from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables" / "20260824_demo2_delivery"
OUTPUT_PATH = OUTPUT_DIR / "RehabAffect_Demo2_阶段整改与后续方向_20260824.docx"
SCREENSHOT_DIR = ROOT / "output" / "playwright" / "demo2_v9_final_release"

NAVY = "17365D"
BLUE = "2E74B5"
INK = "1F2933"
MUTED = "5B6872"
LINE = "D7E0E8"
LIGHT_BLUE = "EDF4FB"
LIGHT_GRAY = "F4F6F8"
LIGHT_GREEN = "EAF4F0"
LIGHT_RED = "FFF1EE"
WHITE = "FFFFFF"
FONT = "Microsoft YaHei"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **edges: dict[str, str]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in edges.items():
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), value)


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        table_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = table_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        table_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            tc_mar = tc_pr.find(qn("w:tcMar"))
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for edge, value in (("top", 90), ("bottom", 90), ("start", 120), ("end", 120)):
                item = tc_mar.find(qn(f"w:{edge}"))
                if item is None:
                    item = OxmlElement(f"w:{edge}")
                    tc_mar.append(item)
                item.set(qn("w:w"), str(value))
                item.set(qn("w:type"), "dxa")


def set_run(run, *, size: float = 10.5, bold: bool = False, color: str = INK, italic: bool = False) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc: Document, text: str, *, bold: bool = False, color: str = INK, size: float = 10.5,
             after: float = 6, before: float = 0, align=WD_ALIGN_PARAGRAPH.LEFT,
             italic: bool = False, keep_next: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.keep_with_next = keep_next
    set_run(paragraph.add_run(text), size=size, bold=bold, color=color, italic=italic)


def add_bullet(doc: Document, text: str, *, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.left_indent = Inches(0.5 if level == 0 else 0.75)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    set_run(paragraph.add_run(text), size=10.2)


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.1
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    set_run(paragraph.add_run(text), size=10.2)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = False
    run = paragraph.add_run(text)
    set_run(run, size={1: 16, 2: 13, 3: 11.5}[level], bold=True, color=BLUE if level < 3 else NAVY)


def add_callout(doc: Document, title: str, body: str, *, fill: str = LIGHT_BLUE, accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, left={"val": "single", "sz": "18", "color": accent},
                    top={"val": "nil"}, bottom={"val": "nil"}, right={"val": "nil"})
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(title), size=10.5, bold=True, color=NAVY)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    set_run(p2.add_run(body), size=10.1, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(header), size=9.5, bold=True, color=NAVY)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "FAFBFC")
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_run(p.add_run(value), size=9.1, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_picture(doc: Document, name: str, caption: str, *, width: float = 6.35) -> None:
    path = SCREENSHOT_DIR / name
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    set_run(cap.add_run(caption), size=8.8, color=MUTED)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(paragraph.add_run("第 "), size=8.5, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    set_run(paragraph.add_run(" 页"), size=8.5, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for level, size, before, after in ((1, 16, 16, 8), (2, 13, 12, 6), (3, 11.5, 8, 4)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE if level < 3 else NAVY)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run("RehabAffect | Demo2 阶段整改与交付"), size=8.5, bold=True, color=MUTED)
    set_cell_border_like_paragraph(p, color=LINE)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def set_cell_border_like_paragraph(paragraph, color: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_cover(doc: Document) -> None:
    add_text(doc, "内部阶段汇报", bold=True, color=BLUE, size=10.5, after=9, before=8)
    add_text(doc, "辅读学校课堂学习证据 Demo2", bold=True, color=NAVY, size=24, after=2)
    add_text(doc, "阶段整改、前端优化、数据更新与后续方向", bold=True, color=INK, size=17, after=14)
    add_table(doc, ["项目", "内容"], [
        ["汇报范围", "2026年8月20日至2026年8月24日"],
        ["当前版本", "Demo2 本地可展示版本"],
        ["交付内容", "汇报文档 + 可运行源码与数据更新包"],
        ["使用对象", "公司内部负责人、学校项目沟通人员、研发人员"],
        ["数据级别", "含学生个体化资料，仅限授权范围内内部使用"],
    ], [1800, 7560])
    add_callout(
        doc,
        "本轮结论",
        "Demo2 已从分散的信息展示页整理为一条可复核的课堂证据链：课程目标进入系统后，关联课堂视频、学生表现、教师支持、后续变化、教师复核和下一步行动。当前适合公司内部展示和学校小范围影子测试，不应表述为已经达到全校自动化投放或真实模型准确率达标。",
        fill=LIGHT_GREEN,
        accent="2F6F60",
    )
    add_heading(doc, "领导可重点查看", 2)
    for item in [
        "第2部分：本次整改要求如何落到页面和功能。",
        "第5部分：本次数据与源文件更新了什么。",
        "第7部分：从演示原型走向学校试点、App与全校部署的路线。",
        "附录：本次压缩包包含内容，以及未放入压缩包的1.44GB原始视频清单。",
    ]:
        add_bullet(doc, item)
    add_text(doc, "编制日期：2026年8月24日", color=MUTED, size=9.2, before=8)
    doc.add_page_break()


def build() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "1. 本轮整改范围与当前产品定位", 1)
    add_text(doc, "本轮工作以学校提出的课程确认、三学生同步分析、师生互动、教师复核、学习掌握判断、课程反馈和个体成长跟踪为主线，同时处理了前端视觉、移动端适配、打印输出和运行性能。")
    add_callout(doc, "当前产品定位", "面向辅读学校的课堂学习证据与个体化支持工作台。系统帮助教师从一节课中定位值得复核的关键时段，整理学生在何种支持下出现何种学习表现，并形成下一节课可执行的支持建议。")
    add_heading(doc, "1.1 当前已经形成的完整流程", 2)
    for item in [
        "课前确认：选择课程，录入或确认课程名称、课程类型、学校标准、考核口径、教学目标和输入资料。",
        "课中证据整理：同步关联学生机与教师机视频，以固定座位ROI和30秒窗口组织三名学生的可见表现。",
        "课后初评：结合课程目标、课堂阶段、学生回应、教师提示和个体化资料形成掌握度与参与度初评。",
        "教师复核：只把关键片段交给教师确认，记录正确性、提示等级、教师调整、学生后续变化和画面外情况。",
        "输出与留档：形成个体页面、状态时间线、学习判断、支持工具、动作建议、课程分析、深度报告、CSV和JSON。",
    ]:
        add_number(doc, item)
    add_heading(doc, "1.2 当前不能夸大的部分", 2)
    add_bullet(doc, "页面上的分数和置信度是演示性初评，不等同于经过校方双人标注验证的模型准确率。")
    add_bullet(doc, "当前不能依据一节课判断学生能力退步、教师教学好坏或医学诊断。")
    add_bullet(doc, "当前仍是离线MP4分析和本地演示，不是希沃数据自动同步，也不是实时全校监控系统。")

    add_heading(doc, "2. 上周四至今的整改内容", 1)
    add_table(doc, ["整改要求", "本轮实现", "当前边界"], [
        ["课前课程与考核标准确认", "七步课程确认入口，一页一个问题，确认后进入工作台", "正式使用还需接入学校课程模板与权限"],
        ["三名学生同步分析", "同一课堂画面中提供群像/个体视图和三学生独立结论页", "身份由校方确认座位ROI，不做人脸或障碍类型推断"],
        ["师生互动细节", "按学生表现、教师支持、学生后续变化组成互动链", "说话人和画面外行为仍需教师确认"],
        ["教师复核标注", "关键片段折叠呈现，保留六步复核和CSV清单", "教师只复核系统筛选片段，不逐帧标注"],
        ["学习掌握判断", "按课程目标、机会次数、独立/提示等级分别展示", "未观察到机会时标记U，不直接判定不会"],
        ["课程反馈与万字报告", "提供约千字课程分析与本地万字深度报告、打印/PDF", "报告仍需教师复核后定稿"],
        ["个体成长方向", "五领域情境趋势和长期追踪入口", "目前是探索性展示，未训练真实长期预测模型"],
    ], [2200, 4200, 2960])
    add_picture(doc, "00_entry_viewport_desktop.png", "图1  七步课程确认入口：先确认目标与评价口径，再进入视频分析。")

    add_heading(doc, "3. 本次前端优化", 1)
    add_text(doc, "前端重构目标不是增加装饰，而是让学校人员更快看懂证据、边界和下一步操作。当前版本使用低饱和白色与矿物色保证阅读稳定；按最终品牌要求，下一轮将把主视觉统一校正为儿童友好、简洁的蓝白体系。")
    add_heading(doc, "3.1 已完成的视觉与交互调整", 2)
    for item in [
        "重建产品头部、课程入口、四阶段证据流程、课程基线、课堂证据区和分析导航。",
        "十四个分析页面使用统一的信息层级、字体、间距、线性图标、按钮和状态表达。",
        "移除Emoji和混杂图标；教师待复核使用唯一强调色，避免用彩色卡片制造虚假确定感。",
        "移动端课堂画面改为全宽16:9；360、390、768、1024、1366和1440像素视口均无页面横向溢出。",
        "所有可见按钮和表单控件不低于40像素；补充键盘导航、ARIA状态和减少动画偏好。",
        "课程分析和深度报告具备打印与A4 PDF输出；JSON限制在可滚动代码区，不再拉长整页。",
    ]:
        add_bullet(doc, item)
    add_picture(doc, "03_workspace_desktop.png", "图2  课堂学习证据台：课程边界、课堂画面、三学生状态和个体证据在同一工作流中。")
    add_heading(doc, "3.2 自动化验收结果", 2)
    add_table(doc, ["验收项", "结果"], [
        ["响应式视口", "1440×1000、1366×768、1024×900、768×900、390×844、360×800"],
        ["横向溢出 / 坏图 / 未命名控件", "0 / 0 / 0"],
        ["小于40px的可见按钮或表单", "0"],
        ["关键文字最低对比度", "4.68:1，通过WCAG AA"],
        ["本地页面传输", "Gzip后约134KB，原始HTML约1.46MB"],
        ["自动测试", "38 passed，5 skipped，无失败"],
    ], [3100, 6260])

    add_heading(doc, "4. 教师真正使用的核心页面", 1)
    add_heading(doc, "4.1 教师复核", 2)
    add_text(doc, "系统不是把每个动作都列出来，而是优先筛出和课程目标、学习机会、教师支持或证据不足有关的关键片段。默认只展开一个片段，其余片段保留课程、学生、时间和复核状态摘要，减少教师浏览负担。")
    add_picture(doc, "06_teacher_review_desktop.png", "图3  教师复核页：时间窗口、系统解释、教师可执行动作与关键片段复核在同一页完成。")
    add_heading(doc, "4.2 学习判断与支持动作", 2)
    add_bullet(doc, "掌握度与参与度分开：安坐、看向材料不直接等于已经学会。")
    add_bullet(doc, "教师提示等级单独记录：独立完成、简单提示、示范或部分协助、完整协助、无法判断。")
    add_bullet(doc, "下一步建议落到可执行任务：缩短指令、视觉提示、等待6秒、替代表达、安排一次独立成功机会。")
    add_bullet(doc, "机器人动作只作为未来接口展示，限制为低刺激、无接触、可关闭并由教师确认。")
    add_picture(doc, "09_support_tools_desktop.png", "图4  个体化支持工具：短指令、等待时间、替代表达、提示顺序和安全约束。")

    add_heading(doc, "5. 本次数据更新与原始资料整理", 1)
    add_heading(doc, "5.1 已纳入分析的数据", 2)
    add_table(doc, ["数据层", "文件/内容", "本次用途"], [
        ["学校原始资料", "2份课程教案、3份学生IEP", "课程目标、考核口径、个体优势、支持需要和安全边界"],
        ["课堂视频", "数学与语文各1组学生机/教师机视频，共4段", "课堂阶段、座位ROI、30秒窗口、师生互动候选证据"],
        ["处理配置", "session_config.json", "课程、学生、ROI、目标、上次掌握度和文件路径"],
        ["分析结果", "analysis.json", "两门课、三名学生、课堂窗口、学习初评和建议动作"],
        ["视觉证据", "两门课群像图和6张学生个体图", "身份复核、页面演示和关键证据定位"],
        ["教师反馈", "teacher_reviews.json", "保存教师对关键片段的复核结果"],
        ["报告输出", "最新TXT与JSON深度报告", "领导查看、学校复盘和后续结构化存储"],
    ], [1750, 3400, 4210])
    add_callout(doc, "隐私说明", "压缩包包含真实学生姓名、IEP摘要、课堂截图和学校原始文档，只能在学校或公司授权范围内使用。不得上传公开网盘、公开代码仓库或对外演示。正式App必须增加角色权限、访问日志、加密存储、脱敏导出和到期删除。", fill=LIGHT_RED, accent="C65B4B")
    add_heading(doc, "5.2 原始视频为何不直接放入微信包", 2)
    add_text(doc, "4段原始视频合计约1.44GB，且属于高敏感课堂数据。本次轻量交付包保留完整视频文件名、大小、原路径和用途说明，但不重复拷贝视频。需要迁移到其他电脑时，应通过公司加密硬盘或学校认可的受控传输方式单独交接。")

    add_heading(doc, "6. 源代码与交付包构成", 1)
    add_table(doc, ["目录", "内容", "负责人如何使用"], [
        ["src/rehab_affect", "Demo2页面、报告生成、本地服务及相关运行模块", "研发继续修改和部署"],
        ["data/processed/seewo_demo2", "两门课三学生处理配置", "替换新课程、新学生和ROI"],
        ["outputs/seewo_demo2", "当前分析结果、图像、教师复核和最新报告", "查看本次演示真实输出"],
        ["source_documents", "3份IEP和2份课程教案", "核对课程与个体支持依据"],
        ["scripts + tests", "启动、UI自动检查和Demo2测试", "复现页面并验证改动"],
        ["design-system + docs", "前端规则、验收记录和说明", "约束后续页面风格与交付口径"],
        ["screenshots", "入口、证据台、教师复核、支持工具、报告和移动端", "不启动程序也可查看当前界面"],
    ], [2100, 3800, 3460])
    add_heading(doc, "6.1 在其他电脑上的启动方式", 2)
    for item in [
        "安装Python 3.11或3.12，并按pyproject.toml安装项目依赖。",
        "将原始视频放到video_manifest.csv记录的对应位置，或修改session_config.json中的路径。",
        "双击start_school_demo2.cmd；程序启动后打开本机地址 http://127.0.0.1:8791/seewo_student_demo2。",
        "首次处理原始视频需要OpenCV；只查看已生成analysis.json时可直接加载当前结果。",
    ]:
        add_number(doc, item)
    add_callout(doc, "本地地址边界", "127.0.0.1只在启动程序的电脑上可访问。若要让校内多台电脑或公网访问，需要部署服务器、域名/HTTPS、账号权限、数据库和隐私安全措施。")

    add_heading(doc, "7. 后续方向与实施优先级", 1)
    add_heading(doc, "7.1 下一轮：形成可稳定试用的蓝白产品版", 2)
    add_table(doc, ["优先级", "工作", "完成标准"], [
        ["P0", "主视觉统一为儿童友好、简洁的蓝白体系", "入口、工作台、14个页面和移动端使用同一品牌规范"],
        ["P0", "真正的课程/视频上传与任务队列", "教师选择课程、上传文件、查看进度、失败可重试"],
        ["P0", "教师标注和审核闭环", "复核可保存、撤回、导出并用于误差统计"],
        ["P0", "学生、课程、片段、复核数据库", "不再依赖单个JSON文件，支持多班、多课和版本记录"],
        ["P0", "隐私与权限", "教师、管理员、研发账号分权；日志、加密和删除流程可审计"],
    ], [1150, 3650, 4560])
    add_heading(doc, "7.2 模型和Agent方向", 2)
    add_bullet(doc, "视觉模型：学生检测/跟踪、姿态与头部朝向、动作时序特征，先做可见行为，不把微表情解释为情绪诊断。")
    add_bullet(doc, "音频模型：VAD、课堂ASR和说话人分离；远场归属不确定时保持U，不强绑定某个学生。")
    add_bullet(doc, "多模态融合：课程阶段、学习机会、视觉线索、音频、教师提示、IEP支持策略和历史记录共同生成human_state与学习证据。")
    add_bullet(doc, "报告Agent：读取结构化证据和教师复核，生成课程报告、个体变化摘要和下一课建议；Agent不能改动原始证据或替代教师结论。")
    add_bullet(doc, "评估：以教师双人复核建立金标准，计算关键片段Precision/Recall、状态一致性、提示等级误差和教师复核耗时。")
    add_heading(doc, "7.3 App、全校与机器人方向", 2)
    add_bullet(doc, "App/小程序优先承担任务提交、进度查看、关键片段复核和报告阅读；大文件上传与模型推理放在Web/服务器端。")
    add_bullet(doc, "全校部署需要班级/学生主数据、课程体系、数据生命周期、GPU任务队列、监控告警和多校区权限，而不是简单复制当前页面。")
    add_bullet(doc, "机器人接口保留屏幕、语音、动作和教师端四类输出；动作由安全策略和教师确认门控制，优先低刺激屏幕/语音反馈，再进入仿真和真机。")
    add_picture(doc, "10_action_plan_desktop.png", "图5  动作方案页：把状态证据转为屏幕、语音、机器人动作和教师端建议。")

    add_heading(doc, "8. 建议的近期安排", 1)
    add_table(doc, ["时间", "建议任务", "主要产出"], [
        ["第1周", "蓝白视觉统一、上传入口、任务状态和失败提示", "可由非研发人员完整走通的产品版"],
        ["第2周", "教师复核数据落库、导出、权限和日志", "学校可执行的影子测试流程"],
        ["第3-4周", "选3名学生、2-4节课进行双人标注和误差分析", "真实Precision/Recall、一致性和复核耗时"],
        ["第5-6周", "视觉/音频基线模型与多模态融合", "可与当前规则基线对比的模型报告"],
        ["第7-8周", "服务器部署、App/小程序复核端和机器人接口仿真", "受控试点版本与下一阶段预算依据"],
    ], [1300, 4300, 3760])
    add_callout(doc, "需要领导确认的三件事", "1）学校是否允许以当前3名学生做2-4节课影子测试；2）原始视频与IEP采用何种授权、加密和删除规则；3）近期交付优先级是学校试点、App/小程序，还是机器人仿真接口。", fill=LIGHT_GREEN, accent="2F6F60")
    add_heading(doc, "8.1 最终判断", 2)
    add_text(doc, "当前版本已经能够清楚展示产品方法和完整业务闭环，也具备继续迭代的代码与数据基础。下一阶段最重要的不是继续堆页面，而是用教师复核数据验证关键片段是否找得准、学习判断是否可解释、教师每节课需要花多少时间，并完成隐私和权限条件。在这些条件成立后，再扩展App、全校系统、Agent和机器人动作接口，路线会更稳。")

    doc.add_page_break()
    add_heading(doc, "附录A：本次交付包清单", 1)
    add_table(doc, ["交付文件", "说明"], [
        ["RehabAffect_Demo2_阶段整改与后续方向_20260824.docx", "本汇报文档"],
        ["RehabAffect_Demo2_源代码与数据更新包_20260824.zip", "代码、配置、处理数据、学校原始文档、测试、截图和说明"],
        ["README_交付说明.md", "启动、目录、隐私边界和原始视频处理方式"],
        ["video_manifest.csv", "4段原始视频的名称、大小、路径和是否纳入压缩包"],
    ], [4300, 5060])
    add_heading(doc, "附录B：本次未重复打包的原始视频", 1)
    add_table(doc, ["课程", "视角", "文件", "大小"], [
        ["得数是3的加法", "学生机", "得数是3的加法 Student D-学生机.mp4", "528.0MB"],
        ["得数是3的加法", "教师机", "得数是3的加法 Student D-教师机.mp4", "275.9MB"],
        ["操场上（第2课时）", "学生机", "3.操场上4.28-学生机.mp4", "463.3MB"],
        ["操场上（第2课时）", "教师机", "3.操场上4.28-教师机.mp4", "199.9MB"],
    ], [1900, 1400, 4260, 1800])
    add_text(doc, "原始视频合计约1.44GB。建议通过加密硬盘或学校认可的受控方式单独交接，不建议通过普通微信群或公开网盘传播。", bold=True, color="9B1C1C", size=9.8, before=4)

    doc.core_properties.title = "RehabAffect Demo2 阶段整改、前端优化、数据更新与后续方向"
    doc.core_properties.subject = "辅读学校课堂学习证据 Demo2 阶段交付"
    doc.core_properties.author = "RehabAffect 项目组"
    doc.core_properties.keywords = "Demo2, 辅读学校, 课堂学习证据, 前端优化, 数据更新, 源代码"
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build())
