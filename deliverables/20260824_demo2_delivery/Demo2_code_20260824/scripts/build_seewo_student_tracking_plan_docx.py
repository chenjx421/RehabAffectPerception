from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "Seewo学生个体状态追踪与轻量训练方案_v2.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(24, 31, 42)
MUTED = RGBColor(92, 102, 112)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_GREEN = "EAF6EF"
PALE_YELLOW = "FFF7E1"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_margins(table, top=80, bottom=80, start=120, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text: str = "", *, size=11, bold=False, color=INK, after=6, before=0,
             align=None, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text: str, *, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_number(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_heading(doc, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color=BLUE)
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=BLUE)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        set_run_font(run, size=12, bold=True, color=DARK_BLUE)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)


def add_simple_table(doc, headers: list[str], rows: list[list[str]], widths_dxa: list[int],
                     header_fill: str = LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    set_cell_margins(table)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=10.5, bold=True, color=INK)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=10, color=INK)
    add_para(doc, "", after=4)
    return table


def add_callout(doc, title: str, body: str, fill: str = PALE_GREEN) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    set_cell_margins(table, top=140, bottom=140, start=180, end=180)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    add_para(doc, "", after=4)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("Seewo 学生个体状态追踪方案")
    set_run_font(r, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("内部方案草案 | 2026-07-01")
    set_run_font(r, size=9, color=MUTED)


def build_doc() -> None:
    doc = Document()
    configure_styles(doc)

    add_para(doc, "方案设计", size=11, bold=True, color=MUTED, after=0)
    add_para(
        doc,
        "基于 Seewo 课堂视频与评课报告的学生个体状态追踪与轻量训练方案",
        size=23,
        bold=True,
        color=RGBColor(0, 0, 0),
        after=4,
    )
    add_para(
        doc,
        "面向辅读学校 AI 助教系统：从单个学生基线建立，到全班长期追踪扩展",
        size=13,
        color=MUTED,
        after=14,
    )
    meta_rows = [
        ["项目定位", "基于课堂视频、声音、Seewo 评课报告和学生测评档案的个体化课堂状态追踪系统"],
        ["第一阶段目标", "先选择 1 名学生，建立个人行为基线，生成课堂状态时间线、趋势报告和动作建议"],
        ["数据来源", "Seewo 下载课堂视频、课堂声音、Word 评课详情、学生匿名测评记录、已有专业评估与行为指导建议"],
        ["训练策略", "低清视频友好的 ROI 行为特征 + 音频特征 + 报告结构特征 + 教师抽样校正"],
        ["边界说明", "不做医学诊断，不承诺单节课训练出泛化大模型，输出仅作为教师观察和干预参考"],
    ]
    add_simple_table(doc, ["项目", "说明"], meta_rows, [1800, 7560], header_fill=LIGHT_BLUE)

    add_callout(
        doc,
        "核心判断",
        "一节 30 多分钟 Seewo 课堂视频足够完成 Demo、建立学生 A 的初始行为基线和生成第一份追踪报告；"
        "但不足以训练强泛化模型。正确路线是先做个体化追踪，再通过多节课、多学生和教师抽样校正逐步训练轻量模型。",
        fill=PALE_GREEN,
    )

    add_heading(doc, "1. 项目主线", 1)
    add_para(
        doc,
        "当前方案建议从“全班状态泛化识别”收敛为“单个学生的个体化课堂状态追踪”。"
        "原因是辅读学校学生差异大、视频清晰度有限、早期样本量少，直接训练通用模型风险高；"
        "先建立一个学生的个人基线，更容易解释、展示和积累真实数据。",
    )
    for item in [
        "从 Seewo 选择一节可下载课堂视频和对应 Word 评课详情，作为第一条真实课程样本。",
        "由老师或项目人员在视频首帧框选目标学生座位区域，系统只追踪该学生。",
        "系统按时间切片，融合视频动作、课堂声音、Seewo 报告字段和测评档案，生成学生状态时间线。",
        "老师只抽样校正确认关键片段，形成可训练标签，不要求学校承担数据处理工作。",
        "当该学生积累 3-5 节课后开始看趋势；积累更多学生后扩展为全班追踪系统。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2. Seewo 可利用数据", 1)
    add_para(
        doc,
        "根据学校提供的 Seewo 后台截图、样例《小雨滴的旅行》课堂智能反馈报告，以及 Seewo 帮助中心说明，"
        "每节课可作为一个独立课程样本导入。系统不依赖学校手工处理数据，而是由我们解析视频、音频和 Word 报告。",
    )
    data_rows = [
        [
            "课堂视频",
            "整节课 MP4、教师机/学生机视角、可能的课堂切片视频",
            "抽帧、学生 ROI 追踪、动作幅度、静止时长、身体偏离、离座倾向",
        ],
        [
            "课堂声音",
            "MP4 内置音轨、教师与学生发言、环境噪声",
            "VAD 语音活动检测、语速/音量、说话人粗分、问题-等待-回答片段；ASR 作为可选增强",
        ],
        [
            "Word 评课详情",
            "课程元数据、课堂时间分配、抬头率、举手率、参与度、应答时长、教师语速、问答模式、提问类型、弗兰德斯指标",
            "提供课堂上下文、弱标签、时间结构和训练特征，辅助解释学生状态变化原因",
        ],
        [
            "学生测评档案",
            "注意力、语言理解、动作协调、感知敏感性、老师观察记录",
            "建立个人基线，判断进步/退步，并个性化调整动作建议",
        ],
        [
            "专业评估与行为指导资料",
            "在学校和家长授权范围内可使用的既有诊断相关建议、康复/心理行为指导、个别化教育计划、禁忌刺激和支持策略",
            "作为学生画像、安全边界和动作策略约束；系统不新增医学诊断，只引用既有建议辅助个体化支持",
        ],
    ]
    add_simple_table(doc, ["数据类别", "可获得内容", "进入模型/数据库的方式"], data_rows, [1500, 3000, 4860], header_fill=LIGHT_BLUE)

    add_heading(doc, "3. 评课报告字段如何进入训练", 1)
    add_para(
        doc,
        "样例 Word 报告显示，Seewo 已经给出大量课堂层面的结构化统计。虽然这些统计不是针对单个学生的精确标签，"
        "但可以成为模型理解课堂阶段、教师节奏和学生整体参与状态的重要上下文。",
    )
    report_rows = [
        ["课程基础信息", "课程名称、学段学科、授课教师、授课时间、课程总时长", "生成 course_id，关联视频、音频、状态观察和测评记录"],
        ["课堂时间分配", "教师讲授、师生问答、学生观察、巩固检测等时长", "判断当前片段属于讲授、问答、练习、观察等场景"],
        ["学生参与指标", "平均抬头率、平均举手率、平均参与度", "作为全班参与背景，辅助判断个体状态是否偏离课堂整体节奏"],
        ["学生应答", "应答时长、回答建构分类", "识别“问题后无反应/反应慢/理解困难”的训练信号"],
        ["教师语音", "讲授字数、平均语速、教师提问和候答时间", "与音频特征融合，判断是否因为语速、长指令或等待时间影响学生状态"],
        ["问答结构", "IRE 问答模式、布鲁姆提问类型、4MAT 提问类型、提问列表", "为状态原因分析提供课堂任务难度和认知负荷背景"],
        ["师生互动", "I/D、PSSR、CCR、PIR、TQR 等指标", "用于课后趋势解释和跨课程对比，不直接替代个体状态标签"],
    ]
    add_simple_table(doc, ["字段组", "样例字段", "使用方式"], report_rows, [1700, 3300, 4360], header_fill=LIGHT_GRAY)

    doc.add_page_break()
    add_heading(doc, "4. 专业评估与行为指导如何进入系统", 1)
    add_para(
        doc,
        "辅读学校学生可能已有来自医院、康复机构、心理行为支持人员或学校个别化教育计划的相关建议。"
        "这些资料对系统很有价值，但属于高度敏感信息，必须在学校和家长授权范围内使用。"
        "系统应把它们作为个体化支持背景和安全约束，而不是作为自动诊断依据。",
    )
    guidance_rows = [
        ["可用内容", "已有诊断相关建议、行为指导、感统/康复建议、个别化教育目标、老师长期观察记录"],
        ["进入方式", "结构化为 support_profile：能力特点、触发因素、有效提示方式、禁忌刺激、适合任务粒度"],
        ["模型作用", "为状态解释提供背景，例如长指令理解困难、声音敏感、需要视觉提示等"],
        ["动作约束", "决定哪些动作可以做、哪些不能做，例如避免突然大声、强光、靠近或复杂连续指令"],
        ["合规边界", "只引用既有授权资料，不让系统生成新的医学诊断；对外展示时使用匿名编号和概括描述"],
    ]
    add_simple_table(doc, ["项目", "说明"], guidance_rows, [1800, 7560], header_fill=LIGHT_BLUE)

    add_heading(doc, "5. 视频清晰度差时的读取策略", 1)
    add_para(
        doc,
        "低清课堂视频不适合做精细表情和眼神识别。第一版应主动避开“微表情情绪识别”，"
        "改做低清条件下更可靠的课堂行为状态追踪。",
    )
    for item in [
        "采用目标学生座位 ROI 追踪，第一版不强制全自动多人识别。",
        "每秒抽 1-3 帧，优先看动作变化、静止持续、身体偏离、趴桌/低头、离座、手部任务动作是否停止。",
        "每 10-30 秒生成一个状态片段，避免单帧误判。",
        "为每个片段记录 video_quality、遮挡、可见身体部位和可判断程度；低质量片段输出 uncertain。",
        "判断标准以该学生自己的历史基线为主，不用普通学生标准硬套辅读学生。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "6. 声音如何进入训练", 1)
    add_para(
        doc,
        "课堂声音必须进入训练，但第一阶段不把准确 ASR 作为刚性前提。教室噪声、多人同时说话和设备拾音位置会影响识别，"
        "因此优先使用更稳的音频特征，再逐步加入 ASR 文本。",
    )
    audio_rows = [
        ["VAD 语音活动", "检测什么时候有人说话、沉默、连续噪声", "识别教师指令后学生是否有回应、是否存在长时间无响应"],
        ["说话人粗分", "教师/学生/多人混声的粗粒度分段", "与 Seewo 问答和候答时间结合，定位问题-等待-回答结构"],
        ["语速与音量", "教师语速、学生回答强弱、音量突变", "分析长指令、过快讲授、嘈杂环境对个体学生状态的影响"],
        ["ASR 文本", "关键词、任务指令、学生回答内容", "作为第二阶段增强；在嘈杂场景下只做辅助证据，不单独决定状态"],
        ["声音事件", "叫喊、持续噪声、突然安静", "辅助判断课堂环境变化和是否需要老师关注"],
    ]
    add_simple_table(doc, ["音频特征", "含义", "训练/分析用途"], audio_rows, [1900, 2700, 4760], header_fill=LIGHT_BLUE)

    add_heading(doc, "7. 样本少时的训练路线", 1)
    add_callout(
        doc,
        "训练原则",
        "第一阶段不训练大模型，而是使用预训练/规则特征提取 + 轻量模型 + 教师抽样校正。"
        "模型目标不是诊断情绪，而是判断学生是否脱离当前课堂任务、可能原因是什么、应该给出什么低刺激引导动作。",
        fill=PALE_YELLOW,
    )
    volume_rows = [
        ["1 节课", "30-40 分钟，切成约 60-240 个片段，但高度相关", "完成 Demo、初始化个人基线、生成第一份课堂状态报告"],
        ["3-5 节课", "同一学生不同课堂", "观察初步趋势，区分该学生的正常参与、停滞和偏离模式"],
        ["10-20 节课", "同一学生跨学科/跨任务", "训练个人状态分类 baseline，评估时间平滑和教师校正效果"],
        ["30-50 节课", "多个学生/一个班级", "训练班级级轻量模型，并加入个体化校准"],
        ["100+ 节课", "多班级、多学科、多教师", "再考虑更强泛化模型和动作推荐策略优化"],
    ]
    add_simple_table(doc, ["数据量", "含义", "合理目标"], volume_rows, [1600, 3100, 4660], header_fill=LIGHT_GRAY)

    add_heading(doc, "8. 训练数据生成流程", 1)
    steps = [
        "导入一节 Seewo 课程：MP4 视频、音轨、Word 评课详情、课程元数据。",
        "老师或项目人员选择目标学生，手动框选座位 ROI，并绑定匿名 student_id。",
        "系统每 10-30 秒切片，抽帧并提取视频特征，同时从音频中提取语音活动和音量/语速特征。",
        "解析 Word 评课报告，提取课堂时间分配、参与度、抬头率、应答时长、问答模式和提问列表等字段。",
        "融合测评档案与既有专业评估/行为指导建议，建立学生个人初始基线和动作安全边界。",
        "系统自动生成状态初判：engaged、low_attention、confused_or_stuck、restless、needs_help、uncertain。",
        "老师只抽样校正不确定片段和关键异常片段，形成弱监督训练集。",
        "训练轻量分类模型，并用滑动窗口/EMA/HMM 做时间平滑。",
        "输出学生状态时间线、进步趋势、可解释证据和屏幕/机器人动作建议。",
    ]
    for step in steps:
        add_number(doc, step)

    add_heading(doc, "9. 数据库设计", 1)
    db_rows = [
        ["students", "student_id、匿名编号、年级、能力画像、注意事项、刺激禁忌", "保存长期个人画像，不存真实姓名"],
        ["assessments", "student_id、测评日期、注意力、语言理解、动作协调、感知敏感、老师备注", "用于个体基线和长期进步/恶化判断"],
        ["support_profiles", "student_id、既有专业建议摘要、触发因素、有效支持方式、禁忌刺激、授权范围", "保存行为指导和个体化教育支持信息，约束动作策略"],
        ["courses", "course_id、课程名、学科、教师、时间、时长、Seewo 来源", "关联每节 Seewo 课堂"],
        ["lesson_reports", "course_id、课堂时间分配、参与指标、语速、问答结构、报告文件路径", "保存 Word 报告解析结果"],
        ["video_clips", "clip_id、course_id、起止时间、ROI、视频质量、音频质量", "训练样本基础单位"],
        ["features", "clip_id、motion_energy、stillness、body_shift、vad、speech_rate、scene_features", "模型输入特征"],
        ["observations", "student_id、clip_id、state、confidence、evidence、teacher_label", "状态识别和教师校正记录"],
        ["interventions", "student_id、clip_id、建议动作、是否采用、后续状态是否改善", "形成动作推荐闭环"],
    ]
    add_simple_table(doc, ["数据表", "关键字段", "作用"], db_rows, [1800, 4200, 3360], header_fill=LIGHT_BLUE)

    add_heading(doc, "10. 模型输出与动作建议", 1)
    add_para(
        doc,
        "最终输出必须落到学生受益和低刺激引导动作上。动作可以先由屏幕或老师端执行，后续再接桌面机器人、实体机器人或 ROS2/WebSocket 接口。",
    )
    action_rows = [
        ["engaged", "继续参与", "保持当前节奏，不额外打断"],
        ["low_attention", "注意力下降", "屏幕轻微高亮当前任务，短句提示“我们先看这里”"],
        ["confused_or_stuck", "疑似困惑/停滞", "重复一步指令，降低任务粒度，必要时提醒老师靠近确认"],
        ["restless", "明显躁动", "减少刺激，暂停强交互，老师端提示关注"],
        ["needs_help", "需要帮助", "老师端推送片段证据和建议，不由系统强制干预"],
        ["uncertain", "无法判断", "标记为需人工确认，不输出强动作"],
    ]
    add_simple_table(doc, ["状态", "解释", "建议动作"], action_rows, [1800, 2100, 5460], header_fill=LIGHT_GRAY)
    add_para(
        doc,
        "动作策略需要读取 support_profile 中的禁忌刺激和有效支持方式。例如学生对突然声音敏感时，系统应优先使用屏幕高亮或老师端提醒；"
        "学生对视觉提示反应更好时，系统可优先推荐图像/步骤卡，而不是连续语音提示。",
        color=INK,
    )

    add_para(doc, "示例 JSON 输出：", bold=True, color=DARK_BLUE, after=4)
    code = """{
  "student_id": "student_A",
  "course_id": "seewo_lesson_001",
  "time_range": [480, 500],
  "state": "confused_or_stuck",
  "confidence": 0.64,
  "evidence": [
    "task_related_hand_motion_stopped",
    "body_orientation_away",
    "teacher_question_waiting_segment"
  ],
  "context": {
    "scene": "teacher_instruction",
    "report_features": ["high_wait_time", "application_question"],
    "audio_features": ["teacher_speech_active", "student_response_absent"]
  },
  "recommended_action": {
    "screen": "highlight_current_step",
    "speech": "我们先看这里，只做第一步。",
    "robot_motion": "slow_turn_to_screen",
    "teacher_alert": true
  }
}"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    for line in code.splitlines():
        r = p.add_run(line + "\n")
        set_run_font(r, size=9, color=RGBColor(20, 40, 55))

    doc.add_page_break()
    add_heading(doc, "11. 第一版交付物", 1)
    deliverable_rows = [
        ["Seewo 导入器", "读取课程 MP4、音频和 Word 评课详情，生成统一 course manifest"],
        ["单学生追踪器", "手动框选目标学生 ROI，生成行为特征和状态时间线"],
        ["学生画像数据库", "SQLite 起步，保存学生匿名档案、测评记录、专业评估/行为指导摘要、课程记录和状态观察"],
        ["轻量训练模块", "基于教师校正片段训练 Logistic Regression / Random Forest / LightGBM baseline"],
        ["学生追踪面板", "展示今日课堂、历史趋势、片段证据、状态解释和动作建议"],
        ["机器人/屏幕接口", "输出 WebSocket/ROS2-ready JSON，支持后续机器人动作执行"],
        ["技术报告", "说明数据来源、模型边界、评估指标、隐私处理和下一步扩展"],
    ]
    add_simple_table(doc, ["交付物", "说明"], deliverable_rows, [2200, 7160], header_fill=LIGHT_BLUE)

    doc.add_page_break()
    add_heading(doc, "12. 开发计划", 1)
    roadmap_rows = [
        ["第 1 周", "Seewo Word 报告解析、课程 manifest、数据库表、Demo 页面改为单学生追踪结构"],
        ["第 2 周", "MP4 读取、抽帧、音频提取、ROI 框选、视频/声音基础特征生成"],
        ["第 3 周", "教师抽样校正工具、轻量 baseline 训练、时间平滑、状态时间线评估"],
        ["第 4 周", "测评档案与行为指导融合、趋势图、动作建议输出、PPT/Word/现场演示材料"],
        ["后续", "多节课累积、多学生扩展、云端批处理、机器人/屏幕执行端联调"],
    ]
    add_simple_table(doc, ["阶段", "重点任务"], roadmap_rows, [1600, 7760], header_fill=LIGHT_GRAY)

    add_heading(doc, "13. 学校侧需要配合的内容", 1)
    for item in [
        "提供 1 节可下载 Seewo 课堂 MP4 和对应 Word 评课详情，优先选择画面中目标学生位置稳定的一节课。",
        "指定 1 名匿名试点学生，提供座位位置和少量基础测评信息；不需要提供真实姓名。",
        "如学校和家长授权，可提供该学生已有的行为指导建议、个别化教育目标、禁忌刺激和有效支持方式；不要求提供完整病历。",
        "允许在校内或授权设备上进行视频处理；如隐私要求较高，原始视频可不离校，只导出匿名特征。",
        "老师每节课抽样确认 5-10 个关键片段，选择“对/不对/不确定”和可能原因。",
        "确认哪些提示方式对学生安全、可接受，哪些声音、光效或动作应避免。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "14. 风险与应对", 1)
    risk_rows = [
        ["视频清晰度低", "不做表情识别，改用 ROI 行为、动作持续性、离座/趴桌/静止等粗粒度特征"],
        ["样本量少", "第一节课只做初始化；通过多节课和教师抽样校正逐步训练个人模型"],
        ["音频嘈杂", "优先 VAD、语速、音量、问答结构，ASR 只作为增强证据"],
        ["教师没有时间标注", "系统只推送不确定和关键异常片段，每节课控制在 5-10 分钟确认量"],
        ["隐私和合规", "匿名 student_id、原始视频可校内处理、专业评估信息只保留必要摘要、不做医学诊断、输出仅作辅助观察"],
        ["敏感资料误用", "专业评估/行为指导只作为安全边界和个体化支持依据，不作为模型自动贴标签或对外展示内容"],
        ["动作误触发", "所有动作低刺激、可关闭；高风险状态优先提醒老师，不直接强干预学生"],
    ]
    add_simple_table(doc, ["风险", "应对策略"], risk_rows, [2200, 7160], header_fill=LIGHT_GRAY)

    add_heading(doc, "15. 结论", 1)
    add_para(
        doc,
        "在当前条件下，最稳的路线不是承诺用少量数据训练通用大模型，而是把 Seewo 课堂视频、课堂声音、评课报告、学生测评档案和既有专业行为指导整合起来，"
        "先完成单个学生的个体化状态追踪。该路线能用一节课启动 Demo，用 3-5 节课形成初步趋势，用更多课程和教师校正逐步训练轻量模型，"
        "最终扩展为全班学生的长期追踪系统，并为屏幕、老师端和机器人动作输出提供可解释依据。",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "附录：Seewo 样例报告可提取字段清单", 1)
    add_para(
        doc,
        "以下字段来自学校提供的《小雨滴的旅行》课堂智能反馈报告样例，后续可开发 Word 报告解析器自动抽取。",
    )
    appendix_rows = [
        ["课程元数据", "课程名、教师、学段学科、授课时间、学校、总时长"],
        ["课程结构", "教师讲授、师生问答、学生观察、巩固检测等时长"],
        ["学生指标", "平均抬头率、平均举手率、平均参与度、提问举手人数、回答人数"],
        ["学习行为", "听讲、阅读、视听、演示、讨论、实践、教给他人"],
        ["声音和问答", "讲授字数、平均语速、核心提问次数、评价次数、应答时长、候答时间"],
        ["认知负荷", "布鲁姆提问类型、4MAT 提问类型、学生回答建构分类"],
        ["教师反馈", "简单肯定、针对肯定、激励、否定、重复等理答类型"],
        ["互动分析", "I/D、PSSR、CCR、PIR、TQR、S-T/Rt-Ch 教学分析"],
        ["提问列表", "每个问题的分类与文本，可用于任务难度和课堂阶段分析"],
        ["专业支持资料", "既有诊断相关建议、行为指导、个别化教育目标、禁忌刺激、有效提示方式、家校沟通注意事项"],
    ]
    add_simple_table(doc, ["字段组", "具体字段"], appendix_rows, [2200, 7160], header_fill=LIGHT_BLUE)

    add_para(
        doc,
        "参考来源：学校提供的 Seewo 后台截图与样例评课详情；Seewo 帮助中心关于 Word 报告下载和课堂实录/切片视频下载的说明。",
        size=9.5,
        color=MUTED,
        after=0,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_doc()
