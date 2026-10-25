from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_seewo_demo2_intro_docx as base


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SHOTS = ROOT / "output" / "playwright" / "v5"
TEACHER_DOC = DOCS / "Demo2_教师关键片段复核与标注操作说明.docx"
APP_DOC = DOCS / "Demo2_App化产品方向与实施条件说明.docx"

# Design preset: compact_reference_guide. First-page pattern: customer_pack.
NAVY = "10233F"
BLUE = "1E40AF"
TEAL = "0F766E"
AMBER = "B45309"
MUTED = "5E6C80"
PALE_BLUE = "EFF6FF"
PALE_TEAL = "ECFDF5"
PALE_GOLD = "FFF7E8"


def set_header_footer(doc: Document, label: str) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    base.set_run_font(header.add_run(f"RehabAffectPerception  |  {label}"), 9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    base.set_run_font(footer.add_run("课堂教育支持原型 · 第 "), 9, color=MUTED)
    base.add_page_field(footer)
    base.set_run_font(footer.add_run(" 页"), 9, color=MUTED)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def add_table(doc: Document, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = str(text)
        shade_cell(cell, "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)
    base.format_table(table, widths, font_size=font_size)
    repeat_header(table.rows[0])
    return table


def add_title_block(doc: Document, kicker: str, title: str, subtitle: str, meta: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(7)
    base.set_run_font(p.add_run(kicker), 9, bold=True, color=TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    base.set_run_font(p.add_run(title), 23, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    base.set_run_font(p.add_run(subtitle), 12.5, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    base.set_run_font(p.add_run(meta), 9.5, color=MUTED)


def add_picture(doc: Document, filename: str, caption: str, width=6.25) -> None:
    path = SHOTS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    inline = p.add_run().add_picture(str(path), width=Inches(width))
    inline._inline.docPr.set("descr", caption)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(8)
    base.set_run_font(cp.add_run(caption), 8.8, color=MUTED, italic=True)


def add_bullet(doc: Document, text: str, prefix: str | None = None) -> None:
    base.add_list_item(doc, text, "List Bullet", bold_prefix=prefix)


def add_numbered_steps(doc: Document, steps: tuple[str, ...]) -> None:
    num_id = base.create_numbering_instance(doc)
    for text in steps:
        prefix = text.split("：", 1)[0] + "：" if "：" in text else None
        base.add_numbered_item(doc, num_id, text, bold_prefix=prefix)


def build_teacher_doc() -> Path:
    doc = Document()
    base.configure_page(doc)
    base.configure_styles(doc)
    set_header_footer(doc, "教师关键片段复核与标注操作说明")
    doc.core_properties.title = "Demo2 教师关键片段复核与标注操作说明"
    doc.core_properties.subject = "系统预筛关键片段后的教师复核SOP、字段口径与定稿规则"
    doc.core_properties.author = "RehabAffectPerception"
    doc.core_properties.comments = "不构成医学诊断、学生能力定级或教师绩效结论。"

    add_title_block(
        doc,
        "校方操作材料 · 教师复核 SOP",
        "Demo2 教师关键片段复核与标注操作说明",
        "系统先找片段，教师补足课堂语义；学校不做算法处理，也不需要逐帧标注",
        f"适用页面：http://127.0.0.1:8791/seewo_student_demo2  |  版本：Demo2 v5  |  日期：{date.today().isoformat()}",
    )
    base.add_callout(
        doc,
        "一句话任务",
        "老师只复核系统筛出的少量关键片段：确认系统描述是否符合、学生回答或操作是否正确、完成时用了哪一级提示，并补充画面外情况。",
        fill=PALE_TEAL,
        border=TEAL,
    )
    add_table(
        doc,
        ("学校需要做", "学校不需要做"),
        (
            ("确认课程、学生、时段和任务；对关键片段作教育语义复核", "不需要剪视频、抽帧、写代码、训练模型或处理音视频"),
            ("按统一口径选结论、正确性和提示等级；必要时写一句补充", "不要求逐帧标注，也不要求给每个普通动作命名"),
            ("对不清楚的内容选择“不确定/U”，不勉强下结论", "不根据外貌、单次表情或障碍名称判断能力、情绪或医学情况"),
        ),
        (4680, 4680),
    )

    base.page_break(doc)
    base.add_heading(doc, "1. 为什么必须由教师复核", 1)
    base.add_para(
        doc,
        "课堂视频只能提供可见动作、姿态、任务活动和部分声音线索。系统无法自动知道画面外是否有人提示、声音来自谁、学生回答是否符合本课题意，也不能仅凭安静、低动作或短暂离开视线判断学生不会。教师复核的作用，是把这些课堂语义补回证据链。",
    )
    for text in (
        "保护学生：证据不足时记为U，不把未观察误写成不会、退步或情绪问题。",
        "保护教师：支持效果指标只复盘本节课的目标、提示和回应，不作为个人绩效评价。",
        "提高报告可信度：约1000字课程分析在关键片段复核后定稿；万字报告保留完整证据和修正记录。",
        "形成训练数据：教师修正可逐步沉淀为高质量标签，但未经授权不得外发或用于其他用途。",
    ):
        add_bullet(doc, text, text.split("：", 1)[0] + "：")
    base.add_heading(doc, "系统会先给老师什么", 2)
    add_table(
        doc,
        ("字段", "示例", "作用"),
        (
            ("课程", "生活数学《得数是3的加法》", "避免脱离本课目标解释动作"),
            ("学生", "Student A / 匿名编号", "确认当前复核对象"),
            ("片段编号", "student_A_clip_01", "用于报告、CSV和后续追溯"),
            ("时间段", "10:00–10:30", "返回原视频核对"),
            ("教学环节", "教师示范 / 个别练习", "理解学生当时应做什么"),
            ("系统判断", "观察到回应或操作线索", "仅描述可见表现，不直接判定学会"),
            ("复核原因", "任务、正确性或提示等级不清", "告诉老师为什么选中该片段"),
        ),
        (1700, 3300, 4360),
        8.9,
    )

    base.page_break(doc)
    base.add_heading(doc, "2. 开始复核前先确认课程", 1)
    base.add_para(
        doc,
        "进入学生结果前，先确认本次课程和学校标准。课程不同，同一个动作的含义可能完全不同；因此课程信息没有确认时，不应形成正式学生结论。",
    )
    add_numbered_steps(
        doc,
        (
            "选择课程：选择与课堂视频对应的课程记录。",
            "确认课程类型与名称：例如生活数学《得数是3的加法》。",
            "核对官方教案与课程标准：写清学生本节课应理解或完成什么。",
            "核对本次执行要求：确认导入、示范、引导练习、个别练习和总结等环节。",
            "核对学校考核标准：统一4级表现量表和U值规则。",
            "填写确认人并勾选边界：确认关键片段仍需教师复核后进入分析。",
        ),
    )
    add_picture(doc, "01_course_entry.png", "图1  课程确认入口：分析必须绑定课程标准、教案要求和学校考核口径")

    base.page_break(doc)
    base.add_heading(doc, "3. 页面内的复核步骤", 1)
    add_numbered_steps(
        doc,
        (
            "进入教师复核：在结果区选择“复核—教师复核”。",
            "确认片段头信息：先看课程、学生、片段编号、时间段和教学环节是否对应。",
            "阅读系统判断与复核原因：系统只描述观察线索，并说明为什么需要教师确认。",
            "必要时定位原时段：点击“定位到完整时间线核对”，结合完整状态窗口和原视频查看上下文。",
            "选择教师结论：在“符合、需修正、不确定”中选择一项。",
            "确认正确性：选择正确、部分正确、不正确、无明确作答机会或无法判断。",
            "确认提示等级：选择4级独立、3级简单提示、2级示范/部分协助、1级完整协助或U无法判断。",
            "补充画面外情况：写明任务、材料、声音归属、教师或同伴提示等关键信息。",
            "保存并继续：点击“保存本片段复核”，完成当前学生后切换另外两名学生。",
            "导出整课清单：点击“导出整课复核清单CSV”，用于教研汇总或校内留档。",
        ),
    )
    add_picture(doc, "04_teacher_review_card.png", "图2  单个关键片段复核卡：定位信息、系统判断、复核原因和四项教师输入")

    base.page_break(doc)
    base.add_heading(doc, "4. 三类结论怎样选择", 1)
    add_table(
        doc,
        ("结论", "什么时候选", "进入报告后的处理"),
        (
            ("符合", "系统对该时段可见表现的描述基本准确", "保留为证据；正确性和提示等级仍按教师选择记录"),
            ("需修正", "系统描述与真实课堂情况不一致，或漏掉关键上下文", "以教师补充说明为准，保留修正记录"),
            ("不确定", "画面、声音、任务或人员归属不足，教师也无法判断", "不计入不会或退步；进入待补证据项"),
        ),
        (1400, 4020, 3940),
    )
    base.add_heading(doc, "正确性选项", 2)
    for text in (
        "正确：学生回应符合本节课当前任务要求。",
        "部分正确：学生只完成部分步骤，或答案包含正确成分但不完整。",
        "不正确：存在明确作答机会，且学生回应不符合当前任务。",
        "无明确作答机会：该时段没有针对该学生的清楚问题或操作机会。",
        "无法判断：声音、画面、材料或任务语义不足。",
    ):
        add_bullet(doc, text, text.split("：", 1)[0] + "：")
    base.add_heading(doc, "提示等级口径", 2)
    add_table(
        doc,
        ("等级", "定义", "常见例子"),
        (
            ("4级 · 独立完成", "给出常规任务后，学生在约定等待时间内独立完成", "无需额外语言、视觉、示范或肢体协助"),
            ("3级 · 简单提示", "在简短语言、关键词、图片或指向提示后完成", "“先看数字卡”、指向材料、流程条提示"),
            ("2级 · 示范/部分协助", "教师示范关键步骤或提供部分操作协助后完成", "先示范一次，再由学生继续；协助摆放部分材料"),
            ("1级 · 完整协助", "在持续示范、逐步带领或完整协助下完成", "教师几乎带领全部步骤"),
            ("U · 无法判断", "证据不足、遮挡、声音归属不清或没有明确机会", "不进入不会、退步或低分统计"),
        ),
        (1800, 3900, 3660),
        8.8,
    )

    base.page_break(doc)
    base.add_heading(doc, "5. 补充说明怎样写才有用", 1)
    base.add_callout(
        doc,
        "写作原则",
        "只补充能改变结论的信息：当时任务、学生实际回应、提示来源、材料变化、声音归属、遮挡原因或异常课堂事件。用可观察事实，不写人格判断和医学推断。",
        fill=PALE_GOLD,
        border=AMBER,
    )
    add_table(
        doc,
        ("推荐写法", "不推荐写法", "原因"),
        (
            ("教师在画面外指向数字卡后，学生选择了3。", "学生不专心。", "前者说明提示与回应；后者是笼统评价"),
            ("声音来自旁边学生，本人没有明确口头回答。", "语音识别错了。", "前者补充说话人归属，便于修正"),
            ("学生先独立摆出1，再经图片提示补充2。", "学生会一点。", "前者能区分独立和提示后完成"),
            ("学生当时在等待材料，没有得到作答机会。", "学生没有反应。", "前者防止把未观察记为不会"),
            ("该片段被教师遮挡，无法判断手部操作。", "看不清。", "前者说明具体证据缺口"),
        ),
        (3720, 2760, 2880),
        8.7,
    )
    base.add_heading(doc, "复核清单的完整字段", 2)
    base.add_para(doc, "定位字段：课程、学生匿名编号/姓名、片段编号、原视频文件、开始秒、结束秒、显示时间段、教学环节、课堂活动。", bold_prefix="定位字段：")
    base.add_para(doc, "判断字段：系统判断、可见证据、复核原因、教师结论、回答/操作正确性、提示等级、补充说明。", bold_prefix="判断字段：")
    base.add_para(doc, "审计字段：复核人、复核时间、系统版本、是否修改、报告定稿状态。正式App中应自动记录审计字段。", bold_prefix="审计字段：")

    base.page_break(doc)
    base.add_heading(doc, "6. 什么时候算复核完成", 1)
    add_table(
        doc,
        ("检查项", "完成标准"),
        (
            ("课程一致", "课程名称、标准、考核口径与视频相符"),
            ("对象一致", "三名学生身份和固定座位已由校方确认"),
            ("片段结论", "每个关键片段已选符合、需修正或不确定"),
            ("学习语义", "有明确机会的片段已确认正确性和提示等级"),
            ("证据缺口", "遮挡、声音归属不清、无机会等均已说明或标U"),
            ("保存状态", "页面显示复核进度完整，必要时已导出CSV"),
            ("报告状态", "约1000字课程分析经教师复核后定稿；深度报告保留修正记录"),
        ),
        (2500, 6860),
    )
    base.add_heading(doc, "建议的校内责任分工", 2)
    for text in (
        "任课教师：确认任务语义、答案正确性和提示等级。",
        "助教或熟悉学生的教师：补充画面外提示、学生表达方式和特殊情况。",
        "教研或项目负责人：抽查复核一致性，确认课程级报告是否可用于校内复盘。",
        "研发人员：完成视频处理、片段预筛、系统维护和数据安全，不代替教师作最终教育判断。",
    ):
        add_bullet(doc, text, text.split("：", 1)[0] + "：")
    add_picture(doc, "03_course_report.png", "图3  一节课、一个学校形成一份约1000字课程分析，复核后定稿")

    base.add_heading(doc, "7. 特殊教育场景的边界与支持工具", 1)
    base.add_para(
        doc,
        "系统面向特殊教育课堂时，重点是尊重学生的表达差异和等待时间。页面保留短指令、6秒等待、替代表达、成功机会、提示顺序和刺激约束等工具，但不自动对学生连续催促，也不主动接触学生。",
    )
    for text in (
        "允许口语、指认、手势、求助卡和实物操作等多种回应，不只认可清晰口语。",
        "提示后至少等待6秒，是否继续、延长、示范或结束由教师决定。",
        "提示后完成必须保留提示等级，不记录为独立完成。",
        "先安排一次可完成的任务，再进入困难目标，减少连续失败。",
        "避免突然提高音量、连续催促、快速靠近和主动身体接触。",
        "任何高影响动作由教师确认，系统可随时暂停。",
    ):
        add_bullet(doc, text)
    add_picture(doc, "05_special_support.png", "图4  个体支持工具：短指令、等待计时、替代表达、成功机会和安全约束")
    base.add_callout(
        doc,
        "使用边界",
        "本系统不构成医学诊断、学生能力定级或教师绩效结论。单节课只能形成当前基线；连续同类任务、统一考核口径和教师复核后，才讨论稳定、进步、退步或知识迁移。",
        fill=PALE_GOLD,
        border=AMBER,
    )

    TEACHER_DOC.parent.mkdir(parents=True, exist_ok=True)
    doc.save(TEACHER_DOC)
    return TEACHER_DOC


def build_app_doc() -> Path:
    doc = Document()
    base.configure_page(doc)
    base.configure_styles(doc)
    set_header_footer(doc, "App化产品方向与实施条件说明")
    doc.core_properties.title = "Demo2 App化产品方向与实施条件说明"
    doc.core_properties.subject = "教师端App、分析服务、学生档案、报告和机器人接口的产品化条件"
    doc.core_properties.author = "RehabAffectPerception"
    doc.core_properties.comments = "产品方向讨论稿，不替代校方的数据、伦理和安全审批。"

    add_title_block(
        doc,
        "产品方向讨论稿 · App化条件",
        "Demo2 App化产品方向与实施条件说明",
        "从本地课堂分析 Demo 走向教师端 App、校内追踪服务与机器人动作接口",
        f"建议载体：教师端平板/电脑 Web App（PWA）+ 校内或合规云端分析服务 + 可选机器人执行端  |  日期：{date.today().isoformat()}",
    )
    base.add_callout(
        doc,
        "推荐的最终载体",
        "第一载体是教师使用的平板或电脑App，完成课程确认、课堂文件导入、三学生复核、报告和跨课追踪；后台负责音视频处理和模型计算；机器人或课堂屏幕是可选执行端，通过WebSocket/ROS2接收低刺激动作，不作为唯一入口。",
        fill=PALE_TEAL,
        border=TEAL,
    )
    add_table(
        doc,
        ("组成", "主要用户", "职责"),
        (
            ("教师端 App / PWA", "任课教师、助教、教研人员", "课程确认、文件导入、片段复核、查看报告与支持卡"),
            ("分析与档案服务", "系统后台", "视频/音频处理、多模态融合、档案更新、报告和审计"),
            ("管理端", "校方授权管理员", "账号、课程、学生匿名编号、授权、保存周期和导出管理"),
            ("具身执行端", "课堂屏幕、桌面机器人或移动机器人", "执行屏幕、语音和动作原语，并回传学生响应"),
        ),
        (2100, 2700, 4560),
    )

    base.page_break(doc)
    base.add_heading(doc, "1. App要解决的完整工作", 1)
    base.add_para(
        doc,
        "App不是把现有页面套进手机壳，而是把“课程建立—文件导入—多模态分析—教师复核—课程报告—学生档案—下一课支持—机器人执行—反馈回流”做成有账号、权限、任务状态和审计记录的产品流程。",
    )
    add_numbered_steps(
        doc,
        (
            "创建课程：填写课程名称、课程标准、本次教案要求和学校考核口径。",
            "选择学生：从授权名单选择本节课重点学生，确认匿名编号和固定座位。",
            "导入资料：上传课堂视频及音频、教案、测评与上次课程记录。",
            "自动分析：后台完成切片、画面/声音质量评估、状态估计、学习机会和关键片段预筛。",
            "教师复核：App推送少量关键片段，教师确认真实性、正确性、提示等级和画面外情况。",
            "形成输出：生成一节课约1000字学校分析、学生个体记录、深度留档和结构化JSON。",
            "执行支持：把低刺激屏幕、语音和动作建议发送给课堂终端或机器人，重要动作需教师确认。",
            "回写档案：保存本课基线、复核、支持效果和下一课目标，形成跨课趋势。",
        ),
    )
    add_picture(doc, "02_three_student_workbench.png", "图1  现有Demo已经具备三学生群像/人像、掌握度、状态和动作建议的统一工作台")

    base.add_heading(doc, "2. 建议的App功能模块", 1)
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    add_table(
        doc,
        ("模块", "核心功能", "最低产品条件"),
        (
            ("课程中心", "课程模板、标准、教案、考核量表、课程版本", "每次分析必须绑定一条课程记录"),
            ("文件导入", "手动上传、批量导入、处理进度、失败重试", "首版支持校方从希沃下载后手动上传；接口接入后再自动同步"),
            ("课堂分析", "三人切换、群像/人像、状态时间线、学习机会", "低清或遮挡时显示证据质量与U值"),
            ("教师复核", "关键片段、四步复核、保存、CSV导出、复核进度", "系统预筛，学校不做逐帧标注"),
            ("课程报告", "约1000字课程分析、学校结论边界、下一课建议", "关键片段复核后标为定稿"),
            ("学生档案", "既有测评、五大领域、课堂记录、趋势和支持偏好", "原始事实、教师结论和模型推断分层保存"),
            ("个体支持", "短指令、等待计时、替代表达、提示顺序、安全约束", "教师控制、低刺激、可关闭、无主动接触"),
            ("机器人接口", "动作审核、发送、回传、终止和接管", "白名单动作、教师确认、急停和审计"),
            ("管理与安全", "账号角色、授权、日志、保留期限、脱敏导出", "每次访问和导出可追溯"),
        ),
        (1800, 4050, 3510),
        8.3,
    )

    base.page_break(doc)
    base.add_heading(doc, "3. 输入、处理和输出条件", 1)
    base.add_heading(doc, "输入条件", 2)
    for text in (
        "课程：课程名称、学科/类型、学校课程标准、本次教案和明确考核口径。",
        "学生：匿名编号、固定座位或教师确认的ROI、上次同类任务表现、可使用的个体教育支持资料。",
        "媒体：校方授权的课堂视频及其中音频；如有教师机视频或独立音轨可提高任务和说话人判断。",
        "复核：至少由熟悉课程和学生的教师处理系统筛出的关键片段。",
        "授权：学校明确数据用途、访问人员、保存位置、保存周期、模型训练是否允许和可否离校。",
    ):
        add_bullet(doc, text, text.split("：", 1)[0] + "：")
    base.add_heading(doc, "媒体质量建议", 2)
    add_table(
        doc,
        ("项目", "推荐条件", "条件不足时的处理"),
        (
            ("画面", "固定视角，优先720p及以上；学生上半身与操作区尽量可见", "仍可分析课堂阶段和较大动作，但细粒度手部、眼神和表情降级"),
            ("帧率", "15 fps及以上即可支持轻量动作与姿态分析", "低帧率增加时间窗口，减少瞬时动作结论"),
            ("声音", "尽量保留原音轨；教师清晰、环境噪声可控", "远场音频不强制绑定学生；说话人不清时记为不确定"),
            ("座位", "重点学生座位相对稳定，由教师确认", "每节课进入前重新确认ROI，不用外貌推断身份或障碍"),
            ("遮挡", "关键学习时段尽量减少教师或同伴长期遮挡", "输出证据质量；遮挡片段不进入能力结论"),
        ),
        (1500, 4090, 3770),
        8.7,
    )
    base.add_heading(doc, "输出条件", 2)
    for text in (
        "一课一校约1000字课程分析：面向学校和任课教师，关键片段复核后定稿。",
        "学生个体记录：掌握初评、状态时间线、提示等级、下一课目标和跨课趋势。",
        "深度报告：约一万字，用于完整证据归档、研发复盘和试点总结。",
        "结构化human_state与动作JSON：供数据库、WebSocket、ROS2和机器人控制端调用。",
        "复核清单：CSV或App内任务，保留课程、学生、片段、教师结论和审计字段。",
    ):
        add_bullet(doc, text)

    base.add_heading(doc, "4. 技术架构与模型条件", 1)
    add_table(
        doc,
        ("层级", "建议组件", "产品职责"),
        (
            ("客户端", "PWA / 平板Web / 桌面Web", "课程配置、上传、复核、报告、支持卡和动作确认"),
            ("媒体服务", "对象存储、转码、切片、音轨提取、质量检测", "统一不同来源视频，生成可追溯时间戳"),
            ("视觉感知", "OpenCV + 人体/姿态/动作轻量模型", "固定ROI、可见性、姿态、动作变化和课堂活动线索"),
            ("语音与文本", "VAD + ASR + 关键词/任务语义；说话人分离按质量启用", "教师指令、课堂阶段、问答机会和音频证据质量"),
            ("多模态融合", "规则/轻量分类器 + 时间平滑 + 置信度校准", "将视频、音频、任务、档案和历史整合为状态与掌握初评"),
            ("报告服务", "证据模板为默认，可选内网/合规LLM整理文字", "保证无LLM也能生成；LLM不得替代证据和教师复核"),
            ("数据层", "关系数据库 + 对象存储 + 版本化知识库", "课程、学生、片段、复核、报告、动作和审计分层保存"),
            ("机器人适配", "WebSocket优先，ROS2 Adapter后续", "发送白名单动作，接收执行状态和学生响应"),
        ),
        (1700, 3600, 4060),
        8.4,
    )
    base.add_callout(
        doc,
        "模型策略",
        "首版不依赖大模型端到端判断。先用可解释的任务机会、轻量视觉/语音线索、时间平滑和教师复核跑通闭环；积累稳定标签后，再训练时序融合模型或引入VLM/VLA做对比，不把模型名称当作产品价值。",
    )

    base.page_break(doc)
    base.add_heading(doc, "5. 算力与部署条件", 1)
    add_table(
        doc,
        ("场景", "建议配置", "能完成的工作"),
        (
            ("教师端浏览与复核", "普通办公电脑/平板，8–16GB内存，稳定浏览器", "课程配置、上传、复核、报告和查看趋势；不承担训练"),
            ("本地轻量离线分析", "4–8核CPU、16GB内存、SSD；可选8–12GB显存GPU", "视频切片、OpenCV、轻量姿态/动作模型、短视频ASR和批量报告"),
            ("校内边缘服务器", "32GB内存、1–2TB加密存储、12–24GB显存GPU", "多课程队列、较快推理、本地存储和统一账号；适合数据不离校"),
            ("合规云端推理/训练", "按需16–24GB显存起步；较大VLM可需24–48GB", "模型微调、批量实验、评估和多学校隔离部署；须先确认数据可上云"),
            ("机器人或课堂终端", "能运行WebSocket客户端的PC/盒子；机器人控制器独立", "接收屏幕、语音和动作原语；高影响动作等待教师确认"),
        ),
        (1800, 3500, 4060),
        8.5,
    )
    base.add_para(
        doc,
        "现阶段结论：普通电脑足以运行页面、规则/轻量分析和复核流程；要批量处理大量视频、训练时序模型或使用较大视觉语言模型，再按授权情况使用校内GPU或云服务器。算力投入应在数据口径和复核流程稳定后扩大。",
        bold_prefix="现阶段结论：",
    )

    base.page_break(doc)
    base.add_heading(doc, "6. 数据库、知识库与学生档案条件", 1)
    base.add_para(
        doc,
        "App需要把可核对事实、教师教育判断、模型推断和支持知识分开保存。混在一个字段里会导致后续无法追踪结论来源，也容易把既往诊断或障碍类型错误地用于能力判断。",
    )
    add_table(
        doc,
        ("数据层", "内容", "约束"),
        (
            ("原始资料层", "课程、教案、授权视频/音频、校方测评原文", "权限最高，保留原始来源和保存期限"),
            ("可观察事实层", "时间段、任务、姿态/动作、ASR文本、证据质量", "只描述观察到的内容，不写医学解释"),
            ("教师复核层", "符合/修正/不确定、正确性、提示等级、补充说明", "保留复核人、时间和修改历史"),
            ("模型推断层", "状态、掌握初评、置信度、风险和动作建议", "保留模型版本，可被教师修正"),
            ("学生追踪层", "跨课目标、趋势、有效支持、刺激约束", "只用于教育支持，趋势必须基于可比任务"),
            ("支持知识库", "短指令、等待、替代表达、提示撤除、安全规则", "不得自动生成医学处置；高影响动作交教师"),
        ),
        (1800, 3900, 3660),
        8.5,
    )
    for text in (
        "学生使用匿名编号；姓名映射由学校单独管理，研发演示尽量不显示真实身份。",
        "五大领域资料不齐时允许缺失，不用其他领域分数推算未记录领域。",
        "病理或医学资料只作为教师和照护者已确认的安全约束来源，不由系统推断诊断。",
        "跨课比较必须记录课程、任务难度、提示口径和证据质量，不能只比一个总分。",
    ):
        add_bullet(doc, text)

    base.page_break(doc)
    base.add_heading(doc, "7. 教师工作流和运营条件", 1)
    add_table(
        doc,
        ("角色", "必须参与的工作", "系统承担的工作"),
        (
            ("任课教师", "确认课程目标、任务语义、正确性和提示等级", "预筛片段、汇总时间线、生成初稿"),
            ("助教/熟悉学生的教师", "补充画面外提示、替代表达和特殊情况", "保存修正并更新档案"),
            ("教研负责人", "抽查口径一致性、批准学校级报告", "统计复核进度、版本和趋势"),
            ("校方管理员", "账号、授权、保存周期和导出审批", "权限控制、日志、数据隔离"),
            ("研发/运维", "模型、服务器、故障、评估和安全维护", "不代替教师作最终教育判断"),
        ),
        (1600, 4060, 3700),
        8.7,
    )
    base.add_heading(doc, "首轮试点必须先确认", 2)
    for text in (
        "是否允许下载课堂视频、是否允许音频、数据是否能离校、是否允许用于模型训练。",
        "三名重点学生由谁确认、固定座位是否稳定、哪类课程先做、怎样算一次明确学习机会。",
        "谁负责关键片段复核、谁批准报告、每节课可接受的复核工作量。",
        "哪些声音、灯光、屏幕和机器人动作对学生不合适，如何一键暂停和教师接管。",
        "试点评价用什么指标：证据可追溯率、复核完成率、报告可用性、教师时间成本和学生支持改善。",
    ):
        add_bullet(doc, text)

    base.page_break(doc)
    base.add_heading(doc, "8. 机器人与具身智能接入条件", 1)
    base.add_para(
        doc,
        "App产品化后，机器人不是用来替代教师判断，而是作为动作执行端。系统把课堂感知和学生档案融合为human_state，再从白名单中选择屏幕、语音和动作原语；教师确认后执行，并把学生后续反应回传，形成感知—状态估计—动作—反馈闭环。",
    )
    add_table(
        doc,
        ("输出通道", "低风险示例", "必须限制"),
        (
            ("屏幕", "单步任务、两个可选答案、视觉流程条、等待倒计时", "避免闪烁、复杂动画和连续弹窗"),
            ("语音", "短句提示、具体表扬、暂停询问", "控制音量、语速和重复次数；不得突然大声"),
            ("机器人动作", "慢速点头、缓慢转向材料、保持位置", "不主动触碰、不快速靠近、不夸张动作"),
            ("教师端", "提醒复核、请求接管、显示动作原因", "教师可拒绝、暂停、改写和急停"),
        ),
        (1600, 4200, 3560),
        8.7,
    )
    base.add_heading(doc, "接口条件", 2)
    for text in (
        "WebSocket优先：适合Windows演示、网页端、屏幕端和机器人Mock快速联调。",
        "ROS2适配：正式机器人接入时发布human_state和action_request，订阅执行结果与安全状态。",
        "动作白名单：action_id映射到经过验证的屏幕、语音和运动参数，禁止模型直接生成底层电机命令。",
        "反馈回流：记录动作是否执行、教师是否否决、学生状态是否改善以及是否需要接管。",
        "仿真验证：先在MuJoCo验证动作选择和安全状态机，再根据硬件条件升级Isaac Sim或真机。",
    ):
        add_bullet(doc, text, text.split("：", 1)[0] + "：")
    add_picture(doc, "05_special_support.png", "图2  教师控制的个体支持工具可直接映射为屏幕、语音与机器人动作原语")

    base.add_heading(doc, "9. 隐私、安全与合规条件", 1)
    for text in (
        "先授权后处理：明确学校、家长/监护人、使用人员、用途、保存周期和退出机制。",
        "最小化收集：只收集当前课程分析必需的数据；能用匿名编号时不显示姓名。",
        "分区存储：原视频、学生档案、模型特征、报告和导出文件设置不同权限与保留期限。",
        "数据不离校优先：校方要求本地保存时，采用校内服务器或边缘设备；上云前重新审批。",
        "全程审计：记录上传、查看、复核、导出、删除、模型版本和机器人动作。",
        "结论边界：不做医学诊断、障碍分类、能力定级、惩罚建议和教师个人绩效排名。",
        "人为接管：所有高影响建议由教师确认，系统提供暂停、撤回和申诉/修正入口。",
    ):
        add_bullet(doc, text, text.split("：", 1)[0] + "：")
    base.add_callout(
        doc,
        "上线门槛",
        "没有明确的数据授权、身份确认、考核口径、教师复核责任和安全动作白名单时，只能作为离线技术演示，不能进入真实课堂自动运行。",
        fill=PALE_GOLD,
        border=AMBER,
    )

    base.page_break(doc)
    base.add_heading(doc, "10. 产品化分阶段路径", 1)
    add_table(
        doc,
        ("阶段", "产品形态", "主要交付", "进入下一阶段条件"),
        (
            ("A. 本地演示", "本地Web Demo", "两节课、三学生、复核、课程报告、深度报告、动作JSON", "校方确认流程和字段有用"),
            ("B. 受控试点", "PWA + 手动上传 + 私有后台", "账号、课程任务、上传队列、复核工作台、审计日志", "授权、角色、保存策略和复核责任明确"),
            ("C. 校内部署", "教师App + 校内服务器", "多班级、多课程、学生档案、模型版本和报表", "视频质量、复核一致性和系统稳定性达标"),
            ("D. 具身接入", "App + 屏幕/机器人", "WebSocket/ROS2、动作白名单、教师确认、反馈回流", "安全评审和仿真验证完成"),
            ("E. 多校扩展", "租户隔离平台", "学校隔离、配置化标准、跨设备运维和模型评估", "不同学校数据和权限完全隔离"),
        ),
        (1400, 1900, 3560, 2500),
        8.2,
    )
    base.add_heading(doc, "试点验收指标", 2)
    for text in (
        "流程：课程确认、文件导入、关键片段复核、课程报告和档案回写能完整跑通。",
        "数据：每条结论能回到课程、学生、片段、系统版本和教师修正。",
        "质量：不确定/U能正确隔离，低清、遮挡和声音归属问题不会被强行解释。",
        "效率：系统筛片后，教师无需观看整段录像即可完成主要复核；实际时间由试点记录。",
        "教育价值：老师能据此调整下一课目标、提示方式和替代表达，而不是只得到状态标签。",
        "具身闭环：至少一类屏幕/语音/动作建议可被安全执行并回传结果。",
    ):
        add_bullet(doc, text, text.split("：", 1)[0] + "：")
    base.add_callout(
        doc,
        "结论",
        "做成App是现实的，但前提不是先堆更大的模型，而是先稳定课程口径、关键片段复核、学生档案分层、隐私权限和动作安全。当前Demo已经覆盖核心业务骨架；产品化新增工作主要集中在上传任务、账号权限、媒体服务、数据库、审计、部署和真实试点评估。",
        fill=PALE_TEAL,
        border=TEAL,
    )

    APP_DOC.parent.mkdir(parents=True, exist_ok=True)
    doc.save(APP_DOC)
    return APP_DOC


if __name__ == "__main__":
    for output in (build_teacher_doc(), build_app_doc()):
        print(output)
