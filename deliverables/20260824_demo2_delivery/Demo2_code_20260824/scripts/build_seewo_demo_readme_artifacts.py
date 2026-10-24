from __future__ import annotations

import base64
import html
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_ROOT / "docs"
SOURCE_MD = DOCS_DIR / "seewo_student_demo_readme_zh.md"
DOCX_OUT = DOCS_DIR / "seewo_student_demo_readme_zh.docx"
HTML_OUT = DOCS_DIR / "seewo_student_demo_readme_zh_embedded.html"


def clean_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    return text.strip()


def add_paragraph_text(document: Document, text: str, *, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    run = paragraph.add_run(clean_inline(text))
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(10.5)


def add_code_block(document: Document, lines: list[str]) -> None:
    if not lines:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.15)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if index < len(lines) - 1:
            run.add_break()


def add_docx_table(document: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for idx, cell_text in enumerate(rows[0]):
        cell = table.rows[0].cells[idx]
        cell.text = clean_inline(cell_text)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(9)
    for row in rows[2:] if _is_separator_row(rows[1]) else rows[1:]:
        cells = table.add_row().cells
        for idx, cell_text in enumerate(row[: len(cells)]):
            cells[idx].text = clean_inline(cell_text)
            for paragraph in cells[idx].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(9)
    document.add_paragraph()


def _is_separator_row(row: list[str]) -> bool:
    return all(set(cell.strip()) <= {"-", ":"} and "-" in cell for cell in row)


def build_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(31, 78, 121)),
        ("Heading 3", 11.5, RGBColor(67, 67, 67)),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = color

    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_docx_table(document, table_rows)
            table_rows = []

    title_done = False
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_table()
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            flush_table()
            continue
        if line.startswith("|") and line.endswith("|"):
            table_rows.append([cell.strip() for cell in line.strip("|").split("|")])
            continue
        flush_table()

        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            alt, rel = image_match.groups()
            image_path = (SOURCE_MD.parent / rel).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(6.6))
            caption = document.add_paragraph(clean_inline(alt))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(10)
            for run in caption.runs:
                run.font.name = "Microsoft YaHei"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(90, 90, 90)
            continue

        if line.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(10)
            run = paragraph.add_run(clean_inline(line[2:]))
            run.bold = True
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor(15, 23, 42)
            title_done = True
            continue
        if line.startswith("## "):
            if title_done:
                document.add_paragraph()
            document.add_heading(clean_inline(line[3:]), level=1)
            continue
        if line.startswith("### "):
            document.add_heading(clean_inline(line[4:]), level=2)
            continue
        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.space_after = Pt(8)
            run = paragraph.add_run(clean_inline(line[2:]))
            run.italic = True
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(31, 78, 121)
            continue
        if re.match(r"^\d+\.\s+", line):
            add_paragraph_text(document, re.sub(r"^\d+\.\s+", "", line), style="List Number")
            continue
        if line.startswith("- "):
            add_paragraph_text(document, line[2:], style="List Bullet")
            continue
        add_paragraph_text(document, line)

    if in_code:
        add_code_block(document, code_lines)
    flush_table()
    document.save(DOCX_OUT)


def md_inline_to_html(text: str) -> str:
    escaped = html.escape(text)
    escaped = re.sub(r"`([^`]+)`", r"<code>\1</code>", escaped)
    escaped = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", escaped)
    return escaped


def image_to_data_uri(path: Path) -> str:
    payload = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:image/png;base64,{payload}"


def build_html() -> None:
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    body: list[str] = []
    in_code = False
    code_lines: list[str] = []
    in_ul = False
    in_ol = False
    table_rows: list[list[str]] = []

    def close_lists() -> None:
        nonlocal in_ul, in_ol
        if in_ul:
            body.append("</ul>")
            in_ul = False
        if in_ol:
            body.append("</ol>")
            in_ol = False

    def flush_table() -> None:
        nonlocal table_rows
        if not table_rows:
            return
        body.append("<table>")
        for row_index, row in enumerate(table_rows):
            if row_index == 1 and _is_separator_row(row):
                continue
            tag = "th" if row_index == 0 else "td"
            body.append("<tr>" + "".join(f"<{tag}>{md_inline_to_html(cell)}</{tag}>" for cell in row) + "</tr>")
        body.append("</table>")
        table_rows = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_table()
            close_lists()
            if in_code:
                body.append("<pre><code>" + html.escape("\n".join(code_lines)) + "</code></pre>")
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            flush_table()
            close_lists()
            continue
        if line.startswith("|") and line.endswith("|"):
            close_lists()
            table_rows.append([cell.strip() for cell in line.strip("|").split("|")])
            continue
        flush_table()
        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            close_lists()
            alt, rel = image_match.groups()
            image_path = (SOURCE_MD.parent / rel).resolve()
            body.append(
                f'<figure><img src="{image_to_data_uri(image_path)}" alt="{html.escape(alt)}">'
                f"<figcaption>{html.escape(alt)}</figcaption></figure>"
            )
            continue
        if line.startswith("# "):
            close_lists()
            body.append(f"<h1>{md_inline_to_html(line[2:])}</h1>")
        elif line.startswith("## "):
            close_lists()
            body.append(f"<h2>{md_inline_to_html(line[3:])}</h2>")
        elif line.startswith("### "):
            close_lists()
            body.append(f"<h3>{md_inline_to_html(line[4:])}</h3>")
        elif line.startswith("> "):
            close_lists()
            body.append(f"<blockquote>{md_inline_to_html(line[2:])}</blockquote>")
        elif re.match(r"^\d+\.\s+", line):
            if not in_ol:
                close_lists()
                body.append("<ol>")
                in_ol = True
            body.append(f"<li>{md_inline_to_html(re.sub(r'^\\d+\\.\\s+', '', line))}</li>")
        elif line.startswith("- "):
            if not in_ul:
                close_lists()
                body.append("<ul>")
                in_ul = True
            body.append(f"<li>{md_inline_to_html(line[2:])}</li>")
        else:
            close_lists()
            body.append(f"<p>{md_inline_to_html(line)}</p>")
    if in_code:
        body.append("<pre><code>" + html.escape("\n".join(code_lines)) + "</code></pre>")
    flush_table()
    close_lists()

    document = f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Seewo 学生个体状态追踪 Demo 截图说明</title>
  <style>
    body {{
      margin: 0;
      background: #f3f6fa;
      color: #172033;
      font-family: "Microsoft YaHei", "Segoe UI", Arial, sans-serif;
      line-height: 1.72;
    }}
    main {{
      max-width: 980px;
      margin: 0 auto;
      padding: 36px 28px 64px;
      background: #fff;
      box-shadow: 0 18px 48px rgba(18, 31, 52, .12);
    }}
    h1 {{ font-size: 32px; line-height: 1.25; margin: 0 0 18px; color: #0f172a; }}
    h2 {{ font-size: 23px; margin: 34px 0 12px; color: #174ea6; border-bottom: 1px solid #dbe3ef; padding-bottom: 6px; }}
    h3 {{ font-size: 18px; margin: 22px 0 8px; color: #243b64; }}
    p {{ margin: 8px 0; }}
    blockquote {{ margin: 16px 0; padding: 12px 14px; border-left: 4px solid #2563eb; background: #eff6ff; }}
    code {{ background: #edf2f7; padding: 1px 5px; border-radius: 4px; }}
    pre {{ background: #111827; color: #e5eef7; padding: 14px; border-radius: 8px; overflow: auto; }}
    table {{ border-collapse: collapse; width: 100%; margin: 12px 0 18px; }}
    th, td {{ border: 1px solid #d8e0ea; padding: 8px 10px; text-align: left; vertical-align: top; }}
    th {{ background: #f1f5f9; }}
    figure {{ margin: 18px 0 28px; }}
    img {{ width: 100%; border: 1px solid #d8e0ea; border-radius: 8px; display: block; }}
    figcaption {{ text-align: center; color: #64748b; font-size: 13px; margin-top: 8px; }}
  </style>
</head>
<body>
<main>
{chr(10).join(body)}
</main>
</body>
</html>
"""
    HTML_OUT.write_text(document, encoding="utf-8")


def main() -> None:
    build_docx()
    build_html()
    print(DOCX_OUT)
    print(HTML_OUT)


if __name__ == "__main__":
    main()
