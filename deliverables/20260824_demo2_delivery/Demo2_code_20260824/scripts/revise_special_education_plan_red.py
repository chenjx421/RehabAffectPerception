from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


INPUT = Path(
    r"C:\Path\To\Developer\Documents\xwechat_files\wxid_2m302j15cyrl22_2ea9\msg\file\2026-06\special_education_embodied_classroom_design_zh.docx"
)
OUTPUT_WECHAT = INPUT.with_name("special_education_embodied_classroom_design_zh_会议后红字修改版.docx")
OUTPUT_PROJECT = (
    Path(__file__).resolve().parents[1]
    / "docs"
    / "special_education_embodied_classroom_design_zh_会议后红字修改版.docx"
)

RED = RGBColor(192, 0, 0)


def set_run_font(run, *, red: bool = False, bold: bool | None = None, size: float | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if red:
        run.font.color.rgb = RED
    if bold is not None:
        run.bold = bold
    if size is not None:
        run.font.size = Pt(size)


def replace_paragraph(paragraph, text: str, *, red: bool = True, bold: bool | None = None) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, red=red, bold=bold)


def append_red_run(paragraph, text: str, *, bold: bool | None = None) -> None:
    run = paragraph.add_run(text)
    set_run_font(run, red=True, bold=bold)


def replace_cell(cell, text: str, *, red: bool = True, bold: bool | None = None) -> None:
    for paragraph in cell.paragraphs:
        paragraph.clear()
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_run_font(run, red=red, bold=bold)


def replace_para_containing(doc: Document, needle: str, text: str, *, red: bool = True) -> None:
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            replace_paragraph(paragraph, text, red=red)
            return
    raise ValueError(f"Paragraph not found: {needle}")


def replace_heading_containing(doc: Document, needle: str, text: str) -> None:
    for paragraph in doc.paragraphs:
        if needle in paragraph.text and paragraph.style.name.startswith("Heading"):
            replace_paragraph(paragraph, text, red=True, bold=True)
            return
    raise ValueError(f"Heading not found: {needle}")


def replace_table_cell_containing(doc: Document, needle: str, text: str, *, red: bool = True) -> None:
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if needle in cell.text:
                    replace_cell(cell, text, red=red)
                    return
    raise ValueError(f"Cell not found: {needle}")


def replace_row_first_cell(doc: Document, first_cell_text: str, row_values: list[str], *, red: bool = True) -> None:
    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == first_cell_text:
                if len(row.cells) < len(row_values):
                    raise ValueError(f"Row has too few cells: {first_cell_text}")
                for cell, value in zip(row.cells, row_values):
                    replace_cell(cell, value, red=red)
                return
    raise ValueError(f"Row not found: {first_cell_text}")


def append_school_questions(doc: Document) -> None:
    heading = doc.add_paragraph(style="Heading 1")
    append_red_run(heading, "17. 需向校方确认的问题（会议后新增）", bold=True)

    intro = doc.add_paragraph()
    append_red_run(
        intro,
        "以下问题用于在与校方线上沟通时确认第一版 Demo 的真实落点，重点围绕学生受益、软件互动、注意力维持、数据边界和设备条件展开。",
    )

    questions = [
        (
            "第一版 Demo 的核心目标应优先落在哪个学生端问题上？",
            "学校更希望先验证“学生更愿意参与互动”“注意力维持时间更长”“分心后更容易回到任务”“生活适应/认知任务理解更好”中的哪一项？是否有最适合试点的一节课、一个年级或一个小组？",
        ),
        (
            "学生最适合使用哪类软件互动任务？",
            "如果第一版使用平板、大屏或网页互动，学校更希望优先做生活适应、认知训练、课堂常规、情境判断还是听指令执行？学生更能接受图片选择、拖拽排序、动画反馈、语音互动还是闯关式游戏？哪些形式可能导致过度兴奋或不适合课堂？",
        ),
        (
            "老师如何判断学生分心、困惑或跟不上？",
            "希望校方举例说明典型表现，例如视线离开、低头、乱动、异常发声、长时间不操作、反复点错、求助、拒绝继续等。还需要区分哪些只是正常波动，哪些需要老师关注，以及是否需要按学生个体建立不同判断标准。",
        ),
        (
            "系统反馈应直接给学生，还是先给老师？",
            "需要确认学校更接受学生端轻提示、教师端提醒、二者同时提示，还是只记录用于课后复盘。也要确认提示频率、持续时间、音量、动画强度、光效限制，以及老师是否必须能够一键关闭所有提示。",
        ),
        (
            "第一版允许使用哪些数据和设备条件？",
            "需确认是否允许摄像头观察头部朝向、坐姿和动作幅度，是否允许麦克风检测音量或异常发声，是否可以保存任务日志或匿名化特征，是否需要家长授权，以及学校现有的大屏、平板、电脑、摄像头、网络条件能支持单人、小组还是整班试用。",
        ),
    ]

    for index, (title, detail) in enumerate(questions, 1):
        p = doc.add_paragraph(style="List Number")
        append_red_run(p, f"{title} ", bold=True)
        append_red_run(p, detail)


def revise_doc() -> None:
    if not INPUT.exists():
        raise FileNotFoundError(INPUT)

    doc = Document(INPUT)

    replace_paragraph(doc.paragraphs[0], "辅读学校 AI 互动助教与学生注意力提升系统设计方案", bold=True)
    replace_paragraph(doc.paragraphs[1], "学生互动任务 → 注意力/学习状态感知 → AI 引导反馈 → 教师复盘")
    replace_paragraph(
        doc.paragraphs[2],
        "产品设计版 | 场景：辅读学校真实课堂 | 第一版策略：软件互动优先、学生受益优先、硬件后置",
    )

    replace_table_cell_containing(
        doc,
        "并把最终决策落到机器人可执行的低刺激动作上",
        "开发一套面向辅读学校课堂的 AI 互动助教系统：通过摄像头、任务表现和教师标注感知学生课堂参与状态，估计学生是否专注、困惑、疲劳、过度活跃、抗拒或需要帮助。第一版优先将反馈落到学生端互动内容、屏幕提示、任务难度调整和教师端提醒上；机器人、智能坐垫等硬件作为后续迭代方向。",
    )

    replace_para_containing(
        doc,
        "本项目定位为辅读学校课堂中的具身智能辅助系统",
        "本项目第一版定位为软件优先的 AI 互动助教系统，核心服务对象是学生，目标是通过更有吸引力的互动任务和轻量状态感知，提升学生课堂参与度和注意力维持时间。教师端主要承担配置、观察和复盘作用，机器人、智能坐垫等硬件设备作为后续增强方向。",
    )
    replace_para_containing(
        doc,
        "系统最终要回答的核心问题是",
        "系统第一版要回答的核心问题是：学生是否愿意参与当前任务？是否出现分心、困惑、停顿或反复错误？系统应通过什么互动内容、提示方式或任务难度调整帮助学生重新进入任务？哪些情况需要教师关注？",
    )

    replace_row_first_cell(
        doc,
        "核心对象",
        ["核心对象", "特殊儿童学生、任课教师、学生端平板/大屏互动界面；机器人和智能坐垫作为后续增强载体。"],
    )
    replace_row_first_cell(
        doc,
        "核心能力",
        ["核心能力", "学生端互动任务、注意力/学习状态感知、AI 引导反馈、教师端复盘。"],
    )
    replace_row_first_cell(
        doc,
        "核心输出",
        ["核心输出", "学生参与状态、注意力/困惑线索、任务表现记录、系统反馈建议、教师关注提示。"],
    )
    replace_row_first_cell(
        doc,
        "展示目标",
        ["展示目标", "形成一个能演示软件闭环的 Demo：学生做互动任务，系统判断参与状态，学生端给出轻量反馈，教师端看到记录和复盘。"],
    )

    replace_table_cell_containing(
        doc,
        "它由学生任务端、教师控制端、多模态状态感知模块、机器人动作决策模块、仿真 Benchmark 和课后报告组成",
        "辅读学校 AI 互动助教系统。它由学生端互动任务、教师控制端、注意力/学习状态感知模块、AI 互动引导策略、课堂数据报告和后续仿真验证组成。第一版先做软件 Demo，硬件不作为初期核心投入。",
    )
    replace_row_first_cell(
        doc,
        "教师控制端",
        ["教师控制端", "班级与学生选择、任务启动、状态提示、人工标注、互动提示方式与强度开关、课后报告。", "教师电脑或平板网页。"],
    )
    replace_row_first_cell(
        doc,
        "机器人决策模块",
        ["AI 互动引导策略模块", "把学生状态转换成软件反馈策略，包含屏幕提示、动画反馈、语音短提示、任务难度调整、教师端提醒和低刺激约束。", "策略服务，支持 WebSocket 接口，后续可接机器人或其他硬件。"],
    )
    replace_row_first_cell(
        doc,
        "仿真 Benchmark",
        ["仿真 Benchmark", "在 MuJoCo 中模拟课堂空间、学生状态变化和反馈策略，用于后续验证而非第一版展示核心。", "可运行仿真脚本、指标报告和场景截图。"],
    )

    replace_para_containing(
        doc,
        "本方案的具身智能重点在",
        "本方案当前阶段的具身智能不等同于必须使用实体机器人，而是强调系统能感知学生在课堂任务中的行为状态，并根据状态改变互动内容、提示方式和教学节奏。第一版优先通过学生端软件互动和教师端提醒形成闭环；实体机器人可作为后续增强载体。",
    )
    replace_row_first_cell(
        doc,
        "Action Suggestion",
        ["AI Feedback Suggestion", "策略根据学生状态、任务阶段、风险等级和教师权限选择反馈方式。", "学生端互动提示、任务难度调整和教师端提醒。"],
    )
    replace_row_first_cell(
        doc,
        "Embodied Action",
        ["Interaction Feedback", "屏幕高亮、动画反馈、语音短提示、分步骤说明、降低任务难度、暂停等待、提醒教师。", "可观察的软件反馈和课堂互动变化，后续可扩展到机器人动作。"],
    )

    replace_para_containing(
        doc,
        "第一版采用“轻量多模态状态估计器 + 行为树机器人策略”",
        "第一版采用“轻量注意力/学习状态估计器 + AI 互动引导策略”。这个选择更适合当前会议共识：先用软件互动证明学生端价值，降低硬件成本和维护风险；后续再把成熟策略迁移到机器人、智能坐垫或更高保真仿真中。",
    )
    replace_row_first_cell(
        doc,
        "动作决策",
        ["互动策略", "规则策略/行为树，显式编码低刺激、教师可控和任务难度调整逻辑。", "后续可升级为个体化推荐、Contextual Bandit 或强化学习策略微调。"],
    )

    replace_para_containing(
        doc,
        "学生端承担两个作用",
        "学生端是第一版方案的核心展示载体：一方面通过更有吸引力的互动任务提升学生参与度，另一方面记录反应时、停顿、错误、求助等行为线索，为注意力和学习状态判断提供依据。",
    )

    replace_heading_containing(doc, "5.4 机器人动作输出", "5.4 AI 互动引导与反馈输出")
    replace_para_containing(
        doc,
        "最终决策必须落在机器人动作上",
        "第一版的最终决策应落在可执行的软件反馈上，而不是强依赖实体机器人。反馈设计遵循低刺激、短时、可中断、教师可控原则，包括屏幕高亮、动画反馈、语音短提示、分步骤提示、降低任务难度、暂停等待和教师端提醒。机器人动作可保留为后续增强方向。",
    )
    replace_cell(doc.tables[8].rows[0].cells[1], "第一版系统反馈", red=True, bold=True)
    for row in doc.tables[8].rows[1:]:
        if row.cells[0].text.startswith("A0"):
            replace_cell(row.cells[1], "保持界面安静，只记录状态，不打断课堂。")
        elif row.cells[0].text.startswith("A1"):
            replace_cell(row.cells[1], "屏幕轻微高亮或动画提示，把注意力指向任务区域。")
        elif row.cells[0].text.startswith("A2"):
            replace_cell(row.cells[1], "出现一句短提示或柔和语音：“我们看第一步”“可以再试一次”。")
        elif row.cells[0].text.startswith("A3"):
            replace_cell(row.cells[1], "将教师预设指令拆成一步一步呈现，减少一次性信息量。")
        elif row.cells[0].text.startswith("A4"):
            replace_cell(row.cells[1], "减少选项、增加示例、切换更简单任务。")
        elif row.cells[0].text.startswith("A5"):
            replace_cell(row.cells[1], "暂停任务，给出“休息一下/换一个”的选择。")
        elif row.cells[0].text.startswith("A6"):
            replace_cell(row.cells[1], "在教师端高亮并请求教师关注，学生端停止继续刺激。")

    replace_para_containing(
        doc,
        "第一版展示不依赖真实学生隐私数据",
        "第一版展示不依赖真实学生隐私数据。系统先用可控的模拟数据、任务交互日志和教师脚本标注跑通“学生互动—状态感知—软件反馈—教师复盘”的闭环，再在获得学校和家长授权后进入试点采集。",
    )
    replace_para_containing(
        doc,
        "第一步使用 MuJoCo 做粗略可行版本",
        "MuJoCo 在本方案中作为后续验证工具，用来模拟课堂空间、学生状态变化和反馈策略效果；第一版 Demo 的重点仍是可直接展示的软件互动体验。仿真用于证明策略可重复评估，为后续机器人或 Isaac Sim 升级预留基础。",
    )

    replace_row_first_cell(
        doc,
        "机器人载体",
        ["机器人载体", "桌面机器人、轮式机器人、机械臂屏幕组合，或仿真机器人。", "后续迭代方向，不作为第一版必需硬件。"],
    )
    replace_table_cell_containing(
        doc,
        "一个可展示的辅读学校具身智能课堂 Demo",
        "第一版最终成品是一个可展示的软件 Demo：学生端互动任务 + 摄像头/任务日志注意力感知 + AI 引导反馈 + 教师端观察复盘。机器人和智能坐垫不作为第一版交付核心，而作为后续硬件增强方向。",
    )
    replace_row_first_cell(
        doc,
        "机器人端",
        ["互动反馈端", "屏幕高亮、动画提示、语音短提示、任务难度调整和教师端提醒；机器人动作作为后续扩展。"],
    )
    replace_row_first_cell(
        doc,
        "接口层",
        ["接口层", "输出 human_state JSON 和 feedback_action JSON，支持 WebSocket，后续预留 ROS2/机器人接入。"],
    )

    replace_row_first_cell(
        doc,
        "robot_action",
        ["feedback_action", "action_type、intensity、target、prompt_text、display_effect、duration、teacher_approval_required", "让学生端和教师端执行具体反馈，后续可映射到机器人动作。"],
    )
    replace_row_first_cell(
        doc,
        "机器人动作",
        ["AI 互动引导策略", "第 4 周", "完成软件反馈策略和教师端提醒，支持 WebSocket/硬件预留接口。", "学生端互动反馈与教师端提醒演示。"],
    )
    replace_row_first_cell(
        doc,
        "闭环完整",
        ["闭环完整", "能从学生任务表现生成状态，再触发学生端互动反馈或教师端提醒，并在教师端留下记录。"],
    )
    replace_row_first_cell(
        doc,
        "具身相关",
        ["学生受益", "系统能通过互动内容、提示方式和任务难度调整，帮助学生更愿意参与课堂或在分心后回到任务。"],
    )
    replace_para_containing(
        doc,
        "本方案把重点放在辅读学校课堂中的具身智能闭环",
        "根据会议讨论，本方案调整为“软件优先、学生受益优先、硬件后置”的第一版 Demo 路线：先证明 AI 互动内容对学生参与度和注意力的价值，再考虑机器人、智能坐垫等硬件增强。系统通过学生端互动任务、轻量状态感知、AI 引导反馈和教师端复盘形成闭环。",
    )

    append_school_questions(doc)

    OUTPUT_WECHAT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PROJECT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_WECHAT)
    shutil.copyfile(OUTPUT_WECHAT, OUTPUT_PROJECT)
    print(OUTPUT_WECHAT)
    print(OUTPUT_PROJECT)


if __name__ == "__main__":
    revise_doc()
