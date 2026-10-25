from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WECHAT_DIR = Path(
    r"C:\Path\To\Developer\Documents\xwechat_files\wxid_2m302j15cyrl22_2ea9\msg\file\2026-06"
)
OUTPUT_PROJECT = PROJECT_ROOT / "docs" / "special_education_ai_tutor_problem_oriented_upgrade_zh.docx"
OUTPUT_WECHAT = WECHAT_DIR / "special_education_ai_tutor_problem_oriented_upgrade_zh.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(30, 38, 45)
MUTED = RGBColor(95, 105, 115)
RED = RGBColor(192, 0, 0)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
RED_FILL = "FCE4E4"


def set_run_font(run, *, size=None, bold=None, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("辅读学校 AI 互动助教方案 - 问题导向升级版")
    set_run_font(r, size=9, color=MUTED)
    return doc


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("辅读学校 AI 互动助教与学生注意力提升系统设计方案")
    set_run_font(r, size=22, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("问题导向升级版 | 学生具体问题 -> 软件互动解决方案 -> 教师复盘 -> 学校低成本试点")
    set_run_font(r, size=11, color=RED, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("本版根据会议反馈升级：不再以机器人能力展示为主线，而是明确第一版具体解决学生、老师、学校的哪些问题。")
    set_run_font(r, size=10.5, color=RED)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_para(doc: Document, text: str, *, red=False, bold=False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, color=RED if red else INK, bold=bold)


def add_bullet(doc: Document, text: str, *, red=False) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, color=RED if red else INK)


def add_number(doc: Document, text: str, *, red=False) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, color=RED if red else INK)


def add_callout(doc: Document, title: str, body: str, *, red=False) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, RED_FILL if red else CALLOUT_FILL)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, bold=True, color=RED if red else DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2, color=RED if red else INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], *, red_body=False) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, bold=True, color=INK)

    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            if i == 0 and len(value) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, color=RED if red_body else INK)

    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def build_doc() -> None:
    doc = setup_document()
    add_title(doc)

    add_callout(
        doc,
        "升级后一句话方案",
        "第一版先做一个低成本、可进课堂试用的软件 Demo：学生在平板/大屏上完成互动任务，系统通过任务行为、摄像头和教师标注判断参与度、分心、困惑和求助需求，再通过屏幕反馈、语音短提示、任务降难和教师端提醒帮助学生重新进入任务。机器人、智能坐垫等硬件只作为后续增强方向。",
        red=True,
    )

    add_heading(doc, "1. 本版核心调整")
    add_table(
        doc,
        ["原方案问题", "升级后处理"],
        [
            ["落点偏技术", "先写清楚学生、老师、学校各自要解决的具体问题，再映射功能。"],
            ["硬件感太重", "第一版只要求电脑/平板或大屏/摄像头，机器人和智能坐垫后置。"],
            ["学生受益不够具体", "把目标拆成参与、理解、注意力回到任务、减少挫败、表达求助五类。"],
            ["教师价值不够明确", "把教师端定位为观察、低打扰提醒、复盘和备课辅助，而不是替代老师管理。"],
            ["学校难判断是否值得做", "加入可量化试点指标：参与率、回到任务率、任务完成率、教师认可度、部署成本。"],
        ],
        [2200, 7160],
        red_body=True,
    )

    add_heading(doc, "2. 第一版具体解决学生什么问题")
    add_table(
        doc,
        ["学生具体问题", "第一版功能怎么解决", "Demo 怎么呈现", "验证指标"],
        [
            [
                "不愿意参与，容易坐着发呆",
                "把课堂任务做成短任务、即时反馈和轻量闯关，让学生每 30-60 秒有一次可完成的小目标。",
                "学生端显示“听指令找物品/分类/场景选择”任务，完成后出现温和动画和正向反馈。",
                "主动点击次数、任务开始率、连续参与时长。",
            ],
            [
                "听不懂多步骤指令",
                "把老师指令拆成一步一步的图片/语音提示，学生卡住时只显示下一步，不一次给太多信息。",
                "同一任务提供普通模式和拆步模式，学生连续错误后自动切换。",
                "连续错误次数下降，提示后完成率提升。",
            ],
            [
                "分心后很难回到任务",
                "通过无操作时长、视线偏离、反复点错等线索判断可能分心，用低刺激方式拉回。",
                "学生 8 秒无操作后，屏幕轻微高亮任务区域；教师端同步显示“可能分心”。",
                "提示后 15 秒内恢复操作比例。",
            ],
            [
                "反复失败后容易挫败或抗拒",
                "系统降低难度、减少选项、给示例，不用反复催促学生。",
                "错误 2 次后从四选一降为二选一，并给出“我们一步一步来”的提示。",
                "放弃次数、抗拒标注次数、降难后完成率。",
            ],
            [
                "不会表达自己需要帮助",
                "提供简单求助按钮和情绪/需求选择，帮助学生用可视化方式表达“不会/想休息/找老师”。",
                "学生端固定提供“我不会”“再说一遍”“休息一下”“找老师”四个按钮。",
                "主动求助次数、教师响应次数、无效等待时长下降。",
            ],
        ],
        [1700, 3200, 2800, 1660],
        red_body=True,
    )

    add_heading(doc, "3. 第一版具体解决老师什么问题")
    add_table(
        doc,
        ["老师具体问题", "第一版功能怎么解决", "老师能看到什么"],
        [
            ["一个班约 10 人，个体差异大，老师顾不过来", "教师端显示每个学生的任务进度、停顿、错误、求助和可能分心状态。", "学生状态列表：正常参与、可能分心、连续错误、需要帮助。"],
            ["不知道学生到底卡在哪一步", "自动记录每个任务步骤的错误点、停顿时长和提示后是否恢复。", "时间线：第几步卡住、错了几次、用了什么提示、是否恢复。"],
            ["课堂提醒容易打断教学", "系统默认低刺激，且由老师控制提示强度；必要时只提醒老师，不直接打扰学生。", "提示模式：只记录、学生轻提示、教师端提醒、同时提醒。"],
            ["备课和分层困难", "任务模板支持难度分级，老师可选低/中/高难度和图片素材。", "任务包配置：选主题、选难度、选提示方式。"],
            ["课后缺少复盘证据", "自动生成课后报告，用数据说明学生参与、卡住点、帮助点和下次建议。", "报告：完成率、错误类型、提示次数、求助次数、建议下次难度。"],
        ],
        [2100, 4300, 2960],
        red_body=True,
    )

    add_heading(doc, "4. 第一版具体解决学校什么问题")
    add_table(
        doc,
        ["学校关注点", "方案回应", "为什么更容易试点"],
        [
            ["预算和维护压力", "第一版不采购机器人、不上智能坐垫，优先使用现有电脑、平板/大屏和普通摄像头。", "成本低、部署快、坏了也不影响课堂主体。"],
            ["担心硬件损耗和安全", "硬件后置，系统不接触学生，不快速移动，不制造强刺激。", "先软件验证，降低安全和维护风险。"],
            ["担心效果说不清", "设定可量化指标：参与时长、任务完成率、回到任务率、教师认可度。", "试点 1-2 节课即可形成初步数据。"],
            ["隐私和授权顾虑", "第一版优先保存任务日志和匿名化特征，不保存可识别原始视频音频。", "更容易通过学校和家长沟通。"],
            ["后续能否扩展", "任务模板可替换，先做一个课堂任务包，后续扩展到生活适应、课堂常规、出行安全等。", "从一个小场景开始，逐步复制。"],
        ],
        [1800, 4300, 3260],
        red_body=True,
    )

    add_heading(doc, "5. 第一版 MVP 功能范围")
    add_table(
        doc,
        ["模块", "第一版必须做", "暂不作为第一版重点"],
        [
            ["学生端互动任务", "3 个可点击任务：听指令、物品分类、情境选择；每个任务支持简单/普通两档难度。", "复杂三维场景、实体机械臂、多人竞技。"],
            ["注意力/学习状态感知", "任务日志 + 摄像头轻量线索 + 教师标注，判断参与、分心、困惑、求助。", "医学诊断、复杂情绪识别、大模型实时视频分析。"],
            ["AI 引导反馈", "屏幕高亮、动画反馈、语音短提示、拆步提示、降难、教师端提醒。", "实体机器人主动干预、智能坐垫震动提醒。"],
            ["教师端复盘", "实时状态卡片、学生时间线、课后报告。", "复杂教务系统、长期成长档案平台。"],
            ["展示与验证", "一条完整课堂 Demo 链路和一份试点指标表。", "大规模真实数据训练。"],
        ],
        [1700, 4400, 3260],
        red_body=True,
    )

    add_heading(doc, "6. 近期 Demo 场景")
    add_callout(
        doc,
        "推荐 Demo 场景：课堂物品整理与听指令任务",
        "老师发出简单指令，学生在平板或大屏上完成“找物品、分类、排序、情境选择”。系统记录学生是否长时间无操作、是否连续错误、是否主动求助；当学生分心或困惑时，学生端给出低刺激提示，教师端同步记录并生成课后报告。",
        red=True,
    )
    add_number(doc, "老师选择任务包：课堂物品整理，难度选择普通。", red=True)
    add_number(doc, "学生完成 5 个短任务，例如“把文具放进书包”“找到上课需要的东西”。", red=True)
    add_number(doc, "系统模拟或接入摄像头线索，识别长时间无操作、连续错误、视线偏离或求助。", red=True)
    add_number(doc, "学生端触发低刺激反馈：高亮任务区域、给出下一步提示、降低选项数量。", red=True)
    add_number(doc, "教师端显示学生状态和时间线，不强制老师立即打断课堂。", red=True)
    add_number(doc, "课后生成报告：完成率、卡住步骤、提示次数、恢复情况、下次难度建议。", red=True)

    add_heading(doc, "7. 关键产品流程")
    add_table(
        doc,
        ["流程", "输入", "系统处理", "输出"],
        [
            ["学生做任务", "点击、拖拽、选择、停顿、求助", "记录行为日志", "lesson_event"],
            ["状态感知", "任务日志、摄像头轻量线索、教师标注", "融合判断参与/分心/困惑/求助", "student_state"],
            ["反馈决策", "student_state、任务阶段、教师设置", "选择低刺激反馈策略", "feedback_action"],
            ["课堂执行", "学生端或教师端", "显示提示、降难、提醒老师", "学生恢复或教师介入"],
            ["课后复盘", "整节课事件流", "聚合指标和卡点", "session_report"],
        ],
        [1300, 2600, 2900, 2560],
        red_body=True,
    )

    add_heading(doc, "8. 需要的硬件与部署条件")
    add_table(
        doc,
        ["设备", "第一版要求", "说明"],
        [
            ["电脑", "1 台普通 Windows 笔记本即可", "运行本地服务、教师端和 Demo。"],
            ["学生端屏幕", "平板、大屏或可触控网页", "承载学生互动任务，是第一版核心载体。"],
            ["摄像头", "普通 USB 摄像头，可选但建议", "只做轻量头部朝向/坐姿/动作幅度观察。"],
            ["麦克风", "可选", "仅用于音量或异常发声线索，不作为第一版必要条件。"],
            ["机器人/智能坐垫", "后续迭代，不作为第一版必需", "待软件 Demo 证明价值后再评估投入。"],
        ],
        [1600, 3600, 4160],
        red_body=True,
    )

    add_heading(doc, "9. 试点验证指标")
    add_table(
        doc,
        ["对象", "指标", "判断方式"],
        [
            ["学生", "主动参与率", "学生是否愿意开始任务，是否持续点击/选择/拖拽。"],
            ["学生", "回到任务率", "分心或停顿后，提示 15 秒内是否重新操作。"],
            ["学生", "任务完成率", "完成任务数量和降难后完成情况。"],
            ["学生", "求助表达次数", "学生是否能通过按钮表达“不会/休息/找老师”。"],
            ["老师", "有效提醒率", "老师认为系统提示是否真的需要关注。"],
            ["老师", "复盘有用度", "课后报告是否能帮助老师知道学生卡在哪。"],
            ["学校", "部署成本", "是否能用现有设备跑起来，是否不增加维护负担。"],
        ],
        [1200, 2200, 5960],
        red_body=True,
    )

    add_heading(doc, "10. 后续硬件与具身智能升级路径")
    add_para(
        doc,
        "硬件不是取消，而是后置。第一版先验证学生端软件互动是否能提高参与和注意力；如果学校认可，再把已验证的 feedback_action 映射到机器人、桌面互动设备、智能坐垫或 Isaac Sim 仿真中。",
        red=True,
    )
    add_table(
        doc,
        ["阶段", "重点", "输出"],
        [
            ["第一阶段", "软件 Demo：学生互动任务 + 状态感知 + AI 反馈 + 教师复盘", "能在学校演示和小范围试用"],
            ["第二阶段", "加入摄像头/麦克风更稳定的状态线索", "提高判断准确性和可解释性"],
            ["第三阶段", "接入桌面机器人或学校现有硬件", "把软件反馈扩展成具身反馈"],
            ["第四阶段", "MuJoCo/Isaac Sim 做策略 Benchmark", "证明反馈策略在不同学生状态下可重复评估"],
        ],
        [1400, 4700, 3260],
        red_body=True,
    )

    add_heading(doc, "11. 需向校方确认的问题")
    questions = [
        (
            "第一版 Demo 的核心目标应优先落在哪个学生端问题上？",
            "学校更希望先验证学生更愿意参与互动、注意力维持时间更长、分心后更容易回到任务，还是生活适应/认知任务理解更好？是否有最适合试点的一节课、一个年级或一个小组？",
        ),
        (
            "学生最适合使用哪类软件互动任务？",
            "学校更希望优先做生活适应、认知训练、课堂常规、情境判断还是听指令执行？学生更能接受图片选择、拖拽排序、动画反馈、语音互动还是闯关式游戏？哪些形式可能导致过度兴奋或不适合课堂？",
        ),
        (
            "老师如何判断学生分心、困惑或跟不上？",
            "希望校方举例说明典型表现，例如视线离开、低头、乱动、异常发声、长时间不操作、反复点错、求助、拒绝继续等。还需要区分哪些只是正常波动，哪些需要老师关注，以及是否需要按学生个体建立不同判断标准。",
        ),
        (
            "系统反馈应直接给学生，还是先给老师？",
            "需要确认学校更接受学生端轻提示、教师端提醒、二者同时提示，还是只记录用于课后复盘。也要确认提示频率、持续时间、音量、动画强度、光效限制，以及老师是否必须能够一键关闭所有提示。",
        ),
        (
            "第一版允许使用哪些数据和设备条件？",
            "需确认是否允许摄像头观察头部朝向、坐姿和动作幅度，是否允许麦克风检测音量或异常发声，是否可以保存任务日志或匿名化特征，是否需要家长授权，以及学校现有的大屏、平板、电脑、摄像头、网络条件能支持单人、小组还是整班试用。",
        ),
    ]
    for title, detail in questions:
        add_bullet(doc, f"{title} {detail}", red=True)

    add_heading(doc, "12. 结论")
    add_callout(
        doc,
        "升级版结论",
        "本方案第一版不追求昂贵硬件形态，而是先用软件互动解决真实课堂中的具体问题：让学生更愿意参与、在分心后更容易回到任务、在困惑时获得拆步帮助，并让老师看到学生卡点和课后数据。只有当软件 Demo 证明学生端价值后，再推进机器人、智能坐垫或 Isaac Sim 等硬件/仿真升级。",
        red=True,
    )

    OUTPUT_PROJECT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PROJECT)
    if WECHAT_DIR.exists():
        shutil.copyfile(OUTPUT_PROJECT, OUTPUT_WECHAT)
    print(OUTPUT_PROJECT)
    if WECHAT_DIR.exists():
        print(OUTPUT_WECHAT)


if __name__ == "__main__":
    build_doc()
