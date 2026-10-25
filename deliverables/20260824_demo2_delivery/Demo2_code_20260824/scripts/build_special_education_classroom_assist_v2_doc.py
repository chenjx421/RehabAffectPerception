from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "docs" / "special_education_classroom_assist_design_v2_zh.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
DARK = RGBColor(30, 38, 45)
MUTED = RGBColor(95, 105, 115)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"


def set_run_font(run, size=None, bold=None, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("辅读学校课堂辅助与生活适应任务系统设计方案")
    set_run_font(r, size=9, color=MUTED)
    return doc


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 5)
    run = p.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_para(doc: Document, text: str, *, bold=False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, bold=bold)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r)


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, bold=True, color=DARK)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            if i == 0 and len(value) <= 10:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def build_doc() -> None:
    doc = setup_document()

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    r = title.add_run("辅读学校课堂辅助与生活适应任务系统设计方案")
    set_run_font(r, size=22, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sr = subtitle.add_run("产品设计版 | 课堂状态辅助观察 + 生活适应任务工具箱 + 低刺激互动")
    set_run_font(sr, size=11, color=MUTED)

    add_callout(
        doc,
        "方案定位",
        "本方案设计一套可复用于多类生活适应课程的课堂辅助系统：帮助教师观察学生状态、组织低刺激互动，并通过任务工具箱训练学生的生活自理和社会适应能力。",
    )

    add_heading(doc, "1. 项目背景与目标", 1)
    add_para(
        doc,
        "辅读学校学生年龄跨度大、能力差异明显。同一课堂中，学生对教师指令、课程内容和生活概念的理解程度可能不同，教师需要同时完成教学组织、状态观察、行为引导和个别化支持。",
    )
    add_para(
        doc,
        "系统目标是围绕真实课堂开发一个“课堂辅助观察 + 生活适应任务工具箱”。任务内容可以覆盖整理物品、认识校园、点餐、乘车、情境选择、社交礼仪和安全规则训练等生活适应场景。",
    )

    add_heading(doc, "2. 要开发什么", 1)
    add_callout(
        doc,
        "拟开发产品",
        "辅读学校课堂辅助与生活适应任务系统。系统由学生端任务界面、教师端观察面板、低刺激提示机制、课后报告和可选机器人互动组成。",
    )
    add_table(
        doc,
        ["组成部分", "开发内容", "作用"],
        [
            ["学生端任务", "在平板或大屏上完成生活适应任务，例如分类、匹配、情境选择、步骤排序、简单购物、整理物品。", "让学生在可控课堂环境中练习真实生活能力。"],
            ["教师端面板", "显示学生任务进度、卡住次数、提示次数、教师标注和课后报告。", "帮助教师观察和复盘，不替代教师判断。"],
            ["状态辅助观察", "根据任务停顿、连续错误、无反应、异常发声、教师标注等信息判断课堂状态。", "优先做轻量辅助观察，降低技术和合规难度。"],
            ["低刺激提示", "用屏幕、语音或可选机器人进行简短、柔和、可关闭的提示。", "帮助学生回到任务，不制造惊吓或过度兴奋。"],
            ["课程任务包", "把生活适应课拆成多个可复用任务包。", "支持公司逐步扩展产品，而不是一次做大系统。"],
        ],
        [1600, 4700, 3060],
    )

    add_heading(doc, "3. 产品建设原则", 1)
    add_table(
        doc,
        ["原则", "设计要求", "目的"],
        [
            ["教师主导", "系统只提供观察、提示、记录和复盘，最终判断由教师完成。", "保证系统是课堂辅助工具。"],
            ["任务驱动", "围绕生活适应任务组织学生操作、状态观察和反馈。", "让功能服务于真实课堂目标。"],
            ["低刺激", "提示方式保持简短、柔和、可关闭、可调节。", "降低学生受惊或过度兴奋的风险。"],
            ["可扩展", "任务包可以逐步扩展到整理、校园、购物、出行、安全等场景。", "避免产品只绑定单一课程内容。"],
            ["数据谨慎", "优先记录任务日志和教师标注，涉及音视频时走授权流程。", "保护学生隐私和学校合规边界。"],
        ],
        [1500, 4300, 3560],
    )

    add_heading(doc, "4. 核心功能设计", 1)
    add_heading(doc, "4.1 学生端生活适应任务工具箱", 2)
    add_para(doc, "学生端是第一版最重要的展示载体。它应像一个可配置的课堂任务工具，而不是单一小游戏。")
    add_table(
        doc,
        ["任务包", "示例任务", "训练目标"],
        [
            ["物品整理", "把书、本子、铅笔放入正确位置；判断上课需要带什么。", "训练分类、顺序和生活自理。"],
            ["认识校园", "选择教室、厕所、食堂、医务室等场景；判断遇到问题找谁。", "训练空间理解和求助意识。"],
            ["情境选择", "看到同学摔倒、老师提问、排队插队等场景，选择合适行为。", "训练社会规则和课堂行为。"],
            ["简单购物", "选择商品、比较数量、确认需求和完成模拟购买流程。", "训练数量理解和生活消费概念。"],
            ["出行安全", "红绿灯、斑马线、等车、上车刷卡等步骤排序。", "训练公共生活安全。"],
        ],
        [1600, 4800, 2960],
    )

    add_heading(doc, "4.2 教师端观察面板", 2)
    add_para(doc, "教师端不做复杂诊断，而是把课堂过程中本来分散的信息变成可查看、可复盘的记录。")
    add_bullet(doc, "实时查看学生是否完成任务、是否连续错误、是否长时间无操作。")
    add_bullet(doc, "教师可以一键标注：注意力下降、困惑、抗拒、需要帮助、异常发声。")
    add_bullet(doc, "系统自动记录提示次数、教师介入次数和任务完成情况。")
    add_bullet(doc, "课后生成简洁报告，用于教学复盘和家校沟通参考。")

    add_heading(doc, "4.3 课堂状态辅助观察", 2)
    add_para(doc, "第一版建议从任务行为和教师标注开始，不急于依赖真实摄像头算法。这样更适合早期进入学校试点，也更容易处理隐私问题。")
    add_table(
        doc,
        ["状态", "第一版线索", "建议响应"],
        [
            ["专注参与", "任务连续完成、反应及时、错误少。", "继续任务，给予正向反馈。"],
            ["注意力下降", "长时间无操作、反复点错、教师标注注意力下降。", "屏幕轻提示或切换更简单任务。"],
            ["困惑", "连续错误、反复请求帮助、停顿时间长。", "分步骤提示，减少一次性指令。"],
            ["抗拒继续", "选择停止、教师标注抗拒、长时间拒绝操作。", "暂停任务，提供选择权。"],
            ["异常发声/过度活跃", "教师标注或麦克风检测到异常音量。", "减少刺激，提示教师关注。"],
            ["需要帮助", "多次卡住或主动求助。", "提示教师介入，记录帮助点。"],
        ],
        [1700, 4400, 3260],
    )

    add_heading(doc, "4.4 低刺激互动", 2)
    add_para(doc, "互动设备可以先从屏幕提示做起，机器人作为展示加分项，而不是第一阶段必需项。")
    add_table(
        doc,
        ["等级", "方式", "示例"],
        [
            ["0", "不打扰", "只记录状态和任务表现。"],
            ["1", "屏幕提示", "出现小图标或一句短提示：“再看一下题目”。"],
            ["2", "语音提示", "轻声慢速提示：“我们一步一步来”。"],
            ["3", "机器人提示", "机器人转向学生或屏幕，短句鼓励，不靠近学生。"],
            ["4", "教师介入", "系统提示教师关注，暂停学生端任务。"],
        ],
        [900, 2500, 5960],
    )

    add_heading(doc, "5. 需要什么硬件", 1)
    add_table(
        doc,
        ["硬件", "第一版建议", "用途", "必要性"],
        [
            ["笔记本/小主机", "普通 Windows 笔记本即可。", "运行教师端、学生端和本地数据记录。", "必需"],
            ["平板/大屏", "触控平板、互动屏或投屏设备。", "学生完成生活适应任务。", "必需"],
            ["教师端设备", "教师电脑或平板。", "查看学生进度、标注状态、导出报告。", "必需"],
            ["麦克风", "普通 USB 麦克风。", "后续检测异常音量或简单语音线索。", "可选"],
            ["摄像头", "普通 USB 摄像头，后续可升级。", "后续观察头部朝向和课堂参与。", "可选"],
            ["桌面机器人", "支持简单语音、转向或表情屏即可。", "展示低刺激互动，不作为第一版核心依赖。", "可选"],
            ["本地网络", "普通路由器或校园内网。", "连接学生端和教师端。", "建议"],
        ],
        [1500, 2600, 3600, 1660],
    )

    add_heading(doc, "6. 最终载体是什么", 1)
    add_callout(
        doc,
        "最终载体",
        "第一阶段的最终载体应是一个可在学校课堂演示和试用的本地化软件系统，而不是昂贵机器人整机。机器人可以作为可选互动终端接入。",
    )
    add_table(
        doc,
        ["载体", "形态", "说明"],
        [
            ["学生端", "网页/平板应用", "承载生活适应任务包，让学生点击、选择、排序、匹配。"],
            ["教师端", "网页/电脑端", "查看任务进度、状态提示、教师标注和课后报告。"],
            ["任务包", "可配置课程内容", "包括整理物品、认识校园、情境选择、简单购物等。"],
            ["互动终端", "屏幕提示/语音/可选机器人", "执行低刺激提示，老师可随时关闭。"],
            ["数据报告", "本地文件或网页报告", "记录任务完成率、错误类型、提示次数和教师介入点。"],
        ],
        [1400, 2400, 5560],
    )

    add_heading(doc, "7. 近期可展示 Demo", 1)
    add_para(doc, "近期 Demo 应重点展示完整教学闭环，而不是炫技算法。建议做一个“生活适应课堂任务演示”。")
    add_number(doc, "教师选择一个任务包，例如“整理书包”或“认识校园”。")
    add_number(doc, "学生在平板上完成分类、匹配或情境选择。")
    add_number(doc, "系统记录错误、停顿、求助和完成情况。")
    add_number(doc, "教师端显示当前状态：专注、困惑、注意力下降、需要帮助。")
    add_number(doc, "屏幕或机器人给出低刺激提示。")
    add_number(doc, "课后生成报告，显示任务完成率、错误类型、提示次数和建议。")
    add_heading(doc, "8. 第一版功能范围", 1)
    add_table(
        doc,
        ["模块", "开发内容", "交付形式"],
        [
            ["学生端任务", "3-5 个生活适应任务包，每个任务支持难度分级。", "平板或大屏网页应用。"],
            ["教师端", "班级列表、任务进度、状态标注、课后报告。", "教师电脑或平板页面。"],
            ["状态观察", "基于任务行为、教师标注和可选传感器线索进行辅助判断。", "实时状态卡片和时间线。"],
            ["互动提示", "屏幕提示、语音提示和可选机器人提示。", "低刺激提示策略配置。"],
            ["数据报告", "本地保存任务日志、状态变化和课后摘要。", "课堂报告网页或导出文件。"],
        ],
        [1500, 4300, 3560],
    )

    add_heading(doc, "9. 实施路径", 1)
    add_table(
        doc,
        ["阶段", "周期", "目标", "输出"],
        [
            ["第一阶段", "1-2 周", "做出可点击 Demo，验证教师是否认可任务形式。", "学生端任务、教师端报告样例。"],
            ["第二阶段", "2-4 周", "进入学校做公益体验课或小范围试课。", "教师反馈、课堂观察记录。"],
            ["第三阶段", "1-2 月", "扩展任务包，加入更稳定的状态记录和低刺激提示。", "可试点 MVP。"],
            ["第四阶段", "后续", "再接入摄像头、麦克风和可选机器人。", "增强版课堂辅助系统。"],
        ],
        [1200, 1300, 4300, 2560],
    )

    add_heading(doc, "10. 安全与伦理边界", 1)
    add_bullet(doc, "不做医学诊断，只做课堂状态辅助观察。")
    add_bullet(doc, "不替代教师，所有判断和干预由教师最终决定。")
    add_bullet(doc, "不主动接触学生，机器人不快速靠近学生。")
    add_bullet(doc, "低刺激原则，避免大声、强光、突然动作和复杂指令。")
    add_bullet(doc, "教师可随时关闭提示、语音和机器人互动。")
    add_bullet(doc, "第一阶段优先采集任务日志和教师标注，不采集可识别原始视频音频。")
    add_bullet(doc, "如需采集图像或声音，必须获得学校和家长授权，并明确用途和留存边界。")

    add_heading(doc, "11. 验收指标", 1)
    add_table(
        doc,
        ["类别", "指标", "意义"],
        [
            ["学生参与", "任务完成率、主动参与次数、放弃次数。", "判断任务是否适合学生。"],
            ["教师价值", "教师查看报告次数、标注次数、有效提示比例。", "判断系统是否减轻教师负担。"],
            ["任务表现", "错误类型、卡住时长、提示后是否恢复。", "判断生活适应训练效果。"],
            ["安全性", "过度刺激事件、教师关闭提示次数、学生抗拒次数。", "判断系统能否安全进课堂。"],
            ["可运营性", "备课时间、任务包复用率、设备部署时间。", "判断系统能否持续交付。"],
        ],
        [1400, 3600, 4360],
    )

    add_heading(doc, "12. 结论", 1)
    add_para(
        doc,
        "本方案的主线是课堂辅助观察和生活适应任务训练。系统通过学生端任务、教师端观察、低刺激互动和课后报告，帮助教师更好地组织课堂、理解学生状态，并逐步支持学生提升生活自理和社会适应能力。",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_doc()
