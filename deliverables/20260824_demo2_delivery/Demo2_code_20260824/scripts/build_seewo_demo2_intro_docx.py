from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ANALYSIS = ROOT / "outputs" / "seewo_demo2" / "analysis.json"
ASSETS = ROOT / "docs" / "assets" / "seewo_demo2_document"
OUTPUT = ROOT / "docs" / "辅读学校课堂学习成效分析_Demo2介绍文档.docx"

BLUE = "2563EB"
TEAL = "087F8C"
NAVY = "17212B"
MUTED = "667085"
LINE = "D7DEE7"
SOFT = "F4F7FA"
PALE_BLUE = "EFF6FF"
PALE_GOLD = "FFF8EB"
WHITE = "FFFFFF"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=11, bold=False, color=NAVY, italic=False, mono=False):
    latin = "Consolas" if mono else "Calibri"
    east_asia = "Microsoft YaHei"
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, color=NAVY, bold=False):
    style.font.name = "Calibri"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = doc.styles[name]
        set_style_font(style, size, color=color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        set_style_font(style, 11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_page(doc: Document):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for node in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(node)
    set_run_font(run, size=9, color=MUTED)


def configure_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("RehabAffectPerception  |  辅读学校课堂学习成效分析 Demo2"), 9, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    set_run_font(fp.add_run("课堂教育支持原型 · 本地授权数据 · 第 "), 9, color=MUTED)
    add_page_field(fp)
    set_run_font(fp.add_run(" 页"), 9, color=MUTED)


def next_numbering_id(numbering):
    ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    return max(ids or [0]) + 1, max(abstract_ids or [0]) + 1


def create_numbering(doc: Document, fmt: str):
    numbering = doc.part.numbering_part.element
    num_id, abstract_id = next_numbering_id(numbering)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{0xA0B0C000 + abstract_id:08X}")
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    tmpl = OxmlElement("w:tmpl")
    tmpl.set(qn("w:val"), f"{0x1A2B3000 + abstract_id:08X}")
    abstract.append(tmpl)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if fmt == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if fmt == "bullet" else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    lvl.append(ppr)
    if fmt == "bullet":
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), "Symbol")
        rfonts.set(qn("w:hAnsi"), "Symbol")
        rfonts.set(qn("w:hint"), "default")
        rpr.append(rfonts)
        lvl.append(rpr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId")
    abs_id.set(qn("w:val"), str(abstract_id))
    num.append(abs_id)
    numbering.append(num)
    return num_id


def create_numbering_instance(doc: Document, style_name="List Number", start_at=1):
    style = doc.styles[style_name]
    style_num_ids = style._element.xpath("./w:pPr/w:numPr/w:numId")
    if not style_num_ids:
        raise ValueError(f"Style {style_name!r} has no numbering definition")
    base_num_id = style_num_ids[0].get(qn("w:val"))
    numbering = doc.part.numbering_part.element
    base_num = next(
        node for node in numbering.findall(qn("w:num")) if node.get(qn("w:numId")) == base_num_id
    )
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    num_id, _ = next_numbering_id(numbering)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), str(start_at))
    override.append(start)
    num.append(override)
    numbering.append(num)
    return num_id


def add_list_item(doc: Document, text: str, style_name: str, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.1875)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix) :]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_numbered_item(doc: Document, num_id: int, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.1875)
    num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(num)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix) :]))
    else:
        set_run_font(p.add_run(text))
    return p


def set_cell_margins(table):
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblCellMar"))
    if old is not None:
        tbl_pr.remove(old)
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tbl_pr.append(mar)


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        existing = tbl_pr.find(qn(tag))
        if existing is not None:
            tbl_pr.remove(existing)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
    set_cell_margins(table)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def format_table(table, widths, header=True, font_size=9.5):
    set_table_geometry(table, widths)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if header and row_idx == 0:
                set_cell_fill(cell, "E8EEF5")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=header and row_idx == 0)
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)


def add_callout(doc: Document, label: str, text: str, fill=PALE_BLUE, border=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.2
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), border)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    set_run_font(p.add_run(f"{label}  "), bold=True, color=NAVY)
    set_run_font(p.add_run(text), color=NAVY)
    return p


def add_para(doc: Document, text: str, bold_prefix: str | None = None, muted=False):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix) :]), color=MUTED if muted else NAVY)
    else:
        set_run_font(p.add_run(text), color=MUTED if muted else NAVY)
    return p


def add_heading(doc: Document, text: str, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def add_picture(doc: Document, filename: str, caption: str, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption.split("  ", 1)[0])
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(8)
    set_run_font(cp.add_run(caption), size=9, color=MUTED, italic=True)


def add_code_block(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.05
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "111827")
    p._p.get_or_add_pPr().append(shd)
    set_run_font(p.add_run(text), size=8.5, color="DBEAFE", mono=True)


def page_break(doc: Document):
    doc.add_page_break()


def build_document():
    payload = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    course = payload["courses"][0]
    student = course["student_results"][0]
    teacher = course["teacher_effectiveness"]

    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    configure_header_footer(doc)
    bullets = "List Bullet"
    numbers = create_numbering_instance(doc)
    doc.core_properties.title = "辅读学校课堂学习成效分析 Demo2 介绍文档"
    doc.core_properties.subject = "多学生课堂状态、学习掌握与具身支持闭环"
    doc.core_properties.author = "RehabAffectPerception"
    doc.core_properties.comments = "不包含身份证号、联系方式或住址等敏感个人信息。"

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("产品说明 · 校方演示版"), 9, bold=True, color=TEAL)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    set_run_font(title.add_run("辅读学校课堂学习成效分析 Demo2"), 24, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    set_run_font(subtitle.add_run("从三人课堂状态感知，到学习掌握评估与低刺激支持动作"), 13, color=MUTED)
    add_callout(
        doc,
        "一句话定位",
        "面向辅读学校课堂的多学生状态感知与学习成效分析原型：连接课程目标、课堂视频、个体化教育资料、教师复核和机器人/屏幕支持动作。",
    )
    meta = doc.add_table(rows=1, cols=2)
    meta.style = "Table Grid"
    meta.rows[0].cells[0].text = "项目"
    meta.rows[0].cells[1].text = "说明"
    for values in (
        ("文档用途", "校方沟通、内部汇报、Demo讲解与需求确认"),
        ("演示对象", "3名校方确认学生；支持数学与语文课程切换"),
        ("本机地址", "http://127.0.0.1:8791/seewo_student_demo2"),
        ("版本日期", f"Demo2 · {date.today().isoformat()}"),
    ):
        row = meta.add_row()
        row.cells[0].text, row.cells[1].text = values
    format_table(meta, [2700, 6660])
    add_picture(doc, "01_course_setup.png", "图1  监测前课程配置与三名学生目标适配建议", width=6.2)

    page_break(doc)
    add_heading(doc, "1. Demo解决什么问题", 1)
    add_para(
        doc,
        "辅读课堂中，同一班级学生的理解、表达、动作和支持需求差异较大。单纯统计抬头、举手或动作次数，无法回答学生是否真正获得学习机会、是否依赖提示、是否能够迁移，也无法直接指导下一步支持动作。",
    )
    add_callout(
        doc,
        "核心闭环",
        "课程目标与个体基线 → 三人同步课堂感知 → human_state状态估计 → 学习掌握初评 → 教师复核 → 屏幕、语音与机器人动作建议 → 下一课基线更新。",
    )
    add_heading(doc, "服务对象与价值", 2)
    for text in (
        "对学生：把“是否参与”与“是否学会”分开，记录独立完成、提示依赖、一致性和迁移表现，形成连续个体追踪。",
        "对教师：自动筛选值得复核的时段，给出课程目标适配和下一课支持建议，减少从整段视频中人工查找证据的时间。",
        "对学校：形成统一、可追溯的课堂观察口径，同时保留教师最终判断权，不把算法结果用作医学诊断或教师绩效结论。",
        "对机器人系统：输出结构化human_state和低刺激动作参数，为后续ROS2、WebSocket或真机执行预留接口。",
    ):
        add_list_item(doc, text, bullets)
    add_heading(doc, "输入与输出一览", 2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("类别", "当前输入", "系统输出")):
        cell.text = text
    rows = (
        ("课程", "课程名称、课程标准、校方4级考核标准", "课程适配分、分层建议、个体目标修改"),
        ("课堂", "学生机视频、教师机视频、远场音频", "30秒状态窗口、关键互动、证据质量"),
        ("学生", "固定座位身份、IEP支持画像、上次课程表现", "掌握度、参与度、跨课追踪、课后反馈"),
        ("执行", "状态、风险、置信度、教师复核结果", "屏幕提示、短句语音、机器人动作、教师交接"),
    )
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    format_table(table, [1500, 3750, 4110])

    page_break(doc)
    add_heading(doc, "2. 系统工作流程", 1)
    steps = (
        "课程配置：教师在监测前录入课程名称、课程标准和校方考核标准。",
        "目标适配：系统读取上次课程表现和个体支持资料，判断本课目标是否适配，并给出拆分或分层建议。",
        "视频处理：对学生机视频做关键帧抽取，在固定座位ROI内计算动作变化和画面证据质量。",
        "状态估计：按30秒窗口形成稳定参与、回应/操作、回应线索较低或证据不足等状态，并进行时间聚合。",
        "学习评估：结合课程目标、互动机会、正确性、提示等级、一致性和迁移形成初评；不把安静或动作多少直接等同于学会。",
        "反馈执行：教师复核关键时段后，系统生成课后反馈、下一课计划，以及屏幕/语音/机器人动作建议。",
    )
    for step in steps:
        add_numbered_item(doc, numbers, step, bold_prefix=step.split("：", 1)[0] + "：")
    add_heading(doc, "当前处理策略", 2)
    for text in (
        "视频以约10秒关键帧采样、30秒状态窗口聚合，首次处理两门真实课堂约15秒；缓存后页面可快速打开。",
        "固定座位ROI来自校方确认的群像截图，不通过障碍类型或外貌推断身份。",
        "远场音频作为课堂活动与互动背景保留；无法可靠区分说话人时，不强行把语音绑定到具体学生。",
        "原始视频和个体资料在本地授权目录处理，页面与文档不输出身份证号、联系方式和住址。",
    ):
        add_list_item(doc, text, bullets)
    add_heading(doc, "结构化输出示例", 2)
    add_code_block(
        doc,
        '''{
  "person_id": "student_A", "state": "low_response",
  "confidence": 0.84, "risk_level": "low",
  "handoff_required": false,
  "recommended_robot_action": "regain_attention_low_stimulation"
}''',
    )

    page_break(doc)
    add_heading(doc, "3. 监测前课程配置与目标适配", 1)
    add_para(
        doc,
        "课程配置是Demo2相对于单纯状态监测的重要扩展。系统先明确“这节课要学什么、学校如何判定掌握”，再分析学生是否获得了与目标匹配的学习机会。",
    )
    add_picture(doc, "01_course_setup.png", "图2  课程名称、课程标准、校方考核标准与三人适配建议")
    add_heading(doc, "适配判断逻辑", 2)
    for text in (
        "比较本课每个目标的要求等级与学生上次课程表现。",
        "结合个体档案中的优势、支持需要、有效提示和安全约束。",
        "输出“适配、需分层、需调整”及量化适配分。",
        "差距较大时，把目标拆成可单独观察的一步任务，并控制指令长度与提示强度。",
    ):
        add_list_item(doc, text, bullets)
    add_callout(
        doc,
        "评价边界",
        "课程目标不会因为障碍名称被自动降低。系统依据已有表现与可观察任务差距提出建议，最终目标由教师和学校确认。",
        fill=PALE_GOLD,
        border="A96300",
    )

    page_break(doc)
    add_heading(doc, "4. 三人同步课堂证据与独立面板", 1)
    add_para(
        doc,
        "群像视图负责确认三名学生在课堂中的位置关系；个人视图负责放大单名学生的状态、掌握度和支持动作。教师可以在同一页面切换学生，不需要重复上传或打开三个系统。",
    )
    add_picture(doc, "02_three_student_evidence.png", "图3  群像身份复核、三人切换、个人掌握度与支持动作")
    add_heading(doc, "每名学生面板包含", 2)
    for text in (
        "固定座位与身份确认状态；个人课堂图像；课程个体目标。",
        "正确性、独立程度、重复一致性和迁移四维掌握指标。",
        "课堂参与指数、主要状态、待复核窗口数和关键互动证据。",
        "屏幕内容、短句语音、机器人动作和是否需要教师确认。",
        "个体化教育支持画像、下一课建议和约1000至1300字课后反馈。",
    ):
        add_list_item(doc, text, bullets)

    page_break(doc)
    add_heading(doc, "5. Demo1状态分析能力完整保留", 1)
    add_para(
        doc,
        "Demo2不是替代或删减Demo1，而是在原有单学生状态分析上增加多学生、课程目标和学习评估。Demo1原页面仍可独立访问，以下能力也已并入Demo2。",
    )
    add_picture(doc, "03_state_action_linkage.png", "图4  完整30秒状态窗口与当前时段动作联动")
    for text in (
        "完整30秒状态时间线、状态分布、平均置信度和证据质量。",
        "每个窗口包含可见证据、解释、风险等级、教师交接和human_state JSON。",
        "每个状态同步映射屏幕、语音、机器人动作与教师端提醒。",
        "支持“全部窗口/只看待复核”筛选，点击状态可直接跳转对应动作。",
        "保留学生追踪、个体档案、课堂汇总、教师支持效果和完整JSON。",
    ):
        add_list_item(doc, text, bullets)

    add_heading(doc, "6. 学习掌握与个体追踪", 1)
    add_para(
        doc,
        "状态分析回答“学生在课堂中呈现了什么线索”，学习评估回答“学生在明确学习机会中完成到什么程度”。两类指标分开呈现，避免把安静、动作或表情直接当作学业结论。",
    )
    add_heading(doc, "4级表现量表", 2)
    rubric = doc.add_table(rows=1, cols=3)
    rubric.style = "Table Grid"
    for cell, text in zip(rubric.rows[0].cells, ("等级", "定义", "记录原则")):
        cell.text = text
    for values in (
        ("4", "独立完成", "无提示完成目标任务"),
        ("3", "简单提示后完成", "短句、视觉或手势提示"),
        ("2", "示范或部分协助", "教师示范、分步协助后完成"),
        ("1", "尚未表现", "已给予清晰机会但未出现目标行为"),
        ("U", "未观察/不确定", "无机会、遮挡、音频归属不清；不计入不会"),
    ):
        cells = rubric.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    format_table(rubric, [900, 3100, 5360])
    add_heading(doc, "四个掌握维度", 2)
    for text in (
        "正确性：回答或操作是否符合课程目标。",
        "独立程度：完成前接受了多少语言、视觉、示范或身体协助。",
        "重复一致性：相同目标在多次机会中的表现是否稳定。",
        "迁移：能否在新材料、新情境或下一节课中使用同类知识。",
    ):
        add_list_item(doc, text, bullets)
    add_picture(doc, "05_student_profile.png", "图5  个体档案把优势、支持需要、有效方式和安全约束转成行动条件", width=5.25)

    add_heading(doc, "7. 教师支持效果：量化但不僵化", 1)
    add_para(
        doc,
        "系统评价的是“本节课对三名重点学生的支持效果”，不评价教师个人能力。结果用于课后复盘和下一课策略调整，不作为教师绩效考核。",
    )
    add_picture(doc, "04_teacher_support.png", "图6  教师支持效果指数、五项量化标准与复核问题", width=5.55)
    metrics = doc.add_table(rows=1, cols=4)
    metrics.style = "Table Grid"
    for cell, text in zip(metrics.rows[0].cells, ("维度", "权重", "含义", "证据来源")):
        cell.text = text
    metric_notes = {
        "student_engagement": "三人参与窗口",
        "recovery_after_challenge": "低回应后60秒恢复",
        "individualized_fit": "上次表现与本课目标",
        "learning_evidence": "正确性、独立性、一致性、迁移",
        "low_disruption_support": "证据不足与干扰窗口",
    }
    metric_meanings = {
        "student_engagement": "三名学生参与与回应的综合表现",
        "recovery_after_challenge": "低回应后能否在60秒内恢复",
        "individualized_fit": "目标是否匹配上次表现与支持资料",
        "learning_evidence": "四维学习证据是否充分且一致",
        "low_disruption_support": "支持是否克制并尽量减少干扰",
    }
    for metric in teacher["metrics"]:
        cells = metrics.add_row().cells
        values = (
            metric["name"],
            f"{metric['weight']}%",
            metric_meanings.get(metric["key"], metric["benchmark"]),
            metric_notes.get(metric["key"], metric["evidence"]),
        )
        for cell, text in zip(cells, values):
            cell.text = text
    format_table(metrics, [1950, 900, 4110, 2400], font_size=8.8)

    page_break(doc)
    add_heading(doc, "8. 从感知结果到机器人动作", 1)
    add_callout(
        doc,
        "具身闭环",
        "摄像头/麦克风/课程任务 → 多模态证据 → human_state → 风险与支持需求 → 屏幕/语音/机器人动作 → 学生反应 → 教师复核 → 更新个体基线。",
    )
    add_heading(doc, "动作输出结构", 2)
    action_table = doc.add_table(rows=1, cols=4)
    action_table.style = "Table Grid"
    for cell, text in zip(action_table.rows[0].cells, ("状态", "屏幕", "语音", "机器人/教师")):
        cell.text = text
    for values in (
        ("稳定参与", "保持当前画面", "具体正反馈", "静止或缓慢点头"),
        ("回应线索较低", "轻微高亮当前一步", "只说一个短指令", "缓慢转向屏幕"),
        ("正在学习/需支持", "减少无关信息", "重复一步并等待", "保持距离，不接触"),
        ("证据不足", "不触发强提示", "不连续提醒", "无动作，交由教师确认"),
    ):
        cells = action_table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
    format_table(action_table, [1650, 2400, 2550, 2760])
    add_heading(doc, "安全规则", 2)
    for text in (
        "低刺激：避免突然大声、闪烁、快速靠近和夸张动作。",
        "无接触：机器人不主动触碰学生。",
        "教师在环：高影响动作、证据不足和连续低回应必须由教师确认。",
        "可撤销：教师可以关闭、修正或标记系统建议为不确定。",
        "结构化接口：动作建议可进一步发布到WebSocket、ROS2 Mock或仿真环境。",
    ):
        add_list_item(doc, text, bullets)

    page_break(doc)
    add_heading(doc, "9. 页面功能与演示顺序", 1)
    add_heading(doc, "启动方式", 2)
    add_para(doc, "双击项目根目录中的 start_school_demo2.cmd，或访问本机地址：http://127.0.0.1:8791/seewo_student_demo2。该地址仅在运行服务的本机可用。")
    add_heading(doc, "建议的5分钟演示流程", 2)
    demo_steps = (
        "切换数学或语文课程，说明系统支持不同课程目标和考核标准。",
        "展示课程配置，点击“生成适配建议”，比较三名学生的目标差距和分层建议。",
        "进入三人同步证据区，在群像中核对身份，再切换个人画像和三名学生面板。",
        "点击“完整状态”，查看30秒窗口、状态分布和只看待复核筛选。",
        "点击某个状态窗口，展示证据、风险和对应屏幕/语音/机器人动作。",
        "进入“证据复核”，由教师选择符合、需修正或不确定，观察复核进度变化。",
        "展示掌握度、跨课追踪、个体档案和课后反馈，说明未观察不计入不会。",
        "最后展示教师支持效果和JSON，说明系统如何接入后续机器人控制端。",
    )
    demo_numbers = create_numbering_instance(doc)
    for step in demo_steps:
        add_numbered_item(doc, demo_numbers, step)
    add_heading(doc, "页面功能分组", 2)
    groups = doc.add_table(rows=1, cols=3)
    groups.style = "Table Grid"
    for cell, text in zip(groups.rows[0].cells, ("学习评估", "状态与行动", "教学复盘")):
        cell.text = text
    row = groups.add_row().cells
    row[0].text = "掌握度、跨课追踪、课后反馈"
    row[1].text = "完整状态、证据复核、动作方案"
    row[2].text = "支持效果、个体档案、课堂汇总、JSON"
    format_table(groups, [3120, 3120, 3120])

    page_break(doc)
    add_heading(doc, "10. 当前完成度、边界与后续方向", 1)
    add_heading(doc, "当前已经完成", 2)
    for text in (
        "两门真实课堂视频的三学生固定座位ROI分析与群像/个人复核。",
        "课程目标录入、上次表现适配、四维掌握初评和三人独立面板。",
        "Demo1完整状态分析、human_state JSON、动作建议、学生追踪和教师支持效果。",
        "教师片段复核、复核进度、约1000至1300字课后反馈和下一课建议。",
        "本地服务器、一键启动脚本、桌面/移动端页面与自动化测试。",
    ):
        add_list_item(doc, text, bullets)
    add_heading(doc, "必须明确的限制", 2)
    for text in (
        "当前状态识别是基于固定ROI、关键帧动作变化和证据质量的可解释原型，不是训练完成的最终深度模型。",
        "低清全景视频适合动作与参与线索分析，不适合精细表情、眼神或医学状态判断。",
        "远场音频已保留，但当前不能稳定完成学生级说话人归属和回答内容识别。",
        "掌握度中的正确性、提示等级和迁移必须由教师抽检关键片段后确认。",
        "单节课只能建立或更新基线，不能据此判断长期进步或恶化。",
    ):
        add_list_item(doc, text, bullets)
    add_heading(doc, "后续可升级方向", 2)
    for text in (
        "建立教师复核标签库，训练学生级时序状态模型并进行置信度校准。",
        "增加课堂ASR、教师提问检测、候答时间和学生回答正确性半自动标注。",
        "引入跨课数据库，按同类目标比较独立完成比例、提示减少和知识迁移。",
        "通过WebSocket/ROS2连接屏幕、语音或机器人，并在仿真中验证低刺激动作策略。",
    ):
        add_list_item(doc, text, bullets)
    add_callout(
        doc,
        "对外表达建议",
        "这是“课堂状态辅助识别与学习支持系统”，不是医学诊断工具；系统帮助教师更快找到证据、形成个体化支持，并把建议落到可控的机器人动作上。",
        fill=PALE_GOLD,
        border="A96300",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
