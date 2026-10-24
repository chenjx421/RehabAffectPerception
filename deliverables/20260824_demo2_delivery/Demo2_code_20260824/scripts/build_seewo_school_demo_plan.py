from __future__ import annotations

from datetime import datetime
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.build_seewo_student_solution_docx import (  # noqa: E402
    BLUE,
    DARK_BLUE,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    add_body,
    add_bullet,
    add_callout,
    add_heading,
    add_kv_table,
    add_matrix,
    add_number,
    add_picture_block,
    configure_document,
    set_paragraph_spacing,
    set_run_font,
)


DOCS_DIR = PROJECT_ROOT / "docs"
SCREENSHOT_DIR = PROJECT_ROOT / "outputs" / "playwright"
OUTPUT_DOCX = DOCS_DIR / "seewo_school_demo_plan_zh.docx"


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=0, after=4, line=1.05)
    r = p.add_run("辅读学校学生课堂状态追踪与个体化支持系统 Demo 方案")
    set_run_font(r, size=21, bold=True)
    r.font.color.rgb = INK

    p2 = doc.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=8, line=1.12)
    r2 = p2.add_run("基于 Seewo 课堂视频、评课文档与学生基础资料的 AI 辅助教学支持原型")
    set_run_font(r2, size=12)
    r2.font.color.rgb = MUTED

    add_kv_table(
        doc,
        [
            ("方案定位", "面向辅读学校课堂的学生状态追踪与个体化支持系统。系统服务对象首先是学生，同时为老师提供可复核的课堂支持依据。"),
            ("展示对象", "以一名学生的一节真实课堂为样例，展示从课堂数据到学生状态、原因分析、支持建议和课后档案的完整流程。"),
            ("演示载体", "本地 Web Demo 页面：/seewo_student_demo。现场展示建议使用运行服务的电脑连接投影或大屏。"),
            ("边界说明", "系统不做医学诊断，不替代老师判断；所有输出只作为教学支持和课后复盘参考。"),
        ],
        widths=(1.45, 5.05),
    )


def add_footer(doc: Document) -> None:
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        r = p.add_run("RehabAffectPerception | 辅读学校学生课堂支持 Demo 方案")
        set_run_font(r, size=8)
        r.font.color.rgb = MUTED


def build_doc() -> None:
    doc = Document()
    configure_document(doc)
    add_footer(doc)
    add_title(doc)

    add_callout(
        doc,
        "核心想法",
        "把学校已经沉淀在 Seewo 中的课堂视频和评课资料，用于建立单个学生的课堂状态追踪档案。系统不是只判断“专注/不专注”，而是进一步回答：学生什么时候掉队、可能为什么掉队、适合用什么方式把他带回学习任务。",
        fill=LIGHT_BLUE,
    )

    add_heading(doc, "1. 要解决的具体问题", 1)
    add_body(
        doc,
        "辅读学校课堂中，学生之间能力差异明显，同一节课里可能同时出现理解慢、注意力短、动作控制弱、语言表达不稳定等情况。老师在上课时很难持续观察每个学生的细微变化，课后也缺少一套围绕单个学生的连续证据。"
    )
    add_body(
        doc,
        "本系统选择从“一个学生的一节课”切入，而不是一开始做全班泛化。先把单个学生的状态变化、课堂环节、教师支持方式和个体资料放在同一条时间线上，再逐步扩展到多节课、多名学生。"
    )
    add_matrix(
        doc,
        ["实际问题", "系统要给出的结果"],
        [
            ["学生可能没有明确表达“我没听懂”或“我累了”。", "识别低参与、困惑、疲劳风险、抗拒苗头等课堂状态，并给出证据片段。"],
            ["老师难以判断某个学生是短暂走神，还是长期跟不上。", "把每节课的状态变化写入学生档案，形成可对比的趋势。"],
            ["课堂提醒过强可能造成反效果。", "根据学生资料生成低刺激、可关闭、可由老师控制的支持建议。"],
            ["Seewo 记录了课堂，但没有直接形成学生个体成长分析。", "在 Seewo 资料之上补充学生视角的状态追踪和个体化支持闭环。"],
        ],
        [3.10, 3.40],
    )

    add_heading(doc, "2. Demo 展示什么", 1)
    add_picture_block(
        doc,
        SCREENSHOT_DIR / "seewo_student_demo_ui_polished.png",
        "图 1  Demo 页面：围绕目标学生展示课堂状态、证据、个体档案、动作建议和课后复盘信息。",
    )
    add_body(
        doc,
        "现场演示时，可以直接打开页面，从左侧的目标学生和课堂信息开始讲，再依次展示中间的状态时间线、证据与建议，最后展示右侧的支持画像和课堂复盘。"
    )
    add_matrix(
        doc,
        ["展示模块", "校方能看到什么"],
        [
            ["目标学生追踪", "系统明确只追踪被选中的学生，避免把老师或其他学生混入分析结果。"],
            ["课堂状态时间线", "按时间片展示学生状态，例如参与、困惑、低参与、疲劳风险，并附带置信度和证据。"],
            ["个体化支持画像", "结合学生基础资料和测评内容，提示这个学生适合怎样的指令、反馈和互动强度。"],
            ["动作建议", "把分析结果落到屏幕提示、语音提示、机器人动作建议或提醒老师复核。"],
            ["课堂支持复盘", "从学生表现反推这节课的支持策略是否有效，给出下一节课可尝试的改进方向。"],
        ],
        [1.65, 4.85],
    )

    add_heading(doc, "3. 数据如何进入 Demo", 1)
    add_callout(
        doc,
        "当前展示方式",
        "展示阶段采用“人工导入文件，系统自动分析”。也就是说，先从 Seewo 下载或导出课堂材料，由项目团队放入 Demo 数据目录，系统启动后自动读取并生成页面结果。",
        fill=LIGHT_GRAY,
    )
    add_matrix(
        doc,
        ["输入文件", "来源", "用途"],
        [
            ["课堂视频 mp4", "Seewo 课堂录制或导出视频。", "用于抽帧、识别目标学生区域、分析动作和课堂参与状态。"],
            ["评课详情 docx", "Seewo 生成的评课详情文档。", "用于读取课程主题、课堂结构、教师活动和课后评价信息。"],
            ["目标学生截图", "从课堂视频中截取。", "用于确认本次追踪对象，避免误识别。"],
            ["学生基础资料", "校方授权提供的测评、能力领域、行为指导或支持建议。", "用于形成个体化支持画像，并约束系统给出的动作建议。"],
        ],
        [1.35, 2.05, 3.10],
    )
    add_body(
        doc,
        "展示现场不需要老师临时上传文件，也不会自动读取学校后台。后续进入试点后，可以增加一个网页上传入口，由授权老师或项目成员上传一节课的数据包；如果学校允许接口对接，再考虑校内服务器或 Seewo 授权同步。"
    )

    add_heading(doc, "4. 系统输出是什么", 1)
    add_number(doc, "学生状态结果：当前时间片的状态、置信度、证据、风险等级和可能原因。")
    add_number(doc, "个体化支持建议：屏幕如何提示、语音如何表达、机器人可做什么动作、是否需要老师介入。")
    add_number(doc, "课后复盘结果：学生在不同课堂环节中的参与变化，以及可能有效或无效的支持方式。")
    add_number(doc, "学生追踪档案：每节课的表现记录、老师复核结果、能力领域变化和后续建议。")
    add_number(doc, "结构化接口数据：以 JSON 形式输出，后续可接 WebSocket、数据库、屏幕端或机器人控制端。")

    add_heading(doc, "5. 课堂支持效果如何量化", 1)
    add_picture_block(
        doc,
        SCREENSHOT_DIR / "seewo_teacher_effectiveness.png",
        "图 2  课堂支持效果复盘：指标围绕学生是否被有效带回学习任务，而不是给老师排名。",
    )
    add_body(
        doc,
        "系统对课堂支持效果的判断，核心依据是学生状态是否改善，而不是单纯评价老师。量化结果用于教研复盘：哪些环节学生能跟上，哪些环节容易掉队，下一节课可以如何调整。"
    )
    add_matrix(
        doc,
        ["指标", "含义"],
        [
            ["参与维持度", "学生在核心教学环节中是否持续看向任务、保持坐姿和回应。"],
            ["困惑恢复度", "出现茫然、低参与或迟疑后，是否能在提示后回到任务。"],
            ["节奏匹配度", "教师讲解、等待、提问和任务切换是否适合该学生反应速度。"],
            ["低刺激支持", "提示方式是否避免突然大声、强光、快速靠近或复杂指令。"],
            ["个体化适配", "教师或系统建议是否参考该学生的基础资料和能力特点。"],
        ],
        [1.55, 4.95],
    )

    add_heading(doc, "6. 和机器人如何结合", 1)
    add_body(
        doc,
        "机器人在本方案中不是一开始必须落地的硬件，而是系统输出动作的执行载体。Demo 阶段先用网页模拟动作建议，等学校确认支持方式有效后，再接入屏幕、音箱、桌面小机器人或移动机器人。"
    )
    add_matrix(
        doc,
        ["输出通道", "动作示例", "控制原则"],
        [
            ["屏幕", "显示简短图形提示、任务步骤、低刺激鼓励语。", "不遮挡老师教学，不用过强动画。"],
            ["语音", "用短句重复任务，例如“请看图片，先选第一个”。", "音量可控，避免突然播放。"],
            ["机器人动作", "轻微转头、点头、后退保持距离、面向学生等待。", "不主动接触学生，不快速靠近。"],
            ["老师提醒", "提示老师某学生可能低参与，建议人工观察。", "最终判断权交给老师。"],
        ],
        [1.25, 2.75, 2.50],
    )

    add_heading(doc, "7. 现场展示流程", 1)
    add_number(doc, "先说明学校已有数据来源：Seewo 课堂视频、评课详情、学生基础资料。")
    add_number(doc, "打开 Demo 页面，指出本次追踪的是同一节课中的同一名学生。")
    add_number(doc, "展示状态时间线，讲清楚系统不是只看一个动作，而是结合课堂环节和个体资料做判断。")
    add_number(doc, "展示动作建议，强调最终要落到屏幕、语音、机器人动作或老师提醒。")
    add_number(doc, "展示课后复盘，说明系统可以长期跟踪这个学生是否进步、退步或需要调整支持方式。")
    add_body(
        doc,
        "试点建议：先选择 1-3 名学生、2-3 节课进行验证。每节课结束后由老师快速复核关键片段，项目团队根据复核结果修正标签和建议库；当单学生流程稳定后，再扩展到多名学生和班级级看板。"
    )

    doc.core_properties.title = "辅读学校学生课堂状态追踪与个体化支持系统 Demo 方案"
    doc.core_properties.subject = "Seewo classroom student support demo plan"
    doc.core_properties.author = "RehabAffectPerception"
    doc.core_properties.created = datetime.now()
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    build_doc()
