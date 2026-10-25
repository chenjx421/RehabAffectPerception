from __future__ import annotations

from datetime import datetime
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts.build_seewo_student_solution_docx import (  # noqa: E402
    BLUE,
    DARK_BLUE,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    add_body,
    add_bullet,
    add_callout,
    add_caption,
    add_heading,
    add_kv_table,
    add_matrix,
    add_number,
    add_picture_block,
    configure_document,
    set_paragraph_spacing,
    set_run_font,
)


DOCS_DIR = PROJECT_ROOT / "docs"
SCREENSHOT_DIR = PROJECT_ROOT / "outputs" / "playwright"

SOLUTION_DOCX = DOCS_DIR / "seewo_school_demo_solution_v4_zh.docx"
IO_DOCX = DOCS_DIR / "seewo_demo_io_flow_for_school_zh.docx"


def add_title(doc: Document, title: str, subtitle: str, meta_rows: list[tuple[str, str]]) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=0, after=4, line=1.05)
    r = p.add_run(title)
    set_run_font(r, size=22, bold=True)
    r.font.color.rgb = INK

    p2 = doc.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=8, line=1.1)
    r2 = p2.add_run(subtitle)
    set_run_font(r2, size=12)
    r2.font.color.rgb = MUTED

    add_kv_table(doc, meta_rows, widths=(1.55, 4.95))


def add_footer_note(doc: Document, text: str) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        r = p.add_run(text)
        set_run_font(r, size=8)
        r.font.color.rgb = MUTED


def build_solution_doc() -> None:
    doc = Document()
    configure_document(doc)
    add_footer_note(doc, "RehabAffectPerception | Seewo 学生状态追踪 Demo 校方展示材料")

    add_title(
        doc,
        "辅读学校 Seewo 学生个体状态追踪与课堂支持 Demo 方案",
        "V4 展示版：把课堂视频、评课文档和学生资料转化为可复核的学生支持建议",
        [
            ("当前版本", "V4 校方展示版，包含学生状态追踪、个体档案、机器人/屏幕/语音动作建议、教师支持效果评估。"),
            ("展示载体", "本地 Web Demo 页面：/seewo_student_demo。127.0.0.1 地址只能在运行 Demo 的电脑上打开。"),
            ("核心定位", "不是替代老师，也不是医学诊断，而是基于课堂数据为学生提供持续、低刺激、可复核的支持建议。"),
            ("适用场景", "辅读学校课堂中，围绕单名学生进行课中状态观察、课后复盘和长期变化追踪。"),
        ],
    )

    add_callout(
        doc,
        "一句话说明",
        "该 Demo 将 Seewo 已有课堂材料向前推进一步：从“课堂记录与评课”变成“学生个体状态追踪 + 支持动作建议 + 教师效果复盘”的闭环。",
        fill=LIGHT_BLUE,
    )

    add_heading(doc, "1. Demo 要解决的具体问题", 1)
    add_matrix(
        doc,
        ["对象", "实际痛点", "Demo 中的对应能力"],
        [
            ["学生", "课堂中可能出现分心、茫然、疲劳、抗拒等状态，但学生不一定能清楚表达。", "从视频片段、声音线索和课堂任务上下文中形成状态判断，并给出低刺激支持建议。"],
            ["老师", "课中很难同时照顾所有学生，课后也缺少针对单个学生的连续证据。", "把一节课拆成时间线，输出状态趋势、可能原因、建议动作和需要人工复核的位置。"],
            ["学校", "Seewo 有课堂记录和评课材料，但难直接沉淀成学生个体成长档案。", "把每节课的观察结果写入学生追踪档案，后续可横向覆盖到全班学生。"],
            ["项目团队", "真实数据少、清晰度不稳定、隐私要求高，不能一开始就重训练大模型。", "先做轻量分析与知识库融合，保留人工复核入口，再逐步积累可训练数据。"],
        ],
        [1.05, 2.25, 3.20],
    )

    add_heading(doc, "2. 当前展示版包含什么", 1)
    add_picture_block(
        doc,
        SCREENSHOT_DIR / "seewo_student_demo_ui_polished.png",
        "图 1  当前 Demo 首页：上方展示学生状态、证据、时间线和支持建议，下方进入个体档案与课堂复盘。",
    )
    add_matrix(
        doc,
        ["模块", "页面中看得到的结果", "面向校方的解释"],
        [
            ["学生状态识别", "专注、困惑、低参与、疲劳风险等状态及置信度。", "先作为教师观察参考，不把结果当成诊断或处分依据。"],
            ["个体档案融合", "匿名学生画像、能力领域、课堂适配建议。", "把学生基础资料变成“支持约束”，避免给出不适合该学生的提醒方式。"],
            ["课堂时间线", "按视频片段整理状态变化、课堂活动和触发线索。", "帮助老师快速回看“什么时候出现变化、可能和什么课堂环节有关”。"],
            ["机器人/屏幕动作建议", "屏幕提示、语音提示、机器人动作、老师提醒四类输出。", "最终落到可执行动作，但演示阶段可先用屏幕和语音模拟机器人输出。"],
            ["教师支持效果评估", "课堂支持指数、分项评分、证据说明和改进建议。", "评价的是教学支持策略是否帮助该学生回到任务，不是给老师排名。"],
        ],
        [1.30, 2.10, 3.10],
    )

    add_heading(doc, "3. 这次更新后的教师支持效果评估", 1)
    add_picture_block(
        doc,
        SCREENSHOT_DIR / "seewo_teacher_effectiveness.png",
        "图 2  教师支持效果评估：把学生表现转化为可讨论的课堂支持指标。",
    )
    add_body(
        doc,
        "该功能的出发点不是“评价老师好坏”，而是从学生表现反推课堂支持是否有效：学生是否能维持参与，是否在困惑后恢复，教师节奏和支架是否匹配，课堂环境是否降低干扰，支持策略是否符合个体资料。"
    )
    add_matrix(
        doc,
        ["指标", "权重", "量化含义"],
        [
            ["参与维持度", "35%", "学生在课程核心环节中保持任务参与的比例和稳定性。"],
            ["困惑/分心恢复度", "22%", "出现低参与或困惑后，是否能在较短时间内回到任务。"],
            ["节奏与提问支架", "18%", "教师讲解、提问和等待时间是否与学生反应节奏匹配。"],
            ["低干扰支持环境", "15%", "是否减少突然、过强或复杂刺激，避免让学生更紧张。"],
            ["个体化适配度", "10%", "是否结合该学生基础资料调整提示方式、任务难度和反馈强度。"],
        ],
        [2.00, 0.80, 3.70],
    )
    add_callout(
        doc,
        "人性化边界",
        "系统只给“支持效果观察”，不输出“老师好/不好”的标签；低分项会转成下一节课可尝试的建议，例如减少一次性指令、增加视觉提示、降低语音刺激或延长等待时间。",
        fill=LIGHT_GRAY,
    )

    add_heading(doc, "4. 从数据到动作的闭环", 1)
    add_number(doc, "Seewo 导出课堂视频和评课详情，项目团队选择一个目标学生作为观察对象。")
    add_number(doc, "系统抽取视频片段、目标学生区域、声音/课堂活动线索，并形成时间线。")
    add_number(doc, "系统结合学生基础资料，判断当前状态、可能原因、风险等级和证据。")
    add_number(doc, "系统输出可执行建议：屏幕怎么提示、语音怎么说、机器人怎么动、是否提醒老师人工介入。")
    add_number(doc, "老师复核后，结果进入该学生档案，用于下一节课的追踪分析。")

    add_heading(doc, "5. 校方展示时建议讲法", 1)
    add_bullet(doc, "今天展示的是“单学生追踪原型”，不是全校部署系统。先证明一个学生、一节课可以跑通。")
    add_bullet(doc, "当前 Demo 数据已由项目团队提前导入，现场不需要老师上传文件，也不会自动读取学校后台。")
    add_bullet(doc, "如果校方认可方向，下一步可以增加校内上传入口，由授权老师或项目成员手动上传课堂包。")
    add_bullet(doc, "如果学校不希望数据出校，可把系统部署在校内电脑或校内服务器，项目方只看脱敏结果。")
    add_bullet(doc, "后续机器人只是执行载体之一，前期先用屏幕、语音和老师提醒验证干预策略是否对学生有帮助。")

    add_heading(doc, "6. 后续试点路线", 1)
    add_matrix(
        doc,
        ["阶段", "目标", "交付物"],
        [
            ["1-2 周", "围绕 1 名学生、2-3 节课，验证状态时间线和老师复核流程。", "单学生追踪报告、状态标签修订表、老师反馈记录。"],
            ["3-4 周", "扩展到 3-5 名学生，建立个体化支持画像。", "学生匿名档案库、课堂支持建议库、复核后的训练样本。"],
            ["5-8 周", "形成班级级别趋势分析，并接入屏幕/语音/机器人模拟输出。", "班级试点看板、动作策略库、可展示 Demo 和阶段报告。"],
        ],
        [1.15, 2.70, 2.65],
    )

    doc.core_properties.title = "辅读学校 Seewo 学生个体状态追踪与课堂支持 Demo 方案"
    doc.core_properties.subject = "Seewo classroom AI support demo V4"
    doc.core_properties.author = "RehabAffectPerception"
    doc.core_properties.created = datetime.now()
    doc.save(SOLUTION_DOCX)


def build_io_doc() -> None:
    doc = Document()
    configure_document(doc)
    add_footer_note(doc, "RehabAffectPerception | Demo 输入输出与文件流转说明")

    add_title(
        doc,
        "Seewo Demo 输入输出与文件流转说明",
        "校方展示版：明确文件从哪里来、怎么导入、系统输出什么",
        [
            ("当前结论", "当前展示版采用“手动导入文件 + 系统自动分析”，不是自动同步 Seewo 后台。"),
            ("现场展示方式", "演示电脑提前准备数据包并启动本地服务，浏览器打开 /seewo_student_demo 查看结果。"),
            ("是否有上传按钮", "当前展示版没有面向学校老师的浏览器上传按钮；后续试点版可以增加。"),
            ("是否自动传输", "当前不会自动从 Seewo 拉取数据，也不会自动把学校数据传到云端。"),
        ],
    )

    add_callout(
        doc,
        "必须说清楚的一句话",
        "现在的 Demo 是“学校/Seewo 导出材料，项目团队手动放入指定目录，系统自动分析并生成页面与 JSON 结果”。这比纯 PPT 更接近产品，但还不是正式上线平台。",
        fill=LIGHT_BLUE,
    )

    add_heading(doc, "1. 当前 Demo 的输入文件", 1)
    add_matrix(
        doc,
        ["输入项", "来源", "当前导入方式", "用途"],
        [
            ["课堂视频 mp4", "Seewo 课后导出的课堂视频。", "项目团队手动放入 data/raw/seewo_xiaoyudi_001/lesson_video.mp4。", "抽帧、定位目标学生、分析状态变化和课堂活动片段。"],
            ["评课详情 docx", "Seewo 评课详情文档。", "项目团队手动放入 data/raw/seewo_xiaoyudi_001/seewo_lesson_report.docx。", "读取课程名称、学科、课堂结构、评价信息，作为状态解释上下文。"],
            ["目标学生截图/ROI", "从课堂视频中截取学生图像。", "项目团队手动放入 data/raw/seewo_xiaoyudi_001/target_student.png，并在程序中标注观察区域。", "告诉系统“这节课重点追踪哪个学生”，避免把老师或其他学生误认为目标对象。"],
            ["学生基础资料", "校方授权提供的测评、能力领域、行为支持建议等资料。", "先人工脱敏整理，再转成 data/processed/seewo_xiaoyudi_001/student_A_support_profile.json。", "生成个体化支持画像，用于约束语音、屏幕和机器人动作建议。"],
            ["教师复核结果", "老师或项目成员对系统判断的校正。", "当前可先线下记录，后续接入页面复核按钮。", "沉淀训练样本，持续改进状态识别和支持策略。"],
        ],
        [1.25, 1.35, 2.15, 1.75],
    )

    add_heading(doc, "2. 当前文件流转步骤", 1)
    add_number(doc, "学校或 Seewo 后台导出一节课的视频和评课详情。")
    add_number(doc, "项目团队拿到授权材料后，在本地电脑建立该节课的数据目录。")
    add_number(doc, "项目团队从视频中确定目标学生，截取参考图，并标注观察区域。")
    add_number(doc, "学生资料先做匿名化和摘要化处理，只保留教学支持需要的字段。")
    add_number(doc, "启动 Demo 后，系统自动读取本地文件、抽取特征、生成状态时间线和建议。")
    add_number(doc, "浏览器页面展示结果，同时在 outputs 目录保存 analysis.json 等结构化输出。")

    add_heading(doc, "3. 自动与手动边界", 1)
    add_matrix(
        doc,
        ["版本", "文件进入方式", "系统自动做什么", "适用场景"],
        [
            ["当前展示版", "手动：项目团队提前把视频、文档、图片和资料放到本地目录。", "自动读取本地文件、分析状态、生成页面和 JSON。", "给校方演示可行性，避免现场依赖账号、网络和学校后台权限。"],
            ["近期试点版", "半手动：增加浏览器上传页面，由授权老师或项目成员上传一节课数据包。", "自动校验文件、分析并进入学生档案，老师可在页面复核。", "校内小范围试点，方便老师实际使用。"],
            ["正式集成版", "自动或准自动：在学校授权和 Seewo 接口允许时，做定时同步或后台对接。", "自动建立课程记录、学生追踪、复核任务和阶段报告。", "长期部署，但需要数据授权、接口权限和隐私合规流程。"],
        ],
        [1.15, 2.10, 2.05, 1.20],
    )

    add_heading(doc, "4. 当前 Demo 输出什么", 1)
    add_matrix(
        doc,
        ["输出", "保存/展示位置", "给谁看", "含义"],
        [
            ["Web 展示页", "http://127.0.0.1:8791/seewo_student_demo，仅演示电脑可直接打开。", "校方、老师、项目团队。", "可视化展示学生状态、证据、建议动作和教师支持效果。"],
            ["状态时间线", "页面时间线模块和 analysis.json。", "老师、项目成员。", "按视频片段记录学生状态、置信度、证据和可能原因。"],
            ["human_state JSON", "outputs/seewo_xiaoyudi_demo/analysis.json。", "研发和后续机器人接口。", "结构化字段包括 timestamp、person_id、state、confidence、evidence、risk_level、recommended_robot_action 等。"],
            ["学生追踪档案", "页面个体档案模块和 processed JSON。", "老师、教研负责人、项目团队。", "把学生基础资料、每节课表现、支持建议和复核记录长期沉淀。"],
            ["教师支持效果评估", "页面教师效果模块和 API 输出。", "老师和校方教研。", "以学生状态变化为依据，量化课堂支持策略是否有效。"],
            ["机器人/屏幕动作建议", "页面动作建议模块和 human_state JSON。", "未来机器人控制端或屏幕端。", "把分析结果落到可执行动作：屏幕提示、语音提示、低刺激动作、提醒老师。"],
        ],
        [1.35, 1.75, 1.20, 2.20],
    )

    add_heading(doc, "5. 校方展示现场怎么说", 1)
    add_bullet(doc, "如果校方问“文件从哪里上传”：当前不让老师现场上传，项目团队提前把 Seewo 导出的文件手动导入本地目录。")
    add_bullet(doc, "如果校方问“是否自动接入 Seewo”：当前没有自动接入，也不会主动读取学校后台；后续需要学校授权和接口条件。")
    add_bullet(doc, "如果校方问“以后能不能自动化”：可以分两步，先做浏览器上传，再做授权接口同步。")
    add_bullet(doc, "如果校方问“数据会不会外传”：展示版跑在本地电脑；正式试点可部署在校内电脑或校内服务器，原始视频不出校。")
    add_bullet(doc, "如果校方问“老师要做多少工作”：当前主要由项目团队整理；试点期老师只需要确认学生、上传课堂包、复核少量关键片段。")

    add_heading(doc, "6. 建议的后续上传入口设计", 1)
    add_matrix(
        doc,
        ["页面步骤", "老师/项目成员操作", "系统处理"],
        [
            ["新建课堂记录", "选择学生、课程、日期和学科。", "建立 lesson_id，并加载该学生历史支持画像。"],
            ["上传课堂包", "上传 mp4 视频、Seewo 评课 docx、目标学生参考图。", "检查文件格式、时长、清晰度和是否缺项。"],
            ["选择观察学生", "在首帧或截图中框选目标学生。", "生成目标学生 ROI，用于后续片段分析。"],
            ["自动分析", "等待系统处理。", "抽帧、音频摘要、状态识别、生成建议和风险提示。"],
            ["老师复核", "确认或修正关键片段标签。", "把复核结果写回学生档案和训练样本库。"],
        ],
        [1.45, 2.30, 2.75],
    )

    add_heading(doc, "7. 隐私与合规边界", 1)
    add_bullet(doc, "所有学生资料默认匿名化，展示材料中不出现真实姓名、身份证号、家庭信息等敏感内容。")
    add_bullet(doc, "系统输出的是课堂支持观察，不做医学诊断，不替代老师或医生判断。")
    add_bullet(doc, "原始视频、声音和个人资料应在校方授权范围内使用；如学校要求，可只在校内设备处理。")
    add_bullet(doc, "公开汇报只展示脱敏截图、模拟数据或获得授权后的样例，不传播学生原始视频。")

    add_heading(doc, "8. 演示前准备清单", 1)
    add_matrix(
        doc,
        ["检查项", "状态要求"],
        [
            ["本地服务", "演示电脑提前启动服务，确认 /seewo_student_demo 页面可打开。"],
            ["数据文件", "确认 lesson_video.mp4、seewo_lesson_report.docx、target_student.png 和学生画像 JSON 均存在。"],
            ["展示口径", "明确说明当前是手动导入，不是自动拉取学校数据。"],
        ],
        [1.75, 4.75],
    )

    doc.core_properties.title = "Seewo Demo 输入输出与文件流转说明"
    doc.core_properties.subject = "Demo input output and school data flow"
    doc.core_properties.author = "RehabAffectPerception"
    doc.core_properties.created = datetime.now()
    doc.save(IO_DOCX)


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    build_solution_doc()
    build_io_doc()
    print(SOLUTION_DOCX)
    print(IO_DOCX)


if __name__ == "__main__":
    main()
