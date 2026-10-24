"""Local school-package intake, isolated from active dossiers and video results.

Only an operator can import a local archive. Receipt is not teacher review or
training consent. Originals and position-preserving extraction are encrypted.
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import re
import stat
import zipfile
from datetime import date, timedelta
from pathlib import Path, PurePosixPath, PureWindowsPath
from uuid import uuid4

from .demo2_dossier_store import DossierError, _file_content
from .demo2_dossier_vault import DossierVault
from .demo2_identity_store import Demo2IdentityStore
from .student_record_store import _iso

DOMAINS = ('感知', '认知', '生活语文', '生活数学', '生活适应', '言语沟通', '运动与保健')
MISSING = [
    '与本批教案对应的课堂视频、录制日期和课次；现有视频不自动关联。',
    '每名学生在对应视频中的位置或校方确认截图；不按IEP姓名猜测人脸。',
    '确认这份学期IEP是否仍适用，并提供当前课的个别化目标、允许的提示与作答方式。',
    '确认资料使用范围与保留期限；如需训练，另行明确训练授权并复核片段。',
    '若要判断进步，需要同一目标的多次观测及教师复核；本批文件不构成纵向课堂样本。',
]


def _json(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':')).encode('utf8')


def read_package(path):
    """Read without extracting ZIP paths; reject oversized or active content."""
    from docx import Document

    path = Path(path)
    if path.stat().st_size > 40 * 1024 * 1024:
        raise ValueError('资料包超过40 MiB')
    raw = path.read_bytes()
    documents = []
    with zipfile.ZipFile(io.BytesIO(raw), metadata_encoding='gbk') as archive:
        entries = archive.infolist()
        if not 0 < len(entries) <= 100 or sum(e.file_size for e in entries) > 64 * 1024 * 1024:
            raise ValueError('资料包条目数或解压大小超限')
        seen = set()
        for entry in entries:
            name = entry.filename.replace('\\', '/')
            parts = PurePosixPath(name)
            if (parts.is_absolute() or '..' in parts.parts or PureWindowsPath(name).drive
                    or ':' in name or entry.flag_bits & 1 or stat.S_ISLNK(entry.external_attr >> 16)):
                raise ValueError('资料包路径或条目类型不安全')
            if entry.is_dir():
                continue
            if (name in seen or parts.suffix.lower() != '.docx' or entry.file_size > 8 * 1024 * 1024
                    or entry.file_size > max(1, entry.compress_size) * 100):
                raise ValueError('仅支持无重复的安全Word资料，单份不超过8 MiB')
            seen.add(name)
            content, filename, _ = _file_content(parts.name, base64.b64encode(archive.read(entry)).decode())
            doc = Document(io.BytesIO(content))
            # Keep every cell/column: a checkmark's position can carry a score.
            tables = [[[(cell.text) for cell in row.cells] for row in table.rows] for table in doc.tables]
            paragraphs = [p.text for p in doc.paragraphs]
            match = re.match(r'^IEP[-－](.+?)[-－]', filename, re.IGNORECASE)
            if not match and '教案' not in filename:
                raise ValueError('无法确认资料类型，请先核对文件命名')
            all_text = '\n'.join(paragraphs + [c for t in tables for r in t for c in r])
            domains = {domain: [i + 1 for i, table in enumerate(tables)
                                if any(domain in cell for row in table for cell in row)] for domain in DOMAINS}
            documents.append(dict(filename=filename, content=content, sha256=hashlib.sha256(content).hexdigest(),
                                  kind='iep' if match else 'lesson_plan', student_name=match.group(1) if match else '',
                                  table_count=len(tables), domains={k: v for k, v in domains.items() if v},
                                  source_school='上海市青浦区辅读学校' if '上海市青浦区辅读学校' in all_text else '',
                                  extraction=dict(paragraphs=paragraphs, tables=tables),
                                  term='2025学年第二学期' if '2025学年第二学期' in filename else '待核对'))
    if not documents:
        raise ValueError('资料包没有可处理的Word文件')
    return hashlib.sha256(raw).hexdigest(), documents


class MaterialIntakeStore:
    def __init__(self, records):
        self.records = records
        self.identities = Demo2IdentityStore(records)
        self.vault = DossierVault(Path(records.upload_dir).resolve().parent / 'school_material_intake')
        with records._connect() as db:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS demo2_material_batches (
                    batch_id TEXT PRIMARY KEY, owner_account_id TEXT NOT NULL,
                    scope_key TEXT NOT NULL, archive_sha256 TEXT NOT NULL,
                    created_at TEXT NOT NULL, retain_until TEXT NOT NULL,
                    status TEXT NOT NULL, student_count INTEGER NOT NULL,
                    document_count INTEGER NOT NULL, metadata_ciphertext BLOB,
                    UNIQUE(owner_account_id,scope_key,archive_sha256)
                );
                CREATE TABLE IF NOT EXISTS demo2_material_documents (
                    document_id TEXT PRIMARY KEY, batch_id TEXT NOT NULL
                    REFERENCES demo2_material_batches(batch_id), student_id TEXT,
                    kind TEXT NOT NULL, source_sha256 TEXT NOT NULL,
                    content_ciphertext BLOB, metadata_ciphertext BLOB,
                    UNIQUE(batch_id,source_sha256)
                );
            ''')

    def cleanup_expired(self):
        today = date.today().isoformat()
        with self.records._connect() as db:
            db.execute('''UPDATE demo2_material_documents SET content_ciphertext=NULL, metadata_ciphertext=NULL
                WHERE batch_id IN (SELECT batch_id FROM demo2_material_batches WHERE retain_until<=?)''', (today,))
            db.execute('''UPDATE demo2_material_batches SET status='expired',metadata_ciphertext=NULL
                WHERE retain_until<=?''', (today,))

    def import_local(self, token, path):
        """Operator-only method; deliberately not exposed as an HTTP endpoint."""
        actor = self.identities._account(token, write=True)
        digest, documents = read_package(path)
        names = [d['student_name'] for d in documents if d['kind'] == 'iep']
        if len(names) != len(set(names)):
            raise ValueError('同名IEP需要先核对版本，不自动合并')
        if len({d['sha256'] for d in documents}) != len(documents):
            raise ValueError('资料包含重复文件内容')
        today, now = date.today(), _iso()
        batch_id = uuid4().hex
        with self.records._connect() as db:
            db.execute('BEGIN IMMEDIATE')
            context = self.identities._context(self.records._school_context_for_account(db, actor['account_id']))
            scope = self.identities._scope(actor['account_id'], context)
            existing = db.execute('''SELECT batch_id FROM demo2_material_batches
                WHERE owner_account_id=? AND scope_key=? AND archive_sha256=?''',
                (actor['account_id'], scope, digest)).fetchone()
            if existing:
                return dict(batch_id=existing['batch_id'], duplicate=True, documents=len(documents), students=len(names))
            mapping = {}
            for name in names:
                # Never merge two people based on a filename alone.
                if db.execute('''SELECT 1 FROM demo2_identity_students
                    WHERE owner_account_id=? AND scope_key=? AND display_name=?''',
                    (actor['account_id'], scope, name)).fetchone():
                    raise ValueError('本校存在同名学生，请人工核对后关联，不自动覆盖')
                student_id = uuid4().hex
                mapping[name] = student_id
                db.execute('''INSERT INTO demo2_identity_students
                    (student_id,owner_account_id,scope_key,school_context_json,display_name,student_code,class_name,note,created_at)
                    VALUES (?,?,?,?,?,'','','校方资料独立入库；班级、校内编号与视频身份待核对',?)''',
                    (student_id, actor['account_id'], scope, _json(context).decode(), name, now))
            metadata = dict(school_context=context, source_school_names=sorted({d['source_school'] for d in documents if d['source_school']}),
                            video_count=0, training_allowed=False, teacher_reviewed=False,
                            association='independent_package_no_existing_video_link', missing=MISSING)
            db.execute('INSERT INTO demo2_material_batches VALUES (?,?,?,?,?,?,?,?,?,?)',
                       (batch_id, actor['account_id'], scope, digest, now, (today + timedelta(days=60)).isoformat(),
                        'received_pending_review', len(names), len(documents), self.vault.seal(_json(metadata))))
            for document in documents:
                meta = {k: v for k, v in document.items() if k not in {'content', 'sha256', 'kind'}}
                db.execute('INSERT INTO demo2_material_documents VALUES (?,?,?,?,?,?,?)',
                           (uuid4().hex, batch_id, mapping.get(document['student_name']), document['kind'], document['sha256'],
                            self.vault.seal(document['content']), self.vault.seal(_json(meta))))
            self.records._audit(db, 'local_project_import', 'school_material_received', 'material_batch', batch_id,
                                dict(students=len(names), documents=len(documents), video_links=0, training_allowed=False))
        return dict(batch_id=batch_id, duplicate=False, documents=len(documents), students=len(names))

    def _readable(self, db, actor, batch):
        context = self.identities._context(self.records._school_context_for_account(db, actor['account_id']), required=False)
        return (actor['role'] == 'school_uploader' and batch['owner_account_id'] == actor['account_id']
                and batch['scope_key'] == self.identities._scope(actor['account_id'], context))

    def workspace(self, token):
        actor = self.identities._account(token)
        self.cleanup_expired()
        with self.records._connect() as db:
            batches = []
            for row in db.execute('SELECT * FROM demo2_material_batches ORDER BY created_at DESC'):
                if actor['role'] != 'company_admin' and not self._readable(db, actor, row):
                    continue
                item = {k: row[k] for k in ('batch_id', 'created_at', 'retain_until', 'status', 'student_count', 'document_count')}
                item['days_remaining'] = max(0, (date.fromisoformat(row['retain_until']) - date.today()).days)
                item['expiry_reminder'] = item['days_remaining'] <= 15
                item['documents'] = []
                if actor['role'] != 'company_admin' and row['metadata_ciphertext']:
                    item.update(json.loads(self.vault.open(row['metadata_ciphertext'])))
                    for doc in db.execute('SELECT * FROM demo2_material_documents WHERE batch_id=? ORDER BY kind,document_id', (row['batch_id'],)):
                        metadata = json.loads(self.vault.open(doc['metadata_ciphertext']))
                        item['documents'].append(dict(document_id=doc['document_id'], student_id=doc['student_id'], kind=doc['kind'],
                                                      **{k: metadata[k] for k in ('filename', 'student_name', 'table_count', 'domains', 'term')}))
                batches.append(item)
            self.records._audit(db, actor['username'], 'material_inventory_read', 'material_batch', '', {'batch_count': len(batches)})
        return dict(batches=batches, admin_metadata_only=actor['role'] == 'company_admin', training_allowed=False)

    def download(self, token, document_id):
        actor = self.identities._account(token, write=True)
        self.cleanup_expired()
        with self.records._connect() as db:
            row = db.execute('''SELECT d.*,b.owner_account_id,b.scope_key,b.status FROM demo2_material_documents d
                JOIN demo2_material_batches b ON b.batch_id=d.batch_id WHERE d.document_id=?''', (document_id,)).fetchone()
            if row is None or not self._readable(db, actor, row):
                raise DossierError(404, 'material_not_found', '资料不存在或无权访问')
            if row['status'] == 'expired' or not row['content_ciphertext']:
                raise DossierError(410, 'material_expired', '资料已到期')
            metadata = json.loads(self.vault.open(row['metadata_ciphertext']))
            content = self.vault.open(row['content_ciphertext'])
            self.records._audit(db, actor['username'], 'material_original_downloaded', 'material_document', document_id, {})
            return content, metadata['filename']
