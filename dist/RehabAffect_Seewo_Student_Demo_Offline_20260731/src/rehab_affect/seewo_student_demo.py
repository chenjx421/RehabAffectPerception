from __future__ import annotations

import json
import math
import re
import subprocess
import wave
from dataclasses import dataclass
from pathlib import Path
from statistics import mean
from typing import Any

import numpy as np


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data" / "raw" / "seewo_xiaoyudi_001"
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed" / "seewo_xiaoyudi_001"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "seewo_xiaoyudi_demo"
VIDEO_PATH = DATA_DIR / "lesson_video.mp4"
REPORT_PATH = DATA_DIR / "seewo_lesson_report.docx"
TARGET_IMAGE_PATH = DATA_DIR / "target_student.png"
STUDENT_SUPPORT_PROFILE_PATH = PROCESSED_DIR / "student_A_support_profile.json"
ANALYSIS_PATH = OUTPUT_DIR / "analysis.json"
DEMO_UI_VERSION = "student_support_v4"

TARGET_ROI = {"x": 485, "y": 225, "w": 80, "h": 125}
CONTEXT_ROI = {"x": 395, "y": 220, "w": 230, "h": 160}
LEFT_ADULT_ROI = {"x": 300, "y": 210, "w": 90, "h": 220}
RIGHT_ADULT_ROI = {"x": 570, "y": 210, "w": 130, "h": 220}

STATE_ZH = {
    "engaged": "参与",
    "low_attention": "注意力下降",
    "confused_or_stuck": "困惑/停滞",
    "uncertain_adult_interference": "成人干扰下不确定",
    "uncertain": "无法判断",
}

ACTION_LIBRARY = {
    "engaged": {
        "action_id": "continue_current_task",
        "screen": "保持当前教学画面，不额外弹窗",
        "speech": "保持当前节奏，必要时给予具体正反馈。",
        "robot_motion": "idle_or_soft_nod",
        "teacher_alert": False,
    },
    "low_attention": {
        "action_id": "regain_attention_low_stimulation",
        "screen": "轻微高亮当前步骤或目标物",
        "speech": "我们先看这里，只做这一步。",
        "robot_motion": "slow_turn_to_screen",
        "teacher_alert": False,
    },
    "confused_or_stuck": {
        "action_id": "repeat_one_step_with_visual_highlight",
        "screen": "只保留当前一步，弱化无关信息",
        "speech": "没关系，我们先看第一步。",
        "robot_motion": "slow_turn_to_screen_and_nod",
        "teacher_alert": True,
    },
    "uncertain_adult_interference": {
        "action_id": "defer_to_teacher_observation",
        "screen": "不触发学生端强提示",
        "speech": "系统记录为需要人工确认的片段。",
        "robot_motion": "no_motion",
        "teacher_alert": True,
    },
    "uncertain": {
        "action_id": "mark_for_review",
        "screen": "不触发自动干预",
        "speech": "该片段画质或遮挡不足，建议人工确认。",
        "robot_motion": "no_motion",
        "teacher_alert": True,
    },
}


@dataclass(frozen=True)
class VideoMeta:
    fps: float
    frame_count: int
    width: int
    height: int
    duration_sec: float

    def to_dict(self) -> dict[str, Any]:
        return {
            "fps": self.fps,
            "frame_count": self.frame_count,
            "width": self.width,
            "height": self.height,
            "duration_sec": self.duration_sec,
            "duration_min": self.duration_sec / 60,
        }


def load_seewo_student_demo(force: bool = False) -> dict[str, Any]:
    if not force and ANALYSIS_PATH.exists():
        try:
            cached = json.loads(ANALYSIS_PATH.read_text(encoding="utf-8"))
            if (
                cached.get("ui_version") == DEMO_UI_VERSION
                and "student_tracking" in cached
                and "teacher_effectiveness" in cached
            ):
                return cached
        except json.JSONDecodeError:
            pass
    return build_seewo_student_demo()


def build_seewo_student_demo() -> dict[str, Any]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    meta = _read_video_meta(VIDEO_PATH)
    report = _parse_report(REPORT_PATH)
    support_profile = _load_student_support_profile()
    samples = _sample_video_features(VIDEO_PATH, meta)
    audio_windows = _extract_audio_windows(VIDEO_PATH, meta.duration_sec)
    timeline = _build_timeline(samples, audio_windows, meta)
    assets = _write_visual_assets(VIDEO_PATH, meta, timeline)
    summary = _summarize_timeline(timeline)
    student_tracking = _build_student_tracking(timeline, summary, report, support_profile)
    teacher_effectiveness = _build_teacher_effectiveness(timeline, summary, report, support_profile)

    payload = {
        "ui_version": DEMO_UI_VERSION,
        "demo_name": "Seewo 真实课堂单学生状态追踪 Demo",
        "data_source": {
            "video": str(VIDEO_PATH),
            "report": str(REPORT_PATH),
            "target_student_image": str(TARGET_IMAGE_PATH),
            "privacy_note": "本 Demo 使用本地授权样例文件；对外展示应使用匿名编号和必要模糊处理。",
        },
        "target_student": {
            "student_id": "student_A",
            "display_name": "目标学生 A",
            "description": "画面中间穿灰色上衣、红领巾的学生；左侧和右侧靠近人员按老师/秩序维护人员处理。",
            "target_roi": TARGET_ROI,
            "context_roi": CONTEXT_ROI,
            "adult_exclusion_rois": [LEFT_ADULT_ROI, RIGHT_ADULT_ROI],
            "support_profile": {
                "available": bool(support_profile),
                "placeholder": "已接入匿名化 IEP 和五大领域评估摘要。" if support_profile else "后续可接入既有专业评估、行为指导、禁忌刺激和有效提示方式。",
            },
            "support_profile_data": support_profile,
        },
        "video_meta": meta.to_dict(),
        "report_summary": report,
        "timeline": timeline,
        "aggregate": summary,
        "student_tracking": student_tracking,
        "teacher_effectiveness": teacher_effectiveness,
        "assets": assets,
        "review_boundary": {
            "video_quality": "全景低清视频，适合做 ROI 行为追踪，不适合做精细表情/眼神识别。",
            "model_status": "当前为可解释特征原型，不是已训练好的最终模型。",
            "teacher_labeling": "后续需老师抽样校正关键片段，形成训练标签。",
            "medical_boundary": "不做医学诊断；专业评估资料只作为安全约束和个体化支持背景。",
        },
    }
    ANALYSIS_PATH.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def _load_student_support_profile() -> dict[str, Any]:
    if not STUDENT_SUPPORT_PROFILE_PATH.exists():
        return {}
    try:
        profile = json.loads(STUDENT_SUPPORT_PROFILE_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    # Keep the web payload focused on anonymous support needs, not raw private files.
    return {
        "student_id": profile.get("student_id", "student_A"),
        "privacy_level": profile.get("privacy_level", "sensitive_private_local_only"),
        "diagnosis_context": profile.get("diagnosis_context", {}),
        "five_domain_assessment": profile.get("five_domain_assessment", {}),
        "classroom_support_implications": profile.get("classroom_support_implications", {}),
        "knowledge_base_topics_to_add": profile.get("knowledge_base_topics_to_add", []),
    }


def render_seewo_student_demo_html() -> str:
    data = load_seewo_student_demo()
    payload = json.dumps(data, ensure_ascii=False)
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Seewo 学生成长追踪 Demo</title>
  <style>
    :root {{
      --bg: #edf2f7;
      --panel: #ffffff;
      --panel-soft: #f8fafc;
      --ink: #101820;
      --muted: #627084;
      --line: #d7dee8;
      --line-strong: #c5d0dd;
      --blue: #2563eb;
      --cyan: #0891b2;
      --green: #16834a;
      --amber: #b76b00;
      --red: #c24135;
      --violet: #6d4aff;
      --shadow: 0 18px 44px rgba(27, 37, 52, .11);
      --soft-shadow: 0 10px 26px rgba(27, 37, 52, .075);
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font-family: Inter, "Microsoft YaHei", "Segoe UI", Arial, sans-serif;
      background:
        linear-gradient(180deg, #fbfcff 0, #f2f6fb 210px, #e9eff6 100%),
        linear-gradient(90deg, rgba(37,99,235,.06), rgba(8,145,178,.04));
      color: var(--ink);
    }}
    header {{
      position: sticky;
      top: 0;
      z-index: 3;
      display: flex;
      justify-content: space-between;
      align-items: center;
      gap: 18px;
      padding: 18px 24px 16px;
      background: rgba(255,255,255,.92);
      border-bottom: 1px solid rgba(199, 210, 224, .85);
      backdrop-filter: blur(12px);
      box-shadow: 0 8px 24px rgba(27, 37, 52, .06);
    }}
    .eyebrow {{
      color: var(--blue);
      font-size: 12px;
      font-weight: 900;
      letter-spacing: .04em;
      text-transform: uppercase;
      margin-bottom: 5px;
    }}
    h1 {{ margin: 0; font-size: 28px; letter-spacing: 0; line-height: 1.12; }}
    .sub {{ margin-top: 5px; color: var(--muted); font-size: 13px; }}
    .badge-stack {{
      display: flex;
      align-items: center;
      gap: 8px;
      flex-wrap: wrap;
      justify-content: flex-end;
    }}
    .badge {{
      display: inline-flex;
      align-items: center;
      min-height: 30px;
      padding: 4px 10px;
      border-radius: 999px;
      font-weight: 800;
      color: #fff;
      background: linear-gradient(135deg, #1f5bd7, #1687c7);
      font-size: 12px;
      box-shadow: 0 8px 16px rgba(37,99,235,.18);
    }}
    .status-pill {{
      display: inline-flex;
      align-items: center;
      gap: 6px;
      min-height: 30px;
      padding: 4px 10px;
      border-radius: 999px;
      border: 1px solid var(--line);
      color: #203044;
      background: linear-gradient(180deg, #fff 0, #f8fbff 100%);
      font-size: 12px;
      font-weight: 900;
      cursor: pointer;
      transition: border-color .12s ease, color .12s ease, box-shadow .12s ease;
    }}
    .status-pill:hover {{
      border-color: var(--blue);
      color: var(--blue);
      box-shadow: 0 8px 18px rgba(25, 91, 215, .12);
    }}
    .status-pill:focus-visible {{
      outline: 3px solid rgba(37, 99, 235, .22);
      outline-offset: 2px;
    }}
    main {{
      display: grid;
      grid-template-columns: minmax(330px, 360px) minmax(560px, 1.08fr) minmax(360px, 390px);
      gap: 18px;
      padding: 18px 24px 30px;
      align-items: start;
    }}
    .panel {{
      background: var(--panel);
      border: 1px solid rgba(199, 210, 224, .92);
      border-radius: 8px;
      box-shadow: var(--shadow);
      overflow: hidden;
    }}
    main > .panel:last-child {{
      position: sticky;
      top: 96px;
    }}
    .panel h2 {{
      margin: 0;
      padding: 15px 15px 12px;
      font-size: 17px;
      border-bottom: 1px solid var(--line);
      background: linear-gradient(180deg, #fff 0, #f7fafc 100%);
      letter-spacing: 0;
    }}
    .body {{ padding: 14px; }}
    .hero-img {{
      width: 100%;
      border: 1px solid var(--line-strong);
      border-radius: 7px;
      display: block;
      background: #f7f9fb;
      box-shadow: 0 10px 22px rgba(27,37,52,.09);
    }}
    .meta-grid {{
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 8px;
      margin-top: 12px;
    }}
    .metric {{
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 10px;
      background: linear-gradient(180deg, #fff 0, #f8fafc 100%);
      min-height: 66px;
    }}
    .metric .k {{ color: var(--muted); font-size: 12px; font-weight: 700; }}
    .metric .v {{ margin-top: 4px; font-size: 20px; font-weight: 950; color: #0f2544; }}
    .note {{
      border-left: 4px solid var(--blue);
      background: #eef5ff;
      border-radius: 6px;
      padding: 11px;
      font-size: 13px;
      line-height: 1.55;
      color: #203044;
      margin-bottom: 12px;
    }}
    .lead {{
      border-left-color: #1f5bd7;
      background: #edf5ff;
      font-size: 14px;
      font-weight: 700;
    }}
    .insight-strip {{
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 10px;
      margin-bottom: 12px;
    }}
    .insight {{
      border: 1px solid var(--line);
      border-top: 4px solid var(--blue);
      border-radius: 7px;
      padding: 11px 12px 12px;
      background: linear-gradient(180deg, #fff 0, #f7fafc 100%);
      box-shadow: var(--soft-shadow);
      min-height: 86px;
    }}
    .insight:nth-child(1) {{ border-top-color: var(--green); }}
    .insight:nth-child(2) {{ border-top-color: var(--amber); }}
    .insight:nth-child(3) {{ border-top-color: var(--violet); }}
    .insight:nth-child(4) {{ border-top-color: var(--cyan); }}
    .insight .k {{
      color: var(--muted);
      font-size: 12px;
      font-weight: 900;
    }}
    .insight .v {{
      margin-top: 5px;
      font-size: 24px;
      font-weight: 950;
      line-height: 1.05;
    }}
    .insight .hint {{
      margin-top: 6px;
      color: var(--muted);
      font-size: 12px;
      line-height: 1.35;
    }}
    .closed-loop {{
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 8px;
      margin-bottom: 12px;
    }}
    .loop-step {{
      position: relative;
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 10px;
      background: linear-gradient(180deg, #fff 0, #fbfcff 100%);
      min-height: 72px;
    }}
    .loop-step::before {{
      content: "";
      position: absolute;
      left: 10px;
      top: 10px;
      width: 4px;
      height: calc(100% - 20px);
      border-radius: 999px;
      background: #dbeafe;
    }}
    .loop-step strong,
    .loop-step span {{
      margin-left: 10px;
    }}
    .loop-step strong {{
      display: block;
      font-size: 13px;
      margin-bottom: 4px;
    }}
    .loop-step span {{
      color: var(--muted);
      font-size: 12px;
      line-height: 1.35;
    }}
    .teacher-effectiveness {{
      display: grid;
      gap: 10px;
      margin-bottom: 12px;
    }}
    .teacher-summary {{
      border: 1px solid #a9d4e4;
      border-left: 5px solid var(--cyan);
      border-radius: 7px;
      padding: 13px;
      background: linear-gradient(135deg, #f7fcff 0, #eef9fc 58%, #f7f4ff 100%);
      box-shadow: var(--soft-shadow);
    }}
    .teacher-summary .score-row {{
      display: flex;
      justify-content: space-between;
      gap: 12px;
      align-items: center;
    }}
    .score-dial {{
      --score: 0;
      width: 78px;
      height: 78px;
      flex: 0 0 auto;
      border-radius: 50%;
      display: grid;
      place-items: center;
      background:
        radial-gradient(circle at center, #fff 0 58%, transparent 59%),
        conic-gradient(var(--cyan) calc(var(--score) * 1%), #d9e7ef 0);
      box-shadow: inset 0 0 0 1px rgba(8,145,178,.10), 0 10px 22px rgba(8,145,178,.14);
    }}
    .score-dial span {{
      font-size: 26px;
      line-height: 1;
      font-weight: 950;
      color: #0f5f75;
    }}
    .score-dial small {{
      display: block;
      margin-top: 2px;
      font-size: 10px;
      font-weight: 800;
      color: var(--muted);
      text-align: center;
    }}
    .teacher-summary .label {{
      font-size: 18px;
      font-weight: 950;
      margin-bottom: 5px;
    }}
    .teacher-metric-grid {{
      display: grid;
      grid-template-columns: repeat(5, minmax(0, 1fr));
      gap: 9px;
    }}
    .teacher-metric {{
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 9px;
      background: linear-gradient(180deg, #fff 0, #fbfcfe 100%);
      min-height: 112px;
      border-top-width: 4px;
      box-shadow: 0 8px 18px rgba(27,37,52,.06);
    }}
    .teacher-metric.good {{ border-top-color: var(--green); }}
    .teacher-metric.watch {{ border-top-color: var(--amber); }}
    .teacher-metric.support {{ border-top-color: var(--violet); }}
    .teacher-metric.good .fill {{ background: var(--green); }}
    .teacher-metric.watch .fill {{ background: var(--amber); }}
    .teacher-metric.support .fill {{ background: var(--violet); }}
    .teacher-metric strong {{
      display: flex;
      justify-content: space-between;
      gap: 6px;
      font-size: 12px;
      margin-bottom: 6px;
    }}
    .teacher-metric .score-mini {{
      font-size: 14px;
      color: #0f5f75;
      font-weight: 950;
    }}
    .timeline {{
      display: grid;
      gap: 8px;
      max-height: 350px;
      overflow: auto;
      padding-right: 4px;
    }}
    .section-title {{
      margin: 14px 0 8px;
      color: #263445;
      font-size: 14px;
      font-weight: 900;
    }}
    .window-list {{
      display: grid;
      grid-template-columns: repeat(auto-fill, minmax(64px, 1fr));
      gap: 6px;
      max-height: none;
      overflow: visible;
      margin-top: 8px;
    }}
    .window-row {{
      border: 1px solid var(--line);
      border-top-width: 4px;
      border-radius: 7px;
      padding: 7px 6px;
      background: #fff;
      cursor: pointer;
      min-height: 58px;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      transition: transform .12s ease, box-shadow .12s ease;
    }}
    .window-row:hover {{ transform: translateY(-1px); box-shadow: 0 8px 18px rgba(25,36,50,.10); }}
    .window-row.active {{ outline: 2px solid rgba(37,99,235,.22); background: #f8fbff; }}
    .window-row.engaged {{ border-top-color: var(--green); }}
    .window-row.low_attention {{ border-top-color: var(--amber); }}
    .window-row.confused_or_stuck {{ border-top-color: var(--violet); }}
    .window-row.uncertain_adult_interference, .window-row.uncertain {{ border-top-color: var(--red); }}
    .window-time {{ font-size: 12px; font-weight: 900; }}
    .window-state {{ font-size: 11px; color: var(--muted); white-space: nowrap; overflow: hidden; text-overflow: ellipsis; }}
    .event {{
      border: 1px solid var(--line);
      border-left-width: 5px;
      border-radius: 7px;
      padding: 10px 12px;
      background: #fff;
      cursor: pointer;
    }}
    .event.active {{ outline: 2px solid rgba(37,99,235,.25); background: #f6f9ff; }}
    .event.engaged {{ border-left-color: var(--green); }}
    .event.low_attention {{ border-left-color: var(--amber); }}
    .event.confused_or_stuck {{ border-left-color: var(--violet); }}
    .event.uncertain_adult_interference, .event.uncertain {{ border-left-color: var(--red); }}
    .event-top {{
      display: flex;
      justify-content: space-between;
      gap: 8px;
      align-items: center;
      font-weight: 900;
    }}
    .tag {{
      color: white;
      background: var(--blue);
      border-radius: 999px;
      padding: 3px 8px 4px;
      font-size: 12px;
      white-space: nowrap;
    }}
    .tag.engaged {{ background: var(--green); }}
    .tag.low_attention {{ background: var(--amber); }}
    .tag.confused_or_stuck {{ background: var(--violet); }}
    .tag.uncertain_adult_interference, .tag.uncertain {{ background: var(--red); }}
    .bar {{
      height: 8px;
      background: #e6edf5;
      border-radius: 999px;
      overflow: hidden;
      margin-top: 7px;
    }}
    .fill {{ height: 100%; background: var(--blue); }}
    .small {{ color: var(--muted); font-size: 12px; line-height: 1.5; }}
    .state-mix {{
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 8px;
      margin: 8px 0 12px;
    }}
    .state-chip {{
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 9px;
      background: #fff;
    }}
    .state-chip strong {{
      display: flex;
      justify-content: space-between;
      gap: 8px;
      font-size: 12px;
      margin-bottom: 7px;
    }}
    .state-chip .bar {{ margin: 0; height: 7px; }}
    .state-chip.engaged .fill {{ background: var(--green); }}
    .state-chip.low_attention .fill {{ background: var(--amber); }}
    .state-chip.confused_or_stuck .fill {{ background: var(--violet); }}
    .state-chip.uncertain_adult_interference .fill, .state-chip.uncertain .fill {{ background: var(--red); }}
    .detail-card {{
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 12px;
      margin-bottom: 10px;
      background: linear-gradient(180deg, #fff 0, #fbfcfe 100%);
    }}
    .detail-card.is-highlighted {{
      border-color: var(--blue);
      box-shadow: 0 0 0 4px rgba(37, 99, 235, .12);
    }}
    .selected-head {{
      display: flex;
      justify-content: space-between;
      gap: 10px;
      align-items: flex-start;
      margin-bottom: 8px;
    }}
    .selected-head h3 {{ margin: 0; font-size: 16px; }}
    .action-hero {{
      border: 1px solid #bfd4ff;
      border-left: 5px solid var(--blue);
      border-radius: 7px;
      padding: 12px;
      margin-bottom: 10px;
      background: linear-gradient(180deg, #f7fbff 0, #eef5ff 100%);
      box-shadow: var(--soft-shadow);
    }}
    .action-hero h3 {{ margin: 0 0 8px; font-size: 16px; }}
    .action-title {{
      font-size: 20px;
      font-weight: 950;
      line-height: 1.2;
      margin-bottom: 6px;
    }}
    .action-grid {{
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 8px;
      margin-top: 10px;
    }}
    .action-cell {{
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 9px;
      background: #fff;
      line-height: 1.35;
      min-height: 72px;
      font-size: 13px;
      overflow-wrap: anywhere;
    }}
    .action-cell strong {{
      display: block;
      font-size: 12px;
      color: var(--muted);
      margin-bottom: 4px;
    }}
    .tracking-card {{
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 12px;
      margin-bottom: 10px;
      background: linear-gradient(180deg, #fff 0, #fbfcfe 100%);
    }}
    .tracking-card strong {{
      display: block;
      font-size: 16px;
      margin-bottom: 5px;
    }}
    .tracking-card.good {{ border-left: 5px solid var(--green); }}
    .tracking-card.watch {{ border-left: 5px solid var(--amber); }}
    .tracking-card.support {{ border-left: 5px solid var(--violet); }}
    .growth-grid {{
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 10px;
    }}
    .support-list {{
      display: grid;
      gap: 8px;
    }}
    .support-item {{
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 9px 10px;
      background: linear-gradient(180deg, #fff 0, #fbfcfe 100%);
      line-height: 1.45;
    }}
    .profile-grid {{
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 8px;
    }}
    .profile-chip {{
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 9px;
      background: linear-gradient(180deg, #fff 0, #fbfcfe 100%);
      min-height: 74px;
    }}
    .profile-chip strong {{
      display: block;
      font-size: 12px;
      color: var(--muted);
      margin-bottom: 4px;
    }}
    .domain-grid {{
      display: grid;
      grid-template-columns: repeat(3, minmax(0, 1fr));
      gap: 8px;
      margin-bottom: 12px;
    }}
    .domain-card {{
      border: 1px solid var(--line);
      border-radius: 7px;
      padding: 10px;
      background: linear-gradient(180deg, #fff 0, #fbfcfe 100%);
      min-height: 104px;
    }}
    .domain-card.missing {{
      background: #f8fafc;
      color: var(--muted);
      border-style: dashed;
    }}
    .domain-card strong {{
      display: block;
      margin-bottom: 5px;
      font-size: 13px;
    }}
    .profile-note {{
      border-left-color: var(--violet);
      background: #f4f1ff;
    }}
    .review-actions {{
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 6px;
      margin-top: 10px;
    }}
    .review-actions button {{
      min-height: 34px;
      padding: 6px 8px;
      font-size: 12px;
    }}
    .detail-card h3 {{ margin: 0 0 8px; font-size: 15px; }}
    ul {{ padding-left: 20px; margin: 6px 0; }}
    li {{ margin: 4px 0; }}
    pre {{
      white-space: pre-wrap;
      overflow: auto;
      max-height: 245px;
      margin: 0;
      padding: 12px;
      border-radius: 7px;
      background: #101820;
      color: #e7f2ec;
      font-size: 12px;
      line-height: 1.45;
    }}
    .tabs {{ display: flex; gap: 6px; flex-wrap: wrap; margin-bottom: 10px; }}
    button {{
      border: 1px solid var(--line);
      background: white;
      border-radius: 7px;
      padding: 8px 10px;
      font-weight: 800;
      cursor: pointer;
    }}
    button:hover {{ border-color: #b6c5d8; background: #f8fbff; }}
    button.active {{ border-color: var(--blue); color: var(--blue); background: #eff6ff; box-shadow: inset 0 0 0 1px rgba(37,99,235,.12); }}
    @media (max-width: 1100px) {{
      main {{ grid-template-columns: 1fr; }}
      main > .panel:last-child {{ position: static; }}
      .timeline {{ max-height: none; }}
      .insight-strip, .closed-loop, .state-mix, .growth-grid {{
        grid-template-columns: 1fr 1fr;
      }}
      .teacher-metric-grid {{ grid-template-columns: 1fr 1fr; }}
      .domain-grid {{ grid-template-columns: 1fr 1fr; }}
    }}
    @media (max-width: 640px) {{
      header {{ align-items: flex-start; flex-direction: column; }}
      .insight-strip, .closed-loop, .state-mix, .growth-grid, .action-grid, .profile-grid, .domain-grid, .teacher-metric-grid {{
        grid-template-columns: 1fr;
      }}
    }}
  </style>
</head>
<body>
  <header>
    <div>
      <div class="eyebrow">Seewo Complementary AI Tutor</div>
      <h1>学生个体状态追踪与课堂支持 Demo</h1>
      <div class="sub">课堂视频/评课报告 → 单学生状态估计 → 低刺激支持动作 → 老师复核反馈</div>
    </div>
    <div class="badge-stack">
      <span class="badge">real-data prototype</span>
      <button class="status-pill" id="localCaseJump" type="button" title="查看本地匿名个体档案和数据使用边界">本地匿名样例</button>
    </div>
  </header>
  <main>
    <section class="panel">
      <h2>学生画像与数据来源</h2>
      <div class="body">
        <div class="note lead" id="studentConclusion"></div>
        <img class="hero-img" src="{data['assets'].get('roi_reference_url', '')}" alt="目标学生 ROI 参考图" />
        <div class="meta-grid" id="metrics"></div>
        <div class="detail-card">
          <h3>学生追踪档案</h3>
          <div id="studentProfile" class="small"></div>
        </div>
        <div class="detail-card" id="profileCard">
          <h3>IEP 与五大领域档案</h3>
          <div id="foundationProfile" class="profile-grid"></div>
        </div>
        <div class="detail-card">
          <h3>课堂资料摘要</h3>
          <div id="reportSummary" class="small"></div>
        </div>
        <div class="detail-card">
          <h3>下一次课堂支持重点</h3>
          <div id="supportPlan" class="support-list"></div>
        </div>
      </div>
    </section>
    <section class="panel">
      <h2>学生支持结论</h2>
      <div class="body">
        <div class="insight-strip" id="insightStrip"></div>
        <div class="closed-loop" id="closedLoop"></div>
        <div class="section-title">教师支持效果评估</div>
        <div id="teacherEffectiveness" class="teacher-effectiveness"></div>
        <div class="section-title">IEP 与五大领域画像</div>
        <div id="domainProfile" class="domain-grid"></div>
        <div id="profileRationale" class="note profile-note"></div>
        <div id="growthIndicators"></div>
        <div class="note">上面是本节课的学生支持结论；下面保留完整 30 秒动作/状态窗口，用于训练数据、教师抽样校正和定位证据。</div>
        <div class="section-title">状态分布</div>
        <div class="state-mix" id="stateMix"></div>
        <div class="section-title">关键复核片段</div>
        <div class="timeline" id="timeline"></div>
        <div class="section-title">完整动作/状态窗口</div>
        <div class="window-list" id="windowList"></div>
      </div>
    </section>
    <section class="panel">
      <h2>复核与执行建议</h2>
      <div class="body">
        <div id="detail"></div>
        <div class="tabs">
          <button class="active" data-tab="event">当前片段</button>
          <button data-tab="action">动作方案</button>
          <button data-tab="tracking">学生追踪</button>
          <button data-tab="teacher">教师效果</button>
          <button data-tab="profile">个体档案</button>
          <button data-tab="summary">课堂汇总</button>
          <button data-tab="full">完整 JSON</button>
        </div>
        <pre id="json"></pre>
      </div>
    </section>
  </main>
  <script>
    const report = {payload};
    let currentIndex = 0;
    let currentMode = 'moment';
    let activeTab = 'event';
    const $ = (id) => document.getElementById(id);
    const fmt = (n) => Number.isFinite(Number(n)) ? Number(n).toFixed(2) : 'n/a';
    const pct = (n) => Number.isFinite(Number(n)) ? `${{Math.round(Number(n) * 100)}}%` : 'n/a';
    const mmss = (sec) => `${{String(Math.floor(sec/60)).padStart(2,'0')}}:${{String(Math.floor(sec%60)).padStart(2,'0')}}`;

    function stateLabel(state) {{
      const match = report.timeline.find((e) => e.state === state);
      return match ? match.state_zh : state;
    }}

    function domainLabel(domain) {{
      const labels = {{
        cognitive: '认知能力',
        language_communication: '语言沟通',
        motor: '运动能力',
        sensory_perception: '感知领域',
        social_adaptation: '社会适应'
      }};
      return labels[domain] || domain;
    }}

    function supportProfile() {{
      return report.target_student?.support_profile_data || {{}};
    }}

    function profileSummary() {{
      return report.student_tracking?.support_profile_summary || {{available: false}};
    }}

    function profileJsonForDisplay() {{
      const profile = supportProfile();
      const summary = profileSummary();
      const implication = profile.classroom_support_implications || {{}};
      return {{
        student_id: profile.student_id || 'student_A',
        privacy_level: profile.privacy_level || 'local_anonymous_demo',
        education_profile: summary.education_profile || profile.diagnosis_context || {{}},
        available_domains: summary.available_domains || [],
        missing_domains: summary.missing_domains || [],
        support_flags: summary.support_flags || {{}},
        individualization_rationale: summary.why || implication.why || [],
        robot_action_constraints: implication.robot_action_constraints || [],
        knowledge_base_topics_to_add: profile.knowledge_base_topics_to_add || [],
        boundary: '演示页只显示匿名摘要；原始学生资料不在页面公开展示。'
      }};
    }}

    function dominantActionTitle() {{
      const counts = {{}};
      report.timeline.forEach((e) => {{
        const id = e.recommended_action?.action_id || 'mark_for_review';
        counts[id] = (counts[id] || 0) + 1;
      }});
      const top = Object.entries(counts).sort((a, b) => b[1] - a[1])[0];
      return top ? actionTitle(top[0]) : '暂无动作';
    }}

    function renderMetrics() {{
      const meta = report.video_meta;
      const agg = report.aggregate;
      const tracking = report.student_tracking;
      const profile = supportProfile();
      const profileInfo = profileSummary();
      const diagnosis = profile.diagnosis_context || {{}};
      const domains = profile.five_domain_assessment || {{}};
      const teacherAlerts = report.timeline.filter((e) => e.recommended_action?.teacher_alert).length;
      const supportWindows = report.timeline.filter((e) => e.state !== 'engaged').length;
      const topAction = dominantActionTitle();
      $('studentConclusion').innerHTML = tracking.lesson_conclusion.summary;
      $('insightStrip').innerHTML = `
        <div class="insight"><div class="k">课堂参与基线</div><div class="v">${{pct(agg.state_ratio?.engaged)}}</div><div class="hint">来自 30 秒窗口聚合，不等同医学判断。</div></div>
        <div class="insight"><div class="k">需要支持窗口</div><div class="v">${{supportWindows}}</div><div class="hint">用于定位可能需要一步提示或观察的片段。</div></div>
        <div class="insight"><div class="k">老师复核建议</div><div class="v">${{teacherAlerts}}</div><div class="hint">成人遮挡、困惑停滞等片段优先复核。</div></div>
        <div class="insight"><div class="k">主要执行动作</div><div class="v" style="font-size:18px">${{topAction}}</div><div class="hint">输出给屏幕、语音、机器人或老师端。</div></div>
      `;
      $('closedLoop').innerHTML = `
        <div class="loop-step"><strong>1 感知</strong><span>Seewo 视频、课堂声音、评课报告与目标学生 ROI。</span></div>
        <div class="loop-step"><strong>2 状态估计</strong><span>按 30 秒窗口生成参与、注意力下降、困惑/停滞等状态。</span></div>
        <div class="loop-step"><strong>3 动作建议</strong><span>转成屏幕高亮、一步语音、机器人慢动作或老师提醒。</span></div>
        <div class="loop-step"><strong>4 复核学习</strong><span>老师确认后形成个体化追踪和下一次课堂支持策略。</span></div>
      `;
      const teacher = report.teacher_effectiveness || {{}};
      const teacherMetrics = teacher.metrics || [];
      $('teacherEffectiveness').innerHTML = `
        <div class="teacher-summary">
          <div class="score-row">
            <div>
              <div class="label">${{teacher.headline || '教师支持效果：待评估'}}</div>
              <div class="small">${{teacher.summary || '需要先完成学生状态分析。'}}</div>
            </div>
            <div class="score-dial" style="--score:${{teacher.overall_score || 0}}">
              <div><span>${{teacher.overall_score ?? '--'}}</span><small>/100</small></div>
            </div>
          </div>
        </div>
        <div class="teacher-metric-grid">
          ${{teacherMetrics.map((item) => `
            <div class="teacher-metric ${{item.level}}">
              <strong><span>${{item.name}}</span><span class="score-mini">${{item.score}}</span></strong>
              <div class="bar"><div class="fill" style="width:${{item.score}}%"></div></div>
              <div class="small">${{item.evidence}}</div>
            </div>
          `).join('')}}
        </div>
        <div class="note">评价边界：${{teacher.boundary || '该结果仅供课堂复盘参考。'}}</div>
      `;
      const orderedStates = ['engaged', 'low_attention', 'confused_or_stuck', 'uncertain_adult_interference']
        .filter((state) => agg.state_counts?.[state]);
      $('stateMix').innerHTML = orderedStates.map((state) => `
        <div class="state-chip ${{state}}">
          <strong><span>${{stateLabel(state)}}</span><span>${{pct(agg.state_ratio?.[state])}}</span></strong>
          <div class="bar"><div class="fill" style="width:${{pct(agg.state_ratio?.[state])}}"></div></div>
        </div>
      `).join('');
      $('metrics').innerHTML = `
        <div class="metric"><div class="k">课程时长</div><div class="v">${{fmt(meta.duration_min)}} 分</div></div>
        <div class="metric"><div class="k">视频规格</div><div class="v">${{meta.width}}×${{meta.height}}</div></div>
        <div class="metric"><div class="k">关键复核片段</div><div class="v">${{tracking.key_moments.length}}</div></div>
        <div class="metric"><div class="k">本节课结论</div><div class="v">${{agg.top_state_zh}}</div></div>
      `;
      $('studentProfile').innerHTML = `
        匿名 ID：<strong>${{tracking.student_profile.student_id}}</strong><br>
        当前阶段：${{tracking.student_profile.baseline_status}}<br>
        个体资料：${{tracking.student_profile.profile_status}}<br>
        支持约束：${{tracking.student_profile.support_constraint}}
      `;
      if (profileInfo.available) {{
        const available = (profileInfo.available_domains || []).map(domainLabel).join('、') || '暂无';
        const missing = (profileInfo.missing_domains || []).map(domainLabel).join('、') || '无';
        $('foundationProfile').innerHTML = `
          <div class="profile-chip"><strong>IEP 教育档案</strong>${{diagnosis.iep_primary_disability_category || '未标注'}}${{diagnosis.iep_primary_disability_degree ? ' / ' + diagnosis.iep_primary_disability_degree : ''}}</div>
          <div class="profile-chip"><strong>使用边界</strong>只作为支持背景，不作为医学诊断输出</div>
          <div class="profile-chip"><strong>已接入领域</strong>${{available}}</div>
          <div class="profile-chip"><strong>待补领域</strong>${{missing}}</div>
        `;
      }} else {{
        $('foundationProfile').innerHTML = `
          <div class="profile-chip"><strong>资料状态</strong>暂未接入 IEP 或五大领域评估。</div>
          <div class="profile-chip"><strong>下一步</strong>补充认知、语言、运动、感知和社会适应资料。</div>
        `;
      }}
      const orderedDomains = ['cognitive', 'language_communication', 'motor', 'sensory_perception', 'social_adaptation'];
      $('domainProfile').innerHTML = orderedDomains.map((domain) => {{
        const item = domains[domain] || {{available: false}};
        if (!item.available) {{
          return `<div class="domain-card missing"><strong>${{domainLabel(domain)}}</strong><div class="small">本样例暂缺，数据库字段已预留。</div></div>`;
        }}
        const weak = (item.low_dimensions || []).slice(0, 2).map(d => d.dimension).join('、') || '无明显低分维度';
        const strong = (item.high_dimensions || []).slice(0, 1).map(d => d.dimension).join('、') || '待复核';
        return `
          <div class="domain-card">
            <strong>${{domainLabel(domain)}}</strong>
            <div class="small">均分：${{fmt(item.avg_score)}}；低分项：${{item.low_item_count}} / ${{item.item_count}}</div>
            <div class="small">薄弱：${{weak}}</div>
            <div class="small">相对优势：${{strong}}</div>
          </div>
        `;
      }}).join('');
      const why = profile.classroom_support_implications?.why || [];
      $('profileRationale').innerHTML = profileInfo.available
        ? `个体化依据：${{why.slice(0, 3).join(' ')}}`
        : '个体化依据：当前仅基于课堂视频和评课资料；接入五大领域后可进一步约束动作策略。';
      const rs = report.report_summary;
      $('reportSummary').innerHTML = `
        课程：<strong>${{rs.title || '未知'}}</strong><br>
        学科：${{rs.metadata?.['学段学科'] || '未知'}}；教师：${{rs.metadata?.['授课教师'] || '未知'}}<br>
        抬头率：${{rs.student_metrics?.['平均抬头率'] || 'n/a'}}；
        参与度：${{rs.student_metrics?.['平均参与度'] || 'n/a'}}；
        平均语速：${{rs.lecture_metrics?.['平均语速'] || 'n/a'}}
      `;
      $('supportPlan').innerHTML = tracking.next_support_plan.map((item) => `
        <div class="support-item">${{item}}</div>
      `).join('');
      $('growthIndicators').innerHTML = `<div class="growth-grid">${{tracking.growth_indicators.map((item) => `
        <div class="tracking-card ${{item.level}}">
          <strong>${{item.name}}</strong>
          <div class="small">${{item.current}}</div>
          <div class="small">${{item.evidence}}</div>
        </div>
      `).join('')}}</div>`;
    }}

    function renderTimeline() {{
      const moments = report.student_tracking.key_moments;
      $('timeline').innerHTML = moments.map((m, i) => {{
        const e = m.event;
        return `
        <div class="event ${{e.state}} ${{currentMode === 'moment' && i === currentIndex ? 'active' : ''}}" onclick="selectEvent(${{i}})">
          <div class="event-top">
            <span>${{mmss(e.start_sec)}}-${{mmss(e.end_sec)}}  ${{m.title}}</span>
            <span class="tag ${{e.state}}">${{e.state_zh}}</span>
          </div>
          <div class="small">${{m.why_it_matters}}</div>
          <div class="small">复核重点：${{m.teacher_review_question}}</div>
          <div class="bar"><div class="fill" style="width:${{Math.round(e.confidence * 100)}}%"></div></div>
        </div>
      `; }}).join('');
      $('windowList').innerHTML = report.timeline.map((e, i) => `
        <div class="window-row ${{e.state}} ${{currentMode === 'window' && i === currentIndex ? 'active' : ''}}" onclick="selectWindow(${{i}})">
          <div class="window-time">${{mmss(e.start_sec)}}</div>
          <div class="window-state">${{e.state_zh}}</div>
          <div class="bar"><div class="fill" style="width:${{Math.round(e.confidence * 100)}}%"></div></div>
        </div>
      `).join('');
    }}

    function renderDetail() {{
      const selected = getSelected();
      const moment = selected.moment;
      const e = selected.event;
      const action = e.recommended_action;
      const profile = supportProfile();
      const implication = profile.classroom_support_implications || {{}};
      const why = implication.why || [];
      $('detail').innerHTML = `
        <div class="action-hero">
          <h3>当前建议执行</h3>
          <div class="action-title">${{actionTitle(action.action_id)}}</div>
          <div class="small">置信度：${{Math.round(e.confidence * 100)}}%；老师端提醒：${{action.teacher_alert ? '开启' : '不开启'}}</div>
          <div class="action-grid">
            <div class="action-cell"><strong>屏幕</strong>${{action.screen}}</div>
            <div class="action-cell"><strong>语音</strong>${{action.speech}}</div>
            <div class="action-cell"><strong>机器人动作</strong>${{robotMotionTitle(action.robot_motion)}}</div>
            <div class="action-cell"><strong>安全边界</strong>${{action.teacher_alert ? '先交给老师确认，避免误打断学生。' : '只做低刺激提示，不做突然动作。'}}</div>
            <div class="action-cell"><strong>个体化依据</strong>${{why[0] || '当前仅基于课堂片段证据。'}}</div>
            <div class="action-cell"><strong>动作约束</strong>${{(implication.robot_action_constraints || ['低刺激', '老师复核']).map(constraintTitle).join(' / ')}}</div>
          </div>
        </div>
        <div class="detail-card">
          <div class="selected-head">
            <h3>${{mmss(e.start_sec)}}-${{mmss(e.end_sec)}}：${{selected.title}}</h3>
            <span class="tag ${{e.state}}">${{e.state_zh}}</span>
          </div>
          <div class="small">${{selected.description}}</div>
          <div class="small">场景：${{e.scene}}；质量：${{e.quality.video_quality}}</div>
          <div class="review-actions">
            <button>符合观察</button>
            <button>需要修正</button>
            <button>不确定</button>
          </div>
        </div>
        <div class="detail-card">
          <h3>底层证据</h3>
          <ul>${{e.evidence.map(x => `<li>${{x}}</li>`).join('')}}</ul>
        </div>
      `;
      renderJson();
    }}

    function renderJson() {{
      const selected = getSelected();
      const moment = selected.moment;
      const e = selected.event;
      const data = activeTab === 'event'
        ? (currentMode === 'moment' ? moment : e)
        : activeTab === 'action'
          ? e.recommended_action
           : activeTab === 'tracking'
             ? report.student_tracking
           : activeTab === 'teacher'
             ? report.teacher_effectiveness
           : activeTab === 'profile'
            ? profileJsonForDisplay()
           : activeTab === 'summary'
             ? report.aggregate
             : report;
      $('json').textContent = JSON.stringify(data, null, 2);
    }}

    function selectEvent(i) {{
      currentMode = 'moment';
      currentIndex = i;
      renderTimeline();
      renderDetail();
    }}

    function selectWindow(i) {{
      currentMode = 'window';
      currentIndex = i;
      renderTimeline();
      renderDetail();
    }}

    function actionTitle(actionId) {{
      const names = {{
        continue_current_task: '保持当前支持节奏',
        regain_attention_low_stimulation: '低刺激拉回注意力',
        repeat_one_step_with_visual_highlight: '一步视觉提示',
        defer_to_teacher_observation: '交由老师观察确认',
        mark_for_review: '标记为人工复核'
      }};
      return names[actionId] || actionId;
    }}

    function robotMotionTitle(motionId) {{
      const names = {{
        idle_or_soft_nod: '保持原位或轻微点头',
        slow_turn_to_screen: '慢速转向屏幕',
        slow_turn_to_screen_and_nod: '慢速转向屏幕并轻微点头',
        no_motion: '不触发机器人动作'
      }};
      return names[motionId] || motionId;
    }}

    function constraintTitle(constraintId) {{
      const names = {{
        low_stimulation_only: '只使用低刺激反馈',
        no_sudden_loud_voice: '不突然大声',
        no_fast_approach: '不快速靠近',
        keep_teacher_in_the_loop: '保持老师在环确认'
      }};
      return names[constraintId] || constraintId;
    }}

    function getSelected() {{
      if (currentMode === 'window') {{
        const event = report.timeline[currentIndex] || report.timeline[0];
        return {{
          event,
          moment: null,
          title: `${{event.state_zh}}窗口`,
          description: event.interpretation || '完整窗口用于训练数据和证据定位。'
        }};
      }}
      const moment = report.student_tracking.key_moments[currentIndex] || report.student_tracking.key_moments[0];
      return {{
        event: moment.event,
        moment,
        title: moment.title,
        description: `追踪意义：${{moment.why_it_matters}}<br>老师复核：${{moment.teacher_review_question}}`
      }};
    }}

    document.querySelectorAll('button[data-tab]').forEach(btn => {{
      btn.addEventListener('click', () => {{
        activeTab = btn.dataset.tab;
        document.querySelectorAll('button[data-tab]').forEach(b => b.classList.toggle('active', b === btn));
        renderJson();
      }});
    }});

    document.getElementById('localCaseJump')?.addEventListener('click', () => {{
      activeTab = 'profile';
      document.querySelectorAll('button[data-tab]').forEach(b => b.classList.toggle('active', b.dataset.tab === 'profile'));
      renderJson();
      const card = document.getElementById('profileCard');
      if (card) {{
        card.scrollIntoView({{behavior: 'smooth', block: 'center'}});
        card.classList.add('is-highlighted');
        window.setTimeout(() => card.classList.remove('is-highlighted'), 1600);
      }}
    }});
    renderMetrics();
    renderTimeline();
    renderDetail();
  </script>
</body>
</html>"""


def _read_video_meta(video_path: Path) -> VideoMeta:
    cv2 = _import_cv2()
    cap = cv2.VideoCapture(str(video_path))
    if not cap.isOpened():
        raise FileNotFoundError(f"Cannot open video: {video_path}")
    fps = float(cap.get(cv2.CAP_PROP_FPS) or 0)
    frame_count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    width = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH) or 0)
    height = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT) or 0)
    duration = frame_count / fps if fps > 0 else 0.0
    cap.release()
    return VideoMeta(fps=fps, frame_count=frame_count, width=width, height=height, duration_sec=duration)


def _sample_video_features(video_path: Path, meta: VideoMeta, step_sec: int = 5) -> list[dict[str, Any]]:
    cv2 = _import_cv2()
    cap = cv2.VideoCapture(str(video_path))
    previous: dict[str, np.ndarray] = {}
    raw_samples: list[dict[str, Any]] = []
    seconds = range(0, max(1, int(meta.duration_sec)), step_sec)

    for sec in seconds:
        cap.set(cv2.CAP_PROP_POS_FRAMES, int(sec * meta.fps))
        ok, frame = cap.read()
        if not ok:
            continue
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        regions = {
            "target": _crop(gray, TARGET_ROI),
            "context": _crop(gray, CONTEXT_ROI),
            "left_adult": _crop(gray, LEFT_ADULT_ROI),
            "right_adult": _crop(gray, RIGHT_ADULT_ROI),
        }
        sample: dict[str, Any] = {
            "t": sec,
            "brightness": float(np.mean(regions["target"])),
            "blur": float(cv2.Laplacian(regions["target"], cv2.CV_64F).var()),
        }
        for name, crop in regions.items():
            if name in previous and previous[name].shape == crop.shape:
                sample[f"{name}_motion_raw"] = float(np.mean(cv2.absdiff(crop, previous[name])) / 255.0)
            else:
                sample[f"{name}_motion_raw"] = 0.0
            previous[name] = crop.copy()
        raw_samples.append(sample)
    cap.release()

    _normalize_samples(raw_samples, "target_motion_raw", "target_motion")
    _normalize_samples(raw_samples, "context_motion_raw", "context_motion")
    _normalize_samples(raw_samples, "left_adult_motion_raw", "left_adult_motion")
    _normalize_samples(raw_samples, "right_adult_motion_raw", "right_adult_motion")
    return raw_samples


def _extract_audio_windows(video_path: Path, duration_sec: float, window_sec: int = 30) -> list[dict[str, Any]]:
    wav_path = OUTPUT_DIR / "lesson_audio_8k_mono.wav"
    try:
        import imageio_ffmpeg

        ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
        subprocess.run(
            [
                ffmpeg,
                "-y",
                "-i",
                str(video_path),
                "-vn",
                "-ac",
                "1",
                "-ar",
                "8000",
                "-loglevel",
                "error",
                str(wav_path),
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
    except Exception:
        return _empty_audio_windows(duration_sec, window_sec)

    try:
        with wave.open(str(wav_path), "rb") as wav:
            rate = wav.getframerate()
            frames = wav.readframes(wav.getnframes())
            audio = np.frombuffer(frames, dtype=np.int16).astype(np.float32) / 32768.0
    except Exception:
        return _empty_audio_windows(duration_sec, window_sec)

    windows: list[dict[str, Any]] = []
    values = []
    for start in range(0, max(1, int(duration_sec)), window_sec):
        seg = audio[int(start * rate) : int(min(start + window_sec, duration_sec) * rate)]
        rms = float(math.sqrt(float(np.mean(seg * seg)))) if seg.size else 0.0
        values.append(rms)
        windows.append({"start_sec": start, "end_sec": min(start + window_sec, duration_sec), "rms_raw": rms})
    scale = _safe_percentile(values, 90) or 1.0
    for item in windows:
        item["audio_activity"] = _clip01(item["rms_raw"] / scale)
    return windows


def _empty_audio_windows(duration_sec: float, window_sec: int) -> list[dict[str, Any]]:
    return [
        {"start_sec": start, "end_sec": min(start + window_sec, duration_sec), "rms_raw": 0.0, "audio_activity": 0.0}
        for start in range(0, max(1, int(duration_sec)), window_sec)
    ]


def _build_timeline(
    samples: list[dict[str, Any]],
    audio_windows: list[dict[str, Any]],
    meta: VideoMeta,
    window_sec: int = 30,
) -> list[dict[str, Any]]:
    timeline: list[dict[str, Any]] = []
    audio_by_start = {int(w["start_sec"]): w for w in audio_windows}
    for start in range(0, max(1, int(meta.duration_sec)), window_sec):
        end = min(start + window_sec, meta.duration_sec)
        group = [s for s in samples if start <= s["t"] < end]
        if not group:
            continue
        audio = audio_by_start.get(start, {"audio_activity": 0.0, "rms_raw": 0.0})
        features = _aggregate_window_features(group, audio)
        state, confidence, evidence, interpretation = _infer_state(features, start, meta.duration_sec)
        timeline.append(
            {
                "clip_id": f"seewo_student_A_{start:04d}_{int(end):04d}",
                "start_sec": start,
                "end_sec": end,
                "scene": _rough_scene(start, meta.duration_sec),
                "state": state,
                "state_zh": STATE_ZH[state],
                "confidence": confidence,
                "features": features,
                "quality": _quality(features),
                "evidence": evidence,
                "interpretation": interpretation,
                "recommended_action": ACTION_LIBRARY[state],
                "human_state_json": {
                    "timestamp_sec": start,
                    "person_id": "student_A",
                    "state": state,
                    "confidence": confidence,
                    "evidence": evidence,
                    "risk_level": "review" if state.startswith("uncertain") else "low",
                    "handoff_required": state in {"confused_or_stuck", "uncertain_adult_interference", "uncertain"},
                    "recommended_robot_action": ACTION_LIBRARY[state]["action_id"],
                },
            }
        )
    return timeline


def _aggregate_window_features(group: list[dict[str, Any]], audio: dict[str, Any]) -> dict[str, float]:
    target = float(mean(s["target_motion"] for s in group))
    context = float(mean(s["context_motion"] for s in group))
    left = float(mean(s["left_adult_motion"] for s in group))
    right = float(mean(s["right_adult_motion"] for s in group))
    adult = _clip01(max(left, right) * 0.75 + max(0.0, context - target) * 0.25)
    return {
        "target_motion": round(target, 3),
        "context_motion": round(context, 3),
        "adult_interference": round(adult, 3),
        "audio_activity": round(float(audio.get("audio_activity", 0.0)), 3),
        "brightness": round(float(mean(s["brightness"] for s in group)), 2),
        "blur": round(float(mean(s["blur"] for s in group)), 2),
    }


def _infer_state(
    features: dict[str, float],
    start_sec: int,
    duration_sec: float,
) -> tuple[str, float, list[str], str]:
    target = features["target_motion"]
    adult = features["adult_interference"]
    audio = features["audio_activity"]
    blur = features["blur"]

    evidence: list[str] = []
    if target >= 0.35:
        state = "engaged"
        evidence.append("目标学生 ROI 内有持续可见动作，说明仍在参与课堂活动。")
        confidence = 0.68 + min(0.17, target * 0.2)
        interpretation = "目标区域动作较活跃，当前更像参与状态。"
    elif adult >= 0.55 and target < 0.22:
        state = "uncertain_adult_interference"
        evidence.append("目标学生附近成人/老师区域运动较强，容易遮挡或影响 ROI 判断。")
        evidence.append("目标学生自身 ROI 运动较低，暂不触发学生端自动干预。")
        confidence = 0.52
        interpretation = "成人靠近或遮挡时，系统应优先记录并交给老师确认。"
    elif target < 0.14 and audio >= 0.35:
        state = "confused_or_stuck"
        evidence.append("课堂声音活动明显，但目标学生区域动作较少。")
        evidence.append("该模式可能对应听指令/问答阶段中的停滞或未响应。")
        confidence = 0.58
        interpretation = "可能不是单纯分心，更像需要一步提示或老师确认。"
    elif target < 0.18:
        state = "low_attention"
        evidence.append("目标学生区域长时间动作较少。")
        evidence.append("在低清全景视频下，该判断只表示参与线索减弱。")
        confidence = 0.55
        interpretation = "建议使用低刺激视觉提示，不直接公开提醒学生。"
    else:
        state = "engaged"
        evidence.append("目标区域有一定动作变化，未达到异常阈值。")
        confidence = 0.60
        interpretation = "当前维持观察。"

    if blur < 12:
        confidence = max(0.35, confidence - 0.08)
        evidence.append("画面清晰度偏低，置信度已下调。")
    if start_sec > duration_sec - 240:
        evidence.append("接近下课或转场阶段，状态解释应结合课堂流程。")
    return state, round(_clip01(confidence), 3), evidence, interpretation


def _quality(features: dict[str, float]) -> dict[str, Any]:
    blur = features["blur"]
    if blur < 8:
        quality = "low"
    elif blur < 22:
        quality = "medium_low"
    else:
        quality = "medium"
    return {
        "video_quality": quality,
        "blur_score": features["blur"],
        "brightness": features["brightness"],
        "audio_available": features["audio_activity"] > 0,
    }


def _rough_scene(start_sec: int, duration_sec: float) -> str:
    ratio = start_sec / max(duration_sec, 1)
    if ratio < 0.12:
        return "导入/教师说明"
    if ratio < 0.38:
        return "节奏练习与师生问答"
    if ratio < 0.66:
        return "学生观察/互动活动"
    if ratio < 0.90:
        return "巩固检测与课堂反馈"
    return "转场/下课整理"


def _summarize_timeline(timeline: list[dict[str, Any]]) -> dict[str, Any]:
    counts: dict[str, int] = {}
    for item in timeline:
        counts[item["state"]] = counts.get(item["state"], 0) + 1
    total = max(1, len(timeline))
    top_state = max(counts, key=counts.get) if counts else "uncertain"
    return {
        "window_count": len(timeline),
        "state_counts": counts,
        "state_ratio": {key: round(value / total, 3) for key, value in counts.items()},
        "top_state": top_state,
        "top_state_zh": STATE_ZH.get(top_state, top_state),
        "avg_confidence": round(mean(item["confidence"] for item in timeline), 3) if timeline else 0.0,
        "review_required_count": sum(1 for item in timeline if item["human_state_json"]["handoff_required"]),
        "next_training_step": "请老师抽样校正 uncertain/confused_or_stuck 片段，形成第一批弱监督标签。",
    }


def _build_student_tracking(
    timeline: list[dict[str, Any]],
    summary: dict[str, Any],
    report: dict[str, Any],
    support_profile: dict[str, Any] | None = None,
) -> dict[str, Any]:
    support_profile = support_profile or {}
    engaged_ratio = float(summary.get("state_ratio", {}).get("engaged", 0.0))
    review_count = int(summary.get("review_required_count", 0))
    low_attention_count = int(summary.get("state_counts", {}).get("low_attention", 0))
    stuck_count = int(summary.get("state_counts", {}).get("confused_or_stuck", 0))
    adult_count = int(summary.get("state_counts", {}).get("uncertain_adult_interference", 0))
    total = max(1, int(summary.get("window_count", len(timeline))))

    if engaged_ratio >= 0.7:
        main_phrase = "多数时间能维持课堂参与线索"
    elif engaged_ratio >= 0.45:
        main_phrase = "参与状态存在波动，需要结合老师观察确认"
    else:
        main_phrase = "参与线索偏弱，需要优先建立个人支持基线"

    summary_text = (
        f"本节课作为目标学生 A 的初始基线样本：{main_phrase}。"
        f"系统从 {total} 个片段中筛出 {review_count} 个建议老师复核的支持时刻，"
        "用于判断学生是否真的需要提示、降阶或老师靠近确认。"
    )
    if support_profile:
        diagnosis = support_profile.get("diagnosis_context", {})
        category = diagnosis.get("iep_primary_disability_category")
        degree = diagnosis.get("iep_primary_disability_degree")
        summary_text += (
            f" 已接入匿名化 IEP 和五大领域评估摘要：教育档案标注主要障碍类别为{category or '未明'}"
            f"{'，程度' + degree if degree else ''}；系统仅将其用于支持策略约束，不作为医学诊断。"
        )

    key_moments = _select_key_moments(timeline)
    profile_status = (
        "已接入课堂视频、评课资料、IEP 摘要以及认知/语言沟通/运动三类评估；感知领域和社会适应字段已预留。"
        if support_profile
        else "已接入课堂视频和评课资料；测评档案、专业行为指导可在下一步补充。"
    )
    support_constraint = (
        "结合 IEP 和五大领域评估，当前优先采用视觉一步提示、短句语音、低刺激动作和老师复核。"
        if support_profile
        else "当前不做强刺激自动干预；所有提示先按低刺激、可关闭、老师可复核原则处理。"
    )
    completeness_current = (
        "已接入 IEP、认知、语言沟通、运动评估；感知领域和社会适应在本样例中暂缺。"
        if support_profile
        else "已接入课堂视频、声音和评课报告；尚缺学生测评档案与既有行为指导。"
    )
    completeness_evidence = (
        "系统已经能把能力画像转成动作约束：一步提示、视觉辅助、短句语音和低刺激机器人动作。"
        if support_profile
        else "补充这些资料后，系统才能给出更个体化、更安全的支持动作。"
    )
    return {
        "student_profile": {
            "student_id": "student_A",
            "baseline_status": "初始基线建立中",
            "profile_status": profile_status,
            "support_constraint": support_constraint,
        },
        "lesson_conclusion": {
            "summary": summary_text,
            "not_a_diagnosis": "这是课堂支持分析，不是医学诊断或能力定级。",
            "data_status": "一节课只能建立初始基线，不能判断长期进步或恶化；需要后续多节课连续追踪。",
        },
        "growth_indicators": [
            {
                "name": "参与稳定性",
                "level": "good" if engaged_ratio >= 0.7 else "watch",
                "current": f"本节课约 {round(engaged_ratio * 100)}% 片段保持参与线索。",
                "evidence": "用于建立个人基线，后续比较同一学生在不同课程、不同任务中的变化。",
            },
            {
                "name": "需要支持的时刻",
                "level": "support" if low_attention_count + stuck_count else "good",
                "current": f"注意力下降 {low_attention_count} 段，疑似困惑/停滞 {stuck_count} 段。",
                "evidence": "这些片段不是为了贴标签，而是为了让老师判断是否需要一步提示、视觉提示或任务降阶。",
            },
            {
                "name": "成人靠近/遮挡影响",
                "level": "watch" if adult_count else "good",
                "current": f"成人干扰下不确定 {adult_count} 段。",
                "evidence": "全景视频中老师或维护人员靠近会影响自动判断，因此系统把这类片段交给人工复核。",
            },
            {
                "name": "资料完整度",
                "level": "support" if support_profile else "watch",
                "current": completeness_current,
                "evidence": completeness_evidence,
            },
        ],
        "key_moments": key_moments,
        "next_support_plan": [
            "请老师优先复核关键片段，判断这些时刻是真分心、没理解、被遮挡，还是正常课堂波动。",
            (
                "继续补充感知领域和社会适应评估，并记录有效提示方式、禁忌刺激和任务粒度。"
                if support_profile
                else "补充学生测评档案和既有行为指导，用于设置禁忌刺激、有效提示方式和任务粒度。"
            ),
            "下一节课继续追踪同一学生，比较参与稳定性、需要支持片段数量和复核结果是否改善。",
            "动作输出先以屏幕弱高亮、一步提示、老师端提醒为主，暂不做强语音或突然机器人动作。",
        ],
        "support_profile_summary": _summarize_support_profile_for_tracking(support_profile),
        "source_context": {
            "lesson_title": report.get("title", ""),
            "class_participation": report.get("student_metrics", {}),
            "lecture_metrics": report.get("lecture_metrics", {}),
        },
    }


def _select_key_moments(timeline: list[dict[str, Any]], limit: int = 6) -> list[dict[str, Any]]:
    preferred_states = ["confused_or_stuck", "low_attention", "uncertain_adult_interference"]
    selected: list[dict[str, Any]] = []
    seen_windows: set[str] = set()

    for state in preferred_states:
        candidates = [event for event in timeline if event["state"] == state]
        candidates = sorted(candidates, key=lambda event: event["confidence"], reverse=True)
        for event in candidates[:2]:
            if event["clip_id"] not in seen_windows:
                selected.append(_moment_from_event(event))
                seen_windows.add(event["clip_id"])
            if len(selected) >= limit:
                return sorted(selected, key=lambda item: item["event"]["start_sec"])

    if not selected and timeline:
        stable = max(timeline, key=lambda event: event["confidence"])
        selected.append(_moment_from_event(stable, title_override="稳定参与样例"))

    return sorted(selected, key=lambda item: item["event"]["start_sec"])


def _summarize_support_profile_for_tracking(support_profile: dict[str, Any]) -> dict[str, Any]:
    if not support_profile:
        return {"available": False}
    diagnosis = support_profile.get("diagnosis_context", {})
    domains = support_profile.get("five_domain_assessment", {})
    available_domains = [key for key, value in domains.items() if value.get("available")]
    missing_domains = [key for key, value in domains.items() if not value.get("available")]
    implications = support_profile.get("classroom_support_implications", {})
    return {
        "available": True,
        "education_profile": {
            "primary_disability_category": diagnosis.get("iep_primary_disability_category", ""),
            "primary_disability_degree": diagnosis.get("iep_primary_disability_degree", ""),
            "boundary": "仅作为教育支持背景，不作为医学诊断输出。",
        },
        "available_domains": available_domains,
        "missing_domains": missing_domains,
        "support_flags": {
            "prefer_visual_step_prompt": bool(implications.get("prefer_visual_step_prompt")),
            "prefer_short_single_step_speech": bool(implications.get("prefer_short_single_step_speech")),
            "avoid_long_multi_instruction_audio": bool(implications.get("avoid_long_multi_instruction_audio")),
            "teacher_review_required_for_uncertain_states": bool(
                implications.get("teacher_review_required_for_uncertain_states")
            ),
        },
        "why": implications.get("why", []),
    }


def _build_teacher_effectiveness(
    timeline: list[dict[str, Any]],
    summary: dict[str, Any],
    report: dict[str, Any],
    support_profile: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Estimate how well this lesson supported the target student.

    This is deliberately framed as lesson support effectiveness, not a personal
    teacher ranking. Every score is tied to observable student response or Seewo
    lesson context and should be reviewed by a human teacher.
    """
    support_profile = support_profile or {}
    total = max(1, int(summary.get("window_count", len(timeline))))
    state_counts = summary.get("state_counts", {})
    engaged_count = int(state_counts.get("engaged", 0))
    low_attention_count = int(state_counts.get("low_attention", 0))
    stuck_count = int(state_counts.get("confused_or_stuck", 0))
    adult_count = int(state_counts.get("uncertain_adult_interference", 0))
    engaged_ratio = float(summary.get("state_ratio", {}).get("engaged", 0.0))
    challenge_count = low_attention_count + stuck_count
    challenge_ratio = challenge_count / total
    adult_ratio = adult_count / total

    recovered_count = _count_recovered_challenges(timeline)
    recovery_rate = recovered_count / challenge_count if challenge_count else 1.0

    lecture_metrics = report.get("lecture_metrics", {})
    question_metrics = report.get("question_metrics", {})
    lesson_minutes = _parse_duration_minutes(lecture_metrics.get("课程时长", ""))
    lecture_minutes = _parse_duration_minutes(lecture_metrics.get("讲授时长", ""))
    lecture_ratio = lecture_minutes / lesson_minutes if lesson_minutes else 0.0
    speech_speed = _parse_number(lecture_metrics.get("平均语速", ""))
    wait_over_5 = _parse_percent(question_metrics.get("等候5秒以上", ""))
    high_order_question_ratio = sum(
        _parse_percent(question_metrics.get(key, ""))
        for key in ("理解型", "应用型", "分析型", "创造型")
    )

    engagement_score = round(engaged_ratio * 100)
    recovery_score = round(recovery_rate * 100) if challenge_count else 90
    rhythm_score = round(
        0.35 * _speech_speed_score(speech_speed)
        + 0.35 * (wait_over_5 * 100 if wait_over_5 else 72)
        + 0.20 * _lecture_balance_score(lecture_ratio)
        + 0.10 * min(100, high_order_question_ratio * 100)
    )
    low_disruption_score = round(max(45, 100 - adult_ratio * 160))
    personalization_score = _personalization_score(
        support_profile=support_profile,
        engaged_ratio=engaged_ratio,
        challenge_ratio=challenge_ratio,
        recovery_rate=recovery_rate,
    )

    metrics = [
        _teacher_metric(
            key="student_engagement",
            name="参与维持度",
            score=engagement_score,
            weight=35,
            benchmark="70 分以上表示该学生多数时间有可见参与线索。",
            evidence=f"{engaged_count}/{total} 个 30 秒窗口处于参与状态，参与比例约 {round(engaged_ratio * 100)}%。",
            interpretation=(
                "学生在本节课中能较长时间保持课堂参与。"
                if engagement_score >= 70
                else "学生参与稳定性偏弱，需要结合课堂任务和老师观察复盘。"
            ),
        ),
        _teacher_metric(
            key="recovery_after_challenge",
            name="困惑/分心恢复度",
            score=recovery_score,
            weight=22,
            benchmark="出现困惑或注意力下降后，后续 60 秒内回到参与的比例越高越好。",
            evidence=f"共 {challenge_count} 个支持需求窗口，其中 {recovered_count} 个在后续两个窗口内回到参与。",
            interpretation=(
                "学生遇到波动后多数能被课堂节奏带回。"
                if recovery_score >= 70
                else "学生出现波动后的恢复不够稳定，建议复盘这些片段中的提示方式和任务难度。"
            ),
        ),
        _teacher_metric(
            key="pace_and_scaffold",
            name="节奏与提问支架",
            score=rhythm_score,
            weight=18,
            benchmark="参考语速、讲授占比、候答时间和高阶问题比例，判断课堂节奏是否给学生留出理解空间。",
            evidence=(
                f"平均语速 {lecture_metrics.get('平均语速', '未知')}，讲授时长占比约 {round(lecture_ratio * 100)}%，"
                f"候答 5 秒以上 {question_metrics.get('等候5秒以上', '未知')}。"
            ),
            interpretation=(
                "课堂节奏整体较友好，给学生留出了反应和理解时间。"
                if rhythm_score >= 75
                else "课堂节奏或提问支架需要复核，尤其要关注目标学生是否来得及理解与回应。"
            ),
        ),
        _teacher_metric(
            key="low_disruption_support",
            name="低干扰支持环境",
            score=low_disruption_score,
            weight=15,
            benchmark="成人靠近或遮挡片段越少，系统越容易稳定判断；此项只提示复核，不等同扣分。",
            evidence=f"成人靠近/遮挡导致不确定 {adult_count}/{total} 个窗口，比例约 {round(adult_ratio * 100)}%。",
            interpretation=(
                "课堂支持环境较稳定，自动分析受成人遮挡影响较小。"
                if low_disruption_score >= 75
                else "成人靠近片段较多，需要区分这是有效支持、秩序维护还是普通遮挡。"
            ),
        ),
        _teacher_metric(
            key="individualized_fit",
            name="个体化适配度",
            score=personalization_score,
            weight=10,
            benchmark="结合学生 IEP/五大领域资料，观察课堂提示是否适合该学生的认知、语言和动作特点。",
            evidence=(
                "已接入匿名化 IEP、认知、语言沟通和运动能力摘要，并用于解释提示方式。"
                if support_profile
                else "当前尚未接入完整个体档案，只能做课堂表现层面的初步判断。"
            ),
            interpretation=(
                "当前课堂表现与个体资料较匹配，适合继续验证视觉一步提示和短句语音。"
                if personalization_score >= 75
                else "需要补充个体资料或复核提示方式，避免把听不懂误判为不配合。"
            ),
        ),
    ]

    overall = round(sum(item["score"] * item["weight"] for item in metrics) / sum(item["weight"] for item in metrics))
    band = _teacher_score_band(overall)
    return {
        "overall_score": overall,
        "band": band,
        "headline": f"对目标学生的课堂支持效果：{band['label']}",
        "summary": _teacher_effect_summary(overall, engaged_ratio, recovery_rate, adult_ratio),
        "metrics": metrics,
        "rubric": {
            "excellent": "85-100：支持效果稳定，学生能持续参与且波动后恢复较快。",
            "good": "70-84：支持效果较好，但仍有若干片段需要复核和优化。",
            "watch": "55-69：有改进空间，需要重点看困惑、低参与或成人遮挡片段。",
            "review": "0-54：需重点复盘，不能直接下结论，应由老师结合课堂情境确认。",
        },
        "strengths": _teacher_strengths(metrics),
        "improvement_focus": _teacher_improvement_focus(metrics),
        "review_questions": [
            "学生注意力下降时，老师当时是否已经提供视觉提示、动作示范或任务降阶？",
            "成人靠近目标学生的片段，是有效支持、秩序维护，还是影响了学生独立参与？",
            "学生出现困惑后，后续 60 秒是否真的回到任务，还是只是画面上暂时恢复？",
        ],
        "boundary": "该评分只评估本节课对目标学生的支持效果，不评价老师个人能力，不替代听课专家或校方判断。",
    }


def _teacher_metric(
    *,
    key: str,
    name: str,
    score: int,
    weight: int,
    benchmark: str,
    evidence: str,
    interpretation: str,
) -> dict[str, Any]:
    score = int(max(0, min(100, score)))
    return {
        "key": key,
        "name": name,
        "score": score,
        "weight": weight,
        "level": _score_level(score),
        "benchmark": benchmark,
        "evidence": evidence,
        "interpretation": interpretation,
    }


def _count_recovered_challenges(timeline: list[dict[str, Any]]) -> int:
    recovered = 0
    challenge_states = {"low_attention", "confused_or_stuck"}
    for idx, event in enumerate(timeline):
        if event.get("state") not in challenge_states:
            continue
        next_events = timeline[idx + 1 : idx + 3]
        if any(item.get("state") == "engaged" for item in next_events):
            recovered += 1
    return recovered


def _speech_speed_score(speed: float) -> float:
    if speed <= 0:
        return 72.0
    if 3.0 <= speed <= 4.0:
        return 88.0
    if 2.5 <= speed < 3.0:
        return 82.0
    if 4.0 < speed <= 4.5:
        return 72.0
    if 2.0 <= speed < 2.5:
        return 68.0
    return 58.0


def _lecture_balance_score(lecture_ratio: float) -> float:
    if lecture_ratio <= 0:
        return 72.0
    return max(55.0, min(95.0, 100.0 - abs(lecture_ratio - 0.40) * 120.0))


def _personalization_score(
    *,
    support_profile: dict[str, Any],
    engaged_ratio: float,
    challenge_ratio: float,
    recovery_rate: float,
) -> int:
    score = 62
    if support_profile:
        score += 13
    if engaged_ratio >= 0.70:
        score += 8
    elif engaged_ratio >= 0.50:
        score += 4
    if challenge_ratio <= 0.15:
        score += 5
    if recovery_rate < 0.50:
        score -= 8
    return int(max(45, min(95, score)))


def _teacher_score_band(score: int) -> dict[str, str]:
    if score >= 85:
        return {"key": "excellent", "label": "支持效果很稳", "tone": "good"}
    if score >= 70:
        return {"key": "good", "label": "支持效果较好", "tone": "good"}
    if score >= 55:
        return {"key": "watch", "label": "有改进空间", "tone": "watch"}
    return {"key": "review", "label": "需重点复盘", "tone": "support"}


def _score_level(score: int) -> str:
    if score >= 75:
        return "good"
    if score >= 60:
        return "watch"
    return "support"


def _teacher_effect_summary(overall: int, engaged_ratio: float, recovery_rate: float, adult_ratio: float) -> str:
    return (
        f"本节课目标学生参与窗口约 {round(engaged_ratio * 100)}%，综合支持指数为 {overall}/100。"
        f"挑战后 60 秒内恢复比例约 {round(recovery_rate * 100)}%，成人靠近/遮挡比例约 {round(adult_ratio * 100)}%。"
        " 该结果适合用于课后复盘和下一节课策略调整，不适合直接作为教师考核结论。"
    )


def _teacher_strengths(metrics: list[dict[str, Any]]) -> list[str]:
    strengths = [item for item in metrics if item["score"] >= 75]
    return [
        f"{item['name']}较好：{item['interpretation']}"
        for item in strengths[:3]
    ] or ["本节课已有可复盘的数据基础，可继续结合老师标注确认有效支持方式。"]


def _teacher_improvement_focus(metrics: list[dict[str, Any]]) -> list[str]:
    focus = sorted(metrics, key=lambda item: item["score"])[:2]
    return [
        f"优先复盘{item['name']}：{item['interpretation']}"
        for item in focus
    ]


def _parse_percent(value: Any) -> float:
    number = _parse_number(value)
    if number <= 0:
        return 0.0
    return _clip01(number / 100.0)


def _parse_duration_minutes(value: Any) -> float:
    text = str(value or "")
    minute_match = re.search(r"(\d+(?:\.\d+)?)\s*分钟", text)
    compact_minute_match = re.search(r"(\d+(?:\.\d+)?)\s*分", text)
    second_match = re.search(r"(\d+(?:\.\d+)?)\s*秒", text)
    minutes = 0.0
    if minute_match:
        minutes += float(minute_match.group(1))
    elif compact_minute_match:
        minutes += float(compact_minute_match.group(1))
    if second_match:
        minutes += float(second_match.group(1)) / 60.0
    if minutes:
        return minutes
    return _parse_number(text)


def _parse_number(value: Any) -> float:
    match = re.search(r"\d+(?:\.\d+)?", str(value or ""))
    return float(match.group(0)) if match else 0.0


def _moment_from_event(event: dict[str, Any], title_override: str | None = None) -> dict[str, Any]:
    state = event["state"]
    if state == "confused_or_stuck":
        title = "疑似需要一步提示"
        why = "课堂声音活跃但目标学生动作较少，适合作为“是否没理解指令”的复核样本。"
        question = "这个片段学生是在思考、没听懂，还是被其他人遮挡？"
    elif state == "low_attention":
        title = "参与线索减弱"
        why = "目标学生区域动作较少，适合让老师判断是否需要轻量视觉提示。"
        question = "当时学生是否仍在看/听/等待，还是已经脱离任务？"
    elif state == "uncertain_adult_interference":
        title = "成人靠近导致不确定"
        why = "老师或维护人员靠近目标学生，系统不应直接判断学生状态。"
        question = "成人靠近是在帮助学生、维持秩序，还是普通走动？"
    else:
        title = title_override or "稳定参与样例"
        why = "该片段可作为学生个人正常课堂参与基线。"
        question = "这个状态是否能代表该学生较自然的课堂参与状态？"
    return {
        "title": title_override or title,
        "why_it_matters": why,
        "teacher_review_question": question,
        "event": event,
    }


def _write_visual_assets(video_path: Path, meta: VideoMeta, timeline: list[dict[str, Any]]) -> dict[str, str]:
    cv2 = _import_cv2()
    frames_dir = OUTPUT_DIR / "frames"
    frames_dir.mkdir(parents=True, exist_ok=True)
    cap = cv2.VideoCapture(str(video_path))
    roi_path = frames_dir / "target_roi_reference.jpg"
    contact_path = frames_dir / "timeline_contact_sheet.jpg"
    sample_times = [0, 180, 540, 900, 1260, max(0, int(meta.duration_sec) - 60)]
    thumbs = []
    for idx, sec in enumerate(sample_times):
        cap.set(cv2.CAP_PROP_POS_FRAMES, int(sec * meta.fps))
        ok, frame = cap.read()
        if not ok:
            continue
        img = _annotated_frame(frame, sec)
        if idx == 2:
            img.save(roi_path, quality=92)
        thumb = img.resize((300, int(img.height * 300 / img.width)))
        thumbs.append(thumb)
    cap.release()
    if thumbs:
        from PIL import Image

        pad = 12
        cols = 2
        rows = math.ceil(len(thumbs) / cols)
        sheet = Image.new("RGB", (cols * 300 + (cols + 1) * pad, rows * thumbs[0].height + (rows + 1) * pad), (238, 242, 246))
        for i, thumb in enumerate(thumbs):
            sheet.paste(thumb, (pad + (i % cols) * (300 + pad), pad + (i // cols) * (thumb.height + pad)))
        sheet.save(contact_path, quality=92)
    if not roi_path.exists() and thumbs:
        thumbs[0].save(roi_path, quality=92)
    return {
        "roi_reference_url": f"/assets/seewo_xiaoyudi_demo/frames/{roi_path.name}",
        "contact_sheet_url": f"/assets/seewo_xiaoyudi_demo/frames/{contact_path.name}",
        "analysis_json_url": "/api/seewo_student_demo",
    }


def _annotated_frame(frame: np.ndarray, sec: int):
    cv2 = _import_cv2()
    from PIL import Image, ImageDraw

    img = Image.fromarray(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
    draw = ImageDraw.Draw(img)
    _draw_roi(draw, CONTEXT_ROI, "student table context", (65, 115, 245))
    _draw_roi(draw, TARGET_ROI, "target student", (255, 40, 40), width=4)
    _draw_roi(draw, LEFT_ADULT_ROI, "teacher/assistant", (255, 180, 0))
    _draw_roi(draw, RIGHT_ADULT_ROI, "teacher/assistant", (255, 180, 0))
    draw.rectangle((8, 508, 118, 536), fill=(255, 255, 255), outline=(20, 20, 20))
    draw.text((14, 514), _format_time(sec), fill=(0, 0, 0))
    return img


def _draw_roi(draw, roi: dict[str, int], label: str, color: tuple[int, int, int], width: int = 3) -> None:
    x, y, w, h = roi["x"], roi["y"], roi["w"], roi["h"]
    draw.rectangle((x, y, x + w, y + h), outline=color, width=width)
    draw.text((x, max(0, y - 20)), label, fill=color)


def _parse_report(report_path: Path) -> dict[str, Any]:
    try:
        from docx import Document

        document = Document(str(report_path))
    except Exception:
        return {"available": False, "title": "", "metadata": {}, "student_metrics": {}, "lecture_metrics": {}}

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    title = paragraphs[0] if paragraphs else ""
    metadata: dict[str, str] = {}
    for text in paragraphs:
        if "授课教师" in text and "授课时间" in text:
            for line in text.splitlines():
                if "：" in line:
                    key, value = line.split("：", 1)
                    metadata[key.strip()] = value.strip()

    student_metrics: dict[str, str] = {}
    lecture_metrics: dict[str, str] = {}
    behavior_distribution: dict[str, str] = {}
    question_metrics: dict[str, str] = {}
    for table in document.tables:
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        headers = rows[0]
        if headers[:3] == ["类目", "数值", "释义"]:
            for row in rows[1:]:
                if len(row) >= 2:
                    student_metrics[row[0]] = row[1]
        elif len(headers) >= 2 and headers[0] == "课堂行为":
            for row in rows[1:]:
                if len(row) >= 2:
                    behavior_distribution[row[0]] = row[1]
        elif len(rows) >= 1 and len(rows[0]) == 2:
            for row in rows:
                if len(row) >= 2 and row[0] in {"课程时长", "讲授时长", "讲授字数", "平均语速"}:
                    lecture_metrics[row[0]] = row[1]
        elif headers[:3] in (["问答模式", "定义", "次数"], ["提问类型", "次数", "占比"], ["候答时间", "次数", "占比"]):
            for row in rows[1:]:
                if len(row) >= 3:
                    question_metrics[row[0]] = row[-1]
    return {
        "available": True,
        "title": title.replace("《", "").replace("》", ""),
        "metadata": metadata,
        "student_metrics": student_metrics,
        "lecture_metrics": lecture_metrics,
        "behavior_distribution": behavior_distribution,
        "question_metrics": question_metrics,
        "source": str(report_path),
    }


def _normalize_samples(samples: list[dict[str, Any]], raw_key: str, out_key: str) -> None:
    values = [float(s.get(raw_key, 0.0)) for s in samples]
    scale = _safe_percentile(values, 90) or 1.0
    for sample in samples:
        sample[out_key] = _clip01(float(sample.get(raw_key, 0.0)) / scale)


def _crop(gray: np.ndarray, roi: dict[str, int]) -> np.ndarray:
    x, y, w, h = roi["x"], roi["y"], roi["w"], roi["h"]
    return gray[y : y + h, x : x + w]


def _safe_percentile(values: list[float], percentile: float) -> float:
    if not values:
        return 0.0
    return float(np.percentile(np.asarray(values, dtype=np.float32), percentile))


def _clip01(value: float) -> float:
    return max(0.0, min(1.0, float(value)))


def _format_time(sec: int | float) -> str:
    sec = int(sec)
    return f"{sec // 60:02d}:{sec % 60:02d}"


def _import_cv2():
    try:
        import cv2
    except ImportError as exc:  # pragma: no cover - optional runtime path
        raise RuntimeError("OpenCV is required for Seewo video analysis. Install opencv-python-headless.") from exc
    return cv2
