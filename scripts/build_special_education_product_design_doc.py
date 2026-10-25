from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "docs" / "special_education_product_design_zh.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
DARK = RGBColor(30, 38, 45)
MUTED = RGBColor(95, 105, 115)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9DEE6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Pt(width / 20)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_run_font(run, size=None, bold=None, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc: Document, text: str = "", style: str | None = None, *, bold=False) -> None:
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_run_font(run, bold=bold)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, bold=True, color=DARK)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            if i == 0 and len(value) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color in [
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("辅读学校具身智能课堂辅助与生活技能训练系统设计方案")
    set_run_font(r, size=9, color=MUTED)
    return doc


def build_doc() -> None:
    doc = setup_document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    r = title.add_run("辅读学校具身智能课堂辅助与生活技能训练系统设计方案")
    set_run_font(r, size=22, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sr = subtitle.add_run("产品设计版 | 面向课堂辅助、低刺激互动与数字支付生活技能训练")
    set_run_font(sr, size=11, color=MUTED)

    add_callout(
        doc,
        "核心定位",
        "开发一套面向辅读学校的具身智能课堂辅助系统，通过多模态感知、教师端提示、低刺激机器人/屏幕互动和数字支付情境训练，帮助教师更早理解学生状态，并支持学生逐步提升面向真实社会的生活自理能力。",
    )

    add_heading(doc, "1. 项目背景与问题", 1)
    add_para(
        doc,
        "辅读学校学生覆盖一年级到九年级，年龄跨度大、能力差异明显。同一课堂中，学生对指令、课程内容、生活概念的理解程度可能完全不同。教师在课堂中不仅要讲授知识，还要持续观察学生是否注意力下降、是否跟得上、是否情绪波动、是否出现怪叫或乱动，以及是否需要立即介入。",
    )
    add_para(
        doc,
        "本次听课内容为三年级“纸币”教学。课堂暴露出一个现实问题：纸币认知仍有教学价值，但现实生活中学生更需要理解纸币、数字金额、扫码支付、虚拟余额和真实购买行为之间的关系。因此，系统设计不应只停留在“知识教学”，而应转向“课堂状态辅助观察 + 生活化任务训练 + 安全可控互动”。",
    )

    add_heading(doc, "2. 要开发什么", 1)
    add_callout(
        doc,
        "拟开发产品",
        "辅读学校具身智能课堂辅助与数字支付生活技能训练系统。系统由课堂状态感知、教师端辅助提示、低刺激互动设备、数字支付训练任务和课后数据报告组成。",
    )
    add_table(
        doc,
        ["组成部分", "开发内容", "作用"],
        [
            ["课堂状态感知", "通过摄像头、麦克风和任务行为记录，辅助识别注意力下降、困惑、抗拒、异常发声、需要帮助等课堂状态。", "帮助教师更早发现学生状态变化。"],
            ["教师端提示", "在平板或电脑端显示学生状态、风险等级、建议关注对象和课堂复盘摘要。", "不替代教师判断，只提供观察参考。"],
            ["低刺激互动", "通过机器人、屏幕或桌面设备进行轻声提示、视觉引导、任务暂停、简化说明等互动。", "帮助学生回到学习任务，同时避免过度刺激。"],
            ["数字支付训练", "构建虚拟小超市、纸币/数字金额匹配、扫码支付、余额变化、错误支付提醒等任务。", "训练学生面向真实生活的支付能力。"],
            ["课后报告", "记录任务完成率、支付错误类型、注意力变化、提示次数和教师介入建议。", "支持教师复盘和个体化教学。"],
        ],
        [1700, 4700, 2960],
    )

    add_heading(doc, "3. 目标用户与使用场景", 1)
    add_para(doc, "目标用户包括辅读学校教师、特教学生、学校管理者和后续研发/运营人员。第一阶段不面向家庭独立销售，优先作为学校课堂辅助与公益课试点系统。")
    add_table(
        doc,
        ["对象", "需求", "系统回应"],
        [
            ["教师", "需要同时管理课堂、观察学生状态、调整教学节奏。", "提供实时状态提示、低刺激互动建议和课后报告。"],
            ["学生", "理解力差异大，注意力时间短，需要生活化、可操作的训练任务。", "通过虚拟超市和支付任务进行分层训练。"],
            ["学校", "希望提升生活适应课程质量，同时保证安全和隐私。", "提供可控设备、授权流程、数据最小化和试点评估。"],
        ],
        [1500, 3850, 4010],
    )

    add_heading(doc, "4. 核心功能设计", 1)
    add_heading(doc, "4.1 课堂状态感知", 2)
    add_para(doc, "系统识别的是课堂状态，不是医学诊断。输出结果只作为教师观察参考。")
    add_table(
        doc,
        ["状态", "典型表现", "建议处理"],
        [
            ["专注参与", "看向老师/屏幕/材料，能跟随任务。", "继续任务并给予正向反馈。"],
            ["注意力下降", "视线游离、长时间无反应、动作慢。", "轻提示或短任务拉回注意力。"],
            ["困惑", "说“不会”“不懂”，连续出错或等待过久。", "慢速重复说明，简化下一步。"],
            ["精力过旺", "频繁乱动、坐不住、动作幅度大。", "降低刺激，给短动作任务，提示教师关注。"],
            ["异常发声", "突然怪叫、音量异常、非任务相关发声。", "暂停语音刺激，提示教师判断。"],
            ["抗拒继续", "说“不想做”“不要”，回避任务。", "暂停任务，确认意愿，不强迫互动。"],
            ["需要帮助", "主动求助、长时间卡住或错误反复。", "提供一步提示，必要时提醒教师介入。"],
        ],
        [1500, 4400, 3460],
    )

    add_heading(doc, "4.2 数字支付生活技能训练", 2)
    add_para(doc, "数字支付训练模块以“虚拟小超市”为核心，把纸币认知过渡到真实生活中的扫码支付、余额判断和金额确认。")
    add_number(doc, "学生选择商品，例如水、面包、文具。")
    add_number(doc, "系统显示商品价格和当前余额。")
    add_number(doc, "学生选择纸币、输入金额或模拟扫码付款。")
    add_number(doc, "系统判断是否少付、多付、余额不足或误扫金额。")
    add_number(doc, "机器人或屏幕给出低刺激反馈，教师端记录结果。")
    add_table(
        doc,
        ["训练任务", "系统判断", "反馈方式"],
        [
            ["商品 2 元，学生支付 2 元", "金额正确", "鼓励继续，进入下一题。"],
            ["商品 5 元，余额 3 元", "余额不足", "提示“还差 2 元”，引导重新选择。"],
            ["商品 6 元，学生支付 10 元", "需要找零", "提示“应找回 4 元”，训练找零理解。"],
            ["商品 5 元，扫码金额 8 元", "误付风险", "停止支付并提示确认金额。"],
        ],
        [2600, 2500, 4260],
    )

    add_heading(doc, "4.3 低刺激互动策略", 2)
    add_para(doc, "互动策略必须可控、可关闭、不过度刺激。机器人或屏幕只做辅助提示，最终判断由教师完成。")
    add_table(
        doc,
        ["等级", "触发情况", "动作"],
        [
            ["0", "状态正常或置信度低", "只记录，不打扰。"],
            ["1", "轻度注意力下降", "屏幕显示简短视觉提示。"],
            ["2", "视线短暂脱离", "机器人转向学生或屏幕，不靠近学生。"],
            ["3", "困惑或任务卡住", "慢速重复说明，简化任务步骤。"],
            ["4", "抗拒、不适、明显异常行为", "暂停互动，提醒教师介入。"],
        ],
        [900, 3600, 4860],
    )

    add_heading(doc, "5. 需要什么硬件", 1)
    add_table(
        doc,
        ["硬件", "第一版配置", "用途", "必要性"],
        [
            ["笔记本/小主机", "Windows 或 Linux 均可，建议 i5/16GB 起步。", "运行课堂任务、状态感知、教师端页面和数据记录。", "必需"],
            ["摄像头", "普通 USB 摄像头，后续可升级 RGB-D。", "观察头部朝向、姿态、参与度和课堂行为。", "必需"],
            ["麦克风", "USB 麦克风或阵列麦克风。", "采集语音关键词、音量变化和异常发声线索。", "建议"],
            ["平板/大屏", "触控平板、互动屏或投屏设备。", "承载虚拟小超市和支付训练界面。", "必需"],
            ["机器人/桌面设备", "可移动或桌面机器人，支持朝向、灯光、语音。", "执行低刺激提示、朝向互动和任务反馈。", "可选但推荐"],
            ["教师端设备", "教师平板或电脑。", "查看实时提示、手动控制互动强度、查看课后报告。", "必需"],
            ["本地网络", "校园内网或路由器。", "连接教师端、互动屏和机器人。", "必需"],
        ],
        [1500, 2600, 3600, 1660],
    )

    add_heading(doc, "6. 最终载体是什么", 1)
    add_para(doc, "最终载体不是单一机器人，而是一套“软件平台 + 互动终端 + 可选机器人 + 课程内容包”的组合。")
    add_table(
        doc,
        ["载体", "形态", "说明"],
        [
            ["教师端", "网页或平板应用", "显示学生状态、任务进度、提示建议、课堂报告。"],
            ["学生端", "平板/大屏互动界面", "完成虚拟超市、扫码支付、金额判断等生活技能任务。"],
            ["机器人端", "桌面机器人或移动机器人", "执行低刺激提示、朝向互动、语音反馈；不主动接触学生。"],
            ["后台数据", "本地数据文件或学校内网服务", "记录任务日志、状态变化和教师标注。"],
            ["课程包", "数字支付训练课件与任务脚本", "支持教师直接用于生活数学或生活适应课程。"],
        ],
        [1500, 2500, 5360],
    )

    add_heading(doc, "7. 第一版 MVP 开发范围", 1)
    add_callout(
        doc,
        "MVP 建议",
        "第一版不要追求全校全场景覆盖。建议聚焦一个班、一节生活数学课、一个数字支付任务系统，先验证学生是否愿意参与、教师是否觉得有帮助、互动是否安全可控。",
    )
    add_table(
        doc,
        ["模块", "第一版要做", "暂不做"],
        [
            ["课堂状态", "注意力下降、困惑、抗拒、需要帮助、异常发声的辅助识别。", "不做医学诊断，不做障碍类型判断。"],
            ["支付训练", "虚拟小超市、商品价格、余额、扫码支付、少付/多付提醒。", "不接真实支付平台，不产生真实交易。"],
            ["互动方式", "屏幕提示、轻声语音、机器人转向、暂停任务、教师提醒。", "不主动触碰学生，不快速靠近。"],
            ["教师端", "实时状态提示、任务完成情况、课后摘要。", "不做复杂教务系统。"],
            ["数据", "优先采任务日志和教师标注。", "未授权前不采集可识别原始视频音频。"],
        ],
        [1600, 4100, 3660],
    )

    add_heading(doc, "8. 安全与伦理边界", 1)
    add_bullet(doc, "不做医学诊断：只提示课堂状态可能变化，不判断疾病或障碍类型。")
    add_bullet(doc, "不替代教师：最终判断、教育决策和干预由教师负责。")
    add_bullet(doc, "不主动接触学生：机器人不得主动碰学生，不快速靠近学生。")
    add_bullet(doc, "低刺激原则：避免突然大声、强光闪烁、快速动作和复杂指令。")
    add_bullet(doc, "教师可控：所有声音、提示和机器人动作必须可暂停、可关闭、可调整强度。")
    add_bullet(doc, "数据最小化：优先采集任务日志和匿名特征，视频音频采集必须获得学校和家长授权。")
    add_bullet(doc, "隐私保护：不公开展示学生原始影像，不把课堂视频随意用于宣传。")

    add_heading(doc, "9. 试点实施路径", 1)
    add_table(
        doc,
        ["阶段", "时间建议", "目标", "输出"],
        [
            ["阶段一", "1-2 周", "公益课和教师访谈，确认真实需求和安全边界。", "需求清单、试点班级、授权要求。"],
            ["阶段二", "2-4 周", "完成数字支付训练 MVP 和教师端初版。", "可演示系统、任务脚本、基础报告。"],
            ["阶段三", "4-8 周", "在一个班小范围试运行。", "课堂记录、教师反馈、学生任务数据。"],
            ["阶段四", "8 周后", "根据数据迭代互动强度、任务难度和报告形式。", "第二版产品方案和扩展计划。"],
        ],
        [1200, 1500, 3600, 3060],
    )

    add_heading(doc, "10. 验收指标", 1)
    add_table(
        doc,
        ["指标类别", "建议指标", "说明"],
        [
            ["教学效果", "支付任务完成率、金额判断正确率、错误类型减少。", "观察生活技能是否提升。"],
            ["课堂辅助", "注意力下降提醒次数、教师确认有效率、教师介入延迟。", "判断系统是否真的帮助老师。"],
            ["互动安全", "过度刺激事件数、学生抗拒次数、教师手动关闭次数。", "必须作为核心安全指标。"],
            ["使用体验", "教师满意度、准备时间、课后报告可读性。", "决定能否长期进入课堂。"],
            ["数据合规", "授权完整性、数据脱敏率、原始数据留存边界。", "保证项目可持续推进。"],
        ],
        [1500, 3400, 4460],
    )

    add_heading(doc, "11. 开发优先级", 1)
    add_number(doc, "先开发虚拟小超市与数字支付训练任务，形成可上课的学生端。")
    add_number(doc, "同步开发教师端状态提示和课后报告，保证教师能看懂、能复盘。")
    add_number(doc, "接入摄像头和麦克风的轻量状态感知，先做辅助观察，不做诊断。")
    add_number(doc, "增加低刺激机器人或桌面设备互动，动作必须可控、可关闭。")
    add_number(doc, "在小范围试点中收集任务日志和教师反馈，再决定是否采集更复杂数据。")

    add_heading(doc, "12. 总结", 1)
    add_para(
        doc,
        "本方案的重点不是让机器人替代教师讲课，而是围绕辅读学校真实课堂中的状态观察、注意力拉回和生活技能训练需求，开发一套可控、安全、可试点的具身智能辅助系统。第一阶段应优先落地数字支付生活技能训练，因为它与学生未来社会自立能力直接相关，也便于通过任务完成率和错误类型进行评估。",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_doc()
