from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "Demo2_云端版本功能与后续发布计划_20261022.docx"
SCREENSHOTS = ROOT / "output" / "playwright" / "tracking-20261022"

FONT = "Microsoft YaHei"
INK = "1F2933"
BLUE = "1F4E79"
TEAL = "087F78"
MUTED = "687784"
LINE = "D9E1E8"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F6F8FA"
WHITE = "FFFFFF"


def set_run(run, *, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_fill(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=LINE, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        item = borders.find(tag)
        if item is None:
            item = OxmlElement(f"w:{edge}")
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = margins.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def set_table_widths(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            width = widths[min(i, len(widths) - 1)]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_borders(cell)


def add_paragraph(doc, text="", *, size=10.5, bold=False, color=INK, before=0, after=6,
                  align=WD_ALIGN_PARAGRAPH.LEFT, italic=False, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep
    if text:
        set_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    set_run(p.add_run(text), size=10.2)


def add_number(doc, label: str, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    set_run(p.add_run(f"{label}  "), size=10.5, bold=True, color=BLUE)
    set_run(p.add_run(text), size=10.5)


def add_heading(doc, text: str, level=1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    set_run(p.add_run(text), size=16 if level == 1 else 12.5, bold=True, color="000000")


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_widths(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run(header), size=9.3, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_fill(cell, LIGHT_GRAY if row_index % 2 else WHITE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            set_run(p.add_run(value), size=9.1, color=INK)
    set_table_widths(table, widths)
    add_paragraph(doc, after=2)


def add_picture(doc, filename: str, caption: str, width=6.35) -> None:
    path = SCREENSHOTS / filename
    if not path.is_file():
        add_paragraph(doc, f"截图缺失：{filename}", size=9, color=MUTED, italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    cap = add_paragraph(doc, caption, size=8.8, color=MUTED, after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    cap.paragraph_format.keep_with_next = False


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(paragraph.add_run("RehabAffect Demo2  |  "), size=8.2, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()._r
    run.append(begin)
    run.append(instr)
    run.append(separate)
    run.append(end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 21), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        p_pr = style.element.find(qn("w:pPr"))
        if p_pr is not None:
            for border in list(p_pr.findall(qn("w:pBdr"))):
                p_pr.remove(border)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def remove_paragraph_borders(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        for border in list(p_pr.findall(qn("w:pBdr"))):
            p_pr.remove(border)


def build() -> None:
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "RehabAffect Demo2 云端版本功能与后续发布计划"
    doc.core_properties.subject = "当前云端测试版功能、验收结果与后续版本路线"
    doc.core_properties.author = "RehabAffect 项目组"

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(4)
    set_run(title.add_run("RehabAffect Demo2 云端版本功能与后续发布计划"), size=21, bold=True, color="000000")
    add_paragraph(doc, "当前交付说明　2026年10月22日", size=10, color=MUTED, after=13)
    add_paragraph(doc, "这份说明面向经理、学校测试人员和后续研发人员，说明本次云端版本已经交付的功能、实际验收结果、当前边界，以及下一步按照什么顺序继续扩展。当前版本的首要目标是让学校可以上传真实获授权的课堂视频，等待后台处理，并清楚看到每个学生人物的课堂状态时间轴。", size=11, after=9)
    p = add_paragraph(doc, "最新测试入口：", size=10.8, bold=True, color=BLUE, after=2)
    add_hyperlink(p, "https://YOUR_SERVER_IP/demo2/login", "https://YOUR_SERVER_IP/demo2/login")
    add_paragraph(doc, "当前版本号：simple_student_seats_v3_1。账号密码不写入本说明，请继续使用项目方单独交付的测试账号。", size=9.5, color=MUTED, after=12)

    add_heading(doc, "一 当前版本先给结论", 1)
    add_paragraph(doc, "本次版本已经把“登录—上传视频—后台处理—学生状态结果—报告下载”这条主线部署到云端，并用实际云端任务完成了浏览器验收。人物时间轴不再使用当前这一个视频的固定结果：系统从每个上传任务实际读取视频时长和采样画面，按画面从左到右建立最多三个人物的独立观察区域，再为每个人物生成自己的截图、状态和时间片。", after=7)
    add_paragraph(doc, "因此，换成其他符合限制的视频，会走同一套处理链路；不会因为文件名、任务编号、38分钟时长或当前三张截图不同而失效。视频中人物严重遮挡、离开固定区域或镜头大幅变化时，系统会显示证据不足或人脸未确认，而不是把旁边学生的结果冒充过来。", after=8)

    add_heading(doc, "二 当前版本已经具备的功能", 1)
    add_table(doc, ["功能模块", "当前状态", "本次交付说明"], [
        ["登录与测试入口", "已部署", "使用云端 HTTPS 入口和现有测试账号登录；账号密码不放入公开文档。"],
        ["视频上传", "已部署并验证", "支持上传课堂视频，页面显示上传与后台处理状态；上传完成后不需要人工一直停留在结果页。"],
        ["媒体处理", "已部署并验证", "按上传任务实际视频解码、抽取画面；有声音可提取音轨，无声音视频也能继续生成视频证据。"],
        ["学生人物分析", "核心已验证", "最多三名人物，按从左到右建立固定观察区域，每个人物独立计算画面变化和状态线索。"],
        ["Demo1 风格时间轴", "已部署并验证", "点击人物1、人物2、人物3即可切换；时间轴显示时间段、人物截图、状态、证据说明和参考活动指数。"],
        ["彩色人物截图", "已修复并验证", "截图使用原始彩色视频画面和本人物区域，不再出现整张课堂图、黑白图或串入其他学生的截图。"],
        ["姓名绑定与本课归档", "已有最小闭环", "结果页可给匿名人物绑定姓名，保存本节课总结；姓名可以跳过，先看匿名人物分析。"],
        ["最小 IEP 支持摘要", "已有入口", "可在归档时填写简短支持摘要，例如视觉提示或分步骤指令；不要求在普通备注中填写诊断原文。"],
        ["报告与后台记录", "已有入口", "可下载当前任务结果，公司后台保留任务、视频和处理状态；本次公网验收已核对下载结果与页面版本一致。"],
        ["帮助界面", "已有入口", "提供当前网站的基本操作、视频要求和结果查看说明。"],
    ], [1900, 1800, 5660])

    add_heading(doc, "三 老师实际操作主线", 1)
    add_number(doc, "1", "登录后进入上传工作台。当前测试不要求老师先学习复杂的后台菜单。")
    add_number(doc, "2", "填写课程名称等少量必要信息，选择获授权的课堂视频并上传。")
    add_number(doc, "3", "等待页面显示后台处理完成。上传进度和分析进度属于同一条任务链，处理期间不用重复上传。")
    add_number(doc, "4", "进入学生状态结果页，点击人物卡片切换人物；每个人物只显示自己的截图和状态时间轴。")
    add_number(doc, "5", "确认画面中的人物后，可以输入姓名并保存本课总结；没有学生档案时可以先创建最小档案，也可以跳过姓名绑定。")
    add_number(doc, "6", "需要留存时下载结果。后续增强版再接入跨课程曲线、更多 IEP 字段和教师评价对比。")

    add_heading(doc, "四 本次云端验收结果", 1)
    add_table(doc, ["验收项目", "结果", "实际证据"], [
        ["公网登录与任务访问", "通过", "邀请入口、学校账号登录、任务结果读取均正常。"],
        ["完整视频采样", "通过", "实际任务读取约2299个采样帧，不再只使用少量示例帧。"],
        ["人物轨迹独立性", "通过", "3个人物各有77个时间片，人物时间轴内容不同。"],
        ["同一时刻人物截图", "通过", "每个时间片都保留对应时间的本人物区域截图；漏检时不借用邻近人物。"],
        ["彩色画面与人物切换", "通过", "人物1、2、3的截图均可加载，保留彩色，点击人物卡片可切换。"],
        ["报告下载一致性", "通过", "下载结果的分析版本与页面一致，均为 simple_student_seats_v3_1。"],
        ["移动端页面", "通过", "390×844 浏览器视口无横向溢出。"],
        ["本地核心回归", "通过", "人物跟踪和人物截图测试 20 passed；媒体解码烟测 20 passed。"],
    ], [2200, 1500, 5660])
    add_paragraph(doc, "说明：历史全量回归中还有旧版媒体基线和上传权限边界测试失败，原因是旧断言要求结果不能增加学生分析字段，或沿用旧模块的权限返回码；这些失败不影响本次云端主线验收，但后续应单独清理旧测试基线。", size=9.4, color=MUTED, italic=True, after=9)

    add_picture(doc, "cloud-person-tabs.png", "图1　云端结果页的人物卡片：人物按独立截图显示，可点击切换。", width=6.25)
    add_picture(doc, "cloud-person1-panel.png", "图2　人物1结果页：视频、人物预览、状态时间轴和本课归档入口。", width=6.25)

    add_heading(doc, "五 当前边界与使用时的注意事项", 1)
    add_table(doc, ["项目", "当前边界"], [
        ["视频大小与时长", "单文件不超过8GB，视频时长不超过3小时；分辨率不超过4096×2160，源视频帧率不超过120fps。实际测试建议使用MP4、固定机位、画面清楚的视频。"],
        ["人物数量", "当前学生状态主线最多处理3名人物，按从左到右建立人物1、人物2、人物3；超过3名学生需要后续扩展。"],
        ["人物身份", "当前是固定观察区域和人脸候选核对，不是生物特征级人脸识别。学生换座、走出区域、严重遮挡或镜头切换时会标记不确定。"],
        ["状态含义", "当前分数是人物局部画面变化形成的参考活动指数，不等同于专注度、情绪、医学诊断、学习掌握度或学校成绩。"],
        ["音频能力", "有音轨时可以保存和回放音频证据；当前不把远场声音直接当作某个学生的发言，也没有把麦克风采集结果当作完整语言能力结论。"],
        ["数据合规", "真实学生视频、姓名、IEP和诊断资料必须在获授权范围内使用。测试数据可以保留在公司后台，但是否长期保留或用于训练仍需单独取得授权。"],
    ], [2100, 7260])

    add_heading(doc, "六 后续功能按顺序推出", 1)
    add_paragraph(doc, "后续不建议一次把所有复杂功能堆进老师的上传页面。每一阶段先完成一个可验收的闭环，再进入下一阶段；主页面保持“登录—上传—结果”简单，扩展功能放到结果页、帮助页或公司后台。", after=7)
    add_table(doc, ["顺序", "计划版本", "主要功能", "进入下一阶段的条件"], [
        ["1", "v3.2 稳定性", "继续优化大视频分片上传、断点续传、失败重试、队列并发、处理进度和错误提示；补齐不同编码、无音频、短视频和多人同时上传测试。", "四个测试账号、20多个视频能够连续跑完；上传、处理、结果和下载都能查询，失败时能明确说明原因。"],
        ["2", "v3.3 档案闭环", "完善人物到姓名的绑定校验、学生档案、每课总结归档、最小 IEP 摘要权限和跨课历史；支持搜索已有学生，搜不到时创建最小档案。", "姓名、人物截图、当前视频、学生档案和本课总结能够互相追溯，并且学校账号之间完全隔离。"],
        ["3", "v3.4 课堂行为", "在已有时间轴上增加更稳定的分心、互动、回应、困惑和需要帮助等课堂观察标签；处理失语、无声和奇怪叫声时保留证据并标记不确定，不强行翻译成语言能力。", "每个标签都有时间片、截图或音频依据，教师能看懂来源，不能只给一个没有解释的分数。"],
        ["4", "v3.5 双轨评价", "接入学校人工评价表作为参考轨，同时保留系统智能轨；显示两套分数差异、权重依据和可读的解释。正常课程分数保持合理区间，明显安全事件单独进入人工核查。", "学校评价指标、近年文献和实际测试数据完成确认；评分理由可复查，不能把演示权重当成科学结论。"],
        ["5", "v3.6 AI 辅助", "在明确授权、成本和数据边界后，再接入多模态模型或 DeepSeek/Claude 等模型做受控的画面摘要、异常片段解释和教师建议；模型只作为证据解释层，不直接替代基础视频分析。", "先完成脱敏、密钥管理、供应商协议、费用上限、失败降级和人工抽查，确保模型不可用时主流程仍能出结果。"],
        ["6", "v4.0 多校云平台", "迁移到正式域名、对象存储和关系数据库；增加租户隔离、备份恢复、审计、保留期限、到期清理、监控告警和按学校统计。", "完成安全评审、数据授权、云资源预算和恢复演练后，再开放多学校长期使用。"],
    ], [1200, 1650, 3450, 3060])

    add_heading(doc, "七 推荐的发布原则", 1)
    add_bullet(doc, "先保证每个新视频都能顺利完成上传、处理、结果展示和下载，再增加更多评分维度。")
    add_bullet(doc, "先展示可回看的截图、时间和证据，再展示解释和分数；任何证据不足都标记不确定，不用默认正常结果填空。")
    add_bullet(doc, "姓名、IEP和诊断资料按最小必要原则保存，学生档案与视频任务必须通过权限和授权期限关联。")
    add_bullet(doc, "AI 模型只在基础链路稳定后接入，并设置费用上限、超时降级和不调用模型时的本地结果。")
    add_bullet(doc, "每次发布都做一组真实获授权视频的浏览器验收，保留页面截图和下载文件，避免只看接口返回。")

    add_heading(doc, "八 交付文件与入口", 1)
    add_paragraph(doc, "最新网页入口：", size=10.5, bold=True, color=BLUE, after=2)
    link = add_paragraph(doc, after=8)
    add_hyperlink(link, "https://YOUR_SERVER_IP/demo2/login", "https://YOUR_SERVER_IP/demo2/login")
    add_paragraph(doc, "本次云端验收截图保存在项目目录 output/playwright/tracking-20261022 中；账号密码和邀请密钥继续通过私密方式交付，不写入 Word、截图或源码包。", size=9.5, color=MUTED, after=4)
    add_paragraph(doc, "最终结论：当前版本已经可以作为学校测试版使用，主线是上传视频并查看学生人物状态结果。后续研发按“稳定上传与处理 → 姓名与档案闭环 → 课堂行为增强 → 双轨评价 → AI 辅助 → 多校云平台”的顺序推进。", size=10.8, bold=True, color=BLUE, before=8, after=0)

    remove_paragraph_borders(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
