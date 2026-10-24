from __future__ import annotations

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SCREENSHOT_DIR = PROJECT_ROOT / "output" / "playwright" / "demo2_doc_20260814"
OUTPUT_PATH = PROJECT_ROOT / "docs" / "Demo2_全部细分页面截图展示_20260814.docx"

PAGE_WIDTH_IN = 11.0
PAGE_HEIGHT_IN = 8.5
MARGIN_IN = 0.25
CONTENT_WIDTH_IN = PAGE_WIDTH_IN - 2 * MARGIN_IN
CONTENT_HEIGHT_IN = PAGE_HEIGHT_IN - 2 * MARGIN_IN
MAX_IMAGE_HEIGHT_IN = 7.55


def set_east_asia_font(run, name: str) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.1)
    section.footer_distance = Inches(0.1)

    normal = document.styles["Normal"]
    normal.font.name = "Noto Sans SC"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_screenshot_page(document: Document, image_path: Path, first: bool) -> None:
    if not first:
        document.add_page_break()
        paragraph = document.add_paragraph()
    else:
        paragraph = document.add_paragraph()

    with Image.open(image_path) as image:
        width_px, height_px = image.size
    aspect = width_px / height_px
    display_width = CONTENT_WIDTH_IN
    display_height = display_width / aspect
    if display_height > MAX_IMAGE_HEIGHT_IN:
        display_height = MAX_IMAGE_HEIGHT_IN
        display_width = display_height * aspect

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(max(0, (CONTENT_HEIGHT_IN - display_height) * 36))
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(display_width), height=Inches(display_height))


def add_definition(document: Document, label: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.35
    label_run = paragraph.add_run(f"{label}　")
    set_east_asia_font(label_run, "Noto Sans SC")
    label_run.bold = True
    label_run.font.size = Pt(12)
    label_run.font.color.rgb = RGBColor(20, 125, 115)
    text_run = paragraph.add_run(text)
    set_east_asia_font(text_run, "Noto Sans SC")
    text_run.font.size = Pt(11)
    text_run.font.color.rgb = RGBColor(45, 61, 57)


def add_summary_page(document: Document) -> None:
    document.add_page_break()
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(34)
    spacer.paragraph_format.space_after = Pt(0)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(7)
    kicker_run = kicker.add_run("REHABAFFECT · CLASSROOM STUDIO")
    set_east_asia_font(kicker_run, "Noto Sans SC")
    kicker_run.bold = True
    kicker_run.font.size = Pt(9)
    kicker_run.font.color.rgb = RGBColor(20, 125, 115)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(15)
    title_run = title.add_run("这个 Demo 能做什么")
    set_east_asia_font(title_run, "Noto Serif SC")
    title_run.bold = True
    title_run.font.size = Pt(25)
    title_run.font.color.rgb = RGBColor(23, 61, 58)

    lead = document.add_paragraph()
    lead.paragraph_format.space_after = Pt(13)
    lead.paragraph_format.line_spacing = 1.35
    lead_run = lead.add_run(
        "它把一节课的课程目标、课堂视频、学生个体资料和教师复核连接起来，"
        "形成从课堂证据到学生支持建议的完整演示闭环。"
    )
    set_east_asia_font(lead_run, "Noto Sans SC")
    lead_run.font.size = Pt(12)
    lead_run.font.color.rgb = RGBColor(74, 91, 86)

    add_definition(
        document,
        "输入",
        "课程名称、课程标准、校方考核口径、课堂视频与音频、学生个体教育资料及历史课程记录。",
    )
    add_definition(
        document,
        "分析",
        "同步观察三名学生，整理课堂阶段、状态时间线和学习机会；把参与状态与学习掌握度分开呈现。",
    )
    add_definition(
        document,
        "复核",
        "系统先筛选关键片段，教师再确认学生回答是否正确、接受了什么提示，以及画面外发生了什么。",
    )
    add_definition(
        document,
        "输出",
        "生成课程级分析、学生跨课追踪、教师支持效果、深度报告、复核清单 CSV 和结构化 JSON。",
    )
    add_definition(
        document,
        "支持动作",
        "根据学生状态提出低刺激的屏幕、语音和机器人动作建议，高影响动作必须由教师确认。",
    )
    add_definition(
        document,
        "边界",
        "系统只做课堂辅助观察，不做医学诊断，不把安静、动作多少或单个片段直接解释为学会或不会；资料仅用于授权展示。",
    )


def main() -> None:
    screenshots = sorted(SCREENSHOT_DIR.glob("[0-9][0-9]_*.png"))
    if len(screenshots) != 16:
        raise RuntimeError(f"Expected 16 screenshots, found {len(screenshots)}")

    document = Document()
    configure_document(document)
    for index, screenshot in enumerate(screenshots):
        add_screenshot_page(document, screenshot, first=index == 0)
    add_summary_page(document)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "Demo2 全部细分页面截图展示"
    document.core_properties.subject = "辅读学校课堂证据工作台界面展示"
    document.core_properties.author = "RehabAffectPerception"
    document.core_properties.keywords = "课堂分析, 特殊教育, 教师复核, 学生追踪"
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    main()
