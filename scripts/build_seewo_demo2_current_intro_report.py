from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import build_seewo_demo2_intro_docx as base


ROOT = Path(__file__).resolve().parents[1]
ANALYSIS = ROOT / "outputs" / "seewo_demo2" / "analysis.json"
REPORTS = ROOT / "outputs" / "seewo_demo2" / "reports"
PLAYWRIGHT = ROOT / "output" / "playwright"
OLD_ASSETS = ROOT / "docs" / "assets" / "seewo_demo2_document"
OUTPUT = ROOT / "docs" / "辅读学校课堂学习成效分析_Demo2_校方介绍报告.docx"

# Design preset: compact_reference_guide. First-page pattern: customer_pack.
BLUE = base.BLUE
TEAL = base.TEAL
NAVY = base.NAVY
MUTED = base.MUTED
PALE_BLUE = base.PALE_BLUE
PALE_GOLD = base.PALE_GOLD


def add_picture_path(doc: Document, path: Path, caption: str, width: float = 6.28):
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption.split("  ", 1)[0])
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(0)
    cp.paragraph_format.space_after = Pt(8)
    base.set_run_font(cp.add_run(caption), size=9, color=MUTED, italic=True)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_status_strip(doc: Document, values: tuple[tuple[str, str], ...]):
    table = doc.add_table(rows=1, cols=len(values))
    table.style = "Table Grid"
    widths = [9360 // len(values)] * len(values)
    widths[-1] += 9360 - sum(widths)
    for cell, (label, value) in zip(table.rows[0].cells, values):
        cell.text = ""
        shade_cell(cell, "F4F7FA")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        base.set_run_font(p.add_run(label + "\n"), size=8.5, color=MUTED)
        base.set_run_font(p.add_run(value), size=12, bold=True, color=NAVY)
    base.set_table_geometry(table, widths)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    return table


def add_table(doc: Document, headers, rows, widths, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            cell.text = str(value)
    base.format_table(table, widths, font_size=font_size)
    return table


def add_compact_bullet(doc: Document, text: str):
    p = base.add_list_item(doc, text, "List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    for run in p.runs:
        base.set_run_font(run, size=9.6)
    return p


def latest_report_payload():
    files = sorted(REPORTS.glob("*.json"), key=lambda item: item.stat().st_mtime, reverse=True)
    return json.loads(files[0].read_text(encoding="utf-8")) if files else {}


def build_document():
    payload = json.loads(ANALYSIS.read_text(encoding="utf-8"))
    course = payload["courses"][0]
    students = course["student_results"]
    teacher = course["teacher_effectiveness"]
    long_report = latest_report_payload()
    char_count = long_report.get("char_count", 10000)
    report_sections = long_report.get("sections", [])

    doc = Document()
    base.configure_page(doc)
    base.configure_styles(doc)
    base.configure_header_footer(doc)
    bullets = "List Bullet"
    numbers = base.create_numbering_instance(doc)
    doc.core_properties.title = "辅读学校课堂学习成效分析 Demo2 校方介绍报告"
    doc.core_properties.subject = "课程确认、三学生课堂分析、教师复核、综合报告与具身动作闭环"
    doc.core_properties.author = "RehabAffectPerception"
    doc.core_properties.comments = "本报告用于产品介绍，不构成医学诊断、学生能力定级或教师绩效结论。"

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(12)
    kicker.paragraph_format.space_after = Pt(8)
    base.set_run_font(kicker.add_run("校方沟通材料 · 产品介绍报告"), 9, bold=True, color=TEAL)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    base.set_run_font(title.add_run("辅读学校课堂学习成效分析 Demo2"), 24, bold=True, color=NAVY)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    base.set_run_font(
        subtitle.add_run("课程确认、三学生课堂感知、教师复核与个体化支持动作"),
        13,
        color=MUTED,
    )
    base.add_callout(
        doc,
        "产品定位",
        "面向辅读学校的课堂教育支持原型。系统把课程标准、课堂视频与音频、学生既有资料、教师复核和后续支持动作连接起来，回答“这节课教了什么、学生在明确机会中表现如何、下一步怎样支持”。",
    )
    add_status_strip(
        doc,
        (
            ("分析对象", "3名学生"),
            ("课堂样例", "数学 / 语文"),
            ("报告长度", f"{char_count}字"),
            ("教师复核", "全程保留"),
        ),
    )
    meta_rows = (
        ("文档用途", "校方演示、需求确认、公司汇报与试点讨论"),
        ("本机页面", "http://127.0.0.1:8791/seewo_student_demo2"),
        ("打开条件", "在演示电脑运行 start_school_demo2.cmd 后访问；当前不是公网地址"),
        ("报告日期", date.today().isoformat()),
    )
    add_table(doc, ("项目", "说明"), meta_rows, (2500, 6860))
    base.page_break(doc)
    base.add_heading(doc, "1. 这个Demo解决什么问题", 1)
    base.add_para(
        doc,
        "辅读课堂中，学生的理解、表达、动作与支持需求差异较大。仅统计抬头、举手、动作次数或情绪标签，无法说明学生是否获得了学习机会、是否正确完成、依赖了什么提示，也无法直接指导下一节课。Demo2把“课堂状态”与“学习掌握”分开，并把结果落到教师可复核、机器人可执行的支持动作上。",
    )
    problem_rows = (
        ("学生", "表现难以被连续记录，安静或动作多容易被误解", "用明确任务机会记录正确性、独立程度、一致性和迁移，不把动作直接当作能力结论"),
        ("教师", "整段录像复盘耗时，关键互动与提示等级难统一", "自动定位关键时段，允许教师确认、修正或标记不确定，并形成下一课建议"),
        ("学校", "课程、视频、测评资料分散，跨课追踪缺少统一口径", "把课程标准、个体资料、课堂证据和报告存入同一学生追踪框架"),
        ("机器人", "识别结果停留在页面，不能转成安全动作", "输出屏幕、短句语音、动作原语、风险等级和教师接管要求"),
    )
    add_table(doc, ("对象", "现实问题", "Demo给出的处理方式"), problem_rows, (1200, 3400, 4760), 8.8)
    base.add_callout(
        doc,
        "核心价值",
        "它不是替教师给学生下结论，而是把零散课堂录像整理成可核对的证据，把教师经验沉淀为可追踪的数据，再把低刺激支持建议交给屏幕、语音或机器人执行。",
        fill=PALE_GOLD,
        border="A96300",
    )

    base.page_break(doc)
    base.add_heading(doc, "2. 输入和输出", 1)
    io_rows = (
        ("课程信息", "课程名称、课程类型、官方教案与课程标准、学校考核口径", "课程目标、适配判断和分析边界"),
        ("课堂数据", "可下载的课堂视频及其中音频；必要时教师机视频", "30秒状态窗口、互动时段、画面与声音证据质量"),
        ("学生资料", "匿名编号、固定座位、上次课程表现、个体化支持资料", "个体基线、优势、支持需要、安全约束和跨课趋势"),
        ("教师复核", "符合 / 需修正 / 不确定、答案正确性、提示等级、补充说明", "正式结论修正、复核进度和可用于后续训练的高质量标签"),
        ("系统输出", "页面、文本报告、结构化JSON", "掌握度、课堂状态、教师支持复盘、下一课计划、具身动作建议"),
    )
    add_table(doc, ("类别", "输入内容", "形成的结果"), io_rows, (1500, 4140, 3720), 8.8)
    base.add_para(
        doc,
        "输入方式：当前演示采用本地授权目录中的课堂视频、音频和学生资料，由研发人员预处理后生成分析结果。学校不需要做算法处理；正式试点可增加“选择课程—上传/导入文件—确认学生座位—开始分析”的受控入口。",
        bold_prefix="输入方式：",
    )
    base.add_para(
        doc,
        "输出方式：页面用于交互复核；万字报告用于课后阅读和留档；JSON用于数据库、WebSocket、ROS2 Mock或机器人控制端。没有清晰证据的内容标记为U，不计入“不会”。",
        bold_prefix="输出方式：",
    )

    base.page_break(doc)
    base.add_heading(doc, "3. 完整工作流程", 1)
    flow_steps = (
        "课程确认：先选择课程，核对课程类型、课程名称、学校官方教案要求、本次执行要求和考核标准。",
        "资料关联：读取三名学生的固定座位、匿名档案、既有测评与上次同类课程表现。",
        "视频与音频处理：按固定时间窗口提取画面动作、姿态变化、课堂活动和远场声音线索，同时评估遮挡与清晰度。",
        "状态估计：将单帧线索汇总为稳定参与、回应/操作、回应线索较低或证据不足等时间状态。",
        "学习掌握初评：只在明确出现的学习机会中评估正确性、独立程度、重复一致性和迁移。",
        "教师复核：教师查看关键时段，补充任务语义、答案正确性和提示等级，系统据此修正结论。",
        "报告与动作：生成三学生个体报告、教师支持复盘、下一课计划，以及屏幕、语音和机器人动作建议。",
        "反馈回流：把教师修正与下一课结果写回个体档案，形成跨课追踪和后续模型训练数据。",
    )
    for step in flow_steps:
        base.add_numbered_item(doc, numbers, step, bold_prefix=step.split("：", 1)[0] + "：")
    base.add_callout(
        doc,
        "具身智能闭环",
        "Perception（摄像头/麦克风/任务）→ State Estimation（human_state与掌握状态）→ Action Suggestion（屏幕/语音/机器人动作）→ Feedback（学生反应与教师复核）→ 更新下一时刻策略。",
    )
    add_picture_path(
        doc,
        PLAYWRIGHT / "demo2_desktop_main.png",
        "图1  课程配置、三学生目标适配与群像/人像证据在同一页面呈现",
        width=6.25,
    )

    base.page_break(doc)
    base.add_heading(doc, "4. 课程确认与课程适配", 1)
    base.add_para(
        doc,
        "页面不会在课程背景不清楚时直接输出学生结论。进入分析前，使用者必须确认课程和学校标准；切换课程时会重新确认。这样可以避免把同一个动作在不同课堂任务中作相同解释。",
    )
    confirm_rows = (
        ("选择已有课程", "将视频、课程目标、学生基线和报告绑定到同一课程记录"),
        ("课程类型与名称", "区分生活数学、生活语文等不同任务结构"),
        ("官方教案与课程标准", "明确学生本节课应该获得的学习机会"),
        ("本次执行要求", "记录导入、示范、引导练习、个别练习和总结等阶段"),
        ("学校考核标准", "统一4级量表与U值规则，减少研发人员自行解释"),
        ("确认人和边界勾选", "留下责任链，确认关键片段仍需教师复核"),
    )
    add_table(doc, ("确认项", "实际作用"), confirm_rows, (3000, 6360))
    base.add_heading(doc, "目标适配不是降低要求", 2)
    for text in (
        "系统比较本课目标与学生上次同类任务表现，输出“适配、需分层或需调整”。",
        "当目标跨度较大时，建议把任务拆成可观察的一步，并控制指令长度与提示强度。",
        "最终课程目标由教师和学校确认，不依据障碍名称自动降低目标。",
        "一节课只能建立或更新基线，连续多节同类任务才可以讨论进步、稳定或退步。",
    ):
        base.add_list_item(doc, text, bullets)
    add_picture_path(
        doc,
        PLAYWRIGHT / "demo2_report_course_entry.png",
        "图2  课程确认入口确保分析建立在校方标准和明确教学任务之上",
        width=6.05,
    )

    base.page_break(doc)
    base.add_heading(doc, "5. 三学生同步分析与课堂状态", 1)
    base.add_para(
        doc,
        "同一节课中，系统保留群像用于身份与座位关系复核，同时为每名学生提供独立面板。使用者可以切换学生，查看同一时间段内的状态、掌握初评、支持动作和档案信息。",
    )
    student_rows = []
    for student in students:
        student_rows.append(
            (
                student["display_name"],
                student["mastery"]["overall_score"],
                student["mastery"]["band"],
                student["state_summary"]["participation_score"],
                student["course_fit"]["fit_score"],
            )
        )
    add_table(
        doc,
        ("学生", "掌握初评", "结论", "参与指数", "课程适配"),
        student_rows,
        (1900, 1500, 2400, 1700, 1860),
    )
    base.add_heading(doc, "课堂状态保留的内容", 2)
    for text in (
        "以30秒为窗口组织课堂状态，避免用一帧或一个动作下结论。",
        "每个窗口保留时间、可见证据、解释、置信度、证据质量和风险等级。",
        "页面可以只看待复核时段，帮助教师从30多分钟视频中快速找到关键片段。",
        "状态与学习掌握分开：动作变化可以提示参与线索，但不能单独证明答题正确或知识掌握。",
        "低清、遮挡、远场音频归属不清时标记证据不足，不把不确定强行分配给某个学生。",
    ):
        base.add_list_item(doc, text, bullets)
    add_picture_path(
        doc,
        OLD_ASSETS / "03_state_action_linkage.png",
        "图3  30秒状态窗口与当时的屏幕、语音、机器人动作联动",
        width=6.15,
    )

    base.page_break(doc)
    base.add_heading(doc, "6. 学习掌握评估与教师复核", 1)
    base.add_para(
        doc,
        "学习掌握不从表情或坐姿直接推断，而是围绕“教师提出了什么任务、学生是否出现回应、答案或操作是否正确、接受了多少提示、能否重复和迁移”建立证据链。系统先给出初评，教师完成关键片段复核后才形成更可靠的正式结论。",
    )
    rubric_rows = (
        ("4", "独立完成", "无提示完成目标任务"),
        ("3", "简单提示后完成", "短句、视觉或手势提示后完成"),
        ("2", "示范或部分协助", "示范、分步或部分协助后完成"),
        ("1", "尚未表现", "已有清晰机会，但未出现目标表现"),
        ("U", "未观察 / 不确定", "无机会、遮挡、声音归属不清；不计入不会"),
    )
    add_table(doc, ("等级", "定义", "记录原则"), rubric_rows, (900, 3000, 5460))
    base.add_heading(doc, "教师复核标注", 2)
    review_rows = (
        ("结论", "符合 / 需修正 / 不确定", "决定系统初评能否被采纳"),
        ("答案正确性", "正确 / 部分正确 / 错误 / 无法判断", "补充视频模型难以获得的课程语义"),
        ("提示等级", "独立 / 简单提示 / 示范协助 / 无法判断", "区分学会与在支持下完成"),
        ("教师说明", "当时任务、画面外情况、材料或人员影响", "作为正式报告的优先解释"),
    )
    add_table(doc, ("复核内容", "可选项", "作用"), review_rows, (1800, 3560, 4000), 8.9)
    add_picture_path(
        doc,
        PLAYWRIGHT / "demo2_report_teacher_review.png",
        "图4  教师在关键互动卡片中确认结论、答案正确性与提示等级",
        width=6.28,
    )
    base.add_callout(
        doc,
        "训练数据来源",
        "教师复核不是额外装饰，而是后续训练的核心标签来源。系统初评负责找片段，熟悉学生的教师只需复核少量高价值时段，逐步形成小样本、个体化、可追溯的训练集。",
        fill=PALE_GOLD,
        border="A96300",
    )

    base.page_break(doc)
    base.add_heading(doc, "7. 一万字综合报告", 1)
    base.add_para(
        doc,
        f"系统可以生成约一万字的课程综合报告。本次实测报告为{char_count}字，共{len(report_sections) or 11}个章节。报告生成前会保存当前教师复核记录，并明确区分系统初评、教师已确认结论与仍不确定内容。",
    )
    section_rows = tuple(
        (str(index), item.get("title", "")) for index, item in enumerate(report_sections, start=1)
    ) or (
        ("1", "报告任务、结论层级与使用边界"),
        ("2", "进入分析前的课程确认记录"),
        ("3", "数据来源、视频质量与证据可靠性"),
        ("4", "三名学生课堂过程总览"),
        ("5-7", "三名学生个体学习表现、支持需要与行动建议"),
        ("8", "教学支持效果与教师复盘建议"),
        ("9", "教师复核标注完成度与结论修正原则"),
        ("10", "屏幕、语音与机器人动作"),
        ("11", "下一课实施计划与学校协作事项"),
    )
    add_table(doc, ("章节", "报告内容"), section_rows, (1100, 8260), 8.8)
    add_picture_path(
        doc,
        PLAYWRIGHT / "demo2_report_long_report.png",
        "图5  万字报告支持章节导航、全文复制、字符统计和复核完成度提示",
        width=6.28,
    )
    base.add_heading(doc, "报告的使用方式", 2)
    for text in (
        "校方管理：查看课程是否按确认标准分析、证据是否充分、哪些结论仍待复核。",
        "任课教师：快速定位三名学生的关键表现、有效支持方式和下一课优先目标。",
        "教研复盘：比较同类目标在多节课中的独立完成比例、提示变化和迁移表现。",
        "研发训练：将教师已确认的片段转成后续时序模型、语音互动模型和策略模型的训练样本。",
    ):
        base.add_list_item(doc, text, bullets)

    base.page_break(doc)
    base.add_heading(doc, "8. 是否需要内置LLM", 1)
    base.add_callout(
        doc,
        "结论",
        "LLM可以接入，但不应承担原始视频判定或最终教育结论。它最适合做“受约束的报告整理器”：在结构化证据已经生成、教师边界已经确认后，帮助组织长文本和表达。",
    )
    llm_rows = (
        ("默认本地模板", "不需要网络或API；依据固定证据结构生成完整报告", "稳定、可复现、适合现场演示", "语言灵活度较低"),
        ("可选LLM辅助", "只接收匿名化结构化证据和章节要求，不接收原始视频/音频", "表达自然，可减少重复", "需要API、费用、隐私审查与输出校验"),
    )
    add_table(doc, ("模式", "工作方式", "优点", "限制"), llm_rows, (1600, 3440, 2160, 2160), 8.5)
    base.add_heading(doc, "当前安全设计", 2)
    for text in (
        "未配置LLM时，系统仍能完整生成报告，不影响现场演示。",
        "配置LLM后，只发送匿名编号和结构化证据，不发送学生姓名、确认人、原始视频或音频。",
        "LLM调用失败、返回过短或格式异常时，自动回退到本地证据模板。",
        "教师复核结果优先于LLM措辞；LLM不能新增视频中不存在的事实。",
        "正式部署应使用学校或公司批准的私有模型/API，并记录模型版本、调用时间和报告来源。",
    ):
        base.add_list_item(doc, text, bullets)
    base.add_code_block(
        doc,
        "REHAB_REPORT_LLM_BASE_URL=<approved endpoint>\n"
        "REHAB_REPORT_LLM_MODEL=<model name>\n"
        "REHAB_REPORT_LLM_API_KEY=<secret>",
    )

    base.add_heading(doc, "9. 从分析结果到机器人动作", 1)
    base.add_para(
        doc,
        "系统最终输出不止是“专注、茫然”等标签，而是结构化human_state和可执行动作建议。机器人、桌面屏幕或语音设备读取同一动作结构，根据风险和教师确认条件执行低刺激支持。",
    )
    action_rows = (
        ("稳定参与", "保持当前步骤与材料", "简短具体正反馈", "保持位置或轻微点头", "低风险，可关闭"),
        ("回应线索较低", "只高亮当前一步", "一个短指令后等待", "缓慢转向目标，不靠近", "教师可随时接管"),
        ("困惑/需帮助", "减少无关信息，显示两个选项", "重复说明或拆成一步", "保持距离，指向屏幕", "建议教师确认"),
        ("证据不足", "不增加刺激", "不连续呼叫", "无动作", "必须复核"),
        ("风险升高", "暂停任务", "克制询问是否需要帮助", "停止移动并保持距离", "通知教师/照护者"),
    )
    add_table(doc, ("human_state", "屏幕", "语音", "动作原语", "控制条件"), action_rows, (1400, 2200, 2100, 2060, 1600), 8.2)
    base.add_heading(doc, "具身接口", 2)
    for text in (
        "WebSocket：优先用于演示，把human_state和动作建议实时发送到网页、屏幕或控制端。",
        "ROS2 Mock：在Windows上验证消息字段、动作队列和教师接管逻辑，不强依赖真实ROS2环境。",
        "MuJoCo / Isaac Sim：先验证机器人距离、速度、转向、等待和停止等安全约束，再接真机。",
        "真机阶段：机器人只执行经过白名单限制的低刺激动作，高影响干预必须由教师确认。",
    ):
        base.add_list_item(doc, text, bullets)

    base.page_break(doc)
    base.add_heading(doc, "10. 现场演示方法", 1)
    base.add_heading(doc, "启动", 2)
    base.add_code_block(
        doc,
        "双击：D:\\GIT\\RehabAffectPerception\\start_school_demo2.cmd\n"
        "访问：http://127.0.0.1:8791/seewo_student_demo2",
    )
    base.add_para(
        doc,
        "注意：127.0.0.1是演示电脑自己的本地地址，并不是任何电脑都能直接打开的公网链接。现场展示应在运行服务的电脑上打开；若需校内多机访问或公网展示，需要部署到局域网服务器或云服务器并配置访问控制。",
    )
    base.add_heading(doc, "建议的6分钟讲解顺序", 2)
    demo_steps = (
        "在课程确认页说明：系统先确认课程和学校标准，不先入为主分析学生。",
        "进入主页面，展示三名学生的上次表现、课程适配和个体目标差异。",
        "在群像与人像之间切换，说明身份来自校方确认和固定座位，不做外貌或障碍推断。",
        "打开完整状态，展示30秒时间窗口、证据质量、风险与human_state。",
        "打开证据复核，演示教师如何确认、修正或标记不确定，并补充正确性和提示等级。",
        "打开万字报告，说明章节、字符数、复核进度和本地/LLM两种生成模式。",
        "最后打开动作方案或JSON，展示屏幕、语音、机器人动作和教师接管条件。",
    )
    demo_numbers = base.create_numbering_instance(doc)
    for step in demo_steps:
        base.add_numbered_item(doc, demo_numbers, step)
    base.add_heading(doc, "校方需要确认的内容", 2)
    school_rows = (
        ("课程", "课程名称、官方教案要求、学校考核标准是否准确"),
        ("学生", "匿名编号、固定座位和允许使用的既有资料范围"),
        ("复核", "由哪位熟悉学生的教师复核关键片段；可接受的复核工作量"),
        ("数据", "视频/音频导出方式、授权范围、保存期限和展示边界"),
        ("支持", "哪些屏幕、语音或机器人动作可接受，哪些刺激必须禁止"),
    )
    add_table(doc, ("确认类别", "需要讨论的问题"), school_rows, (2100, 7260))

    base.page_break(doc)
    base.add_heading(doc, "11. 当前完成度、算力与实施边界", 1)
    base.add_heading(doc, "当前可展示的成品", 2)
    for text in (
        "课程确认、课程切换、学校标准和4级考核口径。",
        "三学生同步证据、独立面板、状态时间线、掌握初评和支持动作。",
        "教师关键片段复核、复核进度、记录持久化与结论修正。",
        f"约一万字综合报告，本次实测{char_count}字，包含课程、数据质量、三学生、教师支持、机器人动作和下一课计划。",
        "可选LLM接口、失败回退、本地匿名化和结构化JSON输出。",
        "本地一键启动与自动化测试；Demo1状态分析能力仍保留。",
    ):
        add_compact_bullet(doc, text)
    base.add_heading(doc, "算力要求", 2)
    compute_rows = (
        ("当前页面与缓存结果", "普通Windows笔记本CPU即可", "现场展示不要求独立显卡"),
        ("视频预处理与轻量特征", "CPU可运行；GPU可缩短批处理时间", "低清全景优先做动作、姿态和互动时序"),
        ("ASR/说话人/视觉模型微调", "建议NVIDIA 12-24GB显存或云GPU", "按阶段租用，不必从项目第一天持续使用"),
        ("LLM报告整理", "API或本地推理服务器", "默认模板模式不需要LLM算力"),
        ("Isaac Sim升级", "建议云端RTX显卡环境", "MuJoCo可先在本地验证动作策略"),
    )
    add_table(doc, ("工作", "建议资源", "说明"), compute_rows, (2400, 2960, 4000), 8.8)
    base.add_heading(doc, "必须明确的边界", 2)
    for text in (
        "当前状态识别基于固定座位ROI、时间聚合和证据质量，是可解释原型，不是已训练完成的通用深度模型。",
        "低清全景适合动作与互动时序；精细眼神、微表情、学生级远场语音归属仍需更清晰数据或教师复核。",
        "系统不输出疾病诊断，不按障碍类型给学生排名，也不把结果直接用于教师绩效评价。",
        "个体进步或恶化必须基于多节同类任务、一致评价口径和教师复核，不能由单节课决定。",
    ):
        add_compact_bullet(doc, text)
    base.add_callout(
        doc,
        "建议对外表述",
        "这是课堂状态辅助识别与学习支持系统。它帮助学校把课程标准、课堂证据、教师经验和个体支持连接起来，并把建议落到安全、克制、可关闭的屏幕、语音与机器人动作上。",
        fill=PALE_GOLD,
        border="A96300",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
