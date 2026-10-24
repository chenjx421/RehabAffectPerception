"""Build a concise Word delivery note and source-linked development proposal."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "deliverables" / "20260908_platform_entry" / "课堂分析平台入口与后续开发方案.docx"
SOURCE = ROOT / "docs" / "demo2_analysis_execution_plan_20260908_zh.md"


def inline(paragraph, text: str) -> None:
    parts = re.split(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*)", text)
    for part in parts:
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link and link[2].startswith("https://"):
            node = OxmlElement("w:hyperlink")
            node.set(qn("r:id"), paragraph.part.relate_to(link[2], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            prop = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "245B8A")
            prop.append(color)
            run.append(prop)
            word = OxmlElement("w:t")
            word.text = link[1]
            run.append(word)
            node.append(run)
            paragraph._p.append(node)
        elif link:
            paragraph.add_run(link[1])
        else:
            paragraph.add_run(part.strip("*").replace("`", "")).bold = part.startswith("**")


def table(doc: Document, rows: list[list[str]]) -> None:
    item = doc.add_table(rows=0, cols=len(rows[0]))
    item.alignment = WD_TABLE_ALIGNMENT.CENTER
    item.autofit = False
    for index, cells in enumerate(rows):
        row = item.add_row()
        for column, value in enumerate(cells):
            cell = row.cells[column]
            cell.width = Inches(4.8 if column == 0 else 2.1)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            prop = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "bottom", "left", "right"):
                node = OxmlElement(f"w:{edge}")
                for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    node.set(qn(f"w:{key}"), val)
                borders.append(node)
            prop.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "bottom", "left", "right"):
                node = OxmlElement(f"w:{edge}")
                node.set(qn("w:w"), "100")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            prop.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            inline(paragraph, value)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E9EFF5")
                prop.append(shade)
                for run in paragraph.runs:
                    run.bold = True
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        cannot_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cannot_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.15
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Title"].font.size = Pt(21)
    for name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("课堂学习分析平台  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.add_paragraph("课堂分析平台入口与后续开发方案", "Title")
    doc.add_paragraph("云图智航  RehabAffectPerception   2026年9月8日")
    doc.add_paragraph("本方案说明当前账号与学校入口的实际交付、可转发页面的使用方式，以及下一轮视频分析原型的开发范围。现阶段账号和上传流程可以演示与验证；新视频自动分析仍需接入后台处理、语音转写与教师复核，不能将已有课例的演示评分当作新视频的真实结论。")
    doc.add_heading("本次入口功能", level=1)
    items = [
        "登录页面采用蓝白色，提供用户与管理员两个入口。登录身份由后台核验，选择管理员按钮不会提升普通账号的权限。",
        "用户账号简化为 t1 至 t4，管理员使用 admin；底层账号和原通道绑定保留。密码单独交付，不放在公开预览文件内。",
        "用户登录后确认城市、区县和学校，信息保存到本地数据库。下一次进入可以沿用；学校名称用于资料归属，不等同于已经完成跨校租户隔离。",
        "会员注册可填写联系人、手机或邮箱与学校信息，提交后进入待审核申请。管理员能在后台查看；尚未实现自助即时开通、自动短信或付费会员。",
        "安全说明保留在链接及可展开设置中。上传使用“确认授权并上传”，课程入口使用“同意并继续”，取消安全协议勾选框并保留确认记录。脱敏结果长期归档默认关闭。",
        "管理员进入后台可查看上传通道、任务、保存期限、学校信息和会员开通申请。原课堂分析、学生档案、教师复核和报告入口继续保留。",
    ]
    for number, text in enumerate(items, 1):
        doc.add_paragraph(f"{number}. {text}")
    doc.add_heading("页面如何交付", level=1)
    doc.add_paragraph("登录截图可直接转发。附带的 login.html、register.html 和 privacy.html 是离线预览，放在同一目录即可点击查看。预览不提交账号、注册信息或学生资料，不能用于实际登录。")
    p = doc.add_paragraph()
    inline(p, "临时在线预览：[打开登录页面](https://example.invalid/login.html)。该地址只连接上述静态页面，本机与隧道停止后会失效，不是正式学校数据站。")
    doc.add_paragraph("本机实际入口为 http://127.0.0.1:8791/demo2/login。领导或校方在其他电脑直接打开这个地址，会访问他们自己的电脑，无法进入当前平台。固定公网入口需要另行部署持续运行的应用服务器、域名或受控地址、HTTPS、数据库及私有文件存储。")
    # The source proposal contains the researched claims and their direct links.
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith("# ") or line.startswith("**RehabAffectPerception"):
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[-: ]+", cell) for cell in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows)
            continue
        if line.startswith("## ") or line.startswith("### "):
            level = 2 if line.startswith("###") else 1
            title = re.sub(r"^[一二三四五六七八九十]+、", "", line.lstrip("# "))
            title = title.replace("：", " ")
            doc.add_heading(title, level=level)
        else:
            inline(doc.add_paragraph(), line)
        i += 1
    for root in (doc.styles.element, doc.element):
        for border in root.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
