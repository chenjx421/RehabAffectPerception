"""One-page teacher quick-start based on the current Demo2 controls."""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/pilot_20261016/课堂分析平台简明操作手册_20261016.docx'


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.7)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    normal = doc.styles['Normal']
    normal.font.name = '宋体'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    title = doc.styles['Title']
    title.font.name = '黑体'
    title._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    title.font.size = Pt(19)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    doc.add_paragraph('课堂分析平台简明操作手册', 'Title')
    p = doc.add_paragraph('适用当前内测版　2026年10月16日')
    p.runs[0].font.size = Pt(9)
    doc.add_paragraph('基本流程：登录 → 选择学校 → 上传视频 → 确认学生 → 补充观察 → 下载报告。')

    def step(label, text):
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        p.add_run(label + '　').bold = True
        p.add_run(text)

    step('一 准备视频', '使用已获授权的课堂录像，建议MP4、1080P、固定机位，尽量拍清学生和教师互动。支持MP4、MOV、AVI、MKV、WebM；单文件不超过5GB，处理时长上限3小时。有音轨时可提取回放，当前不自动转写；没有声音的视频也可上传。')
    step('二 登录并确认学校', '打开项目方发来的完整邀请链接，选择学校用户并使用分配的账号登录。首次使用先选择城市、区县并搜索学校；搜不到时选择“其他”，按提示填写省市、区县、街道、学校名称和地址。核对后保存，不要随意更换学校。')
    step('三 建立学生信息', '在“本校学生名册”中点击“创建学生”，填写姓名，建议同时填写校内编号和班级。IEP等资料如需使用，从该学生的“资料与授权”入口按提示登记来源、授权期限并上传；这不是上传视频的必填步骤，诊断内容不要写进普通备注。')
    step('四 上传课程视频', '进入“提交课堂视频”，填写课程批次号、课程名称，点击“读取课程工作区”，再选择视频，点击“确认授权并上传”。同一节课的不同机位使用负责人统一提供的批次号。上传期间保持页面打开和网络连接；进度100%只表示传输完成，待任务显示“证据已就绪”再查看结果。每课程名额以页面显示为准。')
    step('五 对应画面与姓名', '处理完成后进入人员确认页面，选中画面中的人脸候选，再选择学生姓名并点击“确认姓名”。漏检时用“手动画框”，在图上拖出框后选择姓名保存；看不清时先不确认，不要猜测。确认仅对应当前采样帧，不代表已完成整段视频人物追踪。')
    step('六 查看与下载结果', '从本课程任务进入视频证据或“课堂复核与报告”，查看原片、时间片段、画面及音轨证据。需要学生表现和评课内容时，补充相应片段的教学目标、提问、回应及提示帮助等观察并保存，再查看评分依据和导出按钮。当前不会仅凭上传视频自动判断掌握程度；证据不足时不生成确定分数，不能把“证据已就绪”当成评课完成。')
    step('七 公司后台查询', '公司人员使用管理员账号登录，在上传记录中查看各通道的新任务、处理状态，点击文件打开视频并查看或下载可用结果；较早记录可翻页查找。学校账号不查看其他账号的任务，敏感学生档案也不因管理员身份而默认开放。')
    p = doc.add_paragraph()
    p.add_run('遇到问题：').bold = True
    p.add_run('排队或处理中先等待，不必重复上传；失败时记录任务编号、提示信息和文件格式交给项目方。已上传任务可用原批次号重新读取。原视频默认保存60天、提前15天提醒，请及时下载获准导出的结果。')
    p = doc.add_paragraph('已测情况：4个完整课堂视频、合计约147分钟通过上传及媒体处理验证；不代表所有网络环境和视频编码均保证成功，也不代表教育评分准确性已验证。')
    p.paragraph_format.space_before = Pt(3)
    for r in p.runs:
        r.font.size = Pt(9)

    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)
    doc.core_properties.title = '课堂分析平台简明操作手册'
    doc.core_properties.subject = '学校用户上传视频与查询报告的基本操作'
    for node in list(doc.styles.element.iter(qn('w:pBdr'))) + list(doc.element.iter(qn('w:pBdr'))):
        node.getparent().remove(node)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(str(OUT))


if __name__ == '__main__':
    main()
