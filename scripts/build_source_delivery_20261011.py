"""Create a clean source snapshot and reader-oriented source explanation."""
import ast
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import zipfile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'outputs/delivery_20261011'
STAGE=BASE/'RehabAffect_Demo2_Source_20261011'
CODE=STAGE/'source'
MODULES=[
('登录与四通道','demo2_portal.py / student_record_store.py','登录、学校归属、上传名额、账号状态和公司汇总。角色与数据范围由服务器校验，不能只靠前端隐藏按钮。'),
('HTTP入口','scenario_lab_server.py','集中分派页面、API、会话与权限检查；启动视频处理与保留期限扫描。当前使用Python HTTP服务，不是已部署的高可用生产网关。'),
('视频与音轨','demo2_media_analysis.py / demo2_media_worker.py','分片完成后进入队列，真实解码、采样画面、提取同期音轨和质量信息。evidence_ready表示证据材料已准备好，不是教学结论已经正确。'),
('人物与学生','demo2_identity_store.py / demo2_identity.py','将具体采样帧中的框、人物角色和学生ID关联；支持手动框选、撤回、更正及版本校验。同名用学生编号区分。'),
('敏感资料','demo2_dossier_store.py / demo2_dossier_vault.py','IEP、专业评估和支持资料单独授权、加密、设定期限，课堂引用保存具体版本。管理员统计不自动开放原件正文。'),
('课堂观察','demo2_measurement_store.py','观察记录关联学生、课程、目标、时间段、提示条件和复核状态；为个体、群像及课程报告提供统一数据。'),
('课堂与报告','demo2_course_page.py / demo2_course_analysis.py / demo2_measurement.py','课程工作区展示群像、个体和图文输出；报告依赖已录入或复核的证据，不能用合成样例替代新视频的真实结果。'),
('学生长期测量','demo2_measurement_store.py / assessment_rubric.py','按目标和支持条件保留逐课记录，进行周、月和滚动周期聚合。缺失不是零分，不同条件下的分数不能直接解释为进步或退步。'),
('教师双轨评价','classroom_rubric.py / demo2_evaluation_store.py / demo2_evaluation.py','保留学校评价轨与系统增强轨。普通课基线沿用学校分数，调整需要理由、证据与复核；保留版本、分差解释及试行权重信息。'),
('支持与改进','teaching_support.py / demo2_support_page.py','记录有效支持、替代表达方式和下一课改进计划。学生不说话不能直接视为没有参与，建议不构成诊断。'),
('安全与适用范围','demo2_safety_page.py / video_applicability.py','安全事件独立核查，视频适用性与技术质量分开。当前不是自动判断暴力事实或自动识别所有无关视频。'),
('外部模型','seewo_long_report.py','保留模型解释的相关适配入口；本次交付没有API密钥，也不默认向外部发送学生数据。真实多模态分析仍需独立接入和验证。'),
]

def write(path,text):
    path.parent.mkdir(parents=True,exist_ok=True);path.write_text(text,encoding='utf8')

def build():
    import sys
    if STAGE.exists() and '--refresh' not in sys.argv:
        raise RuntimeError('Version directory already exists; preserve the previous delivery.')
    CODE.mkdir(parents=True,exist_ok=True)
    files=[];changes=[]
    for folder,extensions in [('src',{'.py'}),('tests',{'.py'}),('configs',{'.json'})]:
        for path in sorted((ROOT/folder).rglob('*')):
            if not path.is_file() or path.suffix not in extensions or '__pycache__' in path.parts:continue
            text=path.read_text(encoding='utf-8-sig')
            if re.search(r'sk-[A-Za-z0-9_-]{24,}',text):raise RuntimeError('Potential secret in '+str(path.relative_to(ROOT)))
            # The legacy fixture asserts historical names. Keep it outside the shareable source snapshot.
            if path.name=='test_seewo_multi_student_demo.py':
                changes.append({'excluded':str(path.relative_to(ROOT)),'reason':'Historical named-student fixture; not needed for new-account smoke test.'});continue
            target=CODE/path.relative_to(ROOT);target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(path,target);files.append(target)
    shutil.copy2(ROOT/'pyproject.toml',CODE/'pyproject.toml')
    for path in sorted((ROOT/'scripts').glob('qa_*.py')):
        text=path.read_text(encoding='utf-8-sig')
        if re.search(r'sk-[A-Za-z0-9_-]{24,}|王亿珩|Student A|Student B|Student C',text):continue
        for asset in [path,path.with_suffix('.js')]:
            if asset.exists():
                target=CODE/'scripts'/asset.name;target.parent.mkdir(exist_ok=True);shutil.copy2(asset,target)
    write(CODE/'requirements-runtime.txt','numpy>=1.24,<3\ncryptography>=42\nPillow>=10\nopencv-python>=4.8\nimageio-ffmpeg>=0.5\npython-docx>=1.1\n')
    write(CODE/'requirements-test.txt','-r requirements-runtime.txt\npytest>=8\n')
    write(CODE/'install.ps1', '''$ErrorActionPreference = 'Stop'
Set-Location -LiteralPath $PSScriptRoot
py -3 -m venv .venv
if ($LASTEXITCODE -ne 0) { throw 'Python 3.10+ is required.' }
& ./.venv/Scripts/python.exe -m pip install -r requirements-runtime.txt
if ($LASTEXITCODE -ne 0) { throw 'Dependency installation failed.' }
''')
    write(CODE/'initialize_local.py', '''"""Initialize only a NEW local database and print newly generated credentials."""
import os
import secrets
from pathlib import Path
root=Path(__file__).resolve().parent
db=root/'outputs/seewo_demo2/rehab_affect_demo2.sqlite3'
if db.exists():
    raise SystemExit('Database exists. Initialization refused; do not overwrite existing data.')
os.environ['REHAB_DEMO2_DB_PATH']=str(db)
os.environ['REHAB_DEMO2_UPLOAD_DIR']=str(db.parent/'student_documents')
os.environ['REHAB_DEMO2_ACCOUNT_PASSWORD_PREFIX']=secrets.token_urlsafe(32)
from src.rehab_affect.student_record_store import get_demo2_student_store
store=get_demo2_student_store()
passwords={name:secrets.token_urlsafe(14) for name in ['t1','t2','t3','t4','admin']}
store.set_seeded_pilot_passwords(db,passwords)
print('New local credentials. Store privately; these are NOT the existing online accounts.')
for name,password in passwords.items(): print(name, password)
''')
    write(CODE/'start.ps1', '''$ErrorActionPreference = 'Stop'
Set-Location -LiteralPath $PSScriptRoot
if (-not (Test-Path .venv/Scripts/python.exe)) { throw 'Run install.ps1 first.' }
if (-not (Test-Path outputs/seewo_demo2/rehab_affect_demo2.sqlite3)) {
  & ./.venv/Scripts/python.exe initialize_local.py
  if ($LASTEXITCODE -ne 0) { throw 'Initialization failed.' }
}
$env:REHAB_DEMO2_DB_PATH = Join-Path $PSScriptRoot 'outputs/seewo_demo2/rehab_affect_demo2.sqlite3'
$env:REHAB_DEMO2_UPLOAD_DIR = Join-Path $PSScriptRoot 'outputs/seewo_demo2/student_documents'
$env:REHAB_DEMO2_MEDIA_WORKER = '1'
Write-Host 'Open http://127.0.0.1:8793/demo2/login ; Ctrl+C stops the server.'
& ./.venv/Scripts/python.exe -m src.rehab_affect.scenario_lab_server --host 127.0.0.1 --port 8793 --policy rule
''')
    readme='''# Demo2 当前源码交付 2026年10月11日

这是当前工作目录的源码快照，包含尚未提交Git的实现，不等同于某个已发布Git版本。

## 内容
- source/：Python后端、内嵌HTML/CSS/JS前端、配置、测试与部分隔离浏览器验收脚本。
- 源码与功能解释.docx：系统结构、功能映射、关键代码及交接边界。
- 源码汇编.txt：src目录完整代码，按文件名分隔，供检索；以source中的文件为执行依据。
- manifest.json：逐文件SHA256；exclusions.json：排除范围。
- validation.json：本次干净副本验证结果（若验证尚未执行则不生成）。

## Windows首次启动
1. 安装Python 3.10或以上，在source目录打开PowerShell。
2. 执行 powershell -ExecutionPolicy Bypass -File install.ps1。
3. 执行 powershell -ExecutionPolicy Bypass -File start.ps1。
4. 首次初始化会打印t1至t4及admin的新随机密码，请单独妥善保存。随后访问 http://127.0.0.1:8793/demo2/login 。
5. 若8793端口已占用，停止启动并先确认已有服务，不要强行关闭其他服务；可修改启动脚本的端口。

新安装不会包含现有在线账号密码、已有学生、历史视频、IEP、模型密钥或评分记录。不要复制生产数据库进源码包。
安装依赖需联网；FFmpeg/FFprobe须在PATH可用，或按demo2_media_analysis.py的工具查找逻辑配置。imageio-ffmpeg只保证可提供ffmpeg，不保证ffprobe。
Python基本页面不需要GPU。浏览器自动化另需Node.js、Playwright、sharp及Edge/Chromium；其安装不是运行网页的前提。
Windows敏感档案使用当前系统用户保护的密钥；迁移到Linux需配置独立Fernet密钥并验证权限。备份须协调数据库、文件及密钥，不可只复制SQLite。

## 使用与边界
教师账号：确认学校 → 建立名册和课程 → 手动上传视频 → 查看处理/音轨/片段 → 匹配学生 → 观察复核 → 报告。
公司管理员：查看通道与任务、下载结果、管理名额和保留策略；不默认开放IEP和诊断正文。
evidence_ready仅代表视频证据材料提取完成。当前不能承诺完全自动课堂理解、自动可靠评课或真实准确率。
双轨评分、教师改进和千字图文报告依赖可追溯观察；学校基线与有证据的复核调整分开。
默认原媒体60天、提前15天提醒，服务运行时扫描；长期归档需单独授权，当前不是云端永久备份。
本包没有公网隧道/邀请码，旧临时链接已经到期。正式对外测试需部署稳定HTTPS入口及访问控制，不要直接暴露Python开发服务。

## 测试
pip install -r requirements-test.txt
python -m pytest tests/test_classroom_rubric.py tests/test_evaluation_reference_context.py -q
部分历史测试依赖未交付的历史缓存/可选模型或机器配置，并非所有历史测试都能在空库通过。不得把历史测试数当成本包新跑结果。

## 交接事项
确认四个账号的主体归属、上传授权、公司长期存储范围、托管和备份负责人；如需外部模型再配置端点与密钥并进行数据边界审查。
源码保留服务进口所依赖的历史仿真/机器人模块，以免破坏导入；它们不是本次公司平台上线承诺，也不需要安装MuJoCo才能启动当前网页。
'''
    write(STAGE/'README.md',readme)
    write(STAGE/'exclusions.json',json.dumps({'excluded_roots':['.runtime','outputs','output','data','models','.git','.venv','dist','deliverables','tmp'],'details':changes,'note':'No live databases, students, media, IEP originals, API keys, live passwords or invitation links.'},ensure_ascii=False,indent=2))
    sources=sorted((CODE/'src').rglob('*.py'))
    write(STAGE/'源码汇编.txt','\n\n'.join('===== '+str(f.relative_to(CODE))+' =====\n'+f.read_text(encoding='utf-8-sig') for f in sources))
    d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
    style=d.styles['Normal'];style.font.name='Microsoft YaHei';style.font.size=Pt(10);style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');style.paragraph_format.space_after=Pt(5)
    for name in ['Title','Heading 1','Heading 2']:
        d.styles[name].font.name='Microsoft YaHei';d.styles[name].element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        d.styles[name].font.color.rgb=RGBColor(0,0,0)
    for border in list(d.styles['Title'].element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    d.add_heading('辅读课堂平台源码与功能解释',0)
    d.add_paragraph('交付日期 2026年10月11日')
    d.add_paragraph('本次交付保存当前软件实现，供公司归档、技术接手和后续开发。它包括可执行源码、初始化和启动方式、测试代码、完整文本源码以及本说明。现有学生数据和运行密钥不随源码流转。')
    d.add_heading('一 交付范围与当前结论',1)
    d.add_paragraph(f'运行源码共{len(sources)}个Python文件，按实际文件统计。前端主要内嵌在Python页面模块中，不是单独的React或Vue工程。包内保存当前工作副本，不只保存最后一次Git提交。')
    d.add_paragraph('已经实现账号、四路上传、媒体处理、学生身份关联、敏感档案、观察复核、报告和双轨评价等工作流。材料提取与人工辅助评价可以测试；完全自动理解整堂课、可靠自动评课和全校规模运行仍需后续验证。')
    d.add_heading('二 如何阅读交付包',1)
    for text in ['先阅读README完成环境准备，再启动source中的服务。新安装使用空库和新生成的账号密码，不恢复旧学生资料。','管理人员可先阅读本说明；技术接手人员根据下方文件映射进入模块。源码汇编.txt包含完整运行代码，具体文件以source为准。','manifest.json用于确认收到的文件是否完整；任何后续修改应形成新快照，不能用行数或文件数证明功能质量。']:
        d.add_paragraph(text)
    d.add_heading('三 系统如何完成一次课堂任务',1)
    d.add_paragraph('浏览器提交课程和视频，服务校验登录及通道名额，分片落盘后建立处理任务。后台工作线程解码和提取材料，教师核对学生身份并补充观察。测量与报告根据同一任务、学生、目标和证据生成；公司管理员在自己的权限内汇总和下载。')
    d.add_paragraph('IEP原件与课程视频不混作同一条数据：学生ID负责关联，具体资料版本负责追溯，授权和保留期限负责限制访问。姓名相同不代表同一个学生，画面框也不代表永久身份识别。')
    d.add_heading('四 页面功能与代码对应',1)
    for title,files,description in MODULES:
        d.add_heading(title,2)
        names=d.add_paragraph(files);names.paragraph_format.keep_with_next=True
        d.add_paragraph(description)
    d.add_heading('五 关键实现代码与解释',1)
    for filename,function,explanation in [
        ('student_record_store.py','_initial_password','首次账户使用配置的初始密码前缀；本包初始化脚本另生成随机独立密码，不复制现网测试密码。'),
        ('student_record_store.py','_configured_course_upload_limit','上传名额由配置和服务端规则控制，不靠前端显示数字实现限制。'),
    ]:
        path=CODE/'src/rehab_affect'/filename;text=path.read_text(encoding='utf8');tree=ast.parse(text)
        node=next(n for n in ast.walk(tree) if isinstance(n,(ast.FunctionDef,ast.AsyncFunctionDef)) and n.name==function)
        d.add_heading(function,2)
        explanation_paragraph=d.add_paragraph(explanation);explanation_paragraph.paragraph_format.keep_with_next=True
        para=d.add_paragraph();para.paragraph_format.keep_together=True
        run=para.add_run('\n'.join(text.splitlines()[node.lineno-1:node.end_lineno]));run.font.name='Consolas';run.font.size=Pt(8)
    d.add_paragraph('代码中的函数和状态名是排查问题的入口。完整函数、调用关系和上下文见对应源文件；不要仅复制上述片段作为独立程序。')
    d.add_heading('六 数据保存与安全',1)
    for text in ['数据库为本地SQLite，文件分开保存在video_tasks和敏感档案目录。关闭浏览器不删除数据；本机磁盘故障仍可能造成丢失，当前源码不等于已经完成云端备份部署。','原媒体默认60天，提前15天提醒，服务运行时执行保留策略扫描。IEP和专业资料独立加密、授权、到期及撤回。公司长期资源只接收另有授权的结构化内容，不默认复制原视频和诊断。','源码包不含学生视频、姓名档案、学校评估原件、数据库、密钥、真实登录密码或旧邀请地址。历史命名学生测试被排除并登记；运行源码未作业务逻辑改写。']:
        d.add_paragraph(text)
    d.add_heading('七 环境启动与迁移',1)
    d.add_paragraph('推荐先在Windows与Edge环境验收。Python 3.10以上，安装requirements-runtime.txt；视频处理还需FFmpeg及FFprobe。运行install.ps1后运行start.ps1。端口默认8793，仅监听本机。正式公网部署应增加受控反向代理、HTTPS、访问控制、备份及运行监控。')
    d.add_paragraph('首次运行生成新数据库及五个随机密码，旧数据不会自动恢复。已有数据库时初始化拒绝覆盖。迁移真实数据需要另行授权，并同时备份数据库、媒体文件和可恢复的加密密钥；不要把生产数据混入本代码压缩包。')
    d.add_heading('八 当前不足与后续开发',1)
    for text in ['分析能力：补充真实音频转写、说话人归属、课堂语义和跨片段人物追踪；必须用独立校方课例评估，不把能调用API视为识别准确。','评分依据：保留学校基线，利用复核证据解释调整。试行权重需要教师一致性和样本校准，参考文献不能直接证明某个具体权重比例。','运行保障：四路完整高清课程并发、弱网中断重试、磁盘容量、到期清理及备份恢复仍需正式部署环境实测。','外部模型：确认供应商、费用和数据流向，使用独立环境变量与密钥管理；先做脱敏结构化解释，不默认上传敏感档案。','交付验证：本包的实际检查以validation.json为准；历史开发中的测试记录不能替代本次验证，也不能等同于真实课堂准确率。']:
        d.add_paragraph(text)
    d.add_heading('九 运行源码索引',1)
    d.add_paragraph('以下为交付的运行模块；历史兼容与仿真模块保留是为了完整导入，并不作为当前公司产品路线。')
    for f in sources:
        line=d.add_paragraph(str(f.relative_to(CODE)),style='Normal')
        line.paragraph_format.space_after=Pt(0)
        line.paragraph_format.line_spacing=Pt(11)
        for run in line.runs:run.font.size=Pt(9)
    foot=sec.footer.paragraphs[0];foot.alignment=2
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    d.save(STAGE/'源码与功能解释.docx')
    print(json.dumps({'stage':str(STAGE),'source_files':len(sources)},ensure_ascii=False))

def finalize():
    entries=[]
    for f in sorted(STAGE.rglob('*')):
        if f.is_file() and f.name!='manifest.json':entries.append({'path':f.relative_to(STAGE).as_posix(),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
    write(STAGE/'manifest.json',json.dumps({'date':'2026-10-11','files':entries},ensure_ascii=False,indent=2))
    target=BASE/(STAGE.name+'.zip')
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
        for f in sorted(STAGE.rglob('*')):
            if f.is_file():z.write(f,STAGE.name+'/'+f.relative_to(STAGE).as_posix())
    with zipfile.ZipFile(target) as z:
        assert z.testzip() is None
        for entry in entries:
            assert hashlib.sha256(z.read(STAGE.name+'/'+entry['path'])).hexdigest()==entry['sha256']
        assert not any(Path(name).suffix.lower() in {'.db','.sqlite3','.mp4','.wav','.enc'} for name in z.namelist())
    print(json.dumps({'zip':str(target),'files':len(entries)+1,'bytes':target.stat().st_size},ensure_ascii=False))

if __name__=='__main__':
    import sys
    finalize() if '--finalize' in sys.argv else build()
