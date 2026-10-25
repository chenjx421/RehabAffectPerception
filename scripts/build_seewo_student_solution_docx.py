from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DOCS_DIR = PROJECT_ROOT / "docs"
OUTPUT_PATH = DOCS_DIR / "seewo_student_ai_support_solution_zh.docx"
SCREENSHOT_DIR = PROJECT_ROOT / "outputs" / "playwright"

SCREENSHOTS = [
    (
        SCREENSHOT_DIR / "manager_01_overview_with_profile.png",
        "图 1  系统总览：从 Seewo 课堂数据进入单学生状态追踪与执行建议",
    ),
    (
        SCREENSHOT_DIR / "manager_02_local_anonymous_profile.png",
        "图 2  个体档案：匿名化 IEP 与五大领域资料进入学生支持画像",
    ),
    (
        SCREENSHOT_DIR / "manager_03_action_constraints.png",
        "图 3  动作建议：个体资料约束屏幕、语音、机器人动作和老师复核",
    ),
]


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 32, 48)
MUTED = RGBColor(85, 85, 85)
LIGHT_BLUE = "EAF3FF"
LIGHT_GRAY = "F4F6F9"
MID_GRAY = "DADCE0"
WHITE = "FFFFFF"


def set_run_font(run, name: str = "Microsoft YaHei", size: int | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_paragraph_spacing(paragraph, before: int = 0, after: int = 6, line: float = 1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = MID_GRAY, size: str = "6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths_inches: list[float]):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_inches[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            for p in cell.paragraphs:
                set_paragraph_spacing(p, before=0, after=2, line=1.08)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p.style = doc.styles[style]
    run = p.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], bold=True)
    run.font.color.rgb = BLUE if level in (1, 2) else DARK_BLUE
    set_paragraph_spacing(p, before={1: 16, 2: 12, 3: 8}[level], after={1: 8, 2: 6, 3: 4}[level], line=1.1)
    return p


def add_body(doc: Document, text: str, bold_lead: str | None = None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=4, line=1.15)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_number(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=4, line=1.15)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_callout(doc: Document, title: str, body: str, fill: str = LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True)
    r.font.color.rgb = DARK_BLUE
    p2 = cell.add_paragraph()
    set_paragraph_spacing(p2, after=2, line=1.15)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10)
    doc.add_paragraph()
    return table


def add_caption(doc: Document, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=2, after=8, line=1.05)
    run = p.add_run(caption)
    set_run_font(run, size=9)
    run.font.color.rgb = MUTED


def add_picture_block(doc: Document, path: Path, caption: str):
    if not path.exists():
        add_callout(doc, "截图缺失", f"未找到截图文件：{path}", fill="FFF4E5")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=4, after=2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.35))
    add_caption(doc, caption)


def add_kv_table(doc: Document, rows: list[tuple[str, str]], widths=(1.65, 4.85)):
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_width(table, list(widths))
    for i, (k, v) in enumerate(rows):
        left, right = table.rows[i].cells
        set_cell_shading(left, LIGHT_GRAY)
        left.paragraphs[0].add_run(k)
        right.paragraphs[0].add_run(v)
        for cell in (left, right):
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r, size=10, bold=(cell is left))
    doc.add_paragraph()
    return table


def add_matrix(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_width(table, widths)
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        set_cell_shading(cell, "E8EEF5")
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            run = cell.paragraphs[0].add_run(value)
            set_run_font(run, size=9)
    doc.add_paragraph()
    return table


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    for style_name in ("Normal", "List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
    for style_name, size, color in (
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 12, DARK_BLUE),
    ):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, before=0, after=3, line=1.05)
    run = p.add_run("辅读学校学生个体状态追踪与课堂支持系统方案")
    set_run_font(run, size=22, bold=True)
    run.font.color.rgb = INK

    p2 = doc.add_paragraph()
    set_paragraph_spacing(p2, before=0, after=8, line=1.1)
    run2 = p2.add_run("基于 Seewo 课堂视频、评课报告与学生基础资料的 AI 辅助支持原型")
    set_run_font(run2, size=12)
    run2.font.color.rgb = MUTED

    add_kv_table(
        doc,
        [
            ("方案定位", "面向辅读学校课堂的单学生状态追踪、个体化支持建议与具身动作输出系统"),
            ("当前载体", "本地可运行 Web Demo；后续可部署为校内服务器/云服务器 + 教师端页面 + 机器人/屏幕执行端"),
            ("演示数据", "Seewo 导出的课堂视频、评课详情、目标学生匿名化 IEP 与五大领域评估摘要"),
            ("使用边界", "仅做教育支持与教师复核参考，不做医学诊断，不替代教师判断"),
        ],
    )


def build_doc():
    doc = Document()
    configure_document(doc)
    add_title(doc)

    add_callout(
        doc,
        "一句话方案",
        "在 Seewo 已有课堂数据基础上，为每名学生建立可持续更新的个体支持档案；系统分析课堂视频、声音与评课信息，识别学生参与、困惑、注意力下降等状态，并输出屏幕、语音、机器人动作和老师端复核建议。",
    )

    add_heading(doc, "1. 建设目标", 1)
    add_body(
        doc,
        "本方案面向辅读学校真实课堂，不以“替代老师管理学生”为目标，而是服务于学生的学习参与和课堂支持。系统通过课堂视频、声音、评课报告与学生基础资料，形成单学生的持续追踪档案，并把分析结果转化为低刺激、可复核、可执行的支持动作。",
    )
    add_bullet(doc, "帮助学生：在困惑、跟不上或注意力下降时获得更适合自己的提示方式。")
    add_bullet(doc, "帮助老师：快速定位值得复看的课堂片段，减少从整节课视频中人工翻找的成本。")
    add_bullet(doc, "帮助学校：把 Seewo 课堂记录进一步沉淀为学生个体发展画像和支持策略库。")
    add_bullet(doc, "帮助后续硬件落地：将支持建议标准化为屏幕、语音、机器人动作和教师端提醒。")

    add_heading(doc, "2. 最终产品载体", 1)
    add_kv_table(
        doc,
        [
            ("教师端 Web 页面", "查看单学生课堂状态、关键片段、个体画像、支持建议和复核结果。"),
            ("学生支持数据库", "按学生、课程、时间片段保存状态记录、复核结果、有效提示方式和禁忌刺激。"),
            ("动作执行接口", "输出给屏幕提示、语音播报、小型机器人动作或老师端提醒。"),
            ("部署形态", "第一阶段本地 Demo；第二阶段校内电脑/云服务器；第三阶段接入机器人或互动屏。"),
        ],
    )

    add_heading(doc, "3. 数据来源与处理方式", 1)
    add_body(doc, "系统优先利用学校已经具备的数据资产，不要求一开始额外安装复杂硬件。当前原型已经围绕 Seewo 数据和校方学生资料建立了可展示链路。")
    add_matrix(
        doc,
        ["数据类型", "来源", "进入系统后的作用"],
        [
            ["课堂视频", "Seewo 导出课堂录像", "定位目标学生，按 30 秒窗口分析参与、动作变化、成人遮挡等线索。"],
            ["课堂声音", "视频自带音频", "估计课堂互动强度，辅助判断学生是否在问答/指令阶段出现停滞。"],
            ["评课详情", "Seewo 评课文档", "提供课程主题、课堂节奏、抬头率、参与度、语速等课堂背景。"],
            ["学生基础资料", "IEP、认知、语言沟通、运动等评估", "形成个体化支持约束，如一步视觉提示、短句语音、低刺激动作。"],
            ["老师复核结果", "教师在系统中确认/修正", "反向更新学生画像，形成更贴合该学生的支持策略。"],
        ],
        [1.18, 1.55, 3.77],
    )
    add_callout(
        doc,
        "数据合规边界",
        "演示和对外沟通时仅展示匿名化学生样例，不公开学生姓名、家庭信息、联系方式和原始评估细项。IEP 中的障碍类别只作为教育支持背景，不作为医学诊断输出。",
        fill="FFF4E5",
    )

    add_heading(doc, "4. 核心功能设计", 1)
    add_matrix(
        doc,
        ["功能模块", "功能说明", "当前 Demo 体现"],
        [
            ["单学生定位", "从全景课堂视频中确定目标学生区域，围绕一个学生进行追踪。", "已使用目标学生 ROI 和课堂截图建立演示。"],
            ["课堂切片分析", "将 30 多分钟课程切成连续时间窗口，输出每段状态和证据。", "已生成 77 个时间窗口和关键复核片段。"],
            ["状态估计", "识别参与、注意力下降、疑似困惑/停滞、成人遮挡不确定等课堂状态。", "已在页面展示状态分布和关键片段。"],
            ["个体画像", "结合 IEP 与五大领域资料，形成该学生的支持背景。", "已接入认知、语言沟通、运动能力摘要，并预留感知和社会适应字段。"],
            ["支持策略生成", "根据状态和个体画像生成屏幕、语音、机器人动作和老师复核建议。", "已输出一步视觉提示、短句语音、慢速动作、安全边界。"],
            ["长期追踪", "每节课生成 session record，累积为 student profile。", "当前已展示单节课基线，下一阶段扩展多节课趋势。"],
        ],
        [1.25, 2.45, 2.8],
    )

    add_heading(doc, "5. 系统工作流程", 1)
    for step in [
        "导入 Seewo 课堂视频、音频和评课详情。",
        "选定目标学生，建立匿名 ID 和初始支持档案。",
        "系统按时间窗口分析课堂状态，并筛出值得老师复核的片段。",
        "结合学生 IEP 与五大领域资料，判断更适合该学生的提示方式。",
        "输出屏幕提示、语音提示、机器人动作和老师端提醒。",
        "老师确认、修正或标记不确定，结果进入学生长期追踪数据库。",
    ]:
        add_number(doc, step)

    add_heading(doc, "6. 当前 Demo 成品展示", 1)
    add_body(doc, "当前 Demo 已经可以作为对内汇报和校方初步沟通材料。页面分为三栏：左侧为学生画像与数据来源，中间为学生支持结论，右侧为复核与执行建议。")
    for image_path, caption in SCREENSHOTS:
        add_picture_block(doc, image_path, caption)

    doc.add_page_break()
    add_heading(doc, "7. 机器人与具身动作结合", 1)
    add_body(
        doc,
        "本方案中的“具身智能”不只体现在识别状态，而是体现在系统把学生状态和个体资料转化为可执行动作，并约束动作的强度、安全边界和执行对象。",
    )
    add_matrix(
        doc,
        ["输出端", "动作示例", "设计原则"],
        [
            ["屏幕端", "只保留当前一步、弱化无关信息、局部高亮任务目标。", "降低认知负荷，不制造额外干扰。"],
            ["语音端", "使用短句提示，如“没关系，我们先看第一步”。", "避免长指令和连续追问。"],
            ["机器人端", "慢速转向屏幕、轻微点头、保持距离、不突然靠近。", "低刺激、非接触、老师在环确认。"],
            ["教师端", "提醒老师复核关键片段，标记符合观察/需要修正/不确定。", "最终判断权保留给老师。"],
        ],
        [1.25, 2.65, 2.6],
    )

    add_heading(doc, "8. 模型与训练路线", 1)
    add_body(doc, "由于真实辅读学校数据样本少、视频清晰度有限，第一阶段不追求训练大模型，而是采用轻量可落地路线：先用视频切片、运动线索、音频活动、评课信息和个体档案形成弱标签，再通过老师复核逐步修正。")
    add_bullet(doc, "短期：基于启发式特征和轻量模型，完成关键片段筛选和状态初判。")
    add_bullet(doc, "中期：积累多节课复核结果，训练学生级/班级级状态分类 baseline。")
    add_bullet(doc, "长期：引入多模态模型或 VLA/VLM 能力，用于更复杂的课堂场景理解和动作策略生成。")
    add_bullet(doc, "算力：本地电脑可完成 demo 推理、特征提取和小模型训练；云服务器主要用于批量视频处理、多学生数据管理和后续模型微调。")

    doc.add_page_break()
    add_heading(doc, "9. 硬件与部署需求", 1)
    add_matrix(
        doc,
        ["阶段", "硬件/环境", "说明"],
        [
            ["Demo 阶段", "普通 Windows 电脑 + Seewo 导出视频", "用于展示页面、视频分析结果和个体化动作建议。"],
            ["试点阶段", "校内电脑或云服务器 + 数据库 + 教师端浏览器", "用于处理多节课数据、保存学生档案和老师复核记录。"],
            ["互动阶段", "互动屏/平板/小型机器人/音箱", "用于执行低刺激提示动作，暂不主动接触学生。"],
            ["扩展阶段", "GPU 云服务器", "用于批量视频处理、多模态模型训练和更复杂的个体化策略学习。"],
        ],
        [1.15, 2.35, 3.0],
    )

    doc.add_page_break()
    add_heading(doc, "10. 阶段性交付规划", 1)
    add_matrix(
        doc,
        ["阶段", "交付内容", "验收标准"],
        [
            ["第一阶段：可展示 Demo", "单学生课堂视频分析、IEP/五大领域接入、动作建议页面。", "能向经理和校方说明系统价值，完成一名学生一节课的追踪样例。"],
            ["第二阶段：多节课追踪", "同一学生多节课记录、趋势对比、支持策略变化。", "能回答学生是否进步、是否更稳定、哪类提示更有效。"],
            ["第三阶段：教师复核闭环", "教师标注入口、复核结果入库、自动更新学生画像。", "能形成可训练数据集和个体化支持策略库。"],
            ["第四阶段：机器人/屏幕联动", "将动作建议输出到屏幕、语音或机器人 mock 接口。", "能展示从感知到动作执行的具身智能闭环。"],
        ],
        [1.4, 2.55, 2.55],
    )

    add_heading(doc, "11. 对学校的实际价值", 1)
    add_bullet(doc, "对学生：获得更适合自身能力特点的课堂提示，减少因听不懂、跟不上而被误判为不配合。")
    add_bullet(doc, "对老师：从整节课视频中快速定位关键片段，形成可复盘、可校正的观察记录。")
    add_bullet(doc, "对学校：把已有 Seewo 数据转化为学生长期发展档案，而不是只停留在课堂评课报表。")
    add_bullet(doc, "对后续项目：为机器人、互动屏和多模态模型训练提供真实业务闭环与数据基础。")

    add_callout(
        doc,
        "当前版本结论",
        "本系统当前已经具备“单学生课堂追踪 + 个体档案接入 + 动作建议输出”的最小可展示闭环。下一阶段重点不是继续堆功能，而是围绕同一学生导入多节课数据，验证系统能否稳定形成长期画像和有效支持策略。",
        fill="EAF7EF",
    )

    # Footer
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RehabAffectPerception | Seewo Complementary AI Tutor")
    set_run_font(r, size=9)
    r.font.color.rgb = MUTED

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    output = build_doc()
    print(output)
