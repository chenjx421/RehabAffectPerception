"""Build the scoped Chinese pilot proposal; never open operational databases.

Run with the bundled documents Python runtime. requests may be supplied through
a temporary PYTHONPATH. Source receipts and rendering files stay in TEMP.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import tempfile
from datetime import datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs/pilot_20261016/学校课堂视频分析试点方案_20261016.docx"
QA = Path(tempfile.gettempdir()) / "pilot_20261016_docqa"
SOURCES = {
    "oss": ("阿里云 OSS 计费概述", "https://help.aliyun.com/zh/oss/product-overview/billing-overview"),
    "sqlite": ("SQLite 适用场景", "https://www.sqlite.org/whentouse.html"),
    "pg": ("PostgreSQL 行级安全", "https://www.postgresql.org/docs/current/ddl-rowsecurity.html"),
    "backup": ("PostgreSQL 备份与恢复", "https://www.postgresql.org/docs/current/backup.html"),
    "law": ("中华人民共和国个人信息保护法", "https://www.gov.cn/xinwen/2021-08/20/content_5632486.htm"),
    "tunnel": ("Cloudflare Quick Tunnels", "https://developers.cloudflare.com/cloudflare-one/networks/connectors/cloudflare-tunnel/do-more-with-tunnels/trycloudflare/"),
}


def fetch_sources():
    QA.mkdir(parents=True, exist_ok=True)
    receipts = {}
    try:
        import requests
    except ImportError:
        requests = None
    for key, (title, url) in SOURCES.items():
        item = {"title": title, "url": url,
                "accessed": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(timespec="seconds")}
        if requests is None:
            item["status"] = "未抓取 requests 不可用"
        else:
            try:
                response = requests.get(url, timeout=(10, 25), headers={"User-Agent": "PilotProposalSourceCheck/1.0"})
                item.update(http_status=response.status_code, final_url=response.url,
                            sha256=hashlib.sha256(response.content).hexdigest())
                response.raise_for_status()
                item["status"] = "已抓取 HTTP 200" if response.status_code == 200 else f"HTTP {response.status_code}"
                (QA / f"source_{key}.html").write_bytes(response.content)
            except Exception as exc:
                item["status"] = "抓取失败 待人工核验"
                item["error_type"] = type(exc).__name__
        receipts[key] = item
        print(key, item["status"], flush=True)
    (QA / "sources.json").write_text(json.dumps(receipts, ensure_ascii=False, indent=2), encoding="utf-8")
    return receipts


def font(style, size, bold=False):
    style.font.name = "宋体"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def paragraph(doc, text, style=None, small=False):
    p = doc.add_paragraph(text, style)
    if small:
        p.paragraph_format.space_after = Pt(4)
        for run in p.runs:
            run.font.size = Pt(9)
    return p


def heading(doc, text):
    doc.add_heading(text, level=2)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for cell, title in zip(t.rows[0].cells, headers):
        cell.text = title
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, row in enumerate(t.rows):
        cant = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant)
        for j, cell in enumerate(row.cells):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + edge)
                for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    el.set(qn("w:" + key), val)
                borders.append(el)
            pr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                el = OxmlElement("w:" + edge)
                el.set(qn("w:w"), "85")
                el.set(qn("w:type"), "dxa")
                margins.append(el)
            pr.append(margins)
            if i == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E9EEF2")
                pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    run.font.size = Pt(9.5)
                    run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def source(doc, receipts, key):
    s = receipts[key]
    paragraph(doc, f"资料  {s['title']}  |  访问 {s['accessed'][:10]}  |  {s['status']}\n{s['url']}", small=True)


def page(doc, title, first=False):
    if not first:
        doc.add_page_break()
    doc.add_heading(title, level=1)


def build(receipts):
    doc = Document()
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.75), Cm(1.65)
    sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.8)
    for name, size, bold in (("Normal", 10.5, False), ("Title", 20, True),
                             ("Heading 1", 15, True), ("Heading 2", 11.5, True)):
        font(doc.styles[name], size, bold)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.23
    normal.space_after = Pt(6)
    for name in ("Heading 1", "Heading 2"):
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after = Pt(7), Pt(6)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "学校课堂视频分析试点方案"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.subject = "套餐口径 成本核算 部署授权与试点验收"

    doc.add_paragraph("学校课堂视频分析试点方案", "Title")
    paragraph(doc, "2026 年 10 月 16 日    面向校方负责人 教师与项目实施团队", small=True)
    paragraph(doc, "本轮试点以可回看的视频证据、教师复核和授权管理为交付重点。校方据此验证教学工作流程，再决定半年或年度采购；报价按实际用量、人工投入和部署范围核定。当前未实现自动教育评分，评分须由教师依据证据完成。")
    heading(doc, "一 当前状态与能力边界")
    table(doc, ["核对项", "维护前基线与能力范围"], [
        ("任务基线", "维护前基线为 28 个任务：25 个 evidence_ready，2 个 uploaded_pending_worker 历史占位文件各 29 字节，1 个已取消。后续验收新增合成任务，当前总量随测试变化。"),
        ("资料基线", "维护前基线中，范围内学生记录、dossier 文档及旧 student_documents 均为 0。合成验收数据与真实学校资料分别统计。"),
        ("实际分析", "源码链路包含 FFprobe 元数据、FFmpeg 解码、每秒 1 帧采样、画面亮度及清晰度和运动量、音轨与能量证据。音频能量不是语音识别，画面变化不是学生专注度。"),
        ("评分责任", "媒体处理输出技术证据，实名绑定由教师确认；情绪和掌握程度不由该提取器判定。教育评分与解释由教师依据片段完成。"),
    ], [2.4, 15.0])
    heading(doc, "二 本轮试点安排")
    paragraph(doc, "试点到期时间已设为 2026-11-06 23:59:59（Asia/Shanghai）。新 gateway 与 keeper 已启动并读取该配置，原访问入口及既有密钥配置保留。管理员沿用原 account_id，登录凭据已更新，22 个旧会话已撤销；学校账号记录按字节比对未变。")
    paragraph(doc, "现网四账号合成素材流程及隔离四校学生绑定已通过验收。指定资料另测 4 个完整视频，约 1.54 GB、147 分钟：上传、解码、保存、管理员读取及 8 次浏览器播放通过，耗时 13 分 26 秒；5 份 DOCX 可解析。未将真实 IEP 绑定人像，也未验证自动教育评分。")
    paragraph(doc, "现网 t2、t3、t4 尚缺学校元数据。实际接入需校方提供真实学校资料、名册与有效授权，由教师确认人物关联。后续补齐到期访问、报告及保留策略执行，结束时交付用量、复核结果和续约建议；历史占位任务不计成功处理量。")
    paragraph(doc, "状态依据：2026 年 10 月 16 日现场维护与验收反馈；维护前数据库基线位于 outputs/seewo_demo2/rehab_affect_demo2.sqlite3。能力依据：demo2_media_analysis.py、demo2_media_worker.py 及相关存储源码。", small=True)

    page(doc, "三 月度 半年与年度套餐")
    paragraph(doc, "三个套餐采用相同的功能及隐私底线，以合同周期和服务频次区分，不把租户隔离、删除能力或授权控制作为付费升级项。下表是签约结构，不是已发布价目表；M、N、S、H 等额度由校方排课和实测结果共同填入合同。")
    table(doc, ["比较项", "月度试用包", "半年教学包", "年度校级包"], [
        ("周期与场景", "1 个月\n验证单校流程", "6 个月\n连续跟踪一个教学阶段", "12 个月\n持续使用及年度复盘"),
        ("分析分钟", "每月 M 分钟", "每月 M 分钟\n合同名义总额 6M", "每月 M 分钟\n合同名义总额 12M"),
        ("视频任务", "每月 N 个有效任务", "每月 N 个\n名义总额 6N", "每月 N 个\n名义总额 12N"),
        ("同时在库容量", "S GB", "S GB\n不自动乘以 6", "S GB\n不自动乘以 12"),
        ("原视频保留", "每条默认 60 天", "每条默认 60 天", "每条默认 60 天"),
        ("教师及服务", "教师自行复核\n支持 H 小时每月", "同左\n增加阶段复盘约定", "同左\n增加年度汇总约定"),
        ("部署与并发", "约定账号与并发额度", "按容量验收后扩展", "按峰值和学校数量核定"),
        ("收费", "P1 待书面报价", "P6 待书面报价", "P12 待书面报价"),
    ], [2.6, 4.93, 4.93, 4.94])
    heading(doc, "计量和超额处理")
    paragraph(doc, "分钟按成功生成可复核证据的原视频时长计量，建议累加秒数后除以 60。一次分析任务对应一个视频版本；同一视频多机位上传分别计量。平台失败后的自动重试不重复扣学校额度，但其实际计算开销进入内部成本。教师主动重新分析须先展示预计用量并取得确认。")
    paragraph(doc, "视频条数控制任务管理量，分钟控制处理负载，容量控制原片及衍生文件的占用，三项不可互相替代。学校用量面板同时显示已用、剩余与预计到期容量；任一额度将满时提醒，超额默认暂停接收新任务，不自动产生未授权费用。单条时长、文件大小、分辨率和并发上限以兼容性测试为准。")
    heading(doc, "合同需填清的边界")
    paragraph(doc, "需确定 M、N、S、H、账号数、并发数、服务时段、是否含人工报告、数据地域及税费。按月发放额度默认不结转，寒暑假调配或跨月池须另约；半年和年度总额只是名义额度，不代表可在首月全部消耗。长期合同不延长每条视频保留期，也不默认包含硬件和现场驻校。")
    paragraph(doc, "不预设半年或年度折扣，不以 6P1 或 12P1 冒充实际报价。终止时另行约定有效授权内的导出窗口和退款规则，合同期与数据授权期分别管理。")

    page(doc, "四 可核算的成本与报价依据")
    paragraph(doc, "先记录一轮真实试点账单和处理日志，再决定额度及价格。以下公式不填虚构云价；同一地域、服务等级、税费口径下使用厂商有效报价或账单单价。当前本地证据提取链路不调用云模型，不应先加入并不存在的模型推理费用。")
    table(doc, ["成本项", "核算公式与必须采集的量"], [
        ("处理计算", "C计算 = Σ(实例实际占用小时 × 对应每小时单价)。预测小时 = V × r / 60，V 为视频分钟，r 为处理分钟与视频分钟比；并发及失败重试用实际资源小时校准。"),
        ("对象存储", "C存储 = Σ(厂商计费 GB月 × 分层单价) + 读写请求费 + 取回费。GB月按小时或日占用折算；原片、帧、音轨、结果、版本副本及备份分别计入，不能仅按上传原片计算。"),
        ("传输", "C传输 = Σ(实际计费出网 GB × 对应单价) + CDN 或专线费。回看、下载、跨地域复制需分开统计；同一流量不能在两项重复记账。"),
        ("基础服务", "C基础 = 数据库 + 队列 + 应用服务 + 日志监控 + 备份 + 域名证书等实际费用。共享成本按透明的租户占用分摊；不要同时计入整机费与已包含的 CPU 费。"),
        ("人力与硬件", "C人力 = Σ(实施、支持、人工报告小时 × 完全工时成本)。C硬件月 = (购置与安装成本 - 残值) / 使用月数 + 维护；教师校内复核时间单列，不冒充平台支出。"),
    ], [2.5, 14.9])
    heading(doc, "容量和带宽估算")
    paragraph(doc, "原片大小 GB ≈ 视频分钟 × 60 × 平均码率 Mbps / 8 / 1000，采用十进制 GB；厂商若用 GiB，应按字节换算。仅作容量示例：600 分钟、总码率 4 Mbps，原片约 18 GB；每月同量上传、滚动保留 60 天时，稳定期原片约 36 GB，尚未加入帧、音轨和备份。这是规划假设，不是现网用量。")
    paragraph(doc, "上传耗时分钟 ≈ 文件 GB × 8000 / 实测有效上行 Mbps / 60。并发上传以学校出口总带宽核算，不能把标称带宽当作持续吞吐。视频处理时延需另计队列等待及解码时间。")
    heading(doc, "报价与盈亏口径")
    paragraph(doc, "单租户月成本 C = C计算 + C存储 + C传输 + C基础分摊 + C人力 + C硬件月。同一处理资源在云租赁费与本地主机折旧中择一归集，备份存储也只计一次；本地电费按实测功耗、时长及电价另计。若目标毛利率为 g，未税参考价 P = C / (1 - g)，0 ≤ g < 1；这不是已确定售价。半年及年度逐月累计成本，单列一次性上线及约定服务，不虚设风险费或固定折扣。")
    source(doc, receipts, "oss")

    page(doc, "五 从本地内测迁移到多租户服务")
    paragraph(doc, "当前内测使用本地 SQLite、文件目录和单机媒体工作进程。多学校扩展采用下表的 PostgreSQL、对象存储与队列方案，按租户隔离、并发处理和恢复能力逐项验收。云端改造、保留策略及报告授权到期的完整执行仍是后续工作，与本轮已完成的入口到期配置和管理员维护分开验收。")
    table(doc, ["模块", "目标实现", "上线证据"], [
        ("PostgreSQL", "学校 tenant_id 贯穿课程、任务、学生、授权和报告；外键及唯一键带租户约束。应用鉴权配合 RLS，服务角色不得绕过隔离。", "跨租户列表、详情、下载、导出及后台任务均有负向测试。"),
        ("对象存储", "原片、帧、音轨按租户分区，存储桶私有；短时签名链接、校验和、加密及生命周期策略。对象前缀不能替代权限检查。", "越权直链被拒绝；上传完整性、过期访问、衍生文件删除可验证。"),
        ("队列与工作进程", "上传完成后投递任务，幂等键绑定租户、源版本和分析版本；租约、重试退避、死信及取消。持久任务状态可恢复。", "重启、超时、重复投递和取消测试，不重复扣额或覆盖新版本。"),
        ("审计与监控", "记录操作者、租户、对象、动作、时间及结果；追踪授权变化、查看、导出、评分修改与删除。日志不放原始 IEP 正文。", "可查询完整事件链，能告警磁盘、队列积压、失败和权限异常。"),
        ("备份与恢复", "PostgreSQL 定期备份及按需 WAL 恢复，备份加密、分离权限，恢复后重放删除清单。对象版本同样受保留策略约束。", "签约前确定 RPO/RTO 并完成恢复演练；有备份不等于可恢复。"),
    ], [2.4, 9.1, 5.9])
    heading(doc, "迁移与切换顺序")
    paragraph(doc, "先导出脱敏结构和数量摘要，建立 tenant_id 映射，再迁移数据库及对象清单，逐项核对记录数、关系、哈希和授权。短暂停写后增量校验并切换；保留受控回退窗口，避免两个数据库继续独立写入。历史占位任务先隔离标记，不作为正常媒体迁移。迁移后用合成数据做完整验收，再接收真实授权数据。")
    paragraph(doc, "当前继续使用原同源快速隧道入口，不提供 SLA 或全时在线承诺。生产部署需配置受控域名、HTTPS、访问控制与监控，并在合同中约定服务时段及故障响应。")
    source(doc, receipts, "sqlite")
    source(doc, receipts, "pg")
    source(doc, receipts, "backup")
    source(doc, receipts, "tunnel")

    page(doc, "六 授权 数据保留与退出机制")
    paragraph(doc, "真实视频与 IEP 等资料应分项告知处理目的、范围、保存期限、使用者、处理地点及撤回方式。儿童及健康相关资料需要更严格保护，依适用法律取得有效授权；涉及不满十四周岁未成年人时取得监护人同意并制定专门处理规则。学校组织试点不等于已获得训练模型或公开展示的许可。")
    table(doc, ["对象或控制", "当前事实", "试点要求"], [
        ("原视频", "默认保存 60 天，到期前 15 天提醒。", "以明确的入库时间计算期限；第 45 天开始提醒，对校方和项目管理员记录送达及处理结果，到期清理并留删除回执。"),
        ("旧资料入口", "旧敏感资料路由已由前缀守卫退役；历史 student_documents 默认 365 天，维护前文档数 0。", "新资料统一进入加密 dossier 授权流程。保留旧路由拒绝访问的回归检查，历史默认值不作为新资料的保留规则。"),
        ("新加密档案", "dossier 最大授权期限 60 天，维护前基线文档数 0。", "记录资料范围、授权版本、期限和撤回状态；到期或撤回即停止使用。密钥与数据库及备份分离，按角色控制访问。"),
        ("账号与会话", "入口到期已配置，管理员 22 个旧会话已撤销；普通会话 8 小时，学校账号记录未变。", "入口到期与账号生命周期分别管理；继续核验到期上传、访问及导出拦截，明确存量会话和下载授权的处理方式。"),
        ("结构化结果", "报告保留及授权到期执行仍有待补齐项。", "为报告、评分、人物关联和审计设定必要期限，落实到期访问控制及清理；覆盖衍生文件、导出缓存与备份恢复，并留验收记录。"),
    ], [2.6, 5.0, 9.8])
    heading(doc, "IEP 使用范围与人员权限")
    paragraph(doc, "任课教师只访问被授权班级及学生，校方管理员管理本校授权和成员，项目运维默认不读取 IEP 内容。临时排障访问需审批、限时并留痕。IEP 只用于授权的教学复核和目标对照；向外部模型发送资料、跨校共享、训练或二次研究须另行审查，不能沿用一般上传同意。")
    heading(doc, "提醒与删除的可验证规则")
    paragraph(doc, "60 天是上限口径，不是必须保留满 60 天。撤回或更短授权应提前终止；提前 15 天提醒应按各对象实际到期日计算，短于 15 天的授权在生效时提示。不得通过重新上传、重新分析或恢复备份悄悄重置期限。若依法需要留存特定记录，应限定范围、隔离访问并告知依据。")
    paragraph(doc, "退出时先核验导出者权限和授权状态，再导出允许范围，生成删除清单与处理回执。备份可按约定轮换窗口失效，但必须披露窗口、禁止业务恢复已删除资料，并在恢复时重放删除标记；不能把前台看不到当作全部物理副本已删除。")
    source(doc, receipts, "law")

    page(doc, "七 硬件兼容性与验收安排")
    paragraph(doc, "优先复用现有摄像设备和教学电脑，先用不含真实学生的样片验证。最低要求应由拍摄质量、上传网络和处理负载决定，不绑定未验证的设备品牌，也不把机器人、独立显卡或专用摄像头列为开展基础试点的前提。")
    table(doc, ["类别", "必需或可选", "兼容性检查"], [
        ("拍摄与收音", "能导出视频的现有手机或摄像机必需；涉及声音证据时需可用音轨。外置麦克风和支架可选。", "建议先测 MP4 容器及 H.264/AAC 编码；检查旋转、音画同步、遮挡、逆光和远距离收音。兼容性以完整解码实测为准。"),
        ("教师访问端", "可用浏览器的电脑及网络必需；大屏、平板为可选访问端。", "实测学校使用的 Edge 或 Chrome 版本，验证分片上传、断线重试、视频定位及报告导出。手机浏览器与设备锁屏上传需单测。"),
        ("本地处理主机", "本地部署必需；云端部署可由服务商资源代替。", "FFmpeg、ffprobe、OpenCV 可运行且有足够磁盘空间；CPU、内存和临时盘容量由最长视频及并发压测决定。当前证据提取不强制 GPU。"),
        ("校网与存储", "稳定上行链路、可控存储空间必需；独立 NAS、UPS、备用网络可选。", "校方确认外网及隧道政策，测实用上行带宽；禁止直接暴露数据库与文件目录。备份与隐私控制必需，但不一定购买独立硬件。"),
        ("增强设备", "多机位、深度相机、机器人均可选。", "各自增加同步、校准、视频分钟和容量开销。未完成接口与安全测试前，不承诺即插即用或自动课堂干预。"),
    ], [2.5, 6.0, 8.9])
    heading(doc, "截至 11 月 6 日的建议节奏")
    paragraph(doc, "10 月 16 至 20 日：已完成入口到期配置、管理员维护、旧敏感路由退役和上述两组流程验收；继续补齐 t2、t3、t4 真实学校元数据。10 月 21 至 30 日：校方提供名册与授权，教师确认人物关联后开展小范围复核，记录时延、成本及耗时。10 月 31 日至 11 月 6 日：核验到期访问、保留与报告策略，完成删除恢复演练和续约评估。")
    heading(doc, "验收交付与决策条件")
    paragraph(doc, "流程验收记录上传成功率、证据生成成功率、端到端 P50/P95 时延、教师复核分钟及成本每视频分钟。维护前 28 个任务只作基线；新增合成任务、真实试点任务分组统计，成功率注明分母，失败、取消和历史占位分别列示。采样前由校方与项目组确认阈值、范围和责任人。")
    paragraph(doc, "教学试点由两名教师独立复核授权样本，记录评分依据、分歧和一致性。自动教育评分尚未实现；后续若引入模型，再单独验证识别及评分效果。生产上线须完成跨租户访问、授权撤回、到期拦截和删除恢复验收，云端改造与保留、报告执行缺口另列交付清单。")
    paragraph(doc, "签约前确认：额度与报价、地域与服务时段、IEP 授权范围、60 天及提前 15 天提醒、试点截止和账号控制、退出导出、RPO/RTO、硬件清单。官方资料访问记录见各节，HTTP 成功仅代表页面取得；动态价格或适用版本仍以签约当日官方报价和条款为准。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(str(OUTPUT), flush=True)
    print("sha256=" + hashlib.sha256(OUTPUT.read_bytes()).hexdigest(), flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sources-only", action="store_true")
    parser.add_argument("--cached-sources", action="store_true")
    args = parser.parse_args()
    receipts = (json.loads((QA / "sources.json").read_text(encoding="utf-8"))
                if args.cached_sources else fetch_sources())
    if not args.sources_only:
        build(receipts)


if __name__ == "__main__":
    main()
