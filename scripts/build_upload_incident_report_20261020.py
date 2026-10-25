"""Build the evidence-grounded Chinese upload incident report."""
from pathlib import Path
import json

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "课堂分析平台_上传问题修复与提速方案_20261020.docx"
RUNTIME = ROOT / ".runtime"


def load(relative):
    return json.loads((RUNTIME / relative).read_text(encoding="utf-8-sig"))


def font(run, size=11, bold=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")


def paragraph(doc, text, bold_prefix=None, space=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    if bold_prefix and text.startswith(bold_prefix):
        font(p.add_run(bold_prefix), bold=True)
        font(p.add_run(text[len(bold_prefix):]))
    else:
        font(p.add_run(text))
    return p


def heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f"Heading {level}")


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for key, value in {"val": "single", "sz": "5", "color": "D9D9D9"}.items():
            el.set(qn(f"w:{key}"), value)
        borders.append(el)
    pr.append(borders)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for values in [headers, *rows]:
        row = t.rows[0] if values is headers else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if values is headers:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for i, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_pr = cell._tc.get_or_add_tcPr()
            margin = OxmlElement("w:tcMar")
            for edge, value in (("top", 80), ("bottom", 80), ("left", 100), ("right", 100)):
                el = OxmlElement(f"w:{edge}")
                el.set(qn("w:w"), str(value))
                el.set(qn("w:type"), "dxa")
                margin.append(el)
            cell_pr.append(margin)
            if values is headers:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EAF0F5")
                cell_pr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(15)
            if values is headers:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(text), 10.5, values is headers)
    return t


def main():
    large = load("upload_reliability_20261020/run_20261020_143137_a3d70ce4/report.json")
    real = load("upload_reliability_20261020/run_20261020_143918_9f19ee5e/report.json")
    browser = load("full_entry_acceptance_20261020_deployed.json")
    assert large["status"] == real["status"] == "passed"
    assert browser["passed"] and len(browser["accounts"]) == 4

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(0.63), Inches(0.63)
    sec.left_margin, sec.right_margin = Inches(0.72), Inches(0.72)
    sec.header_distance, sec.footer_distance = Inches(0.28), Inches(0.28)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Header", "Footer"):
        s = doc.styles[name]
        s.font.name = "Arial"
        s.font.color.rgb = RGBColor(0, 0, 0)
        s._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        s.paragraph_format.widow_control = True
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    body = doc.styles["Normal"]
    body.font.size = Pt(11)
    body.paragraph_format.line_spacing = Pt(16)
    body.paragraph_format.space_after = Pt(5)
    for name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
        s = doc.styles[name]
        s.font.size = Pt(size)
        s.font.bold = True
        s.paragraph_format.line_spacing = Pt(20 if name == "Heading 1" else 17)
        s.paragraph_format.space_before = Pt(9)
        s.paragraph_format.space_after = Pt(5)
        s.paragraph_format.keep_with_next = True
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.line_spacing = Pt(28)
    doc.styles["Title"].paragraph_format.space_after = Pt(5)
    hp = sec.header.paragraphs[0]
    hp.paragraph_format.line_spacing = Pt(12)
    font(hp.add_run("辅读学校课堂分析平台 Demo2"), 9)
    fp = sec.footer.paragraphs[0]
    fp.paragraph_format.line_spacing = Pt(12)
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(fp.add_run("2026年10月20日  |  "), 9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)

    doc.add_paragraph("视频上传问题处理与提速方案", style="Title")
    paragraph(doc, "处理日期  2026年10月20日    适用版本  Demo2 当前学校测试站", space=7)
    paragraph(doc, "本次针对学校反馈的视频上传缓慢、中途停滞问题进行了排查、修复和实际验收。上传可靠性修复已部署，原完整邀请链接、四个学校账号及已有业务数据保持不变。当前可以继续开展受控测试，但临时公网通道仍偏慢，后续需通过云端直传改善速度。")

    heading(doc, "一 问题情况与排查结论")
    paragraph(doc, "学校反馈约670 MB的视频上传约半小时后，进度仍停留在14%至20%左右，随后取消任务。排查发现，一个未取消任务实际已保存约92 MiB，说明并非完全没有接收到视频，而是在传输过程中发生了中断或停滞。已取消任务不能承诺恢复原进度。")
    table(doc, ["排查发现", "对使用体验的影响"], [
        ("原上传方式把视频分片转为文本", "Base64编码使传输数据量增加约三分之一，同一视频需要发送更多数据。"),
        ("上传网关的等待时间偏短", "原请求体总等待约30秒，低速或波动网络下容易超时；前端缺少有效的自动重试与续传。"),
        ("临时公网通道存在明显速度瓶颈", "同一16 MiB基准上传，旧协议公网约60.17秒、隔离本地约0.92秒，表明公网路径是重要瓶颈，不能仅归因于学校网络。"),
        ("周边流程存在可用性问题", "录像保存页与后端分片规格不一致；学校信息不完整时仍可开始上传，容易在后续归属与档案流程中卡住。"),
    ], [2.0, 5.06])

    heading(doc, "二 已完成的修复")
    for lead, text in [
        ("传输方式：", "改为直接发送二进制视频分片，取消Base64额外开销；保留旧接口兼容。"),
        ("断网恢复：", "增加超时控制、自动重试、暂停和断点续传；重新选择原文件后校验已收内容，继续同一任务，避免重复占用名额。"),
        ("保存一致性：", "重复提交同一分片不会重复写入；增加写盘校验、空间预留和存储不足提示，降低传完才发现文件损坏或磁盘不足的风险。"),
        ("进度与错误提示：", "按服务器已确认的数据量显示进度、速度和预计时间，明确区分上传完成与后台处理完成；失败后给出可继续操作的提示。"),
        ("流程补齐：", "修复录像保存和身份写入后的异常返回；新建上传前校验学校信息；同步更新网关状态监测，避免服务正常却显示故障。"),
    ]:
        paragraph(doc, lead + text, bold_prefix=lead, space=4)
    paragraph(doc, "文件上传限制：单文件不超过1 GiB、单节课不超过60分钟；支持MP4、MOV、WebM，推荐MP4（H.264视频与AAC音频）、720p或1080p。不建议上传4K原始大文件。", bold_prefix="文件上传限制：")

    doc.add_page_break()
    heading(doc, "三 本次实际测试与结果")
    paragraph(doc, "本次分别验证公网传输、真实课堂视频处理、浏览器操作与数据隔离。公网测试使用明确标记的合成视频；真实课堂文件在隔离本地环境中处理，未发送到新的第三方模型接口。MiB与GiB为二进制容量单位，与部分设备显示的MB、GB略有差别。")
    table(doc, ["测试项目", "测试方法与范围", "实际结果"], [
        ("大文件公网上传", "670 MiB（约703 MB）合成MP4，经当前公网入口传输；包含中断、响应丢失及恢复验证。", "上传30分42秒完成；全部168个分片收到。"),
        ("大文件完整性与处理", "完整下载上述文件，比较SHA-256校验值，并核对处理结果对应的源文件。", "上传、保存、下载校验一致；后台完成基础证据提取。"),
        ("真实课堂视频", "用户提供的约528 MiB、37分38秒课堂视频，在隔离本地环境验证上传、保存、下载和完整解码。", "基础处理约4分18秒；解码67,726帧、采样2,258帧；检测到音轨。"),
        ("浏览器故障恢复", "公网网页上传16 MiB合成文件；模拟请求失败、限速、暂停、刷新，再选择原文件续传。", "自动重试成功；同一任务续传，文件一致、无重复任务。"),
        ("录像保存", "向录像页面注入合成视频文件，模拟首次请求失败，验证保存与重试。", "4 MiB分片保存成功；无重复任务。本项不等于真实摄像头采集已验收。"),
        ("四账号与公司后台", "四个学校账号登录与隔离检查；合计30次学校页面访问、4个后台子页面及4路原视频下载。", "页面可打开、后台可查询下载；跨通道与未授权访问被拒绝。"),
        ("结果页面与报告", "新上传任务的实际播放、群体与个体页、复核页、报告页，以及JSON和HTML报告下载、刷新。", "结果与任务对应；刷新后保留；未发现浏览器脚本错误。"),
        ("学生档案与IEP关联", "隔离环境下，四通道分别使用合成学校、学生、画面绑定、授权资料、测量与报告数据。", "关联与授权资料读写通过；24次跨归属资料访问被拒绝。"),
        ("异常情况与回归", "校验无音轨、损坏文件、上传限制、学校前置条件、网关和存储边界等。", "专项回归505项通过、1项跳过、0项失败；损坏视频不生成正常结果。"),
    ], [1.32, 3.48, 2.26])
    heading(doc, "测试结果应如何理解", 2)
    paragraph(doc, "670 MiB文件由6秒有效视频加标准容器填充构成，验证的是大文件传输与保存，不能代替长课分析性能测试。37分38秒的真实课堂测试补充验证了解码与基础证据提取，但未验证学生身份自动识别率、教学评分准确率或临床判断。")
    paragraph(doc, "四账号的基本访问及隔离已在线验证；完整扩展页面集中验收了当前学校信息完整的账号，其余账号仍需先补全学校信息。真实学生档案未在公网逐人验收；隔离测试也记录到人工评课记录尚未完全结构化关联到具体任务和学生，应作为后续补齐项。")
    paragraph(doc, "唯一跳过项为Windows当前账号无创建符号链接权限。以上计数为本轮专项回归，不代表整个仓库的全部测试；工程流程通过也不等于教育评价有效性已完成验证。", space=0)

    doc.add_page_break()
    heading(doc, "四 后续提速方案与实施顺序")
    paragraph(doc, "本次670 MiB公网文件实测平均约0.364 MiB/秒，约3.05 Mbps。上传已能够完整完成，但仍需约半小时。因此，本次整改主要解决稳定性和可恢复性，不能承诺已经实现高速上传，也不能据此保证其他学校网络下的速度。")
    table(doc, ["方案", "具体做法", "作用与限制"], [
        ("当前立即使用", "保留现有入口，使用二进制分片、自动重试与续传；规范单节课、文件大小和编码格式。", "无需新增云资源，可继续受控测试；仍受学校网络、临时通道及本机网络限制。"),
        ("优先实施云存储直传", "浏览器取得短时、限定文件的上传凭证，直接向境内对象存储分片上传；后端确认完成并提交分析任务。", "大文件不再经过临时隧道和本机中转，通常更适合多校上传；实际速度仍须在学校实测。"),
        ("正式云端部署", "应用、数据库、分析队列和处理节点迁至同区域云资源，配置固定域名、HTTPS、监控、备份和告警。", "降低本机休眠、断网及临时域名变化的影响；需确认资源预算、运维责任和资料授权。"),
    ], [1.26, 3.55, 2.25])
    heading(doc, "建议的上传与处理流程", 2)
    paragraph(doc, "学校登录并创建课程任务 → 系统签发专属上传凭证 → 浏览器直传对象存储 → 后端核验文件大小、类型、归属及完整性 → 分析队列处理 → 报告关联课程与学生 → 学校及获授权的公司管理员查询和下载。")
    paragraph(doc, "实施时保留现有四账号、课程编号和学生编号，先让新任务走云端直传；旧任务继续可查，逐批核对后再迁移，避免换入口后形成丢失历史资料的新空站。上传凭证不得授予整个存储空间权限，视频与报告下载也必须经过权限校验。")
    heading(doc, "提速目标如何验证", 2)
    paragraph(doc, "以本次670 MiB文件计算，若有效上传带宽分别达到10、20、50 Mbps，纯数据传输理论耗时约为9分22秒、4分41秒、1分52秒。这是按文件字节数÷有效速度计算的参考值，不是云服务商承诺，也不含握手、重试、校验和后台分析时间。")
    paragraph(doc, "上线前应在学校实际网络分别测试100 MB、约500 MB和接近1 GiB的文件，比较单路与四路同时上传；记录实际速度、上传耗时、处理耗时和失败重试情况。还需模拟断网续传、凭证过期、文件超限、磁盘或配额不足、处理服务重启，并逐份核对下载文件与原文件一致。")
    heading(doc, "配套条件与下一步安排", 2)
    paragraph(doc, "第一步先开展学校小批量复测，记录文件大小、时长、出错时间和截图，确认本次修复在学校网络中有效；未取消的旧任务可选择原文件续传，已取消的任务重新上传。第二步确认云账号、部署地域、域名和预算，优先建设云存储直传。第三步补齐四校真实资料关联验收、人工评课记录的任务与学生关联，再扩大学校数量。")
    paragraph(doc, "云端建设需分别核算对象存储、下载流量、请求、应用服务器、分析算力、数据库及备份费用，依据试测量与官方报价确定，本报告不提供未经调研的价格。视频仍遵守现有60天保留及到期前15天提醒规则；IEP、诊断资料及研究使用另按授权处理，不因公司后台可访问就默认永久留存。")
    paragraph(doc, "当前状态：修复已上线，原邀请与账号保持不变，可继续受控测试。云端直传、正式云部署与上述扩展验收尚未完成，不应对外表述为已经上线。", bold_prefix="当前状态：", space=0)

    doc.core_properties.title = "视频上传问题处理与提速方案"
    doc.core_properties.subject = "Demo2上传故障排查、修复、测试及后续建设"
    doc.core_properties.author = "项目研发组"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
