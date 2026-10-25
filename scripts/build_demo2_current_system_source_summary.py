from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_seewo_source_code_docx import (
    BLUE,
    BORDER,
    DARK_BLUE,
    INK,
    LIGHT_BLUE,
    MUTED,
    _add_bullet,
    _add_callout,
    _add_field,
    _add_numbered,
    _add_paragraph,
    _configure_styles,
    _language_for,
    _set_cell_shading,
    _set_cell_text,
    _set_keep_with_next,
    _set_run_font,
    _set_table_borders,
    _set_table_geometry,
    _set_widow_control,
    _update_fields_on_open,
    _xml_safe,
)


ROOT = Path(__file__).resolve().parents[1]
PACKAGE_ROOT = (
    ROOT
    / "deliverables"
    / "20260827_demo2_delivery"
    / "RehabAffect_Demo2_交付包_20260827"
    / "02_源程序"
)
SCREENSHOT_ROOT = ROOT / "output" / "playwright" / "demo2_v15_premium_thousand"
OUTPUT_PATH = ROOT / "docs" / "辅读学校Demo2_系统汇总与源码说明_20260831.docx"


@dataclass(frozen=True)
class SourceFile:
    path: Path
    relative: str
    category: str
    language: str
    lines: int


def _set_current_header_footer(section, *, appendix: bool = False) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.tab_stops.clear_all()
    label = "Demo2 完整源码附录" if appendix else "辅读学校 Demo2 课堂学习支持与追踪系统"
    run = paragraph.add_run(label)
    _set_run_font(run, size=8.5, bold=True, color=MUTED)
    run = paragraph.add_run("\t内部受限资料")
    _set_run_font(run, size=8.5, color=MUTED)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        section.page_width - section.left_margin - section.right_margin,
        WD_TAB_ALIGNMENT.RIGHT,
    )

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run("第 ")
    _set_run_font(run, size=8.5, color=MUTED)
    _add_field(paragraph, "PAGE")
    run = paragraph.add_run(" 页")
    _set_run_font(run, size=8.5, color=MUTED)


def _repeat_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    row_properties.append(node)


def _collect_source_files() -> list[SourceFile]:
    groups = [
        ("核心源代码", PACKAGE_ROOT / "src", {".py"}),
        ("自动测试", PACKAGE_ROOT / "tests", {".py"}),
        ("配置与模型", PACKAGE_ROOT / "configs", {".json"}),
        ("配置与模型", PACKAGE_ROOT / "models", {".json"}),
        ("运行与交付脚本", PACKAGE_ROOT / "scripts", {".py", ".js", ".ps1"}),
    ]
    files: list[SourceFile] = []
    for category, base, extensions in groups:
        if not base.exists():
            continue
        for path in sorted(item for item in base.rglob("*") if item.is_file()):
            if path.suffix.lower() not in extensions or "__pycache__" in path.parts:
                continue
            text = path.read_text(encoding="utf-8", errors="replace")
            files.append(
                SourceFile(
                    path=path,
                    relative=path.relative_to(PACKAGE_ROOT).as_posix(),
                    category=category,
                    language=_language_for(path),
                    lines=len(text.splitlines()),
                )
            )

    for name in [
        "pyproject.toml",
        "requirements.txt",
        "install_environment.cmd",
        "start_school_demo2.cmd",
        "verify_demo.cmd",
    ]:
        path = PACKAGE_ROOT / name
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        files.append(
            SourceFile(
                path=path,
                relative=name,
                category="环境与启动文件",
                language=_language_for(path),
                lines=len(text.splitlines()),
            )
        )
    return sorted(files, key=lambda item: (item.category, item.relative))


def _add_cover(document: Document, files: list[SourceFile]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(24)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run("内部技术交付资料")
    _set_run_font(run, size=11, bold=True, color=BLUE)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run("辅读学校 Demo2 课堂学习支持与追踪系统")
    _set_run_font(run, size=23, bold=True, color=INK)

    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run("系统汇总与源码说明")
    _set_run_font(run, size=17, bold=True, color=DARK_BLUE)

    metadata = [
        ("文档用途", "系统现状汇总、功能边界、运行条件与完整源码归档"),
        ("系统版本", "RehabAffect Demo2 / 2026-08-27 交付版本"),
        ("汇编范围", f"{len(files)} 个代码与配置文件，约 {sum(item.lines for item in files):,} 行"),
        ("生成日期", str(date(2026, 8, 31))),
        ("资料级别", "内部受限，不含原始课堂视频、学生档案原件和历史上传文件"),
    ]
    table = document.add_table(rows=len(metadata), cols=2)
    for index, (label, value) in enumerate(metadata):
        _set_cell_text(table.cell(index, 0), label, bold=True, color=DARK_BLUE)
        _set_cell_text(table.cell(index, 1), value)
        if index % 2 == 0:
            _set_cell_shading(table.cell(index, 0), LIGHT_BLUE)
            _set_cell_shading(table.cell(index, 1), "FAFBFC")
    _set_table_geometry(table, [2700, 6660])
    _set_table_borders(table)

    document.add_paragraph()
    _add_callout(
        document,
        "阅读说明",
        "正文汇总项目定位、处理闭环、页面功能、数据治理、测试边界、下一步整合方向和实体落地条件；附录收录本次交付包的源码与配置，用于技术审查、部署复现和后续交接。",
    )


def _add_toc(document: Document) -> None:
    document.add_heading("目录", level=1)
    paragraph = document.add_paragraph()
    _add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u')
    paragraph.paragraph_format.space_after = Pt(8)
    _add_paragraph(document, "打开 Word 后可按 Ctrl+A、F9 更新目录和页码。", size=9.5, color=MUTED)


def _add_project_summary(document: Document) -> None:
    document.add_heading("一、项目定位与当前系统能力", level=1)
    _add_callout(
        document,
        "系统定位",
        "面向辅读学校课堂的多模态学生状态感知、学习证据整理与支持决策系统。系统将课堂视频、声音、课程标准、师生互动、学生档案和教师复核组织成可追溯证据，并输出个体学习判断、低刺激支持建议、课程报告及后续设备可调用的结构化动作。",
    )

    document.add_heading("1.1 需要解决的课堂问题", level=2)
    for text in [
        "一节课包含多个学生和大量互动，教师难以课后逐段回看每名学生的状态与学习机会。",
        "课堂整体评课无法直接回答某个学生在什么任务上独立完成、在提示后完成或没有获得作答机会。",
        "特殊教育学生的动作、眼神、语言和回应节奏差异明显，不能用单一表情或统一阈值直接判断能力。",
        "分析结果需要回到课程目标和下一步支持，不能停留在动作罗列或笼统的专注度分数。",
        "学生资料、课堂视频、教师标注和长期趋势必须受控保存，并保留教师最终确认权。",
    ]:
        _add_bullet(document, text)

    document.add_heading("1.2 当前已经实现的功能闭环", level=2)
    for text in [
        "课程开始前记录数据安全协议，并确认课程名称、课程标准、校方考核标准、责任人和教学目标。",
        "在同一课堂画面中展示三名学生群像，支持切换到一人一案、历史记录和个体课堂证据。",
        "支持 IEP、五大领域评估和专业资料导入，原文件、解析预览与教师确认摘要分开保存。",
        "展示状态时间线、师生互动、关键片段教师复核和支持前后效果，不以单帧结果替代课堂判断。",
        "依据课程目标和独立完成情况区分独立掌握、支持下掌握、正在学习与证据不足。",
        "形成课程分析、五域趋势和约一千字分章节报告，并保留证据来源、限制条件和下一课建议。",
        "提供四个主体独立视频上传端口、本地 SQLite、60 天保存周期、提前 15 天提醒和授权归档队列。",
        "输出 human_state、风险等级、推荐支持动作和结构化 JSON，为 App、WebSocket、ROS2 或机器人端保留接口。",
    ]:
        _add_bullet(document, text)

    document.add_heading("1.3 当前边界", level=2)
    for text in [
        "当前交付版本适合内部展示、校方受控演示和需求确认，尚未达到全校正式自动评估投放状态。",
        "现有课堂结果包含缓存样例和可解释规则输出；对任意新视频生成可靠结论仍需部署并验证视觉、音频和时序模型。",
        "尚未直连 Seewo 官方后台，当前按校方授权后的人工导出、分端上传和本地导入流程运行。",
        "尚未形成具有统计意义的准确率、召回率和跨学生泛化结论，页面分数不应作为学生定级或教师绩效依据。",
        "机器人动作目前是结构化建议与接口协议，尚未完成真实机器人安全联调。",
        "系统不进行医学诊断，不根据病名自动推断能力，不用一次课堂表现定义学生。",
    ]:
        _add_bullet(document, text)


def _add_system_flow(document: Document) -> None:
    document.add_heading("二、系统处理闭环与代码组成", level=1)
    document.add_heading("2.1 从课程确认到下一课支持", level=2)
    for text in [
        "授权与课程确认：记录数据用途、保存期限、课程标准、考核口径和责任人。",
        "数据输入：接收四路课堂视频、随视频音轨、评课材料、学生档案和教师已有观察。",
        "证据整理：进行视频质量检查、时间切片、身份/座位绑定、互动机会和声音事件对齐。",
        "状态估计：综合动作、姿态、回应延迟、语音线索、任务表现、个人基线和历史状态。",
        "学习判断：围绕每个课程目标区分独立完成、提示后完成、未完成和证据不足。",
        "教师复核：教师确认、修正或标记无法判断，补充提示等级与课堂背景。",
        "支持决策：生成等待、重复说明、降低难度、暂停询问、保持距离或通知教师等建议。",
        "报告与追踪：保存单节课报告、结构化 JSON、五域趋势和下一课观察重点。",
    ]:
        _add_numbered(document, text)

    document.add_heading("2.2 主要代码模块", level=2)
    rows = [
        ("页面与接口服务", "scenario_lab_server.py", "提供 Demo2 页面、静态图片、报告、资料、上传、存储与 JSON API"),
        ("三学生课堂分析", "seewo_multi_student_demo.py", "构建群像、个体状态、互动、掌握、支持效果和完整前端"),
        ("千字课程报告", "seewo_long_report.py", "按固定教育证据结构生成约一千字报告，并保留教师复核"),
        ("学生档案与本地库", "student_record_store.py", "管理 SQLite、资料导入、保存周期、审计和四路视频任务"),
        ("感知与状态主流程", "pipeline.py / fusion.py / smoothing.py", "融合多模态证据并进行时间平滑和置信度处理"),
        ("动作与安全策略", "policy.py / closed_loop.py", "把状态转换为低刺激动作并验证动作反馈与安全规则"),
        ("工程接口", "websocket_server.py / ros2_mock_publisher.py", "将 human_state 和 robot_action 发送给外部控制端"),
        ("自动验证", "tests/ / qa_demo2_ui.js", "验证业务逻辑、页面交互、多端适配和关键接口"),
    ]
    table = document.add_table(rows=1, cols=3)
    for index, label in enumerate(["系统部分", "核心文件", "职责说明"]):
        _set_cell_text(table.cell(0, index), label, bold=True, color=DARK_BLUE)
        _set_cell_shading(table.cell(0, index), LIGHT_BLUE)
    _repeat_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value, bold=index == 0)
        if row_index % 2 == 0:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, [1900, 3200, 4260])
    _set_table_borders(table)


def _add_figure(document: Document, filename: str, caption: str, explanation: str) -> None:
    image = SCREENSHOT_ROOT / filename
    if not image.exists():
        raise FileNotFoundError(image)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image), width=Inches(6.45))
    _add_paragraph(
        document,
        caption,
        size=9.3,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=5,
    )
    _add_paragraph(document, explanation, after=10)


def _add_pages(document: Document) -> None:
    document.add_heading("三、当前页面与功能说明", level=1)
    _add_figure(
        document,
        "00_entry_viewport_desktop.png",
        "图 1  数据安全协议与课程进入流程",
        "上传和分析前先明确数据用途、保存期限、下载范围和责任主体。该步骤不是普通弹窗，而是后续审计、到期处理和报告使用边界的来源。",
    )
    _add_figure(
        document,
        "03_workspace_desktop.png",
        "图 2  班级群像、三学生切换与课堂证据工作区",
        "课堂总览保留原画面、群像/个体切换、三名学生指标、课程适配和教师复核进度。切换学生后，左侧档案和后续分析页面同步切换，避免把不同学生证据混在一起。",
    )
    _add_figure(
        document,
        "06_teacher_review_desktop.png",
        "图 3  教师关键片段复核",
        "系统只推荐值得复核的时间窗口，教师可选择确认、需要修正或无法判断，并补充提示等级和备注。复核结果进入后续报告和模型校准，最终解释权保留给教师。",
    )
    _add_figure(
        document,
        "08_mastery_desktop.png",
        "图 4  课程目标掌握判断",
        "掌握判断围绕本节课的目标、学生获得的作答机会、独立完成情况和提示层级组织，不把安静、动作少或单一注意力指数直接等同于学会。",
    )
    _add_figure(
        document,
        "14_deep_report_generated_desktop.png",
        "图 5  约一千字课程深度报告",
        "报告按课程边界、班级证据、个体学习判断、教师支持效果、下一课安排和使用方式分章节生成。当前本地模板无需外部大模型即可运行；后续 Agent 只能在保留来源引用和教师确认的条件下接入。",
    )
    _add_figure(
        document,
        "12a_four_channel_upload_desktop.png",
        "图 6  四个主体独立视频上传",
        "四个通道分别记录上传人、课程、观察对象和批次号，采用分片上传并返回独立任务编号。未经授权的任务不进行跨通道合并，防止不同主体的数据混淆。",
    )
    _add_figure(
        document,
        "12b_storage_governance_desktop.png",
        "图 7  本地数据库、保存周期与归档队列",
        "结构化元数据保存在本地 SQLite，原文件保存在受控目录。原始视频默认保存 60 天并提前 15 天提醒；脱敏结构化结果只有获得单独授权才进入长期资源队列。",
    )


def _add_data_governance(document: Document) -> None:
    document.add_heading("四、输入输出、数据存储与权限", level=1)
    document.add_heading("4.1 输入与输出", level=2)
    rows = [
        ("课程信息", "课程名称、学科、课程标准、校方考核标准、责任人", "课程确认记录、分析口径"),
        ("课堂视频", "四个主体上传的视频及随视频音轨", "质量预检、时间切片、课堂证据"),
        ("学生资料", "IEP、五大领域评估、专业评估和教师摘要", "个体基线、优势、支持需要与安全边界"),
        ("教师复核", "关键切片标签、提示等级、备注和无法判断项", "可追溯修正、报告依据和训练标签"),
        ("分析结果", "状态、置信度、互动、掌握、支持效果和风险", "页面、千字报告、JSON、下载文件和动作建议"),
    ]
    table = document.add_table(rows=1, cols=3)
    for index, label in enumerate(["数据类别", "主要输入", "系统输出"]):
        _set_cell_text(table.cell(0, index), label, bold=True, color=DARK_BLUE)
        _set_cell_shading(table.cell(0, index), LIGHT_BLUE)
    _repeat_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value, bold=index == 0)
        if row_index % 2 == 0:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, [1600, 3880, 3880])
    _set_table_borders(table)

    document.add_heading("4.2 当前存储设计", level=2)
    for text in [
        "本地 SQLite 保存学生、课程、资料元数据、教师复核、保存策略、上传任务和审计记录。",
        "原始资料和视频文件保存在受控文件目录，不把大文件直接塞入数据库。",
        "原始视频默认保存 60 天，到期前 15 天提醒，可下载、延期或按策略删除。",
        "脱敏结构化结果进入长期资源库前需要单独授权；原始学生视频默认不进入公司长期数据库。",
        "正式部署需要按学校、班级、角色和学生划分权限，并记录查看、下载、修改、归档和删除操作。",
    ]:
        _add_bullet(document, text)

    document.add_heading("4.3 安全与教育边界", level=2)
    for text in [
        "不输出医学诊断，不根据病理名称自动判断学生能力或课程难度。",
        "不跨学生排名，不用单节课堂分数评价教师个人绩效。",
        "状态和掌握结论必须显示证据、不确定性与教师复核状态。",
        "对疼痛风险、抗拒和明显不适，优先停止动作、保持距离并交由教师处理。",
        "机器人或互动设备不得主动触碰学生，动作、声音和光效需要白名单、限幅和一键关闭。",
    ]:
        _add_bullet(document, text)


def _add_status(document: Document) -> None:
    document.add_heading("五、开发状态、测试结果与准确性边界", level=1)
    rows = [
        ("安全协议与课程确认", "已实现", "进入分析前记录用途、标准、责任人和审计信息"),
        ("三学生群像与个体切换", "已实现", "后续分析页面按所选学生同步切换"),
        ("学生资料导入", "已实现原型", "支持原文件、预览与教师确认摘要分开保存"),
        ("状态时间线与互动", "已实现原型", "当前样例可展示，新视频模型仍需验证"),
        ("教师复核与支持效果", "已实现", "支持确认、修正、无法判断和提示等级"),
        ("课程目标掌握判断", "已实现原型", "需要更多教师复核样本校准规则和阈值"),
        ("课程分析与千字报告", "已实现", "本地模板可生成约一千字分章节报告"),
        ("四路视频上传", "已实现基础设施", "分片上传、任务隔离、质量预检和结果下载"),
        ("本地数据库与保存周期", "已实现", "SQLite、60 天/15 天策略和归档队列"),
        ("结构化动作接口", "已实现", "JSON、WebSocket 和 ROS2 mock 可继续对接"),
        ("Seewo 自动同步", "未实现", "当前依赖校方授权后人工导出"),
        ("任意新视频自动分析", "待模型部署", "需要身份跟踪、动作、音频和时序融合模型"),
        ("真实机器人联调", "未实现", "当前仅提供动作协议、白名单策略和仿真基础"),
        ("全校生产投放", "条件未满足", "需要权限、模型评估、运维和正式试点"),
    ]
    table = document.add_table(rows=1, cols=3)
    for index, label in enumerate(["项目项", "当前状态", "说明"]):
        _set_cell_text(table.cell(0, index), label, bold=True, color=DARK_BLUE)
        _set_cell_shading(table.cell(0, index), LIGHT_BLUE)
    _repeat_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            color = INK
            if index == 1 and value.startswith("已实现"):
                color = "16783A"
            elif index == 1:
                color = "B54708"
            _set_cell_text(cells[index], value, bold=index < 2, color=color)
        if row_index % 2 == 0:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, [2600, 1800, 4960])
    _set_table_borders(table)

    document.add_heading("5.1 已完成的软件验证", level=2)
    for text in [
        "Python 自动测试：43 项通过，5 项因可选环境跳过。",
        "浏览器检查覆盖 1440 桌面、1366 学校笔记本、1024/768 平板和 390/360 手机宽度。",
        "检查结果为 0 横向溢出、0 破图、0 未命名控件和 0 控制台错误。",
        "交付压缩包已在独立目录解压并启动，Demo 页面和课堂图片资源返回 HTTP 200。",
        "交付包不包含原始课堂视频、学生档案原件、历史上传文件或本机绝对路径。",
    ]:
        _add_bullet(document, text)

    document.add_heading("5.2 准确性状态", level=2)
    _add_callout(
        document,
        "当前结论",
        "现阶段可以证明软件流程、页面和数据闭环可运行，但不能据此宣称学生状态识别或学习掌握判断已经达到可投放精度。准确性必须在教师复核标注的数据集上，按学生隔离和课程隔离方式评估，并报告状态 F1、关键事件召回率、掌握判断一致率、校准误差、误报率和处理延迟。",
    )


def _add_integration_and_conditions(document: Document) -> None:
    document.add_heading("六、下一步整合方向", level=1)
    stages = [
        ("阶段一：统一数据入口", "把 Seewo 导出视频、评课报告、学生档案、课程标准和四路上传任务统一到课程批次；建立匿名编号、文件质量检查、权限和保存期限。"),
        ("阶段二：教师复核闭环", "由系统推荐关键切片，教师只复核与课程目标直接相关的片段；把确认、修正、无法判断、提示等级和学生回应整理为可训练标签。"),
        ("阶段三：模型服务化", "将人员跟踪、姿态/动作、VAD/ASR、课堂事件识别和时序多模态融合拆成可替换服务；保留规则基线作为对照和安全兜底。"),
        ("阶段四：学校平台化", "增加学校、班级、课程、学生、教师、管理员等角色，形成任务队列、报告中心、到期提醒、统计看板和审计后台。"),
        ("阶段五：具身设备联动", "把已复核状态和风险转成机器人/屏幕/扬声器动作协议；先在仿真和桌面设备验证，再进入真实课堂，以教师确认为执行门控。"),
    ]
    for title, text in stages:
        _add_callout(document, title, text)

    document.add_heading("七、实际落地实施条件", level=1)
    rows = [
        ("校方授权", "学校、监护人和项目方明确视频、音频、档案、分析、下载和保存范围；完成数据安全协议。"),
        ("课程与学生范围", "先选 1 个班、3-5 名学生、1-2 类稳定课程任务，避免一开始覆盖全校全部场景。"),
        ("标注与复核人员", "至少 1 名熟悉学生的特教教师作为标签负责人，研发人员只整理切片和工具，不代替教育判断。"),
        ("试点数据量", "首轮建议积累 20-30 节可用课堂及约 1,000-3,000 个经教师复核的互动片段；该范围用于启动评估，需按实际标签分布调整。"),
        ("采集质量", "优先使用 1080p 固定机位，保证目标学生可见；声音至少能区分教学阶段和明显回应，低清或严重遮挡片段标记为证据不足。"),
        ("上传端硬件", "普通 i5/R5、16GB 内存和现代浏览器即可；上传电脑不强制独立显卡。"),
        ("试点服务端", "建议 8-16 核 CPU、32GB 内存、1TB NVMe；四路并发上传建议校内总上行不少于 40Mbps。"),
        ("模型工作节点", "完整视觉/音频模型建议使用 12-24GB 显存 NVIDIA GPU；训练和批量重分析可使用云 GPU。"),
        ("账号与权限", "至少区分教师、教研负责人、学校管理员和研发运维；学生数据按学校、班级和课程隔离。"),
        ("验收指标", "同时考核数据可用率、关键片段召回率、教师修正率、掌握判断一致率、报告节省时间、延迟和安全误触发。"),
        ("运维与应急", "明确存储容量、备份、到期删除、失败重试、日志审计、模型版本回滚、人工停用和课堂异常处理流程。"),
        ("机器人安全", "真实机器人接入前完成速度、距离、动作白名单、急停、教师确认和不主动触碰学生等安全测试。"),
    ]
    table = document.add_table(rows=1, cols=2)
    for index, label in enumerate(["条件类别", "具体要求"]):
        _set_cell_text(table.cell(0, index), label, bold=True, color=DARK_BLUE)
        _set_cell_shading(table.cell(0, index), LIGHT_BLUE)
    _repeat_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        _set_cell_text(cells[0], values[0], bold=True)
        _set_cell_text(cells[1], values[1])
        if row_index % 2 == 0:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, [2200, 7160])
    _set_table_borders(table)


def _add_source_index(document: Document, files: list[SourceFile]) -> None:
    document.add_heading("八、完整源码索引", level=1)
    _add_paragraph(
        document,
        "下表和附录按 2026-08-27 交付包中的实际源码、测试、配置、模型和启动脚本汇编。运行生成的分析 JSON、SQLite 数据库、图片、PDF、DOCX 和上传文件不重复粘贴为源码。",
    )
    table = document.add_table(rows=1, cols=4)
    for index, label in enumerate(["序号", "文件路径", "类型", "行数"]):
        _set_cell_text(table.cell(0, index), label, bold=True, color=DARK_BLUE)
        _set_cell_shading(table.cell(0, index), LIGHT_BLUE)
    _repeat_header(table.rows[0])
    for number, item in enumerate(files, start=1):
        cells = table.add_row().cells
        values = [str(number), item.relative, item.language, f"{item.lines:,}"]
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value, bold=index == 1)
            if index in {0, 2, 3}:
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if number % 2 == 0:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, [720, 6100, 1300, 1240])
    _set_table_borders(table, color=BORDER)


def _add_code_appendix(document: Document, files: list[SourceFile]) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.22)
    _set_current_header_footer(section, appendix=True)

    document.add_heading("附录 A：完整源程序代码", level=1)
    _add_paragraph(
        document,
        "本附录按交付包中的实际文件逐一汇编。左侧四位数字为文档阅读辅助行号，不属于原始代码内容。",
        size=9.5,
        color=MUTED,
    )
    current_category = None
    for item in files:
        if item.category != current_category:
            current_category = item.category
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.page_break_before = True
            paragraph.paragraph_format.space_after = Pt(6)
            run = paragraph.add_run(current_category)
            _set_run_font(run, size=15, bold=True, color=BLUE)

        heading = document.add_paragraph(style="Source File Heading")
        run = heading.add_run(f"{item.relative}  |  {item.language}  |  {item.lines:,} 行")
        _set_run_font(run, size=10.5, bold=True, color=DARK_BLUE)
        _set_keep_with_next(heading)

        raw_lines = item.path.read_text(encoding="utf-8", errors="replace").splitlines()
        numbered = [
            f"{line_number:04d}  {line.expandtabs(4)}"
            for line_number, line in enumerate(raw_lines, start=1)
        ] or ["0001  "]
        for start in range(0, len(numbered), 80):
            chunk = _xml_safe("\n".join(numbered[start : start + 80]))
            paragraph = document.add_paragraph(style="Source Code")
            _set_widow_control(paragraph, False)
            run = paragraph.add_run(chunk)
            _set_run_font(run, name="Consolas", size=7.2, color=INK)


def build() -> Path:
    if not PACKAGE_ROOT.exists():
        raise FileNotFoundError(PACKAGE_ROOT)
    files = _collect_source_files()
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    _configure_styles(document)
    _set_current_header_footer(section)
    _update_fields_on_open(document)

    _add_cover(document, files)
    document.add_page_break()
    _add_toc(document)
    _add_project_summary(document)
    _add_system_flow(document)
    _add_pages(document)
    _add_data_governance(document)
    _add_status(document)
    _add_integration_and_conditions(document)
    _add_source_index(document, files)
    _add_code_appendix(document, files)

    properties = document.core_properties
    properties.title = "辅读学校 Demo2 系统汇总与源码说明"
    properties.subject = "课堂学习支持、学生追踪、数据治理与完整源码归档"
    properties.author = "RehabAffectPerception 项目组"
    properties.keywords = "辅读学校, Demo2, 学生状态, 学习掌握, 教师复核, 数据治理, 源码"
    properties.comments = "内部受限资料，不含原始学生视频和档案原件"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    print(f"files={len(files)}")
    print(f"lines={sum(item.lines for item in files)}")
    print(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    build()
