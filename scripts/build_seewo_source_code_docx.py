from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
PACKAGE_ROOT = (
    PROJECT_ROOT
    / "dist"
    / "RehabAffect_Seewo_Student_Demo_Offline_20260727"
)
OUTPUT_PATH = (
    PROJECT_ROOT
    / "docs"
    / "辅读学校Demo_源码汇编与通俗说明_20260731.docx"
)

BODY_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17212B"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "F5F6F8"
BORDER = "C9D2DC"
WHITE = "FFFFFF"


@dataclass(frozen=True)
class SourceFile:
    path: Path
    relative: str
    category: str
    language: str
    lines: int


def _set_run_font(
    run,
    *,
    name: str = BODY_FONT,
    size: float = 10.5,
    bold: bool = False,
    color: str = INK,
    italic: bool = False,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(
    cell,
    *,
    top: int = 80,
    bottom: int = 80,
    start: int = 120,
    end: int = 120,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (
        ("top", top),
        ("bottom", bottom),
        ("start", start),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table, color: str = BORDER, size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def _set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            _set_cell_margins(cell)


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _set_repeat_table_header(row) -> None:
    _repeat_table_header(row)


def _set_keep_with_next(paragraph, enabled: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = enabled


def _set_widow_control(paragraph, enabled: bool) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    node = ppr.find(qn("w:widowControl"))
    if node is None:
        node = OxmlElement("w:widowControl")
        ppr.append(node)
    node.set(qn("w:val"), "1" if enabled else "0")


def _add_field(paragraph, instruction: str) -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instruction_run._r.append(instr)

    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    paragraph.add_run("1")

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def _set_header_footer(section, *, appendix: bool = False) -> None:
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.clear_all()
    label = "完整源码附录" if appendix else "辅读学校 Seewo 学生状态追踪 Demo"
    r = p.add_run(label)
    _set_run_font(r, size=8.5, bold=True, color=MUTED)
    r = p.add_run("\t内部受限资料")
    _set_run_font(r, size=8.5, color=MUTED)
    p.paragraph_format.tab_stops.add_tab_stop(
        section.page_width - section.left_margin - section.right_margin,
        WD_TAB_ALIGNMENT.RIGHT,
    )

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("第 ")
    _set_run_font(r, size=8.5, color=MUTED)
    _add_field(fp, "PAGE")
    r = fp.add_run(" 页")
    _set_run_font(r, size=8.5, color=MUTED)


def _configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    code_style = styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = CODE_FONT
    code_style._element.rPr.rFonts.set(qn("w:ascii"), CODE_FONT)
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), CODE_FONT)
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code_style.font.size = Pt(7.2)
    code_style.font.color.rgb = RGBColor.from_string(INK)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    code_style.paragraph_format.line_spacing = Pt(8.6)
    code_style.paragraph_format.keep_together = False
    code_style.paragraph_format.keep_with_next = False

    file_style = styles.add_style("Source File Heading", WD_STYLE_TYPE.PARAGRAPH)
    file_style.font.name = BODY_FONT
    file_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    file_style.font.size = Pt(10.5)
    file_style.font.bold = True
    file_style.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    file_style.paragraph_format.space_before = Pt(8)
    file_style.paragraph_format.space_after = Pt(3)
    file_style.paragraph_format.keep_with_next = True


def _add_paragraph(
    doc: Document,
    text: str,
    *,
    bold_prefix: str | None = None,
    size: float = 10.5,
    color: str = INK,
    after: float = 6,
    align=WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        _set_run_font(first, size=size, bold=True, color=color)
        rest = p.add_run(text[len(bold_prefix) :])
        _set_run_font(rest, size=size, color=color)
    else:
        run = p.add_run(text)
        _set_run_font(run, size=size, color=color)


def _add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    _set_run_font(r, size=10.5)


def _add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    _set_run_font(r, size=10.5)


def _add_callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    _set_table_geometry(table, [9360])
    _set_table_borders(table, color="B8CCE4", size="6")
    cell = table.cell(0, 0)
    _set_cell_shading(cell, PALE_BLUE)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    _set_run_font(r, size=11, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    r = p2.add_run(text)
    _set_run_font(r, size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _set_cell_text(cell, text: str, *, bold: bool = False, color: str = INK) -> None:
    p = cell.paragraphs[0]
    p.clear()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    _set_run_font(r, size=9.5, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _collect_source_files() -> list[SourceFile]:
    roots: list[tuple[str, Path, tuple[str, ...]]] = [
        ("核心源代码", PACKAGE_ROOT / "src", (".py",)),
        ("自动测试", PACKAGE_ROOT / "tests", (".py",)),
        ("配置文件", PACKAGE_ROOT / "configs", (".json",)),
        ("运行与验证工具", PACKAGE_ROOT / "tools", (".py",)),
    ]
    files: list[SourceFile] = []
    for category, root, extensions in roots:
        for path in sorted(item for item in root.rglob("*") if item.is_file()):
            if path.suffix.lower() not in extensions:
                continue
            text = path.read_text(encoding="utf-8", errors="replace")
            files.append(
                SourceFile(
                    path=path,
                    relative=path.relative_to(PACKAGE_ROOT).as_posix(),
                    category=category,
                    language=_language_for(path),
                    lines=len(text.splitlines()),
                )
            )

    top_level = [
        "pyproject.toml",
        "requirements.txt",
        "requirements-dev.txt",
        "requirements-reanalysis.txt",
        "install_environment.ps1",
        "install_environment.cmd",
        "start_demo.ps1",
        "start_demo.cmd",
        "verify_demo.cmd",
    ]
    for name in top_level:
        path = PACKAGE_ROOT / name
        text = path.read_text(encoding="utf-8", errors="replace")
        files.append(
            SourceFile(
                path=path,
                relative=name,
                category="环境与启动脚本",
                language=_language_for(path),
                lines=len(text.splitlines()),
            )
        )
    return sorted(files, key=lambda item: (item.category, item.relative))


def _language_for(path: Path) -> str:
    return {
        ".py": "Python",
        ".json": "JSON",
        ".ps1": "PowerShell",
        ".cmd": "Batch",
        ".toml": "TOML",
        ".txt": "Text",
    }.get(path.suffix.lower(), "Text")


def _add_cover(doc: Document, total_files: int, total_lines: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("内部技术交付资料")
    _set_run_font(r, size=11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("辅读学校 Seewo 学生状态追踪 Demo")
    _set_run_font(r, size=24, bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("源码汇编与通俗说明")
    _set_run_font(r, size=17, bold=True, color=DARK_BLUE)

    meta = [
        ("文档用途", "面向非技术管理者的项目理解与完整源码归档"),
        ("系统版本", "student_support_v4 / Offline Demo"),
        ("汇编范围", f"{total_files} 个代码与配置文件，约 {total_lines:,} 行"),
        ("生成日期", str(date(2026, 7, 31))),
        ("资料级别", "内部受限，不含原始课堂视频和原始学生档案"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    for index, (label, value) in enumerate(meta):
        _set_cell_text(table.cell(index, 0), label, bold=True, color=DARK_BLUE)
        _set_cell_text(table.cell(index, 1), value)
        if index % 2 == 0:
            _set_cell_shading(table.cell(index, 0), LIGHT_BLUE)
            _set_cell_shading(table.cell(index, 1), "FAFBFC")
    _set_table_geometry(table, [2700, 6660])
    _set_table_borders(table)

    doc.add_paragraph()
    _add_callout(
        doc,
        "领导阅读建议",
        "阅读前面的“项目说明、页面功能、测试结果和风险边界”即可了解系统。后半部分是完整源码附件，主要用于技术归档、审查和后续交接，不要求非技术人员逐行阅读。",
    )


def _add_toc(doc: Document) -> None:
    doc.add_heading("目录", level=1)
    p = doc.add_paragraph()
    _add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    p.paragraph_format.space_after = Pt(8)
    _add_paragraph(
        doc,
        "提示：在 Word 中按 Ctrl+A 后按 F9，可更新目录页码。",
        size=9.5,
        color=MUTED,
    )


def _add_executive_summary(doc: Document) -> None:
    doc.add_heading("一、给管理者看的项目说明", level=1)
    _add_callout(
        doc,
        "一句话说明",
        "系统利用学校已有的 Seewo 课堂视频、课堂声音、评课报告和匿名化学生支持档案，定位单个学生在课堂中可能需要帮助的时间段，并给出屏幕、语音、机器人和老师端的低刺激支持建议。",
    )

    doc.add_heading("1.1 这个系统解决什么问题", level=2)
    _add_bullet(doc, "一节课三十多分钟，老师很难课后逐帧回看某个学生的状态变化。")
    _add_bullet(doc, "Seewo 主要提供课堂整体评课信息，缺少面向单个学生的持续时间线。")
    _add_bullet(doc, "特殊学生的表现受任务、声音、老师靠近和个体能力影响，不能只靠一个表情下结论。")
    _add_bullet(doc, "系统需要把“观察结果”继续转换为克制、可关闭、可由老师确认的支持动作。")

    doc.add_heading("1.2 当前已经做到什么", level=2)
    achievements = [
        "读取约 38.3 分钟课堂样例并形成 77 个 30 秒分析窗口。",
        "区分参与、注意力下降、困惑/停滞、成人遮挡下不确定等粗粒度状态。",
        "结合课堂声音、目标区域动作、成人区域遮挡、画面质量和评课背景解释结果。",
        "接入匿名化 IEP 与认知、语言沟通、运动能力摘要，限制提示方式和动作强度。",
        "输出学生状态时间线、关键复核片段、教师课堂支持效果和结构化 JSON。",
        "为屏幕、语音、机器人和老师端生成动作建议，并保留人工复核入口。",
        "提供 Windows 离线启动脚本、API、测试与交付包。",
    ]
    for item in achievements:
        _add_bullet(doc, item)

    doc.add_heading("1.3 当前没有做到什么", level=2)
    boundaries = [
        "没有直连 Seewo 官方后台，当前由人工导出资料后导入。",
        "不是医学诊断系统，也不能给学生贴固定标签。",
        "当前状态估计以可解释规则和轻量特征为主，不是训练完成的大型多模态模型。",
        "当前机器人动作是结构化建议和接口输出，还没有完成真实机器人安全联调。",
        "教师支持指数只用于本节课复盘，不用于教师个人绩效或排名。",
    ]
    for item in boundaries:
        _add_bullet(doc, item)


def _add_architecture(doc: Document) -> None:
    doc.add_heading("二、系统怎么工作", level=1)
    doc.add_heading("2.1 从输入到输出", level=2)
    steps = [
        "输入：Seewo 导出视频及音轨、评课报告、目标学生参考区域和匿名化支持档案。",
        "切片：将课堂按 30 秒切分，形成可追踪、可复核的时间窗口。",
        "感知：计算学生区域动作、上下文动作、成人靠近、声音活动和画面质量。",
        "状态估计：结合课堂背景，输出状态、置信度、证据和不确定性。",
        "动作建议：转换为屏幕弱高亮、短句语音、机器人慢动作或老师提醒。",
        "教师复核：老师选择符合观察、需要修正或不确定，作为后续校准依据。",
    ]
    for item in steps:
        _add_numbered(doc, item)

    doc.add_heading("2.2 主要代码模块", level=2)
    rows = [
        ("页面与服务", "scenario_lab_server.py", "启动本地网页、API 和图片资源服务"),
        ("课堂分析", "seewo_student_demo.py", "视频/音频/报告处理、状态窗口、教师效果和页面生成"),
        ("数据结构", "schema.py", "定义 human_state、robot_action 等结构化对象"),
        ("状态主流程", "pipeline.py", "组织识别、时间平滑、多模态融合和动作输出"),
        ("融合与平滑", "fusion.py / smoothing.py", "减少单帧误判并汇总多模态证据"),
        ("动作与安全", "policy.py", "把状态转换为可控动作并设置安全优先级"),
        ("仿真闭环", "sim_runner.py / closed_loop.py", "验证动作执行后的环境与状态反馈"),
        ("测试与交付", "tests/、tools/、启动脚本", "验证页面、API、策略、数据结构和离线运行"),
    ]
    table = doc.add_table(rows=1, cols=3)
    for index, title in enumerate(("系统部分", "核心文件", "通俗说明")):
        _set_cell_text(table.cell(0, index), title, bold=True, color=DARK_BLUE)
        _set_cell_shading(table.cell(0, index), LIGHT_BLUE)
    _set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value, bold=index == 0)
        if row_index % 2 == 0:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, [1900, 3000, 4460])
    _set_table_borders(table)


def _add_management_readout(doc: Document) -> None:
    doc.add_heading("三、页面上能看到什么", level=1)
    screenshot = (
        PACKAGE_ROOT / "docs" / "screenshots" / "demo_overview.png"
    )
    if screenshot.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(screenshot), width=Inches(6.5))
        p.paragraph_format.space_after = Pt(4)
        _add_paragraph(
            doc,
            "图 1  Demo 页面总览：左侧是学生画像与数据来源，中间是整节课结论，右侧是当前片段的证据、动作方案和教师复核。",
            size=9.5,
            color=MUTED,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            after=10,
        )

    doc.add_heading("3.1 点击时间段有什么作用", level=2)
    _add_paragraph(
        doc,
        "下方时间段不是普通装饰，也不是当前版本的视频播放按钮。每一个时间段都对应课堂中的一个 30 秒分析窗口。点击后，右侧会切换到该窗口的状态、置信度、底层证据、动作建议和 JSON，并允许老师对这一窗口进行复核。",
    )
    _add_paragraph(
        doc,
        "它把“整节课一个总分”下钻成“具体什么时候值得老师看一眼、系统为什么这样判断、下一步可以怎样支持”。后续接入视频播放器后，还可以根据 start_sec 和 end_sec 自动跳转到原视频对应位置。",
    )

    doc.add_heading("3.2 教师支持效果怎么看", level=2)
    teacher_image = (
        PACKAGE_ROOT
        / "docs"
        / "screenshots"
        / "teacher_effectiveness.png"
    )
    if teacher_image.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(teacher_image), width=Inches(6.5))
        p.paragraph_format.space_after = Pt(4)
        _add_paragraph(
            doc,
            "图 2  教师支持效果用于课后复盘该节课对目标学生的支持情况，不评价教师个人能力。",
            size=9.5,
            color=MUTED,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )


def _add_delivery_status(doc: Document) -> None:
    doc.add_heading("四、开发、测试与交付状态", level=1)
    status_rows = [
        ("离线页面", "已完成", "本机浏览器直接查看，不依赖公网"),
        ("学生状态时间线", "已完成", "77 个 30 秒窗口"),
        ("关键片段复核", "已完成", "支持片段切换和三类人工反馈"),
        ("个体支持档案", "已完成原型", "接入匿名化 IEP 与三类能力摘要"),
        ("教师支持效果", "已完成原型", "用于课后复盘，不用于绩效"),
        ("结构化接口", "已完成", "HTTP JSON，可继续接 WebSocket/ROS2"),
        ("自动测试", "已完成", "28 passed，5 skipped"),
        ("Seewo 自动同步", "未完成", "当前为人工导出与本地导入"),
        ("真实机器人联调", "未完成", "当前输出动作协议与建议"),
        ("训练模型", "待积累标签", "当前为可解释 baseline"),
    ]
    table = doc.add_table(rows=1, cols=3)
    for index, title in enumerate(("项目项", "状态", "说明")):
        _set_cell_text(table.cell(0, index), title, bold=True, color=DARK_BLUE)
        _set_cell_shading(table.cell(0, index), LIGHT_BLUE)
    _set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(status_rows, start=1):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            color = INK
            if index == 1 and value.startswith("已完成"):
                color = "16783A"
            elif index == 1 and value.startswith("未完成"):
                color = "B54708"
            _set_cell_text(cells[index], value, bold=index < 2, color=color)
        if row_index % 2 == 0:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, [2400, 1800, 5160])
    _set_table_borders(table)

    doc.add_heading("4.1 测试结果", level=2)
    for item in (
        "自动测试：28 项通过，5 项因可选仿真/训练环境未安装而跳过。",
        "浏览器验证：页面加载、时间段切换、教师复核按钮和教师效果页签正常。",
        "浏览器控制台：0 个错误、0 个警告。",
        "独立压缩包验证：重新解压后，页面、API 和图片资源均返回 HTTP 200。",
        "隐私检查：交付包未包含原始 MP4、原始 IEP、PDF、Excel 和学生真实姓名。",
    ):
        _add_bullet(doc, item)

    doc.add_heading("4.2 领导需要关注的风险边界", level=2)
    for item in (
        "当前结果只能证明技术流程可运行，不能证明对所有学生都准确。",
        "真实落地前需要校方授权、教师标注、跨课程验证和数据访问控制。",
        "页面中的百分比和指数是原型量化结果，必须结合老师观察解释。",
        "低清全景视频不应宣传为精细表情识别或医学状态识别。",
        "任何机器人动作都应经过白名单、安全控制和老师确认。",
    ):
        _add_bullet(doc, item)


def _add_source_index(doc: Document, files: list[SourceFile]) -> None:
    doc.add_heading("五、完整源码索引", level=1)
    _add_paragraph(
        doc,
        "下面列出本 Word 后半部分收录的全部源码、测试、配置与启动脚本。生成的课堂分析 JSON 和缓存数据属于运行结果，不作为源程序重复粘贴；它们仍保留在离线交付包中。",
    )
    table = doc.add_table(rows=1, cols=4)
    for index, title in enumerate(("序号", "文件路径", "类型", "行数")):
        _set_cell_text(table.cell(0, index), title, bold=True, color=DARK_BLUE)
        _set_cell_shading(table.cell(0, index), LIGHT_BLUE)
    _set_repeat_table_header(table.rows[0])
    for number, item in enumerate(files, start=1):
        cells = table.add_row().cells
        values = (str(number), item.relative, item.language, f"{item.lines:,}")
        for index, value in enumerate(values):
            _set_cell_text(cells[index], value, bold=index == 1)
            if index in {0, 2, 3}:
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if number % 2 == 0:
            for cell in cells:
                _set_cell_shading(cell, "FAFBFC")
    _set_table_geometry(table, [720, 6100, 1300, 1240])
    _set_table_borders(table)


def _xml_safe(text: str) -> str:
    return "".join(
        char
        for char in text
        if char in "\t\n\r"
        or 0x20 <= ord(char) <= 0xD7FF
        or 0xE000 <= ord(char) <= 0xFFFD
    )


def _add_code_appendix(doc: Document, files: list[SourceFile]) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.22)
    _set_header_footer(section, appendix=True)

    doc.add_heading("附录 A：完整源程序代码", level=1)
    _add_paragraph(
        doc,
        "本附录按交付包中的实际文件逐一汇编。左侧四位数字为文档阅读辅助行号，不属于原始代码内容。",
        size=9.5,
        color=MUTED,
    )

    current_category = None
    for item in files:
        if item.category != current_category:
            current_category = item.category
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(current_category)
            _set_run_font(r, size=15, bold=True, color=BLUE)

        heading = doc.add_paragraph(style="Source File Heading")
        run = heading.add_run(
            f"{item.relative}  |  {item.language}  |  {item.lines:,} 行"
        )
        _set_run_font(run, size=10.5, bold=True, color=DARK_BLUE)
        _set_keep_with_next(heading)

        raw_lines = item.path.read_text(encoding="utf-8", errors="replace").splitlines()
        numbered = [
            f"{line_number:04d}  {line.expandtabs(4)}"
            for line_number, line in enumerate(raw_lines, start=1)
        ]
        if not numbered:
            numbered = ["0001  "]

        chunk_size = 80
        for start in range(0, len(numbered), chunk_size):
            chunk = _xml_safe("\n".join(numbered[start : start + chunk_size]))
            p = doc.add_paragraph(style="Source Code")
            _set_widow_control(p, False)
            r = p.add_run(chunk)
            _set_run_font(
                r,
                name=CODE_FONT,
                size=7.2,
                color=INK,
            )


def _update_fields_on_open(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build() -> Path:
    if not PACKAGE_ROOT.exists():
        raise FileNotFoundError(f"Package source not found: {PACKAGE_ROOT}")
    files = _collect_source_files()
    total_lines = sum(item.lines for item in files)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    _configure_styles(doc)
    _set_header_footer(section)
    _update_fields_on_open(doc)

    _add_cover(doc, len(files), total_lines)
    doc.add_page_break()
    _add_toc(doc)
    doc.add_page_break()
    _add_executive_summary(doc)
    _add_architecture(doc)
    _add_management_readout(doc)
    _add_delivery_status(doc)
    _add_source_index(doc, files)
    _add_code_appendix(doc, files)

    core_properties = doc.core_properties
    core_properties.title = "辅读学校 Demo 源码汇编与通俗说明"
    core_properties.subject = "Seewo 学生状态追踪 Demo 技术交付"
    core_properties.author = "RehabAffectPerception 项目组"
    core_properties.keywords = "辅读学校, Seewo, 学生状态追踪, 源码汇编"
    core_properties.comments = "内部受限资料，不含原始学生视频和原始档案"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(f"files={len(files)}")
    print(f"lines={total_lines}")
    print(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    build()
