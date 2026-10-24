from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = PROJECT_ROOT / "docs" / "special_education_embodied_classroom_design_zh.docx"
LEGACY_OUTPUT = PROJECT_ROOT / "docs" / "special_education_classroom_assist_design_v2_zh.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(30, 38, 45)
MUTED = RGBColor(95, 105, 115)
TABLE_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
SOFT_BLUE_FILL = "E8EEF5"


def set_run_font(run, size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("辅读学校具身智能课堂辅助系统设计方案")
    set_run_font(r, size=9, color=MUTED)

    return doc


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_para(doc: Document, text: str, *, bold=False, italic=False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, bold=bold, italic=italic)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r)


def add_callout(doc: Document, title: str, body: str, fill: str = CALLOUT_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, fill)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, bold=True, color=DARK_BLUE)

    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, TABLE_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, bold=True, color=INK)

    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            if i == 0 and len(value) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r)

    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_title_block(doc: Document) -> None:
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("辅读学校具身智能课堂辅助系统设计方案")
    set_run_font(r, size=22, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    sr = subtitle.add_run("课堂状态感知 → 学生状态估计 → 机器人动作响应 → 教师复盘")
    set_run_font(sr, size=11, color=MUTED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    mr = meta.add_run("产品设计版 | 场景：辅读学校真实课堂 | 输出：可执行机器人动作与教师辅助决策")
    set_run_font(mr, size=10, color=MUTED)


def build_doc() -> None:
    doc = setup_document()
    add_title_block(doc)

    add_callout(
        doc,
        "一句话方案",
        "开发一套面向辅读学校课堂的具身智能辅助系统：通过摄像头、麦克风、任务表现和教师标注感知学生课堂状态，估计学生是否专注、困惑、疲劳、过度活跃、抗拒或需要帮助，并把最终决策落到机器人可执行的低刺激动作上。",
        SOFT_BLUE_FILL,
    )

    add_heading(doc, "1. 设计定位", 1)
    add_para(
        doc,
        "本项目定位为辅读学校课堂中的具身智能辅助系统。系统具备“看、听、理解状态、执行动作”的闭环能力：面向教师和学生，机器人作为课堂辅助伙伴，承担观察提醒、温和引导、任务陪伴和教师提示等角色。",
    )
    add_para(
        doc,
        "系统最终要回答的核心问题是：在一节真实课堂里，学生当前是否还在参与？是否听懂了？是否需要降低任务难度？是否需要老师介入？机器人此刻应该保持安静、看向学生、给出短句提示、重复说明、暂停任务，还是提醒教师关注？",
    )
    add_table(
        doc,
        ["定位项", "设计结论"],
        [
            ["核心场景", "辅读学校课堂，包括生活适应课、认知训练课、课堂常规训练和融合教育活动。"],
            ["核心对象", "教师、特殊儿童学生、课堂辅助机器人或屏幕互动终端。"],
            ["核心能力", "多模态课堂状态感知、学生状态估计、机器人动作决策、教师端复盘。"],
            ["核心输出", "结构化学生状态、风险等级、证据线索、推荐机器人动作、教师介入建议。"],
            ["展示目标", "形成一个能演示闭环的 Demo：学生做任务，系统判断状态，机器人执行动作，教师端看到记录。"],
        ],
        [1800, 7560],
    )

    add_heading(doc, "2. 要开发什么", 1)
    add_callout(
        doc,
        "拟开发产品",
        "辅读学校具身智能课堂辅助系统。它由学生任务端、教师控制端、多模态状态感知模块、机器人动作决策模块、仿真 Benchmark 和课后报告组成。",
    )
    add_table(
        doc,
        ["组成部分", "开发内容", "交付形态"],
        [
            ["学生任务端", "承载课堂任务，例如听指令、物品整理、校园场景识别、情境选择、出行安全、课堂规则训练。", "平板网页、大屏网页或触控屏。"],
            ["教师控制端", "班级与学生选择、任务启动、状态提示、人工标注、机器人动作开关、课后报告。", "教师电脑或平板网页。"],
            ["状态感知模块", "融合视觉、声音、任务日志、历史状态和教师标注，输出学生课堂状态。", "本地 Python 服务。"],
            ["机器人决策模块", "把学生状态转换成机器人动作，包含低刺激策略、距离约束、动作强度限制和教师介入规则。", "行为树/策略服务，支持 WebSocket 或 ROS2 Mock。"],
            ["仿真 Benchmark", "在 MuJoCo 中模拟课堂空间、机器人、学生状态变化和任务事件，评估机器人策略。", "可运行仿真脚本、指标报告和场景截图。"],
            ["课后报告", "汇总学生参与度、卡住点、提示次数、教师介入点和可改进任务。", "网页报告或导出文档。"],
        ],
        [1600, 5200, 2560],
    )

    add_heading(doc, "3. 具身智能闭环", 1)
    add_para(
        doc,
        "本方案的具身智能重点在“感知到行动”的闭环：机器人根据学生状态、课堂任务、距离和安全边界选择下一步动作，并观察动作后的反馈。这个闭环可以在真实课堂 Demo 中展示，也可以先在 MuJoCo 仿真中做可重复 Benchmark。",
    )
    add_table(
        doc,
        ["环节", "输入/处理", "输出"],
        [
            ["Perception", "摄像头观察头部朝向、身体动作、离座倾向；麦克风观察音量、异常发声、关键词；任务端记录停顿、错误、求助。", "多模态 evidence。"],
            ["State Estimation", "用规则、EMA 平滑、滑动窗口和置信度融合估计学生状态。", "engaged、confused、attention_lost、over_active、refusing、needs_help 等。"],
            ["Action Suggestion", "行为树根据学生状态、任务阶段、风险等级和教师权限选择动作。", "机器人动作指令和教师提示。"],
            ["Embodied Action", "机器人看向学生、转向屏幕、短句鼓励、重复说明、降低任务难度、暂停等待、提醒教师。", "可观察的身体动作和语音/屏幕反馈。"],
            ["Feedback Loop", "观察动作后学生是否恢复操作、错误是否减少、教师是否介入。", "更新状态和后续策略。"],
        ],
        [1700, 5200, 2460],
    )

    add_heading(doc, "4. 具身智能模型选择", 1)
    add_para(
        doc,
        "第一版采用“轻量多模态状态估计器 + 行为树机器人策略”。这个选择适合实习项目和学校展示：可解释、可调参、能在没有大量真实学生数据的情况下运行，也方便在仿真里做策略评估。后续数据增加后，再把部分阈值、融合权重或策略选择升级为可学习模型。",
    )
    add_table(
        doc,
        ["模块", "第一版模型", "后续升级"],
        [
            ["视觉线索", "OpenCV/MediaPipe 获取头部朝向、身体活动量、离座或大幅动作线索。", "更稳定的姿态模型或多目标跟踪。"],
            ["声音线索", "音量、语速近似值、异常发声事件、简单关键词/意图识别。", "ASR + 课堂意图分类 + 音频情绪特征。"],
            ["任务行为", "停顿时长、连续错误、提示后恢复情况、放弃/求助按钮。", "个体化学习曲线和难度自适应。"],
            ["状态融合", "加权融合 + EMA 时间平滑 + 状态机约束。", "置信度校准、HMM 或轻量时序模型。"],
            ["动作决策", "Behavior Tree 行为树，显式编码安全边界和教师权限。", "Contextual Bandit 或强化学习策略微调。"],
            ["仿真验证", "MuJoCo 课堂 Benchmark，学生状态用脚本化 agent 生成。", "Isaac Sim 中加入更真实的视觉、机器人和教室资产。"],
        ],
        [1500, 4200, 3660],
    )

    add_heading(doc, "5. 核心功能设计", 1)
    add_heading(doc, "5.1 学生课堂任务端", 2)
    add_para(
        doc,
        "学生端承担两个作用：一是让学生完成课堂任务，二是为状态估计提供可解释的行为线索。任务以可配置模板承载多类课堂内容，方便教师替换图片、文字、步骤和难度。",
    )
    add_table(
        doc,
        ["任务类型", "示例", "可观测线索"],
        [
            ["听指令执行", "听到“把红色物品放到左边”后完成拖拽或选择。", "反应时、是否反复点错、是否需要重复指令。"],
            ["物品整理", "把书、本子、文具、衣物放入正确位置。", "分类错误、卡住位置、完成顺序。"],
            ["校园场景", "识别教室、厕所、医务室、食堂，判断遇到问题找谁。", "场景理解错误、求助频率。"],
            ["情境选择", "排队、同伴互动、课堂回应、安全规则等情境选择。", "规则理解、冲动选择、教师标注。"],
            ["出行安全", "红绿灯、斑马线、等车、上车、下车等步骤排序。", "步骤遗漏、顺序错误、提示后恢复。"],
        ],
        [1600, 4800, 2960],
    )

    add_heading(doc, "5.2 教师控制端", 2)
    add_bullet(doc, "选择班级、学生、任务包和难度。")
    add_bullet(doc, "实时查看学生状态卡片：参与、困惑、注意力下降、过度活跃、抗拒、需要帮助。")
    add_bullet(doc, "一键标注课堂观察结果，作为状态估计的重要证据。")
    add_bullet(doc, "控制机器人动作等级：静默、屏幕提示、语音提示、转向/点头、暂停、呼叫教师。")
    add_bullet(doc, "课后查看时间线和报告，知道学生在哪些任务点需要更多支持。")

    add_heading(doc, "5.3 多模态状态感知", 2)
    add_table(
        doc,
        ["状态", "可能线索", "置信度处理"],
        [
            ["engaged 参与", "持续操作、反应及时、错误少、头部朝向任务区域。", "短时稳定后提高置信度。"],
            ["attention_lost 注意力下降", "长时间无操作、视线游离、身体活动增多、任务停顿。", "连续窗口触发，避免单帧误判。"],
            ["confused 困惑", "连续错误、反复撤销、求助、提示后仍无法完成。", "结合任务难度和历史错误。"],
            ["over_active 过度活跃", "大幅动作、频繁离座倾向、音量升高、教师标注。", "提高教师提示优先级。"],
            ["refusing 抗拒", "主动停止、长时间拒绝操作、教师标注抗拒。", "降低机器人刺激强度。"],
            ["needs_help 需要帮助", "多次卡住、主动求助、提示后仍未恢复。", "触发教师介入建议。"],
            ["safety_risk 安全风险", "离开座位、靠近危险区域、剧烈动作或教师紧急标注。", "立即停止任务并提醒教师。"],
        ],
        [1800, 5200, 2360],
    )

    add_heading(doc, "5.4 机器人动作输出", 2)
    add_para(
        doc,
        "最终决策必须落在机器人动作上。动作设计遵循低刺激、短时、可中断、教师可控的原则。机器人动作不仅包括移动，也包括朝向、头部姿态、屏幕指向、语音和任务难度调整。",
    )
    add_table(
        doc,
        ["动作等级", "机器人动作", "适用状态"],
        [
            ["A0 观察等待", "保持安静，记录状态，不打断课堂。", "参与稳定、状态不确定。"],
            ["A1 朝向关注", "轻微转头或屏幕高亮，把注意力指向任务区域。", "轻度注意力下降。"],
            ["A2 短句提示", "用慢速短句提示：“我们看第一步”“可以再试一次”。", "困惑、短暂停顿。"],
            ["A3 重复说明", "重复教师预设指令，拆成一步一步。", "听不懂、连续错误。"],
            ["A4 降低难度", "减少选项、增加示例、切换更简单任务。", "持续困惑、挫败苗头。"],
            ["A5 暂停与选择", "暂停任务，给出“休息一下/换一个”的选择。", "抗拒、疲劳、情绪波动。"],
            ["A6 教师介入", "停止自动提示，在教师端高亮并请求教师关注。", "需要帮助或安全风险。"],
        ],
        [1400, 4700, 3260],
    )

    add_heading(doc, "6. 数据从哪里来", 1)
    add_para(
        doc,
        "第一版展示不依赖真实学生隐私数据。系统先用可控的模拟数据、任务交互日志和教师脚本标注跑通闭环，再在获得学校和家长授权后进入试点采集。这样既能快速做出 Demo，也能保护真实学生。",
    )
    add_table(
        doc,
        ["数据来源", "第一版做法", "用途"],
        [
            ["任务日志", "学生端自动记录点击、停顿、错误、求助、完成时长。", "状态估计和课后报告。"],
            ["教师脚本标注", "由演示者或教师在控制端标注注意力下降、困惑、需要帮助等。", "补充真实课堂判断。"],
            ["仿真学生 Agent", "在 MuJoCo 中用状态机生成参与、困惑、过度活跃、抗拒等事件。", "训练和评估机器人策略。"],
            ["非隐私演示素材", "由成年人或团队成员录制声音/动作样例，或使用合成事件流。", "验证麦克风和视觉接口。"],
            ["授权试点数据", "进入学校后，在授权范围内采集匿名化任务日志和必要传感器特征。", "优化阈值、融合权重和个体化策略。"],
        ],
        [1700, 5200, 2460],
    )

    add_heading(doc, "7. 仿真 Benchmark 设计", 1)
    add_para(
        doc,
        "第一步使用 MuJoCo 做粗略可行版本：搭建教室区域、学生座位、任务屏幕和一个简化机器人。学生不是完整真人动画，而是可观测状态 Agent：它会根据脚本产生头部朝向、身体活动量、任务错误、异常发声事件和求助事件。这样可以先验证“状态估计 → 机器人动作 → 学生反馈”的闭环是否成立。",
    )
    add_table(
        doc,
        ["Benchmark 场景", "学生线索生成方式", "机器人目标"],
        [
            ["正常参与", "任务操作连续、头部朝向屏幕、无异常声音。", "保持低打扰，只做正向反馈。"],
            ["注意力下降", "操作停顿、头部偏离、身体活动量升高。", "先朝向/屏幕提示，再短句拉回。"],
            ["困惑卡住", "连续错误、长时间停在同一步、求助事件。", "拆分指令或降低难度。"],
            ["过度活跃", "动作幅度上升、音量事件增加、离座倾向。", "减少刺激，提醒教师关注。"],
            ["抗拒继续", "拒绝操作、停止响应、教师标注抗拒。", "暂停任务，提供选择，避免追问。"],
            ["安全风险", "离开座位或进入危险区域。", "停止任务并请求教师介入。"],
        ],
        [1700, 4200, 3460],
    )
    add_table(
        doc,
        ["指标", "含义"],
        [
            ["动作正确率", "机器人选择的动作是否符合当前学生状态。"],
            ["过度干预率", "学生状态正常时机器人是否过多打断。"],
            ["响应延迟", "从状态变化到机器人动作输出的时间。"],
            ["恢复率", "提示后学生是否重新回到任务。"],
            ["教师介入准确率", "真正需要教师时是否及时提示。"],
            ["低刺激合规率", "动作是否满足声音、距离、速度和频次限制。"],
        ],
        [2200, 7160],
    )

    add_heading(doc, "8. MuJoCo 到 Isaac Sim 的升级路径", 1)
    add_table(
        doc,
        ["阶段", "仿真目标", "输出"],
        [
            ["MuJoCo 第一版", "快速搭建课堂几何、机器人动作、学生状态事件和策略评估。", "可运行 Benchmark、指标表、仿真截图。"],
            ["MuJoCo 完整版", "加入多学生、更多课堂事件、动作代价和教师介入规则。", "策略对比报告和 Demo 视频。"],
            ["Isaac Sim 升级", "加入更真实的教室资产、相机视角、机器人模型、光照和传感器。", "更接近展示级的仿真环境。"],
            ["Sim2Real", "保持 human_state JSON 和 robot_action 接口不变，把仿真策略接到真实摄像头、麦克风和机器人。", "真实 Demo 与仿真 Benchmark 对齐。"],
        ],
        [1700, 5000, 2660],
    )

    add_heading(doc, "9. 需要什么硬件", 1)
    add_table(
        doc,
        ["硬件", "第一版配置", "作用"],
        [
            ["开发电脑", "Windows 笔记本或台式机，建议有独立显卡用于后续 Isaac Sim。", "运行学生端、教师端、状态服务和本地 Demo。"],
            ["摄像头", "普通 USB 摄像头即可，后续可换深度相机。", "采集头部朝向、身体动作和课堂参与线索。"],
            ["麦克风", "USB 麦克风或笔记本麦克风。", "采集音量、异常发声和关键词线索。"],
            ["学生端屏幕", "平板、触控屏、大屏或网页投屏。", "承载课堂任务与低刺激提示。"],
            ["教师端设备", "教师电脑或平板。", "控制任务、查看状态、标注、控制机器人动作。"],
            ["机器人载体", "桌面机器人、轮式机器人、机械臂屏幕组合，或仿真机器人。", "执行朝向、提示、屏幕指引、暂停和教师提醒等动作。"],
            ["云电脑", "从 Isaac Sim 或大规模训练阶段开始需要。", "运行 GPU 仿真、视频渲染和较重模型实验。"],
        ],
        [1600, 4300, 3460],
    )

    add_heading(doc, "10. 最终载体与展示形态", 1)
    add_callout(
        doc,
        "最终成品",
        "一个可展示的辅读学校具身智能课堂 Demo：教师端启动任务，学生端产生任务表现和状态线索，系统输出 human_state，机器人执行低刺激动作，课后生成课堂报告；同时提供 MuJoCo Benchmark 证明机器人策略可以在仿真中反复评估。",
        SOFT_BLUE_FILL,
    )
    add_table(
        doc,
        ["载体", "展示内容"],
        [
            ["学生端网页", "完成课堂任务、产生错误/停顿/求助等状态线索。"],
            ["教师端网页", "展示状态卡片、时间线、教师标注、机器人动作开关和课后报告。"],
            ["机器人端", "执行朝向、短句提示、重复说明、暂停、教师提醒等动作。"],
            ["MuJoCo 仿真", "展示课堂环境、学生 Agent、机器人策略和 Benchmark 指标。"],
            ["接口层", "输出 human_state JSON 和 robot_action JSON，支持 WebSocket，预留 ROS2。"],
            ["演示材料", "8-10 分钟演示脚本、截图、视频和结果报告。"],
        ],
        [1800, 7560],
    )

    add_heading(doc, "11. 近期可展示 Demo 流程", 1)
    add_number(doc, "教师在教师端选择“课堂规则/物品整理/校园场景”等任务包，并选择学生画像。")
    add_number(doc, "学生端开始任务，系统记录点击、停顿、错误、求助和完成情况。")
    add_number(doc, "摄像头和麦克风接口接入演示线索，或由仿真学生 Agent 生成状态事件。")
    add_number(doc, "状态估计器输出 human_state，包括状态、置信度、证据、风险等级和是否需要教师介入。")
    add_number(doc, "机器人动作策略输出 robot_action，例如看向学生、短句提示、重复说明、降低难度或提醒教师。")
    add_number(doc, "教师端显示时间线，课后生成报告，说明什么时候卡住、系统如何响应、学生是否恢复。")
    add_number(doc, "MuJoCo Benchmark 运行同一策略，展示多个课堂状态场景下的动作正确率和过度干预率。")

    add_heading(doc, "12. 接口与输出 JSON", 1)
    add_table(
        doc,
        ["输出", "字段", "用途"],
        [
            ["human_state", "timestamp、student_id、state、confidence、evidence、risk_level、handoff_required", "给教师端、机器人策略和课后报告使用。"],
            ["robot_action", "action_type、intensity、target、speech_text、motion、duration、teacher_approval_required", "让机器人端执行具体动作。"],
            ["lesson_event", "task_id、step_id、event_type、latency、error_count、help_requested", "记录任务过程和状态证据。"],
            ["session_report", "参与度、提示次数、卡住点、教师介入点、建议下次难度", "用于教师复盘。"],
        ],
        [1600, 4400, 3360],
    )

    add_heading(doc, "13. 两个月实施路线图", 1)
    add_table(
        doc,
        ["阶段", "周期", "目标", "交付物"],
        [
            ["MVP", "第 1-2 周", "完成学生任务端、教师端、human_state/robot_action 数据结构。", "可点击课堂 Demo。"],
            ["状态估计", "第 3 周", "接入任务日志、教师标注、简单摄像头/麦克风线索和时间平滑。", "实时状态卡片。"],
            ["机器人动作", "第 4 周", "完成行为树策略和机器人动作输出，支持 WebSocket/ROS2 Mock。", "机器人动作演示。"],
            ["MuJoCo 仿真", "第 5 周", "搭建课堂 Benchmark，模拟学生状态变化和机器人响应。", "仿真场景、截图、指标报告。"],
            ["优化评估", "第 6 周", "做阈值、融合权重、延迟和过度干预率评估。", "评估报告。"],
            ["展示打磨", "第 7-8 周", "整合真实 Demo、仿真 Benchmark、课后报告和演示脚本。", "可展示成品包。"],
        ],
        [1200, 1300, 4200, 2660],
    )

    add_heading(doc, "14. 安全与伦理边界", 1)
    add_bullet(doc, "系统输出是课堂辅助观察结果，不做医学诊断。")
    add_bullet(doc, "教师拥有最终判断权和干预权，机器人动作可随时关闭。")
    add_bullet(doc, "机器人不主动触碰学生，不快速靠近学生，不使用突然大声或强光刺激。")
    add_bullet(doc, "第一版优先使用任务日志、仿真事件和教师标注；真实音视频采集必须获得学校和家长授权。")
    add_bullet(doc, "展示材料不公开真实学生可识别图像、声音和个人信息。")
    add_bullet(doc, "状态标签使用“可能注意力下降”“建议教师关注”等克制表达。")

    add_heading(doc, "15. 验收标准", 1)
    add_table(
        doc,
        ["类别", "验收点"],
        [
            ["闭环完整", "能从学生任务表现生成状态，再触发机器人动作，并在教师端留下记录。"],
            ["具身相关", "机器人动作不是文本建议，而是可执行的朝向、语音、屏幕指引、暂停和教师提醒。"],
            ["仿真可跑", "MuJoCo 至少包含 5 个课堂状态场景，并输出策略指标。"],
            ["数据合理", "第一版不依赖真实学生隐私数据，也能通过模拟和日志跑通。"],
            ["教师可用", "教师能启动任务、标注状态、控制动作、查看课后报告。"],
            ["安全可控", "所有提示可关闭，动作强度可配置，风险状态优先请求教师介入。"],
        ],
        [1800, 7560],
    )

    add_heading(doc, "16. 结论", 1)
    add_para(
        doc,
        "本方案把重点放在辅读学校课堂中的具身智能闭环：系统通过多模态线索理解学生状态，再把决策落到克制、安全、可解释的机器人动作上。近期成品可以先以网页端、WebSocket/ROS2 Mock、MuJoCo Benchmark 和可视化报告完成展示；后续再升级到 Isaac Sim 和真实机器人接入。",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    shutil.copyfile(OUTPUT, LEGACY_OUTPUT)
    print(OUTPUT)
    print(LEGACY_OUTPUT)


if __name__ == "__main__":
    build_doc()
