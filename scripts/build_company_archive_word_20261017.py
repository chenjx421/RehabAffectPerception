"""Create the current company archive note without opening student data."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/pilot_20261017'
QA = OUT / 'word_qa'
NAME = '学校课堂分析平台_公司后台与学生档案说明_20261017.docx'


def font(style, size, bold=False):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')


def para(doc, text, bold_start=None, small=False):
    p = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        p.add_run(bold_start).bold = True
        p.add_run(text[len(bold_start):])
    else:
        p.add_run(text)
    if small:
        for run in p.runs:
            run.font.size = Pt(9.5)
    return p


def heading(doc, text):
    doc.add_heading(text, 2)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text
    for i, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                element = OxmlElement('w:' + edge)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    element.set(qn('w:' + key), value)
                borders.append(element)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), '85')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            pr.append(margins)
            if i == 0:
                fill = OxmlElement('w:shd')
                fill.set(qn('w:fill'), 'EAF1F8')
                pr.append(fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.08
                if j == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = i == 0
    return t


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    QA.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.66), Inches(.65)
    section.left_margin, section.right_margin = Inches(.72), Inches(.72)
    section.header_distance, section.footer_distance = Inches(.25), Inches(.28)
    for name, size, bold in [('Normal', 11, False), ('Title', 20, True),
                              ('Heading 1', 15, True), ('Heading 2', 12, True)]:
        font(doc.styles[name], size, bold)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    normal = doc.styles['Normal'].paragraph_format
    normal.line_spacing, normal.space_after = 1.16, Pt(5)
    normal.widow_control = True
    for name in ('Heading 1', 'Heading 2'):
        pf = doc.styles[name].paragraph_format
        pf.space_before, pf.space_after = Pt(9), Pt(5)
        pf.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = '学校课堂分析平台公司后台与学生档案说明'
    doc.core_properties.subject = '当前存储功能 学生课程归档 使用流程与部署条件'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''

    doc.add_paragraph('公司后台与学生档案说明', 'Title')
    para(doc, '学校课堂分析平台 Demo2    核查日期 2026年10月17日', small=True)
    para(doc, '本说明介绍公司后台的数据保存、学生单独建档及每节课的自动归档方式。当前数据库和归档功能已经实现，部署在现有 Windows 主机上；这不等于已经建成独立云数据库，也不代表真实学生资料已经全部录入。')

    heading(doc, '一 当前实际状态')
    para(doc, '本次只读检查显示：数据库完整性检查通过，本地应用和网关可访问。数据库保留36条视频任务，其中33条为视频证据处理就绪，2条为历史待处理记录，1条已取消。这是任务状态统计，不能视为36节真实课堂样本，更不能视为33条教育评分已获验证。')
    para(doc, '新的学生名册、公司授权学生摘要及自动课程历史当前均为0条。系统功能已经可以使用，但仍需正式建档、确认画面身份和完成相应授权，才能开始积累真实个人档案。已有2条身份绑定记录不代表已建立完整学生档案。')

    heading(doc, '二 数据保存在哪里')
    table(doc, ['数据类别', '保存方式', '后台用途'], [
        ('账号与视频任务', 'SQLite 保存账号权限、任务归属、处理状态及文件索引', '按账号与通道查询，不因选择学校而改变租户权限'),
        ('原视频与派生媒体', '受控文件目录保存视频、音轨和采样帧，不直接塞进数据库', '管理员按权限查看、下载；到期后原媒体不可回看'),
        ('学生名册与身份关联', '学生编号关联本校档案，绑定记录关联具体任务和采样帧', '让学生、课堂视频与相应资料能够对应'),
        ('授权摘要与课程历史', '独立授权管理，加密保存摘要及自动归档内容', '按脱敏学生编号查询、导出课程历史和有效复核测量'),
        ('原文资料与审计记录', 'IEP等原文按资料模块和授权管理；操作另有审计留痕', '不能因后台可访问，就默认复制全部原文或用于训练'),
    ], [1.18, 3.10, 2.78])
    para(doc, '当前数据库位置：项目目录 outputs/seewo_demo2/rehab_affect_demo2.sqlite3。视频等文件由数据库记录受控路径。后续备份或迁移必须同时覆盖数据库、附件及加密所需配置，不能只复制一个数据库文件。', small=True)

    p = doc.add_heading('学生建档与课程归档操作', 1)
    p.paragraph_format.page_break_before = True
    heading(doc, '三 教师第一次使用')
    para(doc, '1. 登录学校账号并核对学校归属，在“本校名册”建立学生档案。学生名册用于统一身份；IEP和诊断资料属于另行授权的补充材料，并非为了进入分析就必须提交完整病历。')
    para(doc, '2. 上传课程视频，等待证据处理完成。在采样画面中选中或手动画框，确认画面人物对应哪名学生；系统识别不清时不能直接把猜测当成正确身份。当前确认的是采样帧身份，不代表整节课持续跟踪已经完成。')
    para(doc, '3. 从学生名册进入“公司留存摘要”，核对授权范围、依据和期限，再开启“自动归入后续课程”。旧授权不会自动打开此功能；需要补录旧课程时，应另外选择补录并确认授权覆盖。')
    para(doc, '4. 后续本学生已绑定身份且处理完成的课程，会自动归入同一授权档案，无需每节课手工填写摘要。没有复核测量时显示待复核；只有有效、适用且已核准的测量才能显示分数。')

    heading(doc, '四 公司后台如何调取')
    para(doc, '管理员在公司后台查看各通道的上传任务及处理状态；学生相关的授权数据从“学生档案摘要库”查看。可查询脱敏学生编号、通道、累计课程数、最近课程和授权截止日，再展开课程历史并导出授权 JSON。视频仍从原有上传任务入口按权限查看，摘要库不额外授予原视频读取权限。')
    para(doc, '归档记录包含课程日期或上传日期、媒体技术信息、复核状态及有效的结构化测量，并保留来源版本和修订信息。重复同步不重复建档；身份关联或测量被撤回后，对应分数会撤下。后台同步失败会提示待重试，原先已经保存的主记录不会因此被误报为未保存。')

    heading(doc, '五 保存期限和授权边界')
    table(doc, ['对象', '当前执行口径'], [
        ('原视频和音频', '原有60天保存规则不变，到期前15天提醒；正常到期清理包含派生媒体。'),
        ('独立授权摘要', '默认60天、最长365天，且不能超过实际授权期限；旧测试数据不自动批量转存。'),
        ('原视频到期后的已核准记录', '仅在独立摘要授权仍有效、正常清理有可信审计且身份与测量依据未变化时，保留先前核准数字，并标明原证据不可回看。'),
        ('授权撤回或到期', '公司不可继续读取，相关加密摘要和课程历史内容被逻辑清除；已下载副本和备份仍需另行按授权管理。'),
    ], [1.65, 5.41])
    para(doc, '后台可以保存和读取，不等于公司可以永久保留、开展科研或训练模型。诊断、IEP原文和可识别身份资料不因开启课程归档而自动获得新的使用许可。旧运营配置中的365天或730天设置，也不能替代具体资料授权及相应清理规则。', small=True)

    p = doc.add_heading('验收结果与实际部署条件', 1)
    p.paragraph_format.page_break_before = True
    heading(doc, '六 已完成的工程验证')
    para(doc, '使用两段生成的非学生视频，经真实浏览器完成上传、FFmpeg处理、身份确认、自动课程归档、公司后台查询与JSON下载。另行验证重复同步不重复归档、有效复核测量自动更新、测量或身份撤回后撤下分数，以及摘要授权撤回后阻止访问。桌面和手机视口均检查了页面显示。')
    para(doc, '归档相关首轮回归通过277项；最终课程历史定向回归通过76项，包含到期先扫描再清理、删除前媒体缺失、同秒启用归档边界等情况。测试集合有重叠，不应相加为独立测试总数。工程测试通过不代表学生状态、掌握度或课堂评分准确率已完成真实校验。', small=True)

    image = ROOT / 'output/playwright/automatic-archive-20261016/company-course-history-desktop.png'
    with Image.open(image) as source:
        # Show the actual course-history region rather than shrinking a full page.
        crop = source.crop((150, 630, 1290, 1305))
        screenshot = QA / 'company_history_crop.png'
        crop.save(screenshot)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(screenshot), width=Inches(6.0))
    para(doc, '图1 公司后台课程历史验收截图。画面使用合成数据，演示同一脱敏学生下的两条课程记录，一条待复核、一条已复核，以及授权JSON导出入口；图中分数不是学生真实成效。', small=True)

    heading(doc, '七 当前边界与后续实施')
    para(doc, '访问状态：原邀请信息、四个学校账号和原有数据保留。本次核查本地后台正常，但原公网入口连接检查失败，因此不能在本说明中承诺外部测试入口此刻可用。邀请配置到2026年11月6日结束，不等于临时隧道有可用性保障。')
    para(doc, '下一步先恢复并验收外部访问，选择获准的学生建立真实档案，再按“上传—绑定—复核—查询—下载”逐通道验收。正式多校部署还需固定域名与服务主机、数据库和附件的备份恢复、权限审计、授权到期清理及容量监控；可再迁移至PostgreSQL和对象存储。目前没有为本次更新购买云资源，也没有把学生资料自动发送给新模型接口。')

    output = OUT / NAME
    doc.save(output)
    print(output)


if __name__ == '__main__':
    build()
