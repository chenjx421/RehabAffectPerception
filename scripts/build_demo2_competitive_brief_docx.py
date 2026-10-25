from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\GIT\RehabAffectPerception")
OUTPUT = ROOT / "docs" / "Demo2_核心竞争力与课堂表现分析维度_20260817.docx"

BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
INK = "173B36"
BODY = "2F4945"
MUTED = "687A77"
TEAL = "0E897D"
TEAL_SOFT = "E8F3EF"
BLUE_SOFT = "E8EEF5"
BORDER = "CBD8D4"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=85, start=140, bottom=85, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_table_borders(table, color=BORDER, size=5):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_run_font(run, size=11, bold=False, color=BODY, name="Noto Sans SC"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_inline_label(paragraph, label, text):
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size=10, bold=True, color=INK)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=10, color=BODY)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DEEP_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Noto Sans SC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name == "Heading 2" else before)
        style.paragraph_format.space_after = Pt(5 if style_name == "Heading 2" else after)


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="B7D6CE", size=4)
    cell = table.cell(0, 0)
    set_cell_shading(cell, TEAL_SOFT)
    set_cell_margins(cell, top=110, start=180, bottom=110, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(text)
    set_run_font(run, size=10, bold=True, color=INK)


def add_advantage_bullets(doc):
    advantages = [
        ("以课程目标为参照。", "先录入课程标准和校方量表，再分析学生是否出现与目标相关的可核验表现，避免把“坐得安静”直接等同于“学会了”。"),
        ("群像与个体同时成立。", "同一段课堂视频可建立多名学生的独立证据链，同时保留教师、材料、同伴和课堂阶段等环境信息。"),
        ("结论可以回到证据。", "每30秒形成时间切片，保留状态、置信度、证据质量、提示等级和待复核片段，教师可以确认、纠正或标记证据不足。"),
        ("分析能进入后续行动。", "课程级报告、个体档案和跨课趋势共同生成下一课建议，并可进一步映射到屏幕提示、语音提示或机器人克制动作。"),
    ]
    for title, body in advantages:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.16
        add_inline_label(paragraph, title, body)


def add_dimension_table(doc):
    data = [
        (("目标适配", "课程目标与学生既有能力是否匹配。"), ("可见参与", "是否持续朝向材料、教师或任务区域。")),
        (("任务正确性", "回答、选择或操作是否符合目标。"), ("独立程度", "完成任务需要何种语言、视觉或动作提示。")),
        (("重复一致性", "同类表现能否在多个时间段稳定复现。"), ("迁移表现", "能否在新材料或新情境中再次完成。")),
        (("支持效果", "教师提示前后，学生响应是否发生积极变化。"), ("发展趋势", "跨课程比较进步、波动、退化风险与证据不足。")),
    ]
    table = doc.add_table(rows=len(data), cols=2)
    set_table_geometry(table, [4680, 4680])
    set_table_borders(table, color=BORDER, size=4)
    for row_index, row_data in enumerate(data):
        for col_index, (title, body) in enumerate(row_data):
            cell = table.cell(row_index, col_index)
            set_cell_shading(cell, WHITE if row_index % 2 else BLUE_SOFT)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            label_run = paragraph.add_run(f"{title}：")
            set_run_font(label_run, size=9.5, bold=True, color=INK)
            body_run = paragraph.add_run(body)
            set_run_font(body_run, size=9.5, color=BODY)


def add_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("RehabAffectPerception  |  课堂学习证据与支持决策 Demo2")
    set_run_font(run, size=8, color=MUTED)


def build_docx():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.46)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)
    configure_styles(doc)
    add_footer(section)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("PROJECT BRIEF  |  DEMO2")
    set_run_font(run, size=8.5, bold=True, color=TEAL)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("课堂学习证据与支持决策 Demo2")
    set_run_font(run, size=20, bold=True, color=INK, name="Noto Serif SC")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("核心竞争力与学生课堂表现分析维度")
    set_run_font(run, size=11, color=MUTED)

    add_callout(
        doc,
        "本Demo面向辅读学校课堂，将低清课堂视频、语音、任务互动、课程标准和学生历史资料组织成可复核证据。核心不是给学生贴“专注或不专注”的标签，而是判断其在具体课程目标上表现了什么、需要什么支持，以及下一节课如何调整。",
    )

    doc.add_heading("一、具有竞争力的功能", level=2)
    add_advantage_bullets(doc)

    doc.add_heading("二、分析一个孩子课堂表现的八个维度", level=2)
    add_dimension_table(doc)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.line_spacing = 1.15
    lead = note.add_run("使用边界：")
    set_run_font(lead, size=8.8, bold=True, color=DEEP_BLUE)
    body = note.add_run("所有结果均为教学观察与辅助决策，不做医学诊断；证据不足时标记不确定，并由教师完成最终复核。")
    set_run_font(body, size=8.8, color=MUTED)

    doc.core_properties.title = "Demo2 核心竞争力与课堂表现分析维度"
    doc.core_properties.subject = "辅读学校课堂学习证据与支持决策"
    doc.core_properties.author = "RehabAffectPerception"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_docx()
