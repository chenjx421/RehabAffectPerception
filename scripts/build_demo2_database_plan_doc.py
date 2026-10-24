from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from build_demo2_phase_delivery_doc import (
    BLUE,
    INK,
    LIGHT_BLUE,
    LIGHT_GREEN,
    LIGHT_RED,
    MUTED,
    NAVY,
    add_bullet,
    add_callout,
    add_heading,
    add_number,
    add_table,
    add_text,
    configure_document,
    set_run,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "docs" / "RehabAffect_Demo2_本地数据库与四路视频上传规划_20260826.docx"


def configure_header(doc: Document) -> None:
    configure_document(doc)
    header = doc.sections[0].header.paragraphs[0]
    header.clear()
    set_run(header.add_run("RehabAffect | 本地数据库与四路视频上传规划"), size=8.5, bold=True, color=MUTED)


def add_masthead(doc: Document) -> None:
    add_text(doc, "技术与产品规划", bold=True, color=BLUE, size=10.5, after=7, before=8)
    add_text(doc, "RehabAffect Demo2", bold=True, color=NAVY, size=24, after=2)
    add_text(doc, "本地数据库、学生资料导入与四路视频上传规划", bold=True, color=INK, size=16.5, after=14)
    for label, value in [
        ("适用场景", "辅读学校课堂学习证据与个体化支持平台"),
        ("当前形态", "本地 SQLite + 受控文件目录 + 浏览器工作台"),
        ("规划范围", "学生资料、四路视频、分析任务、报告、权限、保存期限与具身动作接口"),
        ("版本日期", "2026年8月26日"),
        ("文档状态", "供公司与校方确认试点条件，不替代数据授权协议"),
    ]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2.5)
        set_run(paragraph.add_run(f"{label}："), size=10.5, bold=True, color=NAVY)
        set_run(paragraph.add_run(value), size=10.5, color=INK)
    add_callout(
        doc,
        "核心判断",
        "需求可以落地。四个“端口”应设计为同一平台内四个相互隔离的上传通道，而不是四套系统或四个网络端口。当前本地版已经具备 SQLite 数据层、学生资料上传、保存策略、四路分片上传、独立质量预检和结果下载；正式课堂状态分析仍需接入账号权限、任务队列、GPU模型和教师复核。",
        fill=LIGHT_GREEN,
        accent="2F6F60",
    )
    add_heading(doc, "本文件需要校方确认的重点", 2)
    for item in [
        "四个上传主体的真实身份：学生、教师、班级、设备还是摄像机视角。",
        "四段视频是否属于同一节课，以及是否需要自动时间对齐和多视角融合。",
        "视频数量、时长、码率、并发上传量和学校实际上行带宽。",
        "原视频、IEP、结构化档案和报告各自保存期限、下载范围与删除审批人。",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()


def build() -> Path:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_header(doc)
    add_masthead(doc)

    add_heading(doc, "1. 当前能力与产品边界", 1)
    add_text(doc, "本轮功能不是单纯增加两个上传按钮，而是补齐“资料进入系统、结构化确认、数据库留档、保存期限、四路任务、结果下载”这一条产品链路。")
    add_table(doc, ["能力", "当前已实现", "正式试点仍需补齐"], [
        ["学生资料导入", "IEP、五域、专业评估和安全支持资料上传；原件与结构化摘要分开保存", "账号权限、版本审批、批量导入和更完整的PDF/Excel解析"],
        ["本地数据库", "SQLite保存学生、资料元数据、支持摘要、四路任务、策略和审计", "把现有课程、证据窗口、教师复核和报告从JSON迁入数据库"],
        ["四路视频", "四个隔离通道、4MB分片、SHA-256、质量预检、独立JSON下载", "断点续传、后台队列、完整模型推理和多视角同步"],
        ["数据治理", "显示真实本地路径、空间、保存期限、下载和删除状态", "角色权限、自动到期任务、备份恢复和云端加密"],
    ], [1700, 3800, 3860])
    add_callout(doc, "不可夸大的部分", "当前预检分析判断视频能否被读取、画面质量是否适合后续处理，不等同于已经完成学生课堂状态模型分析。学生掌握度、课堂状态和具身动作建议仍需课程目标、人物/座位绑定、模型版本和教师复核。", fill=LIGHT_RED, accent="C65B4B")

    add_heading(doc, "2. 三条核心工作流", 1)
    add_heading(doc, "2.1 学生资料导入", 2)
    for item in [
        "教师选择学生并确认资料已获得校方授权。",
        "上传原文件，系统记录文件类型、来源、大小、SHA-256、上传时间和到期时间。",
        "DOCX、TXT、CSV和JSON生成文本预览；PDF、Excel和图片保存原件后由教师填写摘要。",
        "教师确认已有专业结论、教育目标、有效支持方式和安全边界。",
        "结构化摘要进入个体支持档案，并约束提示强度、机器人接近距离与交互方式。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2.2 四路视频上传", 2)
    for item in [
        "四个主体分别进入自己的通道，填写上传主体、课程、课堂批次号和观察对象。",
        "浏览器按4MB分片上传，服务端校验顺序、登记大小和最终文件哈希。",
        "四路任务默认相互隔离，各自显示进度、状态、到期时间和结果下载。",
        "上传完成后进行可读性、时长、分辨率、帧率、亮度、清晰度和运动量预检。",
        "相同session_id且获得授权的视频，后续可进入时间同步和多视角融合。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "2.3 课堂证据到具身动作", 2)
    add_text(doc, "产品数据库需要持续支撑以下闭环，而不能停留在文件管理：Perception（视频/音频/课程任务）→ State Estimation（状态与学习证据）→ Action Suggestion（教师、屏幕、语音和机器人动作建议）→ Teacher Confirmation（教师确认）→ Robot Action（执行）→ Feedback（学生后续变化与动作结果）。")

    add_heading(doc, "3. 本地数据库架构", 1)
    add_callout(doc, "当前落点", "SQLite文件：outputs/seewo_demo2/rehab_affect_demo2.sqlite3；学生原件：outputs/seewo_demo2/student_documents/；四路视频与结果：outputs/seewo_demo2/video_tasks/；课程报告：outputs/seewo_demo2/reports/。")
    add_heading(doc, "3.1 数据分层", 2)
    for item in [
        "身份层：学校、账号、角色、班级和学生匿名编号。",
        "教育资料层：IEP、五域评估、专业资料、结构化摘要与版本。",
        "课程层：课程、课堂批次、教学目标、考核标准和四路媒体。",
        "分析层：任务、模型版本、证据窗口、状态初评和置信度。",
        "人工判断层：教师复核、正确性、提示等级、画面外信息和定稿结论。",
        "行动层：教师支持、屏幕/语音/机器人动作、确认和执行反馈。",
        "治理层：保存策略、访问日志、下载、删除和授权状态。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "3.2 数据关系", 2)
    add_text(doc, "school → student → support profile / document；school → course → class session → four video tasks；video task → analysis run → evidence windows → teacher reviews → report；evidence window → recommended action → teacher confirmation → robot execution feedback。")

    doc.add_page_break()
    add_heading(doc, "4. 核心数据表规划", 1)
    add_table(doc, ["数据表", "用途", "实施状态"], [
        ["schools / users / roles", "学校租户、账号、教师/管理员/上传主体权限", "试点版新增"],
        ["students", "学生匿名编号、学校归属和档案状态", "基础表已实现"],
        ["student_documents", "IEP、五域和专业资料原件元数据", "已实现"],
        ["student_support_profiles", "教师确认的结构化支持与安全摘要", "已实现"],
        ["support_profile_versions", "历史版本、生效时间和确认人", "试点版新增"],
        ["courses / class_sessions", "课程目标、考核标准、一节课的session_id", "现有JSON待迁库"],
        ["video_upload_channels", "四个上传通道和主体范围", "已实现"],
        ["video_analysis_tasks", "分片上传、分析状态、结果和到期时间", "基础表已实现"],
        ["media_assets", "原视频、音频、切片、缩略图和对象存储地址", "试点版新增"],
        ["analysis_runs / evidence_windows", "模型版本、参数、证据窗口和状态初评", "现有JSON待迁库"],
        ["teacher_reviews", "片段复核、正确性、提示等级和教师说明", "现有JSON待迁库"],
        ["reports", "课程、个体和管理报告的版本与下载", "现有文件待登记"],
        ["robot_actions", "动作建议、教师确认、执行与学生反馈", "后续具身接口"],
        ["retention_policies / audit_logs", "保存期限、上传、下载、修改和删除留痕", "基础表已实现"],
    ], [2350, 4800, 2210])
    add_text(doc, "全校版应使用学校ID作为所有业务记录的隔离键。学生资料、课程视频和报告不能只靠文件夹名称隔离。", bold=True, color=NAVY, before=4)

    doc.add_page_break()
    add_heading(doc, "5. 保存周期与下载策略", 1)
    add_table(doc, ["数据类型", "默认", "建议范围", "到期处理"], [
        ["原始课堂视频/音频", "60天", "30/60/80/90天", "提醒后删除或审批延长"],
        ["自动生成临时切片", "30天", "7-60天", "报告定稿后自动删除"],
        ["IEP/专业评估原件", "365天", "180-730天", "管理员复核资料有效期"],
        ["脱敏结构化支持档案", "730天", "365天至长期授权", "版本化保留，撤回授权时删除"],
        ["教师复核记录", "365天", "180-730天", "归档或删除"],
        ["课程与个体报告", "365天", "180-730天", "支持下载后归档"],
        ["审计日志", "365天", "365-1095天", "归档保存"],
    ], [2800, 1100, 2400, 3060])
    add_callout(doc, "规则边界", "60天或80天是可配置建议，不是已经得到校方确认的制度。正式启用前必须确认数据是否允许离校、到期提醒时间、删除审批人、授权撤回机制和备份是否同步删除。", fill=LIGHT_BLUE, accent=BLUE)

    add_heading(doc, "6. 四路上传的硬件条件", 1)
    add_table(doc, ["位置", "建议配置", "承担工作"], [
        ["四个上传端", "普通Windows电脑；i5/Ryzen 5；8GB最低、16GB建议；20GB临时空间；Chrome/Edge", "选择课程与学生、分片上传、查看进度和下载结果；无需GPU"],
        ["采集设备", "1080p、H.264、15-25fps、3-5Mbps；固定支架；可选领夹麦/指向麦", "提供稳定画面和可用声音；远场说话人不清时不绑定具体学生"],
        ["学校网络", "单路建议10Mbps上行；四路同时建议总上行40Mbps以上；校内服务器使用千兆局域网", "保证四路上传稳定，减少课堂视频传输时间"],
        ["试点服务器", "8核CPU、32GB内存、1TB NVMe；加密与每日备份", "数据库、文件、四路队列、报告和本地服务"],
        ["推理GPU", "只做上传/数据库不需要；本地多路视觉推理建议12-24GB显存", "视觉、音频和多模态模型推理；Web服务与推理节点分离"],
    ], [1800, 4300, 3260])
    add_heading(doc, "6.1 容量估算", 2)
    add_text(doc, "30分钟、3-5Mbps的1080p H.264视频约为0.68-1.13GB。四路合计约2.7-4.5GB/节课。若每天2节课、原视频保存60天，原视频约324-540GB；再预留30%的数据库、切片、报告和运行空间后约421-702GB。因此单校试点可从1TB可用存储开始，最终按校方真实码率和课程量校准。")

    add_heading(doc, "7. 软件与部署条件", 1)
    add_heading(doc, "7.1 当前本地版", 2)
    for item in ["Python 3.10+、SQLite 3、OpenCV和Chrome/Edge。", "本地HTTP服务、受控文件目录和SQLite数据库。", "无需云账号即可演示；当前数据默认不离开本机授权目录。"]:
        add_bullet(doc, item)
    add_heading(doc, "7.2 学校内网试点版", 2)
    for item in [
        "FastAPI Web API、PostgreSQL、MinIO/学校对象存储。",
        "Redis + Celery/RQ管理四路异步任务，FFmpeg负责转码、音轨提取和时间戳。",
        "Nginx/HTTPS、账号登录、角色权限、访问日志和自动到期任务。",
        "CPU/GPU分析工作节点与页面服务分离，上传失败可续传、分析失败可重试。",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "7.3 云端或多学校版", 2)
    for item in [
        "每所学校独立租户和对象存储前缀，数据库使用租户级隔离。",
        "视频使用服务端加密和临时下载链接，GPU任务按队列扩容。",
        "具备备份恢复、到期删除、审计报表和安全告警。",
        "不同学校之间不能查看彼此的学生、视频和报告。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "8. 实施阶段与验收", 1)
    add_table(doc, ["阶段", "建设内容", "通过标准"], [
        ["A 本地验证", "SQLite、资料导入、四路上传、保存策略和独立预检", "四个主体可独立完成上传、下载，不发生串线"],
        ["B 校内试点", "账号、PostgreSQL、对象存储、任务队列、HTTPS、教师复核", "一个班级2-4节课稳定运行，权限与删除规则确认"],
        ["C 全校服务", "多班多课、模型校准、趋势报告、云/校级部署和具身动作接口", "容量、准确性、教师工作量和安全边界均通过验收"],
    ], [1700, 4300, 3360])
    add_heading(doc, "8.1 功能验收标准", 2)
    for item in [
        "四个通道各自创建任务，文件、身份、进度和结果不串线。",
        "每个任务有唯一编号、上传主体、课堂批次、SHA-256和到期时间。",
        "上传完成后产生可下载结果；完整课堂结论标记模型版本与教师复核状态。",
        "学生原件、结构化摘要和课堂证据可关联，但权限和保存周期相互独立。",
        "上传、查看、下载、修改和删除均可追溯。",
        "未经授权账号不能访问其他主体、班级或学校数据。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "9. 目前仍需校方明确的信息", 1)
    for item in [
        "四个主体是四名学生、四名教师、四个班级、四台设备还是四个摄像机视角？",
        "四段视频是否通常属于同一节课，是否要求自动同步与融合？",
        "典型视频时长、格式、码率、每天上传数量和四路同时上传比例是多少？",
        "每个主体可以查看哪些学生、课程和报告，管理员由谁担任？",
        "下载结果具体包含质量预检、学生状态、课程报告、教师复核表中的哪些内容？",
        "原视频、IEP、结构化档案和报告分别保存多久，是否允许离开学校？",
        "到期删除是自动执行还是管理员审批，需要提前几天提醒？",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "建议会议结论", "硬件与软件条件已经可以给出试点级配置，但采购和服务器容量不能在以上七项未明确时一次定死。最稳妥的下一步是让校方确认四主体定义、视频规模、并发量、权限和保存周期，然后用当前页面进行一次四路真实文件上传演练。", fill=LIGHT_GREEN, accent="2F6F60")

    doc.core_properties.title = "RehabAffect Demo2 本地数据库与四路视频上传规划"
    doc.core_properties.subject = "辅读学校课堂学习证据平台数据与部署规划"
    doc.core_properties.author = "RehabAffect 项目组"
    doc.core_properties.keywords = "SQLite, 本地数据库, 四路视频上传, 辅读学校, 数据治理, 具身智能"
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build())
