from __future__ import annotations

from pathlib import Path
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WECHAT_DIR = Path(
    r"C:\Path\To\Developer\Documents\xwechat_files\wxid_2m302j15cyrl22_2ea9\msg\file\2026-06"
)
OUTPUT_PROJECT = PROJECT_ROOT / "docs" / "special_education_ai_tutor_practical_solution_zh.docx"
OUTPUT_WECHAT = WECHAT_DIR / "special_education_ai_tutor_practical_solution_zh.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(30, 38, 45)
MUTED = RGBColor(95, 105, 115)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"


def set_run_font(run, *, size=None, bold=None, color=None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("辅读学校 AI 互动助教系统落地方案")
    set_run_font(r, size=9, color=MUTED)
    return doc


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("辅读学校 AI 互动助教系统落地方案")
    set_run_font(r, size=22, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("第一版聚焦：课堂常规/生活适应互动任务 + 注意力回流 + 教师复盘")
    set_run_font(r, size=11, color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("目标：先用低成本软件 Demo 证明学生端价值，再讨论机器人、智能坐垫等硬件增强。")
    set_run_font(r, size=10.5, color=INK, bold=True)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)


def add_para(doc: Document, text: str, *, bold=False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    r = p.add_run(text)
    set_run_font(r, color=INK, bold=bold)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, color=INK)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, color=INK)


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade_cell(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, bold=True, color=INK)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            if i == 0 and len(value) <= 12:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def build_doc() -> None:
    doc = setup_document()
    add_title(doc)

    add_callout(
        doc,
        "一句话方案",
        "做一个可在辅读学校课堂中试用的软件型 AI 互动助教：学生通过平板/大屏完成短任务，系统根据任务行为、摄像头轻量线索和教师标注判断学生是否参与、分心、困惑或需要帮助，再用低刺激提示、拆步说明、任务降难和教师端提醒帮助学生回到学习任务。",
    )

    add_heading(doc, "1. 方案定位")
    add_para(
        doc,
        "第一版不做机器人整机方案，也不做复杂硬件系统，而是做一个能进入课堂试用的软件 Demo。它的核心不是替老师管学生，而是让学生在更容易理解、更有反馈的互动任务中获得帮助，同时让老师看到学生卡在哪、什么时候需要关注。",
    )
    add_table(
        doc,
        ["定位项", "落地定义"],
        [
            ["核心用户", "学生是第一受益对象，老师是使用和复盘者，学校是试点和评估方。"],
            ["第一场景", "课堂常规/生活适应课中的 8-10 分钟互动环节。"],
            ["第一任务包", "听指令整理书包、找物品、课堂规则情境选择。"],
            ["第一载体", "学生端平板/大屏网页 + 教师端网页 + 本地 AI 状态服务。"],
            ["第一目标", "提高学生参与、帮助分心学生回到任务、帮助困惑学生理解下一步。"],
            ["不做什么", "第一版不采购机器人、不做智能坐垫、不保存可识别学生音视频。"],
        ],
        [1700, 7660],
    )

    add_heading(doc, "2. 首个落地场景")
    add_callout(
        doc,
        "推荐试点课：课堂物品整理与听指令任务",
        "老师让学生完成“找到上课需要的物品”“把物品放到正确位置”“判断课堂规则是否正确”等短任务。每个任务 30-60 秒，学生通过点击、拖拽或选择完成。系统记录学生是否长时间停顿、反复点错、主动求助或分心，并给出低刺激反馈。",
    )
    add_table(
        doc,
        ["课堂环节", "学生看到什么", "系统记录什么", "系统反馈什么"],
        [
            ["听指令", "屏幕显示 3-4 个物品，老师或系统提示“找到铅笔”。", "反应时、是否点错、是否无操作。", "点对后正向反馈；点错两次后缩小选项。"],
            ["整理物品", "把书、本子、铅笔拖进书包或桌面对应区域。", "拖拽顺序、错误位置、停顿时长。", "卡住时高亮目标区域，给出一步提示。"],
            ["课堂规则", "看图选择“应该怎么做”：举手、等待、看屏幕、找老师。", "选择结果、反复修改、求助按钮。", "错误后展示示例图，不用大声提醒。"],
            ["求助表达", "固定按钮：“我不会”“再说一遍”“休息一下”“找老师”。", "求助类型、触发时间、教师响应。", "教师端显示需要关注，学生端保持低刺激。"],
        ],
        [1400, 3200, 2500, 2260],
    )

    add_heading(doc, "3. 给学生解决的具体问题")
    add_table(
        doc,
        ["学生问题", "方案功能", "为什么能落地"],
        [
            ["不愿意参与", "短任务 + 即时反馈 + 简单成功体验。", "任务不长，学生不需要长时间坚持；完成一次就有反馈。"],
            ["听不懂指令", "把一句话拆成图片、语音和下一步提示。", "降低理解负担，适合能力差异大的班级。"],
            ["分心后回不来", "无操作/视线偏离/反复点错后触发轻提示。", "用屏幕高亮和短提示拉回，不突然打断课堂。"],
            ["反复失败后抗拒", "自动降难，减少选项，给示例。", "避免老师反复催促，也减少学生挫败。"],
            ["不会表达需要帮助", "提供固定求助按钮和需求表达选项。", "学生不用组织复杂语言，也能告诉老师自己卡住了。"],
        ],
        [1900, 3900, 3560],
    )

    add_heading(doc, "4. 给老师解决的具体问题")
    add_table(
        doc,
        ["老师问题", "方案功能", "老师得到什么"],
        [
            ["一节课要照看多个学生", "教师端展示每个学生的任务状态和风险提示。", "知道谁在参与、谁卡住、谁需要关注。"],
            ["不知道学生卡在哪", "记录任务步骤、错误类型、停顿时长和求助。", "课后能看到具体卡点，而不是只凭印象。"],
            ["提醒学生容易打断教学", "提示分级：只记录、学生轻提示、教师端提醒、同时提醒。", "老师可以控制提示强度。"],
            ["学生能力差异大", "任务支持简单/普通两档难度和一键降难。", "同一课堂里可以给不同学生不同难度。"],
            ["课后复盘费时间", "自动生成课堂报告。", "看到完成率、提示次数、求助次数和下次建议。"],
        ],
        [1900, 3900, 3560],
    )

    add_heading(doc, "5. 给学校解决的具体问题")
    add_table(
        doc,
        ["学校问题", "方案回应", "试点方式"],
        [
            ["预算有限", "第一版只用现有电脑、平板/大屏和普通摄像头。", "先做一节课 Demo，不采购复杂硬件。"],
            ["担心维护和损坏", "软件优先，机器人/智能坐垫后置。", "设备少，维护压力低。"],
            ["担心学生隐私", "默认只保存任务日志和匿名化特征，不保存原始音视频。", "试点前确认授权边界。"],
            ["担心效果难评估", "用参与率、回到任务率、完成率、老师认可度衡量。", "1-2 节课可形成初步数据。"],
            ["希望能扩展", "任务模板可替换。", "先做课堂常规，再扩展生活适应、出行安全、社交情境。"],
        ],
        [1800, 4300, 3260],
    )

    add_heading(doc, "6. 产品功能设计")
    add_table(
        doc,
        ["模块", "第一版功能", "交付形态"],
        [
            ["学生端任务", "听指令、物品整理、情境选择、求助表达。", "平板/大屏网页。"],
            ["状态感知", "基于无操作时长、错误次数、求助、摄像头轻量线索判断参与/分心/困惑/求助。", "本地 Python 服务。"],
            ["AI 反馈", "屏幕高亮、动画反馈、短语音、拆步提示、任务降难。", "学生端反馈策略。"],
            ["教师端", "实时状态卡片、提示开关、人工标注、课后报告。", "教师网页。"],
            ["数据报告", "完成率、卡住步骤、提示次数、求助次数、下次难度建议。", "网页报告/导出文件。"],
        ],
        [1500, 5200, 2660],
    )

    add_heading(doc, "7. 第一版 Demo 流程")
    add_number(doc, "老师打开教师端，选择“课堂物品整理”任务包，选择学生或小组。")
    add_number(doc, "学生端出现 5 个短任务，每个任务只需要点击、拖拽或选择。")
    add_number(doc, "系统记录学生操作：反应时、错误次数、停顿、求助、是否完成。")
    add_number(doc, "若学生 8 秒无操作或连续错误 2 次，系统先给学生端轻提示。")
    add_number(doc, "若提示后仍未恢复，教师端出现“建议关注”的提醒。")
    add_number(doc, "课后报告显示：哪个任务卡住、用了什么提示、是否回到任务、下次建议难度。")

    add_heading(doc, "8. 所需硬件与部署")
    add_table(
        doc,
        ["设备", "是否必需", "用途"],
        [
            ["Windows 笔记本", "必需", "运行本地服务、教师端和 Demo。"],
            ["平板/触控屏/大屏", "必需", "学生完成互动任务。"],
            ["普通摄像头", "建议", "观察头部朝向、坐姿、动作幅度等轻量线索。"],
            ["麦克风", "可选", "仅检测音量或异常发声，不作为第一版核心。"],
            ["机器人", "不必需", "后续在软件效果验证后再接入。"],
            ["智能坐垫", "不必需", "后续如学校有预算和维护能力再评估。"],
        ],
        [1900, 1700, 5760],
    )

    add_heading(doc, "9. 数据与隐私边界")
    add_bullet(doc, "第一版优先保存任务日志：点击、错误、停顿、求助、提示、完成情况。")
    add_bullet(doc, "摄像头只用于轻量特征，例如是否长时间偏离屏幕、动作幅度是否明显增加。")
    add_bullet(doc, "默认不保存可识别原始视频和音频。")
    add_bullet(doc, "如果学校希望保存图像或声音，必须先确认学校授权、家长授权、留存周期和使用范围。")
    add_bullet(doc, "系统只输出“可能分心/可能困惑/建议关注”，不做医学判断。")

    add_heading(doc, "10. 试点验收指标")
    add_table(
        doc,
        ["对象", "指标", "合格标准建议"],
        [
            ["学生", "任务开始率", "多数学生愿意开始至少 1 个互动任务。"],
            ["学生", "回到任务率", "轻提示后 15 秒内恢复操作的比例可被记录。"],
            ["学生", "完成率", "简单难度下能完成主要任务。"],
            ["学生", "求助表达", "学生能使用“我不会/再说一遍/找老师”等按钮。"],
            ["老师", "提示有效性", "老师认为系统提示不是乱报，而是有参考价值。"],
            ["老师", "复盘有用度", "课后报告能帮助老师知道学生卡点。"],
            ["学校", "部署可接受", "无需新增昂贵硬件即可跑通。"],
        ],
        [1300, 2300, 5760],
    )

    add_heading(doc, "11. 开发计划")
    add_table(
        doc,
        ["阶段", "周期", "工作内容", "交付物"],
        [
            ["第 1 阶段", "1 周", "完成任务包原型、学生端交互、教师端基础页面。", "可点击 Demo。"],
            ["第 2 阶段", "1 周", "接入任务日志、状态判断、低刺激反馈策略。", "完整课堂闭环。"],
            ["第 3 阶段", "1 周", "加入摄像头轻量线索、教师标注、课后报告。", "可展示 MVP。"],
            ["第 4 阶段", "1 周", "整理演示脚本、试点指标表、校方沟通问题。", "汇报材料和演示视频。"],
            ["后续", "1-2 月", "根据校方反馈扩展任务包、优化状态判断、评估硬件接入。", "试点版产品。"],
        ],
        [1300, 1200, 4600, 2260],
    )

    add_heading(doc, "12. 向校方确认的问题")
    questions = [
        "第一版最希望验证哪个学生效果：更愿意参与、分心后回到任务、听懂指令，还是能主动表达求助？",
        "学校更适合先试哪类任务包：课堂常规、生活适应、物品整理、出行安全，还是情境选择？",
        "老师判断学生分心/困惑/需要帮助时，最看重哪些表现？哪些表现只是正常波动？",
        "反馈应该先给学生还是先给老师？声音、动画、提示频率有没有禁忌？",
        "学校现有设备有哪些？是否有平板/大屏/摄像头/网络？第一版适合单人、小组还是整班试用？",
        "任务日志和摄像头轻量特征是否可以保存？是否需要家长授权？",
    ]
    for question in questions:
        add_bullet(doc, question)

    add_heading(doc, "13. 后续升级方向")
    add_para(
        doc,
        "如果第一版软件 Demo 证明学生端价值，再考虑把反馈策略迁移到硬件。机器人可以承担转向、表情、短句提示；智能坐垫可以作为坐姿和离座趋势补充；MuJoCo 或 Isaac Sim 可以用于策略 Benchmark 和展示视频。但这些都不是第一版成败的前提。",
    )

    add_heading(doc, "14. 结论")
    add_callout(
        doc,
        "落地结论",
        "本方案第一版聚焦一个能进课堂的小闭环：学生做互动任务，系统识别停顿、错误、分心和求助，学生端给出低刺激帮助，教师端看到状态和复盘。它用软件先解决学生参与和注意力回流问题，用数据给老师和学校证明价值，再决定是否投入机器人等硬件。",
    )

    OUTPUT_PROJECT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PROJECT)
    if WECHAT_DIR.exists():
        shutil.copyfile(OUTPUT_PROJECT, OUTPUT_WECHAT)
    print(OUTPUT_PROJECT)
    if WECHAT_DIR.exists():
        print(OUTPUT_WECHAT)


if __name__ == "__main__":
    build_doc()
