"""Rebuild the reference-style explanation and complete source appendix in one Word."""
import hashlib
import json
from pathlib import Path
from copy import deepcopy
import zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from scripts import build_seewo_source_code_docx as old
from scripts.build_source_delivery_20261011 import MODULES

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'outputs/delivery_20261011/RehabAffect_Demo2_Source_20261011/source'
REFERENCE=Path('C:/Path/To/Developer/Documents/xwechat_files/wxid_2m302j15cyrl22_2ea9/msg/file/2026-10/辅读学校Demo_源码汇编与通俗说明_20260731.docx')
OUT=ROOT/'outputs/source_word_20261012'
TARGET=OUT/'辅读学校Demo2_源码汇编与系统说明_20261012.docx'

def paragraph(doc,text):old._add_paragraph(doc,text)

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    files=[]
    categories={'src':'核心源代码','tests':'自动测试','configs':'配置文件','scripts':'隔离验收工具'}
    for path in SOURCE.rglob('*'):
        if path.is_file():
            rel=path.relative_to(SOURCE)
            if path.suffix not in {'.py','.js','.json','.txt','.toml','.ps1'}:raise ValueError(str(rel))
            text=path.read_text(encoding='utf-8-sig')
            files.append(old.SourceFile(path,rel.as_posix(),categories.get(rel.parts[0],'环境与启动脚本'),path.suffix[1:].upper(),len(text.splitlines())))
    order=['核心源代码','自动测试','配置文件','隔离验收工具','环境与启动脚本']
    files.sort(key=lambda f:(order.index(f.category),f.relative))
    for f in files:
        if f.relative.startswith('src/'):
            assert f.path.read_bytes()==(ROOT/f.relative).read_bytes(),'Current source changed: '+f.relative
    total=sum(f.lines for f in files)
    (OUT/'artifact.md').write_text(f'''# Reference contract
Reference: {REFERENCE.name}
SHA256: {hashlib.sha256(REFERENCE.read_bytes()).hexdigest()}
Structure retained: cover, reading guidance, contents, system explanation, architecture, page functions, delivery status, file index, complete landscape source appendix.
Styles retained from the reference document. Title is black without a border. Current claims replace old measurements and obsolete features.
Source scope: all {len(files)} deliverable text files, {total} physical lines, no live datasets or credentials. Line numbers and expanded tabs are reading aids only.
''',encoding='utf8')
    d=Document(REFERENCE)
    body=d._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'):body.remove(child)
    for key,relationship in list(d.part.rels.items()):
        if relationship.reltype.rsplit('/',1)[-1] in {'image','customXml','hyperlink'}:
            d.part.drop_rel(key)
    sec=d.sections[0];sec.orientation=WD_ORIENT.PORTRAIT;sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.left_margin=sec.right_margin=Inches(1);sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.75)
    sec.header_distance=sec.footer_distance=Inches(.3)
    old._set_header_footer(sec)
    sec.header.paragraphs[0].runs[0].text='辅读课堂学生追踪与双轨评价平台'
    title=d.styles['Title'];title.font.color.rgb=RGBColor(0,0,0)
    for border in list(title.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.add_paragraph('内部技术交付资料')
    d.add_paragraph('辅读学校 Demo2',style='Title')
    d.add_paragraph('源码汇编与系统说明',style='Title')
    table=d.add_table(rows=0,cols=2)
    for label,value in [('文档用途','当前系统说明与完整源码归档'),('系统版本','Demo2 当前工作副本'),('汇编范围',f'{len(files)}个文件  约{total:,}行'),('生成日期','2026年10月12日'),('资料级别','内部受限  不含原始学生资料与运行密钥')]:
        cells=table.add_row().cells;old._set_cell_text(cells[0],label,bold=True);old._set_cell_text(cells[1],value)
    old._set_table_geometry(table,[2300,7060]);old._set_table_borders(table)
    d.add_paragraph()
    old._add_callout(d,'阅读建议','前半部分说明当前系统如何工作、页面功能和交付状态；后半部分按文件完整收录源代码。它不是代码节选，也不需要另开TXT才能查看完整代码。')
    paragraph(d,'本次以2026年10月11日干净代码交付包为基础，10月12日再次核对运行源码一致。账号密码、学校视频、IEP原件、数据库和API密钥不在本册中。')
    d.add_page_break();old._add_toc(d);d.add_page_break()
    d.add_heading('一 项目说明',1)
    paragraph(d,'系统将课程视频、学生身份、个别化支持资料和教师观察关联到同一课程任务，帮助学校回看学生表现、整理证据、追踪目标，并比较学校标准与系统增强评价。公司后台负责通道、任务、保留期限及授权范围内的结果管理。')
    d.add_heading('当前已经实现',2)
    for s in ['四个上传账号与公司管理员分权，按学校工作区组织名册、课程和任务。','视频分片上传、排队处理、画面采样、同期音轨提取、媒体回放与结果下载。','自动人物候选入口与手动画框，关联学生ID，支持撤回、更正和版本校验。','IEP、专业评估及支持资料的授权、加密、期限管理和课堂版本引用。','群像与个体观察、教师复核、目标复用、千字图文报告、长期测量与双轨评课。','教学支持、逐课改进记录及独立安全事件人工核查。']:
        old._add_bullet(d,s)
    d.add_heading('需要明确的边界',2)
    paragraph(d,'视频处理成功只表示证据材料准备就绪，不意味着模型已经可靠理解整堂课。当前不能承诺无人参与的学生掌握度判断、医学诊断、自动伤害认定或经过校方验证的评课准确率。语音能量与音轨提取不是语音转写，也不能证明某个孩子回答了问题。')
    paragraph(d,'普通课堂的系统评分以学校评分为基线；提出调整必须有证据、理由和复核记录。不能为了让分数好看统一抬分，也不能将学生需要支持或未说话直接扣到教师身上。')
    d.add_heading('二 系统工作流程与代码模块',1)
    paragraph(d,'登录并确认学校 → 建立课程和学生名册 → 上传视频 → 媒体处理 → 人物与姓名匹配 → 关联授权资料 → 记录或复核片段 → 群像与个体汇总 → 图文报告和双轨评价 → 公司后台查阅与下载。')
    paragraph(d,'前端主要为Python模块内嵌HTML、CSS与JavaScript；后端由Python服务分派HTTP接口，SQLite保存结构化记录，原媒体与加密资料文件分开保存。后台工作线程处理媒体和保留期限扫描。')
    for name,paths,text in MODULES:
        d.add_heading(name,2)
        p=d.add_paragraph(paths);p.paragraph_format.keep_with_next=True
        paragraph(d,text)
    d.add_heading('三 页面功能与使用方式',1)
    for name,text in [
        ('登录 学校与四路上传','学校人员使用用户入口，公司人员使用管理员入口。四账号分别绑定通道，名额由后端控制。视频由使用者从本地手动选择，不会自动从希沃后台取得。短信注册、支付和正式云端部署不属于本次已经验收的完整链路。'),
        ('人物与学生匹配','先建立名册，再将采样画面中的位置对应到学生。自动候选不是最终身份结论；遮挡或检测遗漏时可手动画框。同名通过学生编号区分，具体框和版本需要可追溯。'),
        ('IEP与资料引用','教师在授权范围内导入资料并记录来源和摘要。课堂引用具体资料版本，更新档案不应悄悄改写旧课依据。公司汇总权限不等于自动获得诊断原件阅读权。'),
        ('时间片段与教师复核','点击时间段可定位到相关视频位置，与同期音轨共同回看。复核记录应具体到目标、机会、提示条件和表现；复用教学目标时不把旧课堂证据一并复制为新结论。'),
        ('群像 个体与千字报告','群像用于同一课程的整体浏览，个体页查看某名学生。报告结合指标、证据与支持建议；数据不足时明确留空或待复核。正文使用当前证据模板，并非默认调用外部模型生成诊断。'),
        ('学生进步与教师双轨评价','学生追踪以同一目标和支持条件为比较前提；每课、每周、每月与滚动窗口组织记录。教师评价保留学校轨与增强轨，并展示差异来源。项目参考线和试行权重不能冒充官方验证量表。'),
        ('公司后台与存储','管理员查看任务、通道、到期状态和可下载结果。原媒体默认60天、提前15天提醒；服务运行时扫描。长期结构化资源需要额外授权，当前本地保存不能表述为已经完成公司云端备份。'),
    ]:
        d.add_heading(name,2);paragraph(d,text)
    d.add_heading('四 开发测试与交付状态',1)
    validation=json.loads((SOURCE.parent/'validation.json').read_text(encoding='utf8'))
    paragraph(d,'2026年10月11日已在独立交付副本上验证：空库初始化成功、再次初始化拒绝覆盖已有数据库、登录页面可启动、四个新用户账号和管理员均能登录并打开各自工作区。选取的评分规则与参考上下文测试79项通过。运行源码本次核对未变化。')
    paragraph(d,'上述是指定范围的工程验证，不是所有历史测试均已通过，也不是模型准确率。四路公网高清整课并发、真实课堂语音与学生识别效果、稳定托管、备份恢复和长期运维仍需现场或正式部署环境验收。')
    d.add_heading('安装与交接',2)
    paragraph(d,'代码交付包source目录包含install.ps1、initialize_local.py和start.ps1。安装Python 3.10以上后运行安装脚本，再运行启动脚本；首次生成新数据库和随机密码。视频处理另需FFmpeg及FFprobe；Windows敏感档案密钥受当前系统用户保护，跨机器迁移需单独设计密钥恢复。')
    paragraph(d,'本册是源码阅读和归档载体，行号不是原始程序内容。执行时使用配套源码包，不直接将Word排版文本另存为程序。当前机器密码、旧邀请链接与学生数据不随代码流转。历史仿真模块因导入依赖保留，不构成本次公司的硬件或机器人交付承诺。')
    d.add_heading('五 完整源码索引',1)
    paragraph(d,f'以下{len(files)}个文件全部在后文逐一展开。运行结果、缓存、数据库、个人档案和密钥不在收录范围内；含历史实名学生的旧测试已在代码包排除清单登记。')
    table=d.add_table(rows=1,cols=4)
    for cell,value in zip(table.rows[0].cells,['序号','文件路径','类别','行数']):old._set_cell_text(cell,value,bold=True)
    old._set_repeat_table_header(table.rows[0])
    for i,f in enumerate(files,1):
        cells=table.add_row().cells
        for cell,value in zip(cells,[str(i),f.relative,f.category,str(f.lines)]):old._set_cell_text(cell,value)
    old._set_table_geometry(table,[650,5700,1900,1110]);old._set_table_borders(table)
    sec=d.add_section(WD_SECTION.NEW_PAGE);sec.orientation=WD_ORIENT.LANDSCAPE;sec.page_width=Inches(11);sec.page_height=Inches(8.5)
    sec.left_margin=sec.right_margin=Inches(.45);sec.top_margin=sec.bottom_margin=Inches(.5);sec.header_distance=sec.footer_distance=Inches(.22)
    old._set_header_footer(sec,appendix=True)
    d.add_heading('附录 A 完整源程序代码',1)
    paragraph(d,'按索引顺序完整收录。左侧行号用于定位，制表符按四个空格显示；长行自动折行不代表新增代码行。每个文件的SHA256及收录完整性由配套校验清单记录。')
    inventory=[]
    for number,f in enumerate(files,1):
        heading=d.add_paragraph(style='Source File Heading');heading.paragraph_format.page_break_before=True
        heading.add_run(f'{number:03d}  {f.relative}  |  {f.lines:,} 行')
        lines=f.path.read_text(encoding='utf-8-sig').splitlines()
        for start in range(0,len(lines),60):
            text='\n'.join(f'{i+1:05d}  {lines[i].expandtabs(4)}' for i in range(start,min(start+60,len(lines))))
            p=d.add_paragraph(style='Source Code');p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
            p.paragraph_format.widow_control=False
            p.add_run(text)
        inventory.append({'file':f.relative,'lines':f.lines,'sha256':hashlib.sha256(f.path.read_bytes()).hexdigest()})
    old._update_fields_on_open(d)
    d.core_properties.title='辅读学校Demo2源码汇编与系统说明';d.core_properties.subject='当前源码全文与系统交付说明';d.core_properties.author='RehabAffectPerception'
    d.save(TARGET)
    # Validate every source line from the actual Word paragraph content.
    loaded=Document(TARGET);actual=[]
    for p in loaded.paragraphs:
        if p.style.name=='Source Code':actual.extend(p.text.split('\n'))
    expected=[f'{i+1:05d}  {line.expandtabs(4)}' for f in files for i,line in enumerate(f.path.read_text(encoding='utf-8-sig').splitlines())]
    assert actual==expected,'Source appendix differs from source files'
    report={'docx':str(TARGET),'files':len(files),'lines':total,'all_lines_verified':True,'source_matches_current':True,'inventory':inventory,'validation_snapshot':validation}
    (OUT/'source_manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:v for k,v in report.items() if k not in ['inventory','validation_snapshot']},ensure_ascii=False))

if __name__=='__main__':main()
